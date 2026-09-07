"""Compare official YYH-US Inactive listings with Mabang Excel 7-day sales.

Run from the repository root with uv run --frozen python scripts/check_yyh_inactive_sales.py.
Default: download a fresh source. --source-xlsx uses an explicitly supplied snapshot.
All evidence is written to a new, ignored output directory, never production datasets.
"""
from __future__ import annotations

import argparse
import asyncio
import csv
import hashlib
import json
import os
import re
import shutil
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from services.mabang import official_api
from services.mabang.amazon.fba import store_msku, store_resolver
from services.mabang.amazon.fba.combo_sku import normalize_sku_key
from services.mabang.auth import MabangAuthContext
from shared.infra.net import close_all_network_clients
from saihu_msku_compare import load_env, redact

STORE = 'Amazon-YYH-US'
SID = '2021143528'
SITE = 'us'
WEB_ID = '1039477'


def now():
    return datetime.now(timezone.utc).isoformat()


def save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + '\n')


def clean(value):
    return '' if value is None else str(value).strip()


def integer(value):
    if isinstance(value, bool) or not re.fullmatch(r'\d+', str(value)):
        raise ValueError(f'无效分页整数: {value!r}')
    return int(value)


async def fetch_inactive(post, page_size=1000):
    result, fingerprints = [], set()
    expected = None
    page = 1
    while True:
        payload = await post('listings/search', {
            'shop_id': [SID], 'amazonsite': [SITE], 'pStatus': ['Inactive'],
            'page': str(page), 'pageSize': str(page_size),
        })
        data = payload.get('data')
        if str(payload.get('code')) != '200' or not isinstance(data, dict) or not isinstance(data.get('list'), list):
            official_api.invalid(f'Inactive page={page}', '缺少成功码或 data.list', payload)
        total, pages, current = (integer(data.get(k)) for k in ('total', 'totalPage', 'nowPage'))
        if current != page or pages not in ({0, 1} if total == 0 else {(total + page_size - 1)//page_size}):
            official_api.invalid(f'Inactive page={page}', '分页数量矛盾', payload)
        if expected is None:
            expected = total, pages
        if expected != (total, pages) or len(data['list']) != min(page_size, total-len(result)):
            official_api.invalid(f'Inactive page={page}', '分页总数变化或记录不完整', payload)
        fingerprint = hashlib.sha256(json.dumps(data['list'], sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        if fingerprint in fingerprints:
            official_api.invalid(f'Inactive page={page}', '重复分页', payload)
        fingerprints.add(fingerprint)
        for row in data['list']:
            if not isinstance(row, dict):
                official_api.invalid(f'Inactive page={page}', 'Listing 必须为对象', row)
            shops = {s.strip() for s in clean(row.get('shopIds')).split(',') if s.strip()}
            if shops != {SID} or clean(row.get('amazonsite')).lower() != SITE:
                official_api.invalid(f'Inactive page={page}', '店铺或站点不符', row)
            if clean(row.get('pStatus')).casefold() not in {'inactive', '停售'}:
                official_api.invalid(f'Inactive page={page}', '返回状态不是停售', row)
        result.extend(data['list'])
        print(f'停售 Listing {len(result)}/{total}，page={page}/{pages}', file=sys.stderr, flush=True)
        if len(result) == total:
            return result, page
        page += 1


def read_source(path):
    if not re.fullmatch(r'\d{12}-Amazon-YYH-US_店铺MSKU数据\.xlsx', path.name):
        raise ValueError(f'必须指定 YYH-US 原始 MSKU 源表: {path.name}')
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        values = wb.worksheets[0].iter_rows(values_only=True)
        headers = [clean(x) for x in next(values, ())]
        for field in ['MSKU', 'ASIN', '本地SKU', '7天销量']:
            if headers.count(field) != 1:
                raise ValueError(f'源表缺少或重复列: {field}')
        rows = []
        for number, cells in enumerate(values, 2):
            if not any(v is not None and clean(v) for v in cells):
                continue
            record = dict(zip(headers, cells))
            if clean(record.get('店铺名称')) not in {'', STORE, STORE+'美国'}:
                raise ValueError(f'源表混入其他店铺: 行={number}, 店铺={record["店铺名称"]}')
            rows.append({'msku': record['MSKU'], 'asin': record['ASIN'], 'local_sku': record['本地SKU'], 'sales_7d': record['7天销量'], 'excel_row': number})
        return rows
    finally:
        wb.close()


def sales_number(value):
    if isinstance(value, bool) or value is None or not clean(value):
        raise ValueError('销量为空或类型无效')
    text = clean(value)
    if ',' in text:
        if not re.fullmatch(r'\d{1,3}(,\d{3})+(\.0+)?', text):
            raise ValueError('销量千位分隔格式无效')
        text = text.replace(',', '')
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError('销量不是数值') from exc
    if not number.is_finite() or number < 0 or number != number.to_integral_value():
        raise ValueError('销量必须为有限非负整数')
    return int(number)


def compare(listings, source_rows):
    source, groups = defaultdict(list), defaultdict(list)
    for row in source_rows:
        source[(normalize_sku_key(row['msku']), clean(row['asin']))].append(row)
    for row in listings:
        groups[(normalize_sku_key(row.get('platformSku')), clean(row.get('asin')))].append(row)
    details = []
    for (msku, asin), entries in sorted(groups.items()):
        matches = source.get((msku, asin), []) if msku and asin else []
        reasons = []
        definitions = {(normalize_sku_key(e.get('stockSku')), clean(e.get('stockType'))) for e in entries}
        if not msku or not asin:
            reasons.append('Listing 缺少 MSKU 或 ASIN')
        if len(definitions) != 1:
            reasons.append('Listing 重复绑定冲突')
        if not matches:
            reasons.append('源表没有对应 MSKU+ASIN')
        parsed = []
        for match in matches:
            try:
                parsed.append(sales_number(match['sales_7d']))
            except ValueError as exc:
                reasons.append(f'源表行{match["excel_row"]}: {exc}')
        if len(set(parsed)) > 1 or len({normalize_sku_key(m['local_sku']) for m in matches}) > 1:
            reasons.append('源表重复记录冲突')
        quantity = parsed[0] if parsed and not reasons else None
        details.append({
            'MSKU': msku, 'ASIN': asin,
            'Listing本地SKU': '；'.join(sorted({clean(e.get('stockSku')) for e in entries})),
            '源表本地SKU': '；'.join(sorted({clean(m['local_sku']) for m in matches})),
            '平台状态': '；'.join(sorted({clean(e.get('pStatus')) for e in entries})),
            '近7天销量': quantity,
            '分类': '无法判断' if quantity is None else ('零销量' if quantity == 0 else '非零销量'),
            '原因': '；'.join(dict.fromkeys(reasons)),
            '源表行号': ','.join(str(m['excel_row']) for m in matches),
            '源表原始销量': json.dumps([m['sales_7d'] for m in matches], ensure_ascii=False, default=str),
            'Listing记录数': len(entries),
        })
    counts = Counter(d['分类'] for d in details)
    if counts['非零销量']:
        verdict = '不是全部为零'
    elif not details:
        verdict = '本轮没有停售记录'
    elif counts['无法判断']:
        verdict = '已确认部分均为零，但无法确认全部' if counts['零销量'] else '没有可确认的销量，无法判断'
    else:
        verdict = '本轮停售 MSKU 近7天销量全部为零'
    details.sort(key=lambda d: ({'非零销量': 0, '无法判断': 1, '零销量': 2}[d['分类']], -(d['近7天销量'] or 0), d['MSKU'], d['ASIN']))
    summary = dict(
        verdict=verdict, listing_raw_count=len(listings), inactive_unique_count=len(details),
        inactive_unique_msku_count=len({d['MSKU'] for d in details}),
        source_row_count=len(source_rows), zero_count=counts['零销量'], positive_count=counts['非零销量'],
        unknown_count=counts['无法判断'],
        matched_source_count=sum(bool(d['源表行号']) for d in details),
        listing_duplicate_count=len(listings)-len(details),
        positive_sales_total=sum(d['近7天销量'] for d in details if d['分类'] == '非零销量'),
        positive_examples=[d for d in details if d['分类'] == '非零销量'][:10],
    )
    assert summary['inactive_unique_count'] == sum(counts.values())
    return summary, details


def write_csv(path, details):
    fields = ['MSKU', 'ASIN', 'Listing本地SKU', '源表本地SKU', '平台状态', '近7天销量', '分类', '原因', '源表行号', '源表原始销量', 'Listing记录数']
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fields)
        writer.writeheader()
        for row in details:
            # Keep externally supplied text from becoming an Excel formula.
            writer.writerow({k: "'"+v if isinstance(v, str) and v.lstrip().startswith(('=', '+', '-', '@')) else v for k, v in row.items()})


def install_auth_snapshot(path):
    state = json.loads(path.read_text())
    grouped = defaultdict(list)
    for cookie in state['cookies']:
        grouped[cookie['domain']].append(cookie)
    async def context(**kwargs):
        return MabangAuthContext('', 'explicit_snapshot', dict(grouped), '', '')
    store_msku.get_auth_context = context
    store_resolver.get_auth_context = context


async def run(args):
    out = args.output_dir.resolve()
    out.mkdir(parents=True, exist_ok=False)
    (out/'.gitignore').write_text('*\n')
    if args.env_file:
        load_env(args.env_file)
        os.environ['LXE_DATA_SERVER_URL'] = os.environ.get('LXE_DATA_SERVER_PRIVATE_URL') or os.environ.get('LXE_DATA_SERVER_URL', '')
    if args.auth_state:
        install_auth_snapshot(args.auth_state)
    ledger = []
    async def post(endpoint, body):
        event = dict(endpoint=endpoint, request=body, started=now())
        try:
            payload = await official_api.post_json(endpoint, body, context=f'{STORE} 停售核查 page={body.get("page", "-")}')
            event['response'] = redact(payload, [os.environ.get('LXE_DATA_SERVER_API_KEY', '')])
            return payload
        except Exception as exc:
            event['error'] = official_api.diagnostic(f'{type(exc).__name__}: {exc}')
            raise
        finally:
            event['ended'] = now()
            ledger.append(event)
            save(out/'raw'/f'{len(ledger):03d}.json', event)
    metadata = {'store_name': STORE, 'sid': SID, 'site': SITE, 'started': now()}
    try:
        metadata['source_started'] = now()
        if args.source_xlsx:
            source_path = out/'source'/args.source_xlsx.name
            source_path.parent.mkdir()
            shutil.copyfile(args.source_xlsx, source_path)
            metadata['source_mode'] = 'supplied_snapshot'
        else:
            selected = await store_resolver.resolve_fba_store(STORE)
            if (selected.store.store_name, selected.store.store_id, selected.store.id_type) != (STORE, WEB_ID, 'fbaWarehouseIds[]'):
                raise ValueError(f'网页店铺身份变化: {selected.to_payload()}')
            result = await store_msku.download_store_msku_excel(WEB_ID, 'fbaWarehouseIds[]', store_name=STORE, output_dir=out/'source')
            source_path = Path(result.xlsx_path)
            metadata['source_mode'] = 'fresh_download'
        metadata.update(source_ended=now(), source_xlsx=str(source_path), source_data_time=source_path.name[:12], source_sha256=hashlib.sha256(source_path.read_bytes()).hexdigest())
        rows = read_source(source_path)
        metadata['listing_started'] = now()
        shops = await post('shops/list', {})
        if not isinstance(shops.get('data'), dict):
            official_api.invalid('店铺验证', 'data 必须为对象', shops)
        matches = [r for r in shops['data'].values() if isinstance(r, dict) and r.get('name') == STORE]
        if len(matches) != 1 or str(matches[0].get('sid')) != SID or clean(matches[0].get('amazonsite')).lower() != SITE:
            official_api.invalid('店铺验证', '官方店铺身份变化', shops)
        listings, pages = await fetch_inactive(post)
        metadata.update(listing_ended=now(), listing_pages=pages, logical_official_requests=len(ledger))
        summary, details = compare(listings, rows)
        summary.update(metadata, completed=now())
        write_csv(out/'details.csv', details)
        save(out/'summary.json', summary)
        print(json.dumps(summary, ensure_ascii=False, default=str))
        return 0
    except Exception as exc:
        save(out/'error.json', {**metadata, 'error': official_api.diagnostic(f'{type(exc).__name__}: {exc}')})
        print(official_api.diagnostic(f'{type(exc).__name__}: {exc}'), file=sys.stderr)
        return 1
    finally:
        await close_all_network_clients()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-xlsx', type=Path)
    parser.add_argument('--env-file', type=Path)
    parser.add_argument('--auth-state', type=Path, help='Optional explicit browser state; never copied to outputs')
    parser.add_argument('--output-dir', type=Path, default=Path('outputs')/('yyh-inactive-sales-'+datetime.now().strftime('%Y%m%d-%H%M%S')))
    return asyncio.run(run(parser.parse_args()))


if __name__ == '__main__':
    sys.exit(main())
