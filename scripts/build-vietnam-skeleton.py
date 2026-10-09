"""Rebuild the shipped Vietnam layout from audited parts of a local template.

The source workbook is a maintainer-only reference and must never be committed.
This script creates a new workbook; it does not save or copy the source archive.
Formula/header fingerprints deliberately require review when business rules change.
"""

from __future__ import annotations

from argparse import ArgumentParser
from copy import copy
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.formula import ArrayFormula


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = (
    REPOSITORY_ROOT
    / "python/lxeskill_cli/services/vietnam_replenishment/resources/skeleton.xlsx"
)
SHEETS = (
    ("越南备货清单", 51),
    ("雅仓库存", 13),
    ("雅仓动销", 16),
    ("数据更改", 4),
    ("库存商品信息", 11),
)
FORMULA_COLUMNS = (
    "H",
    *(get_column_letter(column) for column in range(10, 30)),  # J:AC
    "AF", "AG", "AH", "AI",
    "AK", "AL", "AM",
    "AV", "AW", "AX", "AY",
)
DEFAULT_PARAMETERS = (0.8, 0.8, 0, 3900, 0.1, 0.3, 0.6)

# SHA-256 of JSON [(sheet, row-1 values)] and [(column, row-2 formula text)].
# These were audited against the supplied 9-sheet template. Updating either
# fingerprint is a business-rule change, not routine resource regeneration.
AUDITED_HEADER_SHA256 = "acd374ac17a797c729a80b0134c2d8a77af6f8fe7d18e6823df5f39a2c254214"
AUDITED_FORMULA_SHA256 = "6e83bab0064c21fbd2434c811fa09b3d3dac79f96e468c5c7050eccdc96f2d50"


def _fingerprint(value: object) -> str:
    encoded = json.dumps(value, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return sha256(encoded).hexdigest()


def _formula_text(cell: object) -> str:
    if cell.data_type != "f":
        raise ValueError(f"Expected formula at {cell.parent.title}!{cell.coordinate}")
    value = cell.value
    if isinstance(value, ArrayFormula):
        if value.ref != cell.coordinate or not isinstance(value.text, str):
            raise ValueError(f"Unsupported array formula at {cell.parent.title}!{cell.coordinate}")
        return value.text
    if not isinstance(value, str):
        raise ValueError(f"Unsupported formula at {cell.parent.title}!{cell.coordinate}")
    return value


def _copy_cell_style(source: object, target: object) -> None:
    # Copy only these formatting fields. In particular, never carry values,
    # comments, hyperlinks, images, tables, defined names, or relationships.
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format


def build_skeleton(source_path: Path, output_path: Path) -> Path:
    if source_path.resolve() == output_path.resolve():
        raise ValueError("Source template and output skeleton must be different files")
    source = load_workbook(source_path, read_only=False, data_only=False, keep_links=False)
    try:
        missing = [name for name, _ in SHEETS if name not in source]
        if missing:
            raise ValueError(f"Source template is missing required sheets: {missing}")

        headers = [
            (name, [source[name].cell(1, column).value for column in range(1, width + 1)])
            for name, width in SHEETS
        ]
        if _fingerprint(headers) != AUDITED_HEADER_SHA256:
            raise ValueError("Source template headers differ from audited Vietnam layout")

        source_main = source["越南备货清单"]
        formulas = [
            (column, _formula_text(source_main[f"{column}2"]))
            for column in FORMULA_COLUMNS
        ]
        if _fingerprint(formulas) != AUDITED_FORMULA_SHA256:
            raise ValueError("Source template formulas differ from audited Vietnam rules")

        result = Workbook()
        result.remove(result.active)
        result.properties.creator = "LXE Agent"
        result.properties.lastModifiedBy = "LXE Agent"
        result.properties.title = "越南备货清单"
        result.properties.subject = "五表计算骨架"

        for name, width in SHEETS:
            original = source[name]
            target = result.create_sheet(name)
            target.sheet_view.showGridLines = original.sheet_view.showGridLines
            for column_number in range(1, width + 1):
                letter = get_column_letter(column_number)
                dimension = original.column_dimensions.get(letter)
                if dimension is not None and dimension.width is not None:
                    target.column_dimensions[letter].width = dimension.width
                for row_number in (1, 2):
                    _copy_cell_style(original.cell(row_number, column_number), target.cell(row_number, column_number))
                target.cell(1, column_number).value = original.cell(1, column_number).value
            for row_number in (1, 2):
                height = original.row_dimensions[row_number].height
                if height is not None:
                    target.row_dimensions[row_number].height = height
            target.freeze_panes = original.freeze_panes
            if name != "数据更改":
                target.auto_filter.ref = f"A1:{get_column_letter(width)}2"

        target_main = result["越南备货清单"]
        for column in FORMULA_COLUMNS:
            text = _formula_text(source_main[f"{column}2"])
            if column == "AC":
                target_main["AC2"] = ArrayFormula("AC2", text=text)
            else:
                target_main[f"{column}2"] = text

        change = result["数据更改"]
        # The audited template used fixed sales weights. Keep its input
        # fingerprint, then apply this reviewed configurable-weights rule.
        target_main["S2"] = "=K2*数据更改!$E$2/(30+AV2)+L2*数据更改!$F$2/(15+AW2)+M2*数据更改!$G$2/(7+AX2)"
        for column, label in enumerate(("30天销量权重", "15天销量权重", "7天销量权重"), 5):
            change.cell(1, column).value = label
        for column, value in enumerate(DEFAULT_PARAMETERS, 1):
            change.cell(2, column).value = value

        output_path.parent.mkdir(parents=True, exist_ok=True)
        result.save(output_path)
        result.close()
    finally:
        source.close()
    return output_path


def main() -> None:
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path, help="Local nine-sheet business template; read only")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Data-free package resource")
    arguments = parser.parse_args()
    print(build_skeleton(arguments.source, arguments.output))


if __name__ == "__main__":
    main()
