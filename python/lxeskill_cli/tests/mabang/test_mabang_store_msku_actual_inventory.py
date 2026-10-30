from __future__ import annotations

import asyncio
import json
from decimal import Decimal
from http.cookies import SimpleCookie
from pathlib import Path
from unittest.mock import AsyncMock, Mock
from yarl import URL

import pytest
import aiohttp
from aiohttp import web

from mabang_test_helpers import _annotate_active_test_source, _form_value, _form_values, _sheet_names, _xlsx_bytes
from services.mabang.amazon.fba import store_msku_actual_inventory as inv
from services.mabang.auth import MabangAuthContext


async def _async_value(value):
    return value


def _write_xlsx(path: Path, rows: list[dict], *, columns: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_xlsx_bytes(rows, columns=columns))
    return path


def _stock_xlsx_bytes(rows: list[dict]) -> bytes:
    return _xlsx_bytes(rows, columns=["库存SKU编号", "可用库存量"])


SEARCH_TOKEN = "synthetic-search-token"


def test_old_and_clean_store_names_preserve_existing_file_identity(monkeypatch, tmp_path):
    from services.mabang.amazon.fba import combo_sku as combo, store_resolver as stores
    from services.mabang.amazon.fba import store_msku_sales_analysis as sales
    from services.mabang.amazon.fba import store_msku_replenishment as repl
    from services.mabang.amazon.fba.replenishment_template import get_template
    names = ['Amazon-YYH-US', 'Amazon-YYH-US美国']
    # Existing filename normalization already strips the trailing country text.
    # Both names refer to this file; a newer, unrelated shop must not be selected.
    for index, name in enumerate(['Amazon-YYH-US', 'Amazon-OTHER-US']):
        row = dict.fromkeys(inv.SOURCE_COLUMNS, 0)
        row.update({'MSKU': f'M{index}', '父ASIN': 'P', 'ASIN': 'A', '本地SKU': 'S', '商品链接': 'https://www.amazon.com/dp/A', '7天销量': 7, '14天销量': 14, '30天销量': 30, '90天销量': 90, '计划入库': 0})
        _write_xlsx(tmp_path/'source'/f'20260907090{index}-{name}_店铺MSKU数据.xlsx', [row], columns=[*inv.SOURCE_COLUMNS, '90天销量', '计划入库', '单品重量(g)(cm)'])

    async def web_stores():
        return stores.parse_fba_store_options('<li><input name="fbaWarehouseIds[]" value="1039477"><span class="texts">Amazon-YYH-US<span class="shop-country-cn">美国</span></span></li>')
    async def post(endpoint, body, **kwargs):
        if endpoint == 'shops/list':
            return {'code': 200, 'data': {'profile': {'name': names[0], 'sid': 10, 'amazonsite': 'us'}}}
        assert endpoint == 'listings/search' and body['shop_id'] == ['10']
        return {'code': 200, 'data': {'list': [{'platformSku': f'M{i}', 'asin': 'A', 'stockSku': 'S', 'stockType': 1, 'shopIds': '10', 'amazonsite': 'us'} for i in range(2)], 'total': 2, 'nowPage': 1, 'totalPage': 1}}
    stock_path = tmp_path/'stock.xlsx'
    stock_path.write_bytes(_stock_xlsx_bytes([{'库存SKU编号': 'S', '可用库存量': 100}]))
    async def search(skus):
        assert skus == ['S']
        return "", SEARCH_TOKEN
    async def download(**kwargs):
        return stock_path
    monkeypatch.setattr(combo, 'fetch_fba_stores', web_stores)
    monkeypatch.setattr(combo, 'post_json', post)
    monkeypatch.setattr(inv, 'search_warehouse_stock', search)
    monkeypatch.setattr(inv, 'download_warehouse_stock_xlsx', download)
    template = get_template('默认')
    monkeypatch.setattr(repl, 'get_template', lambda name: template)
    _annotate_active_test_source(tmp_path/'source'/'202609070900-Amazon-YYH-US_店铺MSKU数据.xlsx', requested_store_name=names[1])
    for index, name in enumerate(names):
        analysis = sales.analyze_store_msku_sales(name, input_dir=tmp_path/'source', output_dir=tmp_path/'sales')
        actual = asyncio.run(inv.export_store_msku_actual_inventory(name, input_dir=tmp_path/'source', output_dir=tmp_path/'inventory'))
        assert analysis.source_data_time == actual.source_msku_data_time == '202609070900'
        result = repl.calculate_store_msku_replenishment(name, sales_analysis_dir=tmp_path/'sales', actual_inventory_dir=tmp_path/'inventory', output_dir=tmp_path/'replenishment', unlinked_shipments_snapshot_dir=tmp_path/'unlinked')
        assert result.store_name == name
        matched = repl.find_matching_report_files(name, sales_analysis_dir=tmp_path/'sales', actual_inventory_dir=tmp_path/'inventory')
        assert matched.source_data_time == '202609070900'
        assert [r.msku for r in repl.load_inventory_rows(matched.actual_inventory_path)] == ['M0']


class _FakeResponse:
    def __init__(
        self,
        payload: dict | None = None,
        *,
        status: int = 200,
        body: bytes = b"",
        text_body: str | None = None,
        content_type: str = "",
        redirect_count: int = 0,
        sets_cookie: bool = False,
        set_cookie_header: str | None = None,
    ) -> None:
        self.status = status
        self.content_type = content_type
        self.history = (object(),) * redirect_count
        self.cookies = SimpleCookie()
        if set_cookie_header:
            self.cookies.load(set_cookie_header)
        elif sets_cookie:
            self.cookies.load("PHPSESSID=never-log-this-cookie; Path=/")
        self._payload = dict(payload or {})
        self._body = body
        self._text_body = text_body

    async def text(self) -> str:
        if self._text_body is not None:
            return self._text_body
        return json.dumps(self._payload, ensure_ascii=False)

    async def json(self, content_type=None) -> dict:
        if self._text_body is not None:
            raise ValueError("not json")
        return dict(self._payload)

    async def read(self) -> bytes:
        return self._body


class _FakeRequest:
    def __init__(self, response: _FakeResponse) -> None:
        self._response = response

    async def __aenter__(self) -> _FakeResponse:
        return self._response

    async def __aexit__(self, exc_type, exc, tb) -> None:
        return None


class _FakeSession:
    def __init__(self, responses: list[_FakeResponse]) -> None:
        self.responses = list(responses)
        self.calls: list[dict] = []

    def post(self, url: str, **kwargs) -> _FakeRequest:
        self.calls.append({"method": "POST", "url": url, **kwargs})
        return _FakeRequest(self.responses.pop(0))

    def get(self, url: str, **kwargs) -> _FakeRequest:
        self.calls.append({"method": "GET", "url": url, **kwargs})
        return _FakeRequest(self.responses.pop(0))


async def _fake_auth_context(*args, **kwargs) -> MabangAuthContext:
    return MabangAuthContext(
        account="",
        source="test",
        cookies_by_domain={
            ".mabangerp.com": [
                {"name": "PHPSESSID", "value": "sid", "domain": ".mabangerp.com"},
                {
                    "name": "MABANG_ERP_PRO_MEMBERINFO_LOGIN_COOKIE",
                    "value": "memcache-key",
                    "domain": ".mabangerp.com",
                },
                {
                    "name": "MABANG_ERP_PRO_MEMBERINFO_LOGIN_PLUS",
                    "value": "login-plus",
                    "domain": ".mabangerp.com",
                },
                {"name": "signed", "value": "signed-value", "domain": ".mabangerp.com"},
                {"name": "route", "value": "route-value", "domain": ".mabangerp.com"},
            ]
        },
        free_token="",
        wms_cookie_header="",
    )


def _load_records(path: Path, sheet_name: str = "真实库存（深圳仓库）-库存sku") -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in worksheet[1]]
        return [dict(zip(headers, values, strict=False)) for values in worksheet.iter_rows(min_row=2, values_only=True)]
    finally:
        workbook.close()


def _column_number_formats(path: Path, sheet_name: str, column_name: str) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False)
    try:
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in worksheet[1]]
        column_index = headers.index(column_name) + 1
        return [
            cell.number_format
            for cells in worksheet.iter_rows(min_row=2, min_col=column_index, max_col=column_index)
            for cell in cells
        ]
    finally:
        workbook.close()


def _column_fill_colors(path: Path, sheet_name: str, column_name: str) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False)
    try:
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in worksheet[1]]
        column_index = headers.index(column_name) + 1
        return [
            str(cell.fill.fgColor.rgb or "")
            for cells in worksheet.iter_rows(min_row=1, min_col=column_index, max_col=column_index)
            for cell in cells
        ]
    finally:
        workbook.close()


def _assert_standard_dimensions(path: Path, sheet_names: list[str]) -> None:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    try:
        for sheet_name in sheet_names:
            worksheet = workbook[sheet_name]
            assert worksheet.sheet_format.defaultRowHeight == 15
            for row_index in range(1, worksheet.max_row + 1):
                assert worksheet.row_dimensions[row_index].height == 15
            for column_index in range(1, worksheet.max_column + 1):
                assert worksheet.column_dimensions[get_column_letter(column_index)].width == 15
    finally:
        workbook.close()


def test_warehouse_search_form_uses_fixed_warehouse_and_crlf_skus() -> None:
    form = inv._warehouse_search_form_data(["SKU-A", "SKU-B"])

    assert _form_value({"data": form}, "warehouseIds[]") == "1014318"
    assert _form_value({"data": form}, "stockSkuStr") == "SKU-A\r\nSKU-B\r\n"
    assert _form_value({"data": form}, "statusIN[]") == "3"


# Captured 2026-10-14 warehouse export error; only line endings/trailing spaces normalized.
WAREHOUSE_ERROR_HTML = (Path(__file__).parent / "fixtures/warehouse_search_overwritten.html").read_bytes()
WAREHOUSE_ERROR = "查询条件已被新的搜索覆盖，请重新搜索后再导出"


def test_warehouse_html_error_is_not_published_as_xlsx(monkeypatch, tmp_path) -> None:
    path = tmp_path / "202610141356-shop_warehouse_stock.xlsx"
    original = _stock_xlsx_bytes([{"库存SKU编号": "S", "可用库存量": 8}])
    path.write_bytes(original)
    session = _FakeSession([_FakeResponse(body=WAREHOUSE_ERROR_HTML)])
    monkeypatch.setattr(inv, "erp_http_session", session)
    monkeypatch.setattr(inv, "get_auth_context", _fake_auth_context)
    monkeypatch.setattr(inv, "_timestamp_text", lambda: "202610141356")
    with pytest.raises(inv.StoreMskuActualInventoryError, match=WAREHOUSE_ERROR):
        asyncio.run(inv.download_warehouse_stock_xlsx(
            store_name="shop", output_dir=tmp_path, search_token=SEARCH_TOKEN,
        ))
    assert path.read_bytes() == original
    assert len(session.calls) == 1


def test_warehouse_search_failure_stops_before_download(monkeypatch, tmp_path) -> None:
    response = _FakeResponse(payload={"success": False, "message": WAREHOUSE_ERROR})
    _assert_warehouse_pipeline(monkeypatch, tmp_path, [response], succeeds=False, calls=1, expected=WAREHOUSE_ERROR)


def test_warehouse_pipeline_forwards_current_search_token_in_export_url(monkeypatch, tmp_path, caplog) -> None:
    token = "synthetic+token/&=?"
    monkeypatch.setattr(inv, "_warehouse_export_url", lambda: inv.DEFAULT_WAREHOUSE_EXPORT_URL + "&searchToken=stale")
    with caplog.at_level("INFO"):
        session = _assert_warehouse_pipeline(monkeypatch, tmp_path, [
            _FakeResponse(payload={"success": True, "searchToken": token}),
            _FakeResponse(body=_stock_xlsx_bytes([{"库存SKU编号": "S", "可用库存量": 8}])),
        ], succeeds=True, calls=2)
    export_url = URL(str(session.calls[1]["url"]))
    assert export_url.query["searchToken"] == token
    assert export_url.query.getall("searchToken") == [token]
    assert export_url.query["mod"] == "warehouse.doexportwarehousestock"
    assert export_url.query["flag"] == "1"
    assert token not in caplog.text


def test_warehouse_missing_or_invalid_search_token_stops_before_export(monkeypatch, tmp_path) -> None:
    _assert_warehouse_pipeline(monkeypatch, tmp_path, [
        _FakeResponse(payload={"success": True}),
    ], succeeds=False, calls=1, expected="searchToken")


def test_warehouse_http_session_carries_search_cookie_without_shared_cookie_jar(monkeypatch, tmp_path) -> None:
    received: list[tuple[str, bool]] = []
    workbook = _stock_xlsx_bytes([{"库存SKU编号": "S", "可用库存量": 8}])

    async def search(request: web.Request) -> web.Response:
        received.append(("search", "PHPSESSID=sid" in request.headers.get("Cookie", "")))
        response = web.json_response({"success": True, "pageHtml": False, "searchToken": SEARCH_TOKEN})
        response.set_cookie("PHPSESSID", "rotated-test-session", path="/")
        return response

    async def export(request: web.Request) -> web.Response:
        received.append((
            "export",
            "PHPSESSID=rotated-test-session" in request.headers.get("Cookie", "")
            and request.query.get("searchToken") == SEARCH_TOKEN,
        ))
        return web.Response(body=workbook, content_type="application/octet-stream")

    async def run() -> None:
        app = web.Application()
        app.router.add_post("/search", search)
        app.router.add_get("/export", export)
        runner = web.AppRunner(app)
        await runner.setup()
        site = web.TCPSite(runner, "localhost", 0)
        await site.start()
        sockets = list(getattr(site._server, "sockets", []) or [])
        assert sockets
        base_url = f"http://localhost:{int(sockets[0].getsockname()[1])}"
        monkeypatch.setattr(inv, "_warehouse_search_url", lambda: f"{base_url}/search")
        monkeypatch.setattr(inv, "_warehouse_export_url", lambda: f"{base_url}/export")
        monkeypatch.setattr(inv, "get_auth_context", _fake_auth_context)
        try:
            async with aiohttp.ClientSession(cookie_jar=aiohttp.DummyCookieJar()) as session:
                monkeypatch.setattr(inv, "erp_http_session", session)
                cookie_header, search_token = await inv.search_warehouse_stock(["S"])
                output = await inv.download_warehouse_stock_xlsx(
                    store_name="shop", output_dir=tmp_path,
                    cookie_header=cookie_header, search_token=search_token,
                )
                assert output.read_bytes() == workbook
                assert not session.cookie_jar.filter_cookies(f"{base_url}/export")
        finally:
            await runner.cleanup()

    asyncio.run(run())
    assert received == [("search", True), ("export", True)]


def _assert_warehouse_pipeline(
    monkeypatch, tmp_path, responses, *, succeeds, calls, expected="", auth_context=_fake_auth_context,
) -> _FakeSession:
    from test_source_verification import row, source

    source(tmp_path / "source/202610141349-shop_店铺MSKU数据.xlsx", [row(**{
        "父ASIN": "P", "商品链接": "https://example.test/A",
        **dict.fromkeys((*inv.SALES_COLUMNS, *inv.FBA_STOCK_COLUMNS), 0),
    })])
    session = _FakeSession(responses)
    monkeypatch.setattr(inv, "erp_http_session", session)
    monkeypatch.setattr(inv, "get_auth_context", auth_context)

    async def unexpected_refresh(**kwargs):
        pytest.fail("业务错误不得刷新认证或重试")

    monkeypatch.setattr(inv, "refresh_mabang_auth", unexpected_refresh)
    operation = inv.export_store_msku_actual_inventory("shop", input_dir=tmp_path / "source", output_dir=tmp_path / "output")
    if succeeds:
        result = asyncio.run(operation)
        assert Path(result.shenzhen_warehouse_inventory_report_xlsx_path).is_file()
        assert result.source_msku_data_time == "202610141349"
    else:
        with pytest.raises(inv.StoreMskuActualInventoryError, match=expected):
            asyncio.run(operation)
        assert not list((tmp_path / "output").glob("*.xlsx"))
    assert len(session.calls) == calls
    return session


def test_warehouse_error_diagnostic_reuses_redaction_and_explicit_truncation(monkeypatch, tmp_path) -> None:
    secret = "synthetic-warehouse-test-key"
    monkeypatch.setenv("LXE_DATA_SERVER_API_KEY", secret)
    body = f"<html><p>{WAREHOUSE_ERROR}; token={secret}; {'x' * 5000}</p></html>".encode()
    monkeypatch.setattr(inv, "erp_http_session", _FakeSession([_FakeResponse(body=body)]))
    monkeypatch.setattr(inv, "get_auth_context", _fake_auth_context)
    with pytest.raises(inv.StoreMskuActualInventoryError) as error:
        asyncio.run(inv.download_warehouse_stock_xlsx(
            store_name="shop", output_dir=tmp_path, search_token=SEARCH_TOKEN,
        ))
    assert WAREHOUSE_ERROR in str(error.value)
    assert secret not in str(error.value)
    assert "[REDACTED]" in str(error.value)
    assert "[truncated " in str(error.value)
    assert not list(tmp_path.glob("*.xlsx"))


def test_parse_stock_inventory_sums_duplicate_stock_sku_rows(tmp_path) -> None:
    xlsx_path = tmp_path / "stock.xlsx"
    xlsx_path.write_bytes(
        _stock_xlsx_bytes(
            [
                {"库存SKU编号": "STOCK-A", "可用库存量": "2"},
                {"库存SKU编号": "STOCK-A", "可用库存量": "3"},
                {"库存SKU编号": "STOCK-B", "可用库存量": "12"},
            ]
        )
    )

    result = inv.parse_stock_inventory_xlsx(xlsx_path)

    assert result == {
        "STOCK-A": Decimal("5"),
        "STOCK-B": Decimal("12"),
    }


def test_parse_stock_inventory_uses_available_quantity_when_bin_quantity_differs(tmp_path) -> None:
    path = tmp_path / "stock.xlsx"
    path.write_bytes(_xlsx_bytes(
        [
            {"库存SKU编号": "A", "可用库存量": -11499, "仓位库存": 6013},
            {"库存SKU编号": "B", "可用库存量": 0, "仓位库存": 0},
            {"库存SKU编号": "C", "可用库存量": -2, "仓位库存": None},
        ],
        columns=["库存SKU编号", "可用库存量", "仓位库存"],
    ))
    assert inv.parse_stock_inventory_xlsx(path) == {
        "A": Decimal("-11499"), "B": Decimal("0"), "C": Decimal("-2"),
    }


def test_load_store_msku_rows_requires_product_link_column(tmp_path) -> None:
    source_path = tmp_path / "202605251530-Amazon-Test_店铺MSKU数据.xlsx"
    _write_xlsx(
        source_path,
        [{"MSKU": "MSKU-A", "父ASIN": "PARENT-A", "ASIN": "ASIN-A", "本地SKU": "SKU-A"}],
        columns=["MSKU", "父ASIN", "ASIN", "本地SKU"],
    )

    with pytest.raises(inv.StoreMskuActualInventoryError, match="店铺MSKU数据缺少列: 商品链接"):
        inv.load_store_msku_rows(source_path)


def test_find_latest_store_msku_file_accepts_legacy_english_file_name(tmp_path) -> None:
    source_path = tmp_path / "202605251530-Amazon-Test_msku_data.xlsx"
    _write_xlsx(
        source_path,
        [{"MSKU": "MSKU-A", "父ASIN": "PARENT-A", "ASIN": "ASIN-A", "本地SKU": "SKU-A", "商品链接": "https://example.test/a"}],
        columns=["MSKU", "父ASIN", "ASIN", "本地SKU", "商品链接"],
    )

    result = inv.find_latest_store_msku_file("Amazon-Test", input_dir=tmp_path)

    assert result.path == source_path
    assert result.source_data_time == "202605251530"


def test_load_store_msku_rows_requires_sales_and_fba_stock_columns(tmp_path) -> None:
    source_path = tmp_path / "202605251530-Amazon-Test_店铺MSKU数据.xlsx"
    _write_xlsx(
        source_path,
        [
            {
                "MSKU": "MSKU-A",
                "父ASIN": "PARENT-A",
                "ASIN": "ASIN-A",
                "本地SKU": "SKU-A",
                "商品链接": "https://example.test/a",
            }
        ],
        columns=["MSKU", "父ASIN", "ASIN", "本地SKU", "商品链接"],
    )

    with pytest.raises(inv.StoreMskuActualInventoryError, match="7天销量.*调仓中"):
        inv.load_store_msku_rows(source_path)


def test_load_store_msku_rows_allows_missing_optional_name_columns(tmp_path) -> None:
    source_path = tmp_path / "202605251530-Amazon-Test_店铺MSKU数据.xlsx"
    _write_xlsx(
        source_path,
        [
            {
                "MSKU": "MSKU-A",
                "父ASIN": "PARENT-A",
                "ASIN": "ASIN-A",
                "本地SKU": "SKU-A",
                "商品链接": "https://example.test/a",
                "7天销量": 7,
                "14天销量": 14,
                "30天销量": 30,
                "可售": 4,
                "待入库": 3,
                "预留": 2,
                "在途": 1,
                "待调仓": 6,
                "调仓中": 7,
            }
        ],
        columns=[
            "MSKU",
            "父ASIN",
            "ASIN",
            "本地SKU",
            "商品链接",
            "7天销量",
            "14天销量",
            "30天销量",
            "可售",
            "待入库",
            "预留",
            "在途",
            "待调仓",
            "调仓中",
        ],
    )

    rows = inv.load_store_msku_rows(source_path)

    assert len(rows) == 1
    assert rows[0].local_sku_name == ""
    assert rows[0].product_name == ""
    assert rows[0].remark == ""


def test_load_store_msku_rows_parses_invalid_sales_and_fba_stock_as_zero(tmp_path) -> None:
    source_path = tmp_path / "202605251530-Amazon-Test_店铺MSKU数据.xlsx"
    _write_xlsx(
        source_path,
        [
            {
                "MSKU": "MSKU-A",
                "父ASIN": "PARENT-A",
                "ASIN": "ASIN-A",
                "本地SKU": "SKU-A",
                "商品链接": "https://example.test/a",
                "本地SKU名称": "本地产品A",
                "产品名称": "Product A",
                "备注": "清货观察",
                "7天销量": "nan",
                "14天销量": "1,400",
                "30天销量": "not-a-number",
                "可售": "1,000",
                "待入库": "",
                "预留": "nan",
                "在途": "not-a-number",
                "待调仓": "2",
                "调仓中": "bad",
            }
        ],
        columns=[
            "MSKU",
            "父ASIN",
            "ASIN",
            "本地SKU",
            "商品链接",
            "本地SKU名称",
            "产品名称",
            "备注",
            "7天销量",
            "14天销量",
            "30天销量",
            "可售",
            "待入库",
            "预留",
            "在途",
            "待调仓",
            "调仓中",
        ],
    )

    rows = inv.load_store_msku_rows(source_path)

    assert len(rows) == 1
    assert rows[0].local_sku_name == "本地产品A"
    assert rows[0].product_name == "Product A"
    assert rows[0].sales_7d == Decimal("0")
    assert rows[0].sales_14d == Decimal("1400")
    assert rows[0].sales_30d == Decimal("0")
    assert rows[0].fba_sellable == Decimal("1000")
    assert rows[0].fba_inbound == Decimal("0")
    assert rows[0].fba_reserved == Decimal("0")
    assert rows[0].fba_in_transit == Decimal("0")
    assert rows[0].fba_pending_transfer == Decimal("2")
    assert rows[0].fba_transferring == Decimal("0")


def test_calculate_inventory_rows_handles_normal_combo_and_missing_stock() -> None:
    rows = [
        inv.StoreMskuRow(
            "MSKU-A",
            "PARENT-A",
            "ASIN-A",
            "SKU-A",
            "https://example.test/a",
            local_sku_name="单品本地名",
            product_name="Single Product",
            remark="清货单品",
            sales_7d=Decimal("7"),
            sales_14d=Decimal("14"),
            sales_30d=Decimal("30"),
            fba_sellable=Decimal("4"),
            fba_inbound=Decimal("3"),
            fba_reserved=Decimal("2"),
            fba_in_transit=Decimal("1"),
            fba_pending_transfer=Decimal("6"),
            fba_transferring=Decimal("7"),
        ),
        inv.StoreMskuRow(
            "MSKU-C",
            "PARENT-C",
            "ASIN-C",
            "COMBO-A",
            "https://example.test/c",
            local_sku_name="组合本地名",
            product_name="Combo Product",
            remark="组合备注",
            sales_7d=Decimal("70"),
            sales_14d=Decimal("140"),
            sales_30d=Decimal("300"),
            fba_sellable=Decimal("20"),
            fba_inbound=Decimal("15"),
            fba_reserved=Decimal("10"),
            fba_in_transit=Decimal("5"),
            fba_pending_transfer=Decimal("4"),
            fba_transferring=Decimal("6"),
        ),
        inv.StoreMskuRow("MSKU-M", "PARENT-M", "ASIN-M", "SKU-MISSING", "https://example.test/m"),
        inv.StoreMskuRow("MSKU-CM", "PARENT-CM", "ASIN-CM", "COMBO-MISSING", "https://example.test/cm"),
    ]
    combo_map = {
        inv.normalize_sku_key("COMBO-A"): inv.ComboSku(
            combo_sku="COMBO-A",
            components=(
                inv.ComboComponent("STOCK-A", Decimal("1")),
                inv.ComboComponent("STOCK-B", Decimal("2")),
            ),
        ),
        inv.normalize_sku_key("COMBO-MISSING"): inv.ComboSku(
            combo_sku="COMBO-MISSING",
            components=(inv.ComboComponent("STOCK-MISSING", Decimal("1")),),
        ),
    }

    result_rows, missing = inv.calculate_inventory_rows(
        rows,
        combo_map=combo_map,
        stock_quantities={
            "SKU-A": Decimal("9"),
            "STOCK-A": Decimal("10"),
            "STOCK-B": Decimal("12"),
        },
    )

    assert [row.actual_inventory for row in result_rows] == [Decimal("9"), Decimal("6"), None, None]
    assert [row.local_sku_name for row in result_rows] == ["单品本地名", "组合本地名", "", ""]
    assert [row.product_name for row in result_rows] == ["Single Product", "Combo Product", "", ""]
    assert [row.remark for row in result_rows] == ["清货单品", "组合备注", "", ""]
    assert [row.product_link for row in result_rows] == [
        "https://example.test/a",
        "https://example.test/c",
        "https://example.test/m",
        "https://example.test/cm",
    ]
    assert [row.is_combo_sku for row in result_rows] == [False, True, False, True]
    assert result_rows[0].fba_total_inventory == Decimal("23")
    assert result_rows[0].weighted_daily_sales == Decimal("1.0")
    assert result_rows[0].sales_days == Decimal("23")
    assert result_rows[1].fba_total_inventory == Decimal("60")
    assert result_rows[1].weighted_daily_sales == Decimal("10.0")
    assert result_rows[1].sales_days == Decimal("6")
    assert result_rows[1].child_skus == "STOCK-A * 1, STOCK-B * 2"
    assert missing == ["SKU-MISSING", "STOCK-MISSING"]


def test_write_actual_inventory_xlsx_splits_rows_by_inventory_state(tmp_path) -> None:
    output_path = tmp_path / "actual_inventory.xlsx"
    rows = [
        inv.ActualInventoryRow(
            "MSKU-S1",
            "PARENT-S1",
            "ASIN-S1",
            "SKU-S1",
            "https://example.test/stock-1",
            Decimal("10"),
            "",
            fba_total_inventory=Decimal("30"),
            weighted_daily_sales=Decimal("3.456"),
            sales_days=Decimal("8.666"),
            local_sku_name="单品本地1",
            product_name="Stock Product 1",
            remark="清货单品1",
        ),
        inv.ActualInventoryRow(
            "MSKU-S0",
            "PARENT-S0",
            "ASIN-S0",
            "SKU-ZERO",
            "https://example.test/zero",
            Decimal("0"),
            "",
            fba_total_inventory=Decimal("0"),
            weighted_daily_sales=Decimal("0"),
            sales_days=None,
        ),
        inv.ActualInventoryRow(
            "MSKU-C",
            "PARENT-C",
            "ASIN-C",
            "COMBO-A",
            "https://example.test/combo",
            Decimal("6"),
            "STOCK-A * 1",
            is_combo_sku=True,
            fba_total_inventory=Decimal("20"),
            weighted_daily_sales=Decimal("5"),
            sales_days=Decimal("4"),
            local_sku_name="组合本地商品",
            product_name="Combo Product",
            remark="组合备注",
        ),
        inv.ActualInventoryRow("MSKU-N", "PARENT-N", "ASIN-N", "", "https://example.test/no-local", None, ""),
        inv.ActualInventoryRow(
            "MSKU-M",
            "PARENT-M",
            "ASIN-M",
            "SKU-MISSING",
            "https://example.test/missing",
            None,
            "",
        ),
        inv.ActualInventoryRow(
            "MSKU-CM",
            "PARENT-CM",
            "ASIN-CM",
            "COMBO-MISSING",
            "https://example.test/combo-missing",
            None,
            "STOCK-MISSING * 1",
        ),
    ]

    groups = inv.split_inventory_rows(rows)
    inv.write_actual_inventory_xlsx(rows, output_path)

    assert len(groups.combo_inventory_rows) == 1
    assert len(groups.stock_inventory_rows) == 2
    assert len(groups.inventory_rows) == 3
    assert len(groups.no_local_sku_rows) == 1
    assert len(groups.no_inventory_rows) == 2
    assert _sheet_names(output_path) == ["真实库存（深圳仓库）-组合sku", "真实库存（深圳仓库）-库存sku", "无本地SKU", "无库存数据"]
    _assert_standard_dimensions(output_path, _sheet_names(output_path))
    for sheet_name in _sheet_names(output_path):
        records = _load_records(output_path, sheet_name)
        if records:
            expected_columns = (
                list(inv.INVENTORY_OUTPUT_COLUMNS)
                if sheet_name in {"真实库存（深圳仓库）-组合sku", "真实库存（深圳仓库）-库存sku"}
                else list(inv.OUTPUT_COLUMNS)
            )
            assert list(records[0]) == expected_columns

    combo_records = _load_records(output_path, "真实库存（深圳仓库）-组合sku")
    stock_records = _load_records(output_path, "真实库存（深圳仓库）-库存sku")
    no_local_records = _load_records(output_path, "无本地SKU")
    no_inventory_records = _load_records(output_path, "无库存数据")

    assert [record["MSKU"] for record in combo_records] == ["MSKU-C"]
    assert list(combo_records[0]) == list(inv.INVENTORY_OUTPUT_COLUMNS)
    assert combo_records[0]["本地SKU名称"] == "组合本地商品"
    assert combo_records[0]["产品名称"] == "Combo Product"
    assert combo_records[0]["备注"] == "组合备注"
    assert combo_records[0]["商品链接"] == "https://example.test/combo"
    assert combo_records[0]["FBA总库存"] == 20
    assert combo_records[0]["加权日销"] == 5
    assert combo_records[0]["可销售天数"] == 4
    assert combo_records[0]["真实库存（深圳仓库）数量"] == 6
    assert [record["MSKU"] for record in stock_records] == ["MSKU-S1", "MSKU-S0"]
    assert stock_records[0]["本地SKU名称"] == "单品本地1"
    assert stock_records[0]["产品名称"] == "Stock Product 1"
    assert stock_records[0]["备注"] == "清货单品1"
    assert stock_records[0]["商品链接"] == "https://example.test/stock-1"
    assert stock_records[0]["FBA总库存"] == 30
    assert stock_records[0]["加权日销"] == 3.46
    assert stock_records[0]["可销售天数"] == 8.67
    assert stock_records[1]["商品链接"] == "https://example.test/zero"
    assert stock_records[1]["加权日销"] == 0
    assert stock_records[1]["可销售天数"] in (None, "")
    assert stock_records[1]["真实库存（深圳仓库）数量"] == 0
    assert [record["MSKU"] for record in no_local_records] == ["MSKU-N"]
    assert list(no_local_records[0]) == list(inv.OUTPUT_COLUMNS)
    assert no_local_records[0]["商品链接"] == "https://example.test/no-local"
    assert [record["MSKU"] for record in no_inventory_records] == ["MSKU-M", "MSKU-CM"]
    assert list(no_inventory_records[0]) == list(inv.OUTPUT_COLUMNS)
    assert no_inventory_records[0]["商品链接"] == "https://example.test/missing"
    assert no_inventory_records[1]["商品链接"] == "https://example.test/combo-missing"
    assert no_inventory_records[1]["子SKU"] == "STOCK-MISSING * 1"
    assert _column_number_formats(output_path, "真实库存（深圳仓库）-组合sku", "加权日销") == ["0.00"]
    assert _column_number_formats(output_path, "真实库存（深圳仓库）-组合sku", "可销售天数") == ["0.00"]
    assert _column_number_formats(output_path, "真实库存（深圳仓库）-库存sku", "加权日销") == ["0.00", "0.00"]
    assert _column_number_formats(output_path, "真实库存（深圳仓库）-库存sku", "可销售天数") == ["0.00", "0.00"]
    assert _column_fill_colors(output_path, "真实库存（深圳仓库）-组合sku", "MSKU") == ["00FFF2CC", "00FFF2CC"]
    assert _column_fill_colors(output_path, "真实库存（深圳仓库）-组合sku", "真实库存（深圳仓库）数量") == ["00FFF2CC", "00FFF2CC"]
    assert _column_fill_colors(output_path, "真实库存（深圳仓库）-库存sku", "MSKU") == [
        "00FFF2CC",
        "00FFF2CC",
        "00FFF2CC",
    ]
    assert _column_fill_colors(output_path, "真实库存（深圳仓库）-库存sku", "真实库存（深圳仓库）数量") == [
        "00FFF2CC",
        "00FFF2CC",
        "00FFF2CC",
    ]
    assert _column_fill_colors(output_path, "无本地SKU", "MSKU") == ["00000000", "00000000"]
    assert _column_fill_colors(output_path, "无本地SKU", "真实库存（深圳仓库）数量") == ["00000000", "00000000"]
    assert _column_fill_colors(output_path, "无库存数据", "MSKU") == ["00000000", "00000000", "00000000"]
    assert _column_fill_colors(output_path, "无库存数据", "真实库存（深圳仓库）数量") == ["00000000", "00000000", "00000000"]


@pytest.mark.parametrize("explicit_source", [False, True])
def test_export_store_msku_actual_inventory_success_with_fake_network(monkeypatch, tmp_path, explicit_source) -> None:
    source_path = tmp_path / "input" / "202605251530-Amazon-Lerxiuer-FR_店铺MSKU数据.xlsx"
    _write_xlsx(
        source_path,
        [
            {
                "MSKU": "MSKU-A",
                "父ASIN": "PARENT-A",
                "ASIN": "ASIN-A",
                "本地SKU": "SKU-A",
                "商品链接": "https://example.test/a",
                "7天销量": 7,
                "14天销量": 14,
                "30天销量": 30,
                "可售": 4,
                "待入库": 3,
                "预留": 2,
                "在途": 1,
                "待调仓": 6,
                "调仓中": 7,
            },
            {
                "MSKU": "MSKU-C",
                "父ASIN": "PARENT-C",
                "ASIN": "ASIN-C",
                "本地SKU": "COMBO-A",
                "商品链接": "https://example.test/c",
                "7天销量": 70,
                "14天销量": 140,
                "30天销量": 300,
                "可售": 20,
                "待入库": 15,
                "预留": 10,
                "在途": 5,
                "待调仓": 4,
                "调仓中": 6,
            },
            {
                "MSKU": "MSKU-N",
                "父ASIN": "PARENT-N",
                "ASIN": "ASIN-N",
                "本地SKU": "",
                "商品链接": "https://example.test/no-local",
                "7天销量": 1,
                "14天销量": 2,
                "30天销量": 3,
                "可售": 4,
                "待入库": 5,
                "预留": 6,
                "在途": 7,
                "待调仓": 8,
                "调仓中": 9,
            },
        ],
        columns=[
            "MSKU",
            "父ASIN",
            "ASIN",
            "本地SKU",
            "商品链接",
            "7天销量",
            "14天销量",
            "30天销量",
            "可售",
            "待入库",
            "预留",
            "在途",
            "待调仓",
            "调仓中",
        ],
    )
    _annotate_active_test_source(source_path)
    fake_session = _FakeSession(
        [
            _FakeResponse({"success": True, "searchToken": SEARCH_TOKEN}),
            _FakeResponse(
                body=_stock_xlsx_bytes(
                    [
                        {"库存SKU编号": "SKU-A", "可用库存量": 9},
                        {"库存SKU编号": "STOCK-A", "可用库存量": 10},
                        {"库存SKU编号": "STOCK-B", "可用库存量": 12},
                    ]
                )
            ),
        ]
    )
    monkeypatch.setattr(inv, "erp_http_session", fake_session)
    monkeypatch.setattr(inv, "external_http_session", fake_session)
    monkeypatch.setattr(inv, "get_auth_context", _fake_auth_context)
    monkeypatch.setattr(inv, "_timestamp_text", lambda *_args, **_kwargs: "202605271530")
    sleep_calls: list[float] = []

    async def fake_sleep(delay: float) -> None:
        sleep_calls.append(delay)

    monkeypatch.setattr(asyncio, "sleep", fake_sleep)

    source_kwargs = {}
    if explicit_source:
        source_path.with_name("202605261530-Amazon-Lerxiuer-FR_店铺MSKU数据.xlsx").write_bytes(source_path.read_bytes())
        monkeypatch.setattr(inv, "find_latest_store_msku_file", Mock(side_effect=AssertionError("显式源表不得选择历史最新文件")))
        source_kwargs["msku_source_xlsx_path"] = source_path
    result = asyncio.run(
        inv.export_store_msku_actual_inventory(
            "Amazon-Lerxiuer-FR",
            input_dir=tmp_path / "input",
            output_dir=tmp_path / "output",
            **source_kwargs,
        )
    )

    payload = result.to_payload()
    assert payload == {
        "success": True,
        "store_name": "Amazon-Lerxiuer-FR",
        "warehouse_id": "1014318",
        "warehouse_name": "深圳仓库",
        "source_msku_xlsx_path": str(source_path),
        "source_msku_data_time": "202605251530",
        "unique_local_sku_count": 2,
        "detected_combo_sku_count": 1,
        "queried_warehouse_stock_sku_count": 3,
        "matched_warehouse_inventory_msku_row_count": 2,
        "missing_local_sku_msku_row_count": 1,
        "missing_warehouse_inventory_msku_row_count": 0,
        "missing_warehouse_stock_sku_count": 0,
        "missing_warehouse_stock_skus": [],
        "shenzhen_warehouse_inventory_report_xlsx_path": str(tmp_path / "output" / "202605251530-Amazon-Lerxiuer-FR_真实库存（深圳仓库）.xlsx"),
        "result_source": "mabang_store_msku_shenzhen_warehouse_inventory",
        "original_row_count": 3, "binding_verified_row_count": 2, "binding_unverified_row_count": 1,
    }
    report_path = Path(result.shenzhen_warehouse_inventory_report_xlsx_path)
    assert _sheet_names(report_path) == ["真实库存（深圳仓库）-组合sku", "真实库存（深圳仓库）-库存sku", "无本地SKU", "无库存数据", "源数据核验信息"]
    combo_records = _load_records(report_path, "真实库存（深圳仓库）-组合sku")
    stock_records = _load_records(report_path, "真实库存（深圳仓库）-库存sku")
    assert combo_records[0]["真实库存（深圳仓库）数量"] == 6
    assert combo_records[0]["商品链接"] == "https://example.test/c"
    assert combo_records[0]["FBA总库存"] == 60
    assert combo_records[0]["加权日销"] == 10
    assert combo_records[0]["可销售天数"] == 6
    assert combo_records[0]["子SKU"] == "STOCK-A * 1, STOCK-B * 2"
    assert stock_records[0]["真实库存（深圳仓库）数量"] == 9
    assert stock_records[0]["商品链接"] == "https://example.test/a"
    assert stock_records[0]["FBA总库存"] == 23
    assert stock_records[0]["加权日销"] == 1
    assert stock_records[0]["可销售天数"] == 23
    no_local_records = _load_records(report_path, "无本地SKU")
    assert [record["MSKU"] for record in no_local_records] == ["MSKU-N"]
    assert no_local_records[0]["商品链接"] == "https://example.test/no-local"
    assert _load_records(report_path, "无库存数据") == []

    post_calls = [call for call in fake_session.calls if call["method"] == "POST"]
    assert len(post_calls) == 1
    assert sleep_calls == []
    assert _form_value(post_calls[0], "stockSkuStr") == "SKU-A\r\nSTOCK-A\r\nSTOCK-B\r\n"
    assert _form_value(post_calls[0], "warehouseIds[]") == "1014318"


def test_explicit_source_failure_never_uses_latest_or_queries_warehouse(monkeypatch, tmp_path):
    from test_source_verification import source, row

    selected = source(tmp_path / "202610141100-shop_店铺MSKU数据.xlsx", [row()])
    latest = Mock(side_effect=AssertionError("must not substitute historical input"))
    search = AsyncMock(side_effect=AssertionError("must not query after source rejection"))
    monkeypatch.setattr(inv, "find_latest_store_msku_file", latest)
    monkeypatch.setattr(inv, "search_warehouse_stock", search)
    with pytest.raises(Exception, match="店铺不一致"):
        asyncio.run(inv.export_store_msku_actual_inventory(
            "other-shop", msku_source_xlsx_path=selected, output_dir=tmp_path / "out",
        ))
    latest.assert_not_called()
    search.assert_not_called()
    assert not (tmp_path / "out").exists()


def test_shenzhen_pipeline_http_auth_boundary(monkeypatch, tmp_path):
    refresh = AsyncMock()
    from test_source_verification import source, row
    record = row(**dict.fromkeys((*inv.SALES_COLUMNS, *inv.FBA_STOCK_COLUMNS), 0))
    selected = source(tmp_path / "202610141100-shop_店铺MSKU数据.xlsx", [record])
    monkeypatch.setattr(inv, "refresh_mabang_auth", refresh)
    session = _FakeSession([_FakeResponse({"success": True, "searchToken": SEARCH_TOKEN}), _FakeResponse(status=403, text_body="actual rejection", body=b"actual rejection")])
    monkeypatch.setattr(inv, "get_auth_context", _fake_auth_context)
    monkeypatch.setattr(inv, "erp_http_session", session)
    with pytest.raises(inv.MabangRequestError, match="status=403.*actual rejection"):
        asyncio.run(inv.export_store_msku_actual_inventory("shop", msku_source_xlsx_path=selected, output_dir=tmp_path / "out"))
    assert len(session.calls) == 2
    refresh.assert_not_awaited()
    assert not list(tmp_path.rglob("*真实库存*.xlsx"))
