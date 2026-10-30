"""Same-batch CLI integration; model routing and live ERP are separate acceptance."""
from __future__ import annotations

from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from lxeskill import cli as lxeskill
from services.mabang.amazon.fba import replenishment_v2_workbook as v2_workbook
from services.mabang.amazon.fba import source_verification as verification
from services.mabang.amazon.fba import store_msku as msku_export
from services.mabang.amazon.fba import store_msku_actual_inventory as inventory
from services.mabang.amazon.fba import unlinked_shipments as shipments
from services.mabang.amazon.fba.replenishment_v2_formulas import replenishment_v2_formulas
from services.mabang.amazon.fba.sku_catalog import LocalSkuDefinition
from test_replenishment_v2_workbook import (
    SOURCE_HEADERS, STORE, _source_rows, _write_confirmed_empty_snapshot, _write_source,
)


@pytest.mark.parametrize("scenario", ["both-known", "v2-office-failed", "invalid-shipments"])
def test_same_batch_supports_independent_version_cli_results(tmp_path, monkeypatch, capsys, scenario):
    if scenario != "invalid-shipments" and (not os.environ.get("LXE_OFFICE_NODE") or not os.environ.get("LXE_OFFICE_CLI")):
        pytest.skip("Office Kit is unavailable; cross-version workbook acceptance remains NOT VERIFIED")
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path))
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    artifacts = tmp_path / ".lxeagent/artifacts/replenish"
    for setting, directory in {
        "MABANG_STORE_MSKU_OUTPUT_DIR": "store_msku",
        "MABANG_STORE_MSKU_ANALYSIS_OUTPUT_DIR": "sales_analysis",
        "MABANG_STORE_MSKU_INVENTORY_OUTPUT_DIR": "actual_inventory",
        "MABANG_STORE_MSKU_REPLENISHMENT_OUTPUT_DIR": "calculation",
    }.items():
        monkeypatch.setenv(setting, str(artifacts / directory))
    rows = _source_rows()[:3]
    for row in rows:
        row.update({"本地库存": 999, "商品链接": f"https://example.test/{row['MSKU']}"})
    headers = (*SOURCE_HEADERS, "本地库存", "商品链接")
    source = _write_source(
        artifacts / "store_msku/202610131000-Amazon-Test_店铺MSKU数据.xlsx", rows,
        headers=headers,
        local_skus=(
            LocalSkuDefinition("SKU-SEA", 1), LocalSkuDefinition("SKU-CLEAR", 1),
            LocalSkuDefinition("SKU-COMBO", 2, (
                inventory.ComboComponent("CHILD-A", Decimal(2)),
                inventory.ComboComponent("CHILD-B", Decimal(3)),
            )),
        ),
    )
    verified = verification.load_verified_source(source, store_name=STORE)
    report_rows, _ = inventory.calculate_inventory_rows(
        inventory.load_store_msku_rows(source, records=verified.records),
        combo_map=verified.combo_map,
        stock_quantities={"SKU-SEA": Decimal(12), "SKU-CLEAR": Decimal(0),
                          "CHILD-A": Decimal(20), "CHILD-B": Decimal(21)},
        unverified_rows=verification.unverified_reasons(verified.metadata),
    )
    report = inventory.write_actual_inventory_xlsx(
        report_rows, artifacts / "actual_inventory/202610131000-Amazon-Test_真实库存（深圳仓库）.xlsx",
    )
    verification.stamp_report(report, verified.metadata)
    snapshot = _write_confirmed_empty_snapshot(artifacts / "unlinked_shipments_snapshots")
    if scenario == "invalid-shipments":
        raw = tmp_path / "shipments.csv"
        raw.write_text(f"店铺,MSKU,MSKU发货量\n{STORE},SEA-1,17\n{STORE},COMBO-1,23\n", encoding="utf-8")
        query = shipments.StoreUnlinkedShipmentDownloadResult(
            STORE, 401403, "202610131000", [
                shipments.UnlinkedShipmentStatusResult(spec.status_name, int(index == 0), raw_file_path=str(raw) if index == 0 else "")
                for index, spec in enumerate(shipments.UNLINKED_SHIPMENT_STATUS_SPECS)
            ],
        )
        snapshot = Path(shipments.build_store_unlinked_shipments_snapshot(
            [raw], store_name=STORE, query_result=query, snapshot_time="202610131100", output_dir=snapshot.parent,
        ).snapshot_xlsx_path)
        book = load_workbook(snapshot)
        sheet = book[shipments.SNAPSHOT_SUMMARY_SHEET]
        for row in sheet.iter_rows(min_row=2):
            if row[1].value == "SEA-1":
                row[2].value = "invalid"
        book.save(snapshot)
        book.close()
    inputs = [source, report, snapshot]
    hashes = [hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs]

    def unexpected_acquisition(*_args, **_kwargs):
        pytest.fail("both calculators must consume the same acquired batch without ERP or inventory recomputation")

    monkeypatch.setattr(msku_export, "download_store_msku_excel", unexpected_acquisition)
    monkeypatch.setattr(inventory, "calculate_inventory_rows", unexpected_acquisition)
    monkeypatch.setattr(inventory, "search_warehouse_stock", unexpected_acquisition)
    monkeypatch.setattr(inventory, "download_warehouse_stock_xlsx", unexpected_acquisition)

    def execute(skill, arguments):
        monkeypatch.setenv("LXESKILL_SKILL_SCOPE", skill)
        code = lxeskill.main(arguments)
        records = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line.strip()]
        assert len(records) == 1 and records[0]["type"] == "result"
        return code, records[0]

    sales_code, sales = execute("replenishment-sales-analyze", ["replenish", "sales", "analyze", "--store-name", STORE])
    v1 = None
    assert sales_code == 0 and sales["ok"] is True
    assert Path(sales["data"]["source_xlsx_path"]) == source
    v1_code, v1 = execute("replenishment-calculate", [
        "replenish", "calculate", "--store-name", STORE,
        "--unlinked-shipments-snapshot", str(snapshot),
    ])
    if scenario == "invalid-shipments":
        assert v1_code == lxeskill.EXIT_BUSINESS and v1["ok"] is False
        assert "SEA-1" in v1["error"]["message"] and "未关联数量" in v1["error"]["message"]
        assert v1["files"] == []
        assert not list((artifacts / "calculation").glob("*.xlsx"))
    else:
        assert v1_code == 0 and v1["ok"] is True
        assert Path(v1["data"]["sales_analysis_xlsx_path"]) == Path(sales["data"]["report_xlsx_path"])
        assert Path(v1["data"]["actual_inventory_xlsx_path"]) == report
        assert v1["data"]["source_data_time"] == sales["data"]["source_data_time"] == "202610131000"
        assert v1["data"]["row_count"] == 3  # V1 keeps its existing eligibility rules.
        assert v1["files"] == [v1["data"]["report_xlsx_path"]]
        book = load_workbook(v1["files"][0], read_only=True)
        try:
            assert book.sheetnames == ["最终备货意见", "本轮不备货", "链接备货汇总", "源数据核验信息"]
        finally:
            book.close()

    if scenario in {"v2-office-failed", "invalid-shipments"}:
        def fail_office(*_args, **_kwargs):
            raise RuntimeError("fixture V2 Office recalculation failed")
        monkeypatch.setattr(v2_workbook, "_recalculate", fail_office)
    v2_code, v2 = execute("replenishment-workflow-map", [
        "replenish", "v2", "generate", "--store-name", STORE,
        "--msku-source-xlsx-path", str(source),
        "--shenzhen-inventory-report-xlsx-path", str(report),
        "--unlinked-shipments-snapshot-path", str(snapshot),
    ])
    if scenario in {"v2-office-failed", "invalid-shipments"}:
        assert v2_code == lxeskill.EXIT_BUSINESS and v2["ok"] is False
        expected = "未关联数量" if scenario == "invalid-shipments" else "fixture V2 Office recalculation failed"
        assert expected in v2["error"]["message"]
        assert v2["files"] == []
        if scenario == "invalid-shipments":
            assert "SEA-1" in v2["error"]["message"]
            assert not list((artifacts / "final_restock_recommendation").glob("*.xlsx"))
    else:
        assert v2_code == 0 and v2["ok"] is True
        assert v2["data"]["unknown_inventory_count"] == 0
        assert v2["data"]["formula_error_count"] == 0
        assert v2["data"]["row_count"] == len(rows)
        assert v2["files"] == [v2["data"]["output_xlsx_path"]]
        assert "库存待核对" not in Path(v2["files"][0]).name
        book = load_workbook(v2["files"][0], read_only=True, data_only=False)
        try:
            assert book.sheetnames == ["mskulist"] and book.active.max_column == 35
            for index, expected in enumerate([12, 0, 7], 2):
                assert book.active[f"AG{index}"].value == expected
                assert {cell.column_letter for cell in book.active[index] if cell.data_type == "f"} == set(replenishment_v2_formulas(index))
        finally:
            book.close()
    if v1 and v1["ok"]:
        assert Path(v1["files"][0]).is_file()  # V2 failure never replaces or invalidates V1.
        assert set(v1["files"]).isdisjoint(v2["files"])
    assert [hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs] == hashes
