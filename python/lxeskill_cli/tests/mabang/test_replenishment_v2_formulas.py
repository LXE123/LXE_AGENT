"""Formula contract and isolated LibreOffice checks; no ERP or report generation."""
from __future__ import annotations

import shutil
import subprocess

import pytest
from openpyxl import Workbook, load_workbook

from services.mabang.amazon.fba.replenishment_v2_formulas import replenishment_v2_formulas


def test_formula_contract():
    formulas = replenishment_v2_formulas(2)
    assert tuple(formulas) == ("I", "J", "R", "S", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC", "AI")
    assert formulas["J"] == "=E2/7*0.6+F2/14*0.3+G2/30*0.1"
    assert formulas["R"] == "=K2+L2+M2+O2+P2+Q2"
    assert formulas["S"] == '=IF(J2=0,"",R2/J2)'
    # The confirmed XLSX rules guard both shipping days and sea-weight inputs.
    assert formulas["AA"] == '=IF(OR(V2="",W2=""),"",IF((J2*W2-(R2+T2)-IF(Y2="",0,Y2))<=0,"",ROUND(J2*W2-(R2+T2)-IF(Y2="",0,Y2),0)))'
    assert formulas["AC"] == '=IF(OR(AB2="",AD2=""),"",AB2*AD2)'
    assert formulas["AI"] == '=IF(OR(J2="",J2=0),"",(R2+T2+IF(Y2="",0,Y2)+IF(AB2="",0,AB2))/J2)'


def test_seasonal_boost_is_on_by_default_and_can_be_disabled():
    enabled = replenishment_v2_formulas(2)
    disabled = replenishment_v2_formulas(2, peak_boost_enabled=False)
    assert enabled["V"] == replenishment_v2_formulas(2, peak_boost_enabled=True)["V"]
    assert 'SEARCH("热销",AF2)' in enabled["V"]
    assert 'SEARCH("热销",AF2)' not in disabled["V"]
    assert "+0,120" in disabled["V"] and "+0,120" in disabled["W"]
    assert {column: enabled[column] for column in ("I", "J", "R", "S", "U", "X", "Y", "Z", "AA", "AB", "AC", "AI")} == {
        column: disabled[column] for column in ("I", "J", "R", "S", "U", "X", "Y", "Z", "AA", "AB", "AC", "AI")
    }


def test_seasonal_boost_setting_requires_a_boolean():
    with pytest.raises(ValueError, match="peak_boost_enabled"):
        replenishment_v2_formulas(2, peak_boost_enabled=1)


@pytest.mark.parametrize("days", [0, 20, 120])
def test_custom_days_replace_only_the_existing_seasonal_bonus(days):
    default = replenishment_v2_formulas(2)
    custom = replenishment_v2_formulas(2, peak_boost_days=days)
    for column, formula in default.items():
        assert custom[column] == (formula.replace(",30,0))", f",{days},0))") if column in ("V", "W") else formula)
    assert replenishment_v2_formulas(2, peak_boost_days=None) == default
    assert replenishment_v2_formulas(2, peak_boost_days=30) == default
    assert replenishment_v2_formulas(2, peak_boost_enabled=False, peak_boost_days=days) == replenishment_v2_formulas(2, peak_boost_enabled=False)


@pytest.mark.parametrize("days", [-1, 121, 1.5, True, "20"])
def test_formula_rejects_invalid_custom_days(days):
    with pytest.raises(ValueError, match="peak_boost_days"):
        replenishment_v2_formulas(2, peak_boost_days=days)


# Each case overrides cells only to isolate a formula's business boundary.
# None denotes an empty cached result; formula-empty and truly empty inputs
# are deliberately distinguished in the AC cases below.
CASES = [
    ("supply-empty-daily", {"J": None, "R": 100, "T": 20, "Y": 30, "AB": 40}, {"AI": None}),
    ("supply-zero-daily", {"J": 0, "R": 100, "T": 20, "Y": 30, "AB": 40}, {"AI": None}),
    ("supply-empty-advice", {"J": 2, "R": 100, "T": 20, "Y": '=""', "AB": '=""'}, {"AI": 60}),
    ("supply-with-advice", {"J": 2, "R": 100, "T": 20, "Y": 30, "AB": 40}, {"AI": 95}),
    ("supply-unrounded", {"J": 3, "R": 1, "T": 0, "Y": None, "AB": None}, {"AI": 1 / 3}),
    ("sample-9", {"E": 7, "F": 14, "G": 9}, {"I": "样本不足"}),
    ("sample-10", {"E": 0, "F": 0, "G": 10}, {"I": "下降"}),
    ("zero-baseline-sales", {"E": 10, "F": 10, "G": 10}, {"I": "增长"}),
    ("zero-sales-sample-first", {"E": 0, "F": 0, "G": 0}, {"I": "样本不足"}),
    ("trend-085", {"E": 17, "F": 34, "G": 90}, {"I": "下降"}),
    ("trend-above-085", {"E": 17, "F": 34, "G": 89}, {"I": "平稳"}),
    ("trend-12", {"E": 21, "F": 42, "G": 70}, {"I": "增长"}),
    ("trend-below-12", {"E": 21, "F": 42, "G": 71}, {"I": "平稳"}),
    ("daily-weights-ignore-90d", {"E": 7, "F": 28, "G": 90, "H": 999}, {"J": 1.5}),
    ("daily-no-display-rounding", {"E": 1, "F": 0, "G": 0}, {"J": 0.6 / 7}),
    ("fba-components", {"K": 1, "L": 2, "M": 3, "N": 999, "O": 4, "P": 5, "Q": 6}, {"R": 21}),
    ("zero-daily-days", {"J": 0, "R": 10}, {"S": None}),
    ("days-not-floored", {"J": 3, "R": 10}, {"S": 10 / 3}),
    ("clearance-before-hot", {"I": "增长", "J": 11, "S": 80, "AF": "清货，热销，新品"}, {"V": None, "W": None}),
    ("no-replenish-cancels-boost-only", {"I": "平稳", "J": 6, "S": 80, "AF": "不补，热销，出口退税"}, {"V": 75, "W": 90}),
    ("hot-declining", {"I": "下降", "J": 6, "S": 80, "AF": "热销产品"}, {"V": 90, "W": 120}),
    ("peak-steady-no-boost", {"I": "平稳", "J": 11, "S": 80, "AF": "稳定产品"}, {"V": 90, "W": 95}),
    ("peak-declining-no-boost", {"I": "下降", "J": 11, "S": 80, "AF": "稳定产品"}, {"V": 75, "W": 95}),
    ("peak-cap-100", {"I": "平稳", "J": 21, "S": 80, "AF": "热销"}, {"V": 120, "W": 120}),
    ("sample-air-empty", {"I": "样本不足", "J": 1, "S": 80, "AF": "热销"}, {"V": None, "W": None}),
    ("sea-days-70", {"J": 6, "S": 70}, {"W": None}),
    ("sea-daily-5", {"J": 5, "S": 80}, {"W": None}),
    ("sea-days-empty", {"J": 6, "S": '=""'}, {"W": None}),
    ("urgent-at-40-priority", {"I": "平稳", "J": 6, "S": 40, "Y": 10, "AB": 20}, {"U": "空运（急发）"}),
    ("air-above-40", {"I": "平稳", "J": 6, "S": 40.01, "Y": 10, "AB": 20}, {"U": "空运"}),
    ("air-at-70", {"I": "平稳", "J": 6, "S": 70, "Y": 10, "AB": 20}, {"U": "空运"}),
    ("shipping-at-5", {"I": "平稳", "J": 5, "S": 80, "Y": 10, "AB": 20}, {"U": None}),
    ("sea-only", {"I": "平稳", "J": 6, "S": 80, "Y": '=""', "AB": 20}, {"U": "海运"}),
    ("both-above-70", {"I": "平稳", "J": 6, "S": 70.01, "Y": 10, "AB": 20}, {"U": "空运＋海运"}),
    ("shipping-no-sea", {"I": "平稳", "J": 6, "S": 80, "Y": 10, "AB": '=""'}, {"U": None}),
    ("shipping-no-air", {"I": "平稳", "J": 6, "S": 40, "Y": '=""'}, {"U": None}),
    ("shipping-sample", {"I": "样本不足", "J": 6, "S": 40, "Y": 10}, {"U": None}),
    ("shipping-no-days", {"I": "平稳", "J": 6, "S": '=""', "Y": 10}, {"U": None}),
    ("air-plan-no-days", {"V": '=""'}, {"X": None}),
    ("air-plan-zero", {"J": 1, "V": 10, "R": 10, "T": 0}, {"X": None}),
    ("air-plan-negative", {"J": 1, "V": 10, "R": 11, "T": 0}, {"X": None}),
    ("air-plan-deduct-once", {"J": 2, "V": 10, "R": 1, "T": 3}, {"X": 16}),
    ("air-plan-positive-rounds-zero", {"J": 1, "V": 10, "R": 9.9, "T": 0}, {"X": 0, "Y": None}),
    ("air-plan-half-up", {"J": 1, "V": 10, "R": 6.5, "T": 0}, {"X": 4}),
    ("air-weight-no-quantity", {"Y": '=""', "AD": 50}, {"Z": None}),
    ("air-weight-no-weight", {"Y": 10, "AD": None}, {"Z": None}),
    ("air-weight-grams", {"Y": 10, "AD": 2.5}, {"Z": 25}),
    ("sea-plan-no-air-days", {"V": '=""', "W": 90, "J": 6, "R": 480, "T": 0, "Y": '=""'}, {"AA": None}),
    ("sea-plan-truly-empty-air-days", {"V": None, "W": 90, "J": 6, "R": 480, "T": 0, "Y": '=""'}, {"AA": None}),
    ("sea-plan-no-days", {"V": 75, "W": '=""'}, {"AA": None}),
    ("sea-plan-deductions", {"V": 75, "W": 90, "J": 6, "R": 480, "T": 10, "Y": 10}, {"AA": 40}),
    ("sea-plan-empty-air-quantity", {"V": 75, "W": 90, "J": 6, "R": 480, "T": 0, "Y": '=""'}, {"AA": 60}),
    ("sea-plan-zero", {"V": 75, "W": 90, "J": 6, "R": 480, "T": 60, "Y": '=""'}, {"AA": None}),
    ("sea-plan-positive-rounds-zero", {"V": 75, "W": 90, "J": 6, "R": 539.9, "T": 0, "Y": '=""'}, {"AA": 0, "AB": None}),
    ("sea-weight-grams", {"AB": 10, "AD": 2.5}, {"AC": 25}),
    ("sea-weight-formula-empty", {"AB": '=""', "AD": 50}, {"AC": None}),
    ("sea-weight-truly-empty", {"AB": None, "AD": 50}, {"AC": None}),
    ("sea-weight-missing", {"AB": 10, "AD": None}, {"AC": None}),
    ("sea-weight-formula-empty-weight", {"AB": 10, "AD": '=""'}, {"AC": None}),
    ("sea-weight-zero-quantity", {"AB": 0, "AD": 50}, {"AC": 0}),
    ("sea-weight-zero-weight", {"AB": 10, "AD": 0}, {"AC": 0}),
    ("sea-weight-invalid-is-not-hidden", {"AB": "invalid", "AD": 50}, {"AC": "#VALUE!"}),
    ("whole-row-clearance", {"E": 7, "F": 14, "G": 30, "AF": "清货", "AD": 50},
     {"V": None, "W": None, "AA": None, "AB": None, "AC": None}),
    ("whole-row-sea", {"E": 42, "F": 84, "G": 180, "H": 540, "K": 450, "L": 10, "M": 5, "N": 999, "O": 5, "P": 5, "Q": 5, "T": 10, "AD": 100},
     {"I": "平稳", "J": 6, "R": 480, "S": 80, "U": "海运", "V": 75, "W": 90, "X": None, "Y": None, "Z": None, "AA": 50, "AB": 50, "AC": 5000}),
    ("whole-row-combined", {"E": 42, "F": 84, "G": 180, "K": 426, "N": 999, "AD": 100},
     {"I": "平稳", "J": 6, "R": 426, "S": 71, "U": "空运＋海运", "V": 75, "W": 90, "X": 24, "Y": 20, "Z": 2000, "AA": 94, "AB": 90, "AC": 9000}),
]

for trend, days in (("增长", (60, 75, 75, 90, 90, 90)), ("平稳", (45, 60, 60, 75, 75, 90)), ("下降", (45, 45, 45, 60, 60, 75))):
    for daily, expected in zip((2, 2.01, 5, 5.01, 10, 10.01), days):
        CASES.append((f"air-band-{trend}-{daily}", {"I": trend, "J": daily}, {"V": expected}))

for daily, expected in ((5.01, 90), (10, 90), (10.01, 95), (20, 95), (20.01, 100)):
    CASES.append((f"sea-band-{daily}", {"J": daily, "S": 80}, {"W": expected}))

for keyword in ("新品-海运", "一般产品", "透明计划", "稳定产品", "重点产品", "重点关注", "有机会", "出口退税"):
    for trend, air, sea in (("增长", 120, 120),):
        CASES.append((f"peak-{keyword}-{trend}", {"I": trend, "J": 11, "S": 80, "AF": keyword}, {"V": air, "W": sea}))

for remark in ("", "滞销"):
    CASES.append((f"no-peak-{remark}", {"I": "增长", "J": 11, "S": 80, "AF": remark}, {"V": 90, "W": 95}))

for source, target in (("X", "Y"), ("AA", "AB")):
    for quantity, expected in ((4, None), (5, 10), ('=""', None)):
        CASES.append((f"round-{target}-{quantity}", {source: quantity}, {target: expected}))


CASES = [(name, inputs, expected, {}) for name, inputs, expected in CASES]
for enabled, days, remark, expected in [
    (False, 20, "热销", (90, 95)), (True, 20, "热销", (110, 115)),
    (True, 0, "热销", (90, 95)), (True, 120, "热销", (120, 120)),
    (True, 20, "不补，热销", (90, 95)), (True, 20, "清货，热销", (None, None)),
]:
    CASES.append((f"configured-{enabled}-{days}-{remark}",
                  {"I": "平稳", "J": 11, "S": 80, "AF": remark},
                  dict(zip(("V", "W"), expected)),
                  {"peak_boost_enabled": enabled, "peak_boost_days": days}))


@pytest.fixture(scope="module")
def recalculated_cases(tmp_path_factory):
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None:
        pytest.skip("NOT VERIFIED: LibreOffice is unavailable for formula evaluation")
    root = tmp_path_factory.mktemp("replenishment-v2-formulas")
    book = Workbook()
    sheet = book.active
    sheet.title = "formula_cases"
    for row, (name, inputs, _, settings) in enumerate(CASES, 2):
        sheet[f"A{row}"] = name
        for column in ("E", "F", "G", "H", "K", "L", "M", "N", "O", "P", "Q", "T", "AD"):
            sheet[f"{column}{row}"] = 0
        for column, formula in replenishment_v2_formulas(row, **settings).items():
            sheet[f"{column}{row}"] = formula
        for column, value in inputs.items():
            sheet[f"{column}{row}"] = value
    source = root / "cases.xlsx"
    book.save(source)
    book.close()
    output = root / "recalculated"
    output.mkdir()
    # Validation oracle only. Production Office execution belongs to PR2.
    result = subprocess.run(
        [executable, f"-env:UserInstallation={(root / 'profile').as_uri()}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(output), str(source)],
        capture_output=True, text=True, timeout=60,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    recalculated = output / source.name
    assert recalculated.is_file(), result.stdout + result.stderr
    cached = load_workbook(recalculated, data_only=True)
    try:
        assert cached.active.max_row == len(CASES) + 1
        yield {
            name: {column: cached.active[f"{column}{row}"].value for column in expected}
            for row, (name, _, expected, _) in enumerate(CASES, 2)
        }
    finally:
        cached.close()


@pytest.mark.parametrize("name,inputs,expected,settings", CASES, ids=[case[0] for case in CASES])
def test_recalculated_business_boundaries(recalculated_cases, name, inputs, expected, settings):
    actual = recalculated_cases[name]
    for column, value in expected.items():
        if isinstance(value, (int, float)):
            assert actual[column] == pytest.approx(value), (name, column, actual)
        else:
            assert actual[column] == value, (name, column, actual)
