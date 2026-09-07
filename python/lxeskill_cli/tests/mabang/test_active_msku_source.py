from __future__ import annotations

import asyncio
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.mabang.amazon.fba import active_msku_source as active
from services.mabang.amazon.fba import combo_sku as combo
from services.mabang.amazon.fba import store_msku as download
from services.mabang.amazon.fba import store_msku_actual_inventory as inventory
from services.mabang.amazon.fba import store_msku_sales_analysis as sales
from services.mabang.amazon.fba import store_msku_replenishment as replenishment
from services.mabang.amazon.fba.store_resolver import FbaStore
from services.mabang.official_api import OfficialApiError


def row(msku='M', asin='A', sku='S', **extra):
    return {'MSKU': msku, 'ASIN': asin, '本地SKU': sku, **extra}


def binding(msku='M', asin='A', sku='S', kind=1):
    return combo.ListingSkuBinding(msku, asin, sku, kind)


def source(path, records, bindings=(binding(),)):
    headers = list(dict.fromkeys([*inventory.SOURCE_COLUMNS, *sales.REQUIRED_COLUMNS, '本地SKU']))
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.active.append(headers)
    for record in records:
        workbook.active.append([record.get(header) for header in headers])
    workbook.save(path)
    workbook.close()
    if bindings is not None:
        active.annotate_source(path, combo.ActiveListingSnapshot('shop', '10', 'us', bindings), requested_store_name='shop')
    return path


def test_classification_keeps_original_binding_and_only_confirmed_active():
    rows = [row(), row('Amazon.Found.A'), row('inactive'), row('unpaired', sku='')]
    bindings = (binding(), binding(), binding('unpaired', sku='new-live-binding'))
    assert active.classify(rows, bindings) == [('Active', True, ''), ('未确认在售', False, '未匹配本店本站点的 Active Listing'), ('未确认在售', False, '未匹配本店本站点的 Active Listing'), ('Active', True, '')]
    assert rows[3]['本地SKU'] == ''
    assert active.classify([row(' M ', sku=' S ')], (binding(),))[0][1] is True
    assert active.classify([row('m')], (binding(),))[0][1] is False
    assert active.classify([row('Amazon.Found.A')], (binding('Amazon.Found.A'),))[0][1] is True


@pytest.mark.parametrize('records,bindings', [
    ([row()], (binding(asin='changed'),)),
    ([row()], (binding(sku='changed'),)),
    ([row()], (binding(), binding(kind=2))),
    ([row(asin='')], (binding(), binding(asin='B'))),
    ([row(), row('M2')], (binding(), binding('M2', kind=2))),
])
def test_conflicts_do_not_become_exclusions(records, bindings):
    with pytest.raises(OfficialApiError):
        active.classify(records, bindings)


def test_empty_asin_requires_unambiguous_binding():
    assert active.classify([row(asin='')], (binding(), binding())) == [('Active', True, '')]


@pytest.mark.parametrize('mutation', ['sales', 'flag', 'binding', 'store'])
def test_snapshot_rejects_tampering(tmp_path, mutation):
    path = source(tmp_path/'source.xlsx', [row(**{'7天销量': 7})])
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    columns = {cell.value: cell.column for cell in sheet[1]}
    if mutation == 'sales':
        sheet.cell(2, columns['7天销量'], 999)
    elif mutation == 'flag':
        sheet.cell(2, columns['是否参与计算'], False)
    elif mutation == 'binding':
        workbook[active.SHEET].cell(3, 2, workbook[active.SHEET].cell(3, 2).value.replace('"S"', '"changed"'))
    workbook.save(path); workbook.close()
    with pytest.raises(active.ActiveSourceError):
        active.load_active_source(path, store_name='other' if mutation == 'store' else 'shop')


def test_legacy_zero_active_and_report_mix_are_blocked(tmp_path):
    path = source(tmp_path/'legacy.xlsx', [row()], None)
    with pytest.raises(active.ActiveSourceError, match='重新下载'):
        active.load_active_source(path, store_name='shop')
    path = source(tmp_path/'zero.xlsx', [row()], ())
    with pytest.raises(active.ActiveSourceError, match='无符合条件'):
        active.load_active_source(path, store_name='shop')
    a = source(tmp_path/'a.xlsx', [row()]); b = source(tmp_path/'b.xlsx', [row(**{'7天销量': 8})])
    with pytest.raises(active.ActiveSourceError, match='指纹不一致'):
        active.require_matching_reports(a, b, store_name='shop')
    with pytest.raises(active.ActiveSourceError, match='重新下载'):
        active.require_matching_reports(a, tmp_path/'legacy.xlsx', store_name='shop')


def test_end_to_end_active_scope_and_offline_listing_reuse(monkeypatch, tmp_path):
    common = {'父ASIN': 'P', '商品链接': 'https://example.test/A', '7天销量': 7, '14天销量': 14, '30天销量': 30, '90天销量': 90, '可售': 2}
    records = [row(**common), row('COMBO', sku='C', **common), row('no-local', sku='', **common), row('inactive', **{**common, '7天销量': 9000}), row('Amazon.Found.A', **common)]
    bindings = (binding(), binding('COMBO', sku='C', kind=2), binding('no-local', sku='unused-live-sku'))
    path = source(tmp_path/'source'/'202609071200-shop_店铺MSKU数据.xlsx', records, bindings)
    snapshot = active.load_active_source(path, store_name='shop')
    assert snapshot.counts == {'original_row_count': 5, 'active_row_count': 3, 'excluded_row_count': 2}
    queries = []
    async def post(endpoint, body, **kwargs):
        queries.append((endpoint, body))
        assert endpoint == 'combo-skus/search'
        assert body['comboSku'] == 'C'
        return {'code': 200, 'data': {'page': 1, 'rowsPerPage': 20, 'total': 1, 'data': [{'comboSku': 'C', 'comboProductDetail': [{'stockSku': 'S', 'quantity': 2}]}]}}
    async def search(skus):
        assert skus == ['S']
    async def warehouse(**kwargs):
        return tmp_path/'stock.xlsx'
    monkeypatch.setattr(combo, 'post_json', post)
    monkeypatch.setattr(inventory, 'search_warehouse_stock', search)
    monkeypatch.setattr(inventory, 'download_warehouse_stock_xlsx', warehouse)
    monkeypatch.setattr(inventory, 'parse_stock_inventory_xlsx', lambda path: {'S': Decimal(100)})
    analysis = sales.analyze_store_msku_sales('shop', input_dir=tmp_path/'source', output_dir=tmp_path/'sales')
    result = asyncio.run(inventory.export_store_msku_actual_inventory('shop', input_dir=tmp_path/'source', output_dir=tmp_path/'inventory'))
    assert analysis.msku_count == 3
    assert result.matched_warehouse_inventory_msku_row_count == 2
    assert result.missing_local_sku_msku_row_count == 1
    assert result.missing_warehouse_inventory_msku_row_count == 0
    assert len(queries) == 1
    workbook = load_workbook(analysis.report_xlsx_path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            if sheet.title == active.SHEET:
                continue
            values = sheet.values; headers = next(values); rows = [dict(zip(headers, values)) for values in values]
            if rows and '7天销量' in headers:
                assert sum(float(row.get('7天销量') or 0) for row in rows) == 21
    finally:
        workbook.close()
    inputs = replenishment.load_inventory_rows(result.shenzhen_warehouse_inventory_report_xlsx_path)
    assert {item.msku for item in inputs} == {'M', 'COMBO'}
    active.require_matching_reports(Path(analysis.report_xlsx_path), Path(result.shenzhen_warehouse_inventory_report_xlsx_path), store_name='shop')
    # Replaying fixed source/warehouse yields identical rows and no extra Listing request.
    again = asyncio.run(inventory.export_store_msku_actual_inventory('shop', input_dir=tmp_path/'source', output_dir=tmp_path/'inventory2'))
    assert replenishment.load_inventory_rows(again.shenzhen_warehouse_inventory_report_xlsx_path) == inputs


@pytest.mark.parametrize('group', [False, True])
def test_download_publication_and_group_scope(monkeypatch, tmp_path, group):
    async def stores():
        return [FbaStore('shop', '1'), *([FbaStore('child', '2', 'shopId', parent_store_id='1', parent_id_type='fbaWarehouseIds[]')] if group else [])]
    async def pipeline(spec):
        staged = spec.download_file.keywords['output_dir']
        path = source(staged/'202609071200-shop_店铺MSKU数据.xlsx', [row()], None)
        return download.StoreMskuExcelResult('shop', '1', 'fbaWarehouseIds[]', 1, str(path), False, False)
    async def failing_snapshot(name):
        assert not list(tmp_path.glob('*.xlsx')), 'unverified source was published'
        raise OfficialApiError('Listing shop', 'HTTP 401 upstream-denied')
    monkeypatch.setattr(download, 'fetch_fba_stores', stores)
    monkeypatch.setattr(download, 'run_export_pipeline', pipeline)
    monkeypatch.setattr(download, 'fetch_listing_snapshot', failing_snapshot)
    if group:
        result = asyncio.run(download.download_store_msku_excel('1', 'fbaWarehouseIds[]', store_name='shop', output_dir=tmp_path))
        assert Path(result.xlsx_path).exists()
        with pytest.raises(active.ActiveSourceError, match='欧洲整组'):
            active.load_active_source(result.xlsx_path, store_name='shop')
    else:
        with pytest.raises(OfficialApiError, match='upstream-denied'):
            asyncio.run(download.download_store_msku_excel('1', 'fbaWarehouseIds[]', store_name='shop', output_dir=tmp_path))
        assert not list(tmp_path.glob('*.xlsx'))
    assert not list(tmp_path.glob('.active-msku-*'))


@pytest.mark.parametrize('status', ['Inactive', '停售', 'delete', None])
def test_active_query_validates_returned_status(monkeypatch, status):
    async def post(endpoint, body, **kwargs):
        if endpoint == 'shops/list':
            return {'code': 200, 'data': {'profile': {'name': 'shop', 'sid': 10, 'amazonsite': 'us'}}}
        return {'code': 200, 'data': {'list': [{'platformSku': 'M', 'asin': 'A', 'stockSku': 'S', 'stockType': 1, 'shopIds': '10', 'amazonsite': 'us', 'pStatus': status}], 'total': 1, 'totalPage': 1, 'nowPage': 1}}
    monkeypatch.setattr(combo, 'post_json', post)
    with pytest.raises(OfficialApiError, match='平台状态'):
        asyncio.run(combo.fetch_listing_snapshot('shop'))


@pytest.mark.parametrize('extra', [{'店铺名称': 'other'}, {'站点': '德国站'}])
def test_source_scope_is_not_inferred_from_matching_sku(tmp_path, extra):
    path = tmp_path/'source.xlsx'
    workbook = Workbook(); workbook.active.append(['MSKU', 'ASIN', '本地SKU', *extra])
    workbook.active.append(['M', 'A', 'S', *extra.values()]); workbook.save(path); workbook.close()
    with pytest.raises(OfficialApiError, match='其他店铺或站点'):
        active.annotate_source(path, combo.ActiveListingSnapshot('shop', '10', 'us', (binding(),)), requested_store_name='shop')


def test_download_listing_deadline_cancels_without_publication(monkeypatch, tmp_path):
    cancelled = []
    async def stores():
        return [FbaStore('shop', '1')]
    async def pipeline(spec):
        path = source(spec.download_file.keywords['output_dir']/'202609071200-shop_店铺MSKU数据.xlsx', [row()], None)
        return download.StoreMskuExcelResult('shop', '1', 'fbaWarehouseIds[]', 1, str(path), False, False)
    async def forever(name):
        try:
            await asyncio.Event().wait()
        finally:
            cancelled.append(True)
    monkeypatch.setattr(download, 'fetch_fba_stores', stores)
    monkeypatch.setattr(download, 'run_export_pipeline', pipeline)
    monkeypatch.setattr(download, 'fetch_listing_snapshot', forever)
    monkeypatch.setattr(download, 'OFFICIAL_LOOKUP_TIMEOUT_SECONDS', .01)
    with pytest.raises(OfficialApiError, match='TimeoutError'):
        asyncio.run(download.download_store_msku_excel('1', 'fbaWarehouseIds[]', store_name='shop', output_dir=tmp_path))
    assert cancelled == [True]
    assert list(tmp_path.iterdir()) == []


def test_download_cli_business_error_does_not_request_cookie_refresh(monkeypatch):
    from services.agent_cli.mabang import download_store_msku_excel as cli
    async def failure(*args, **kwargs):
        raise OfficialApiError('sid=401403', 'HTTP 401 upstream actual-denial')
    monkeypatch.setattr(cli, 'download_store_msku_excel', failure)
    payload = cli.run({'store_name': 'shop', 'store_id': '1', 'id_type': 'shopId'})
    assert payload['auth_refresh_required'] is False
    assert 'actual-denial' in payload['exception']


@pytest.mark.parametrize("field,value", [("version", 2), ("active_row_count", -1), ("original_row_count", None), ("snapshot_id", ""), ("site", None)])
def test_invalid_metadata_is_rejected_before_calculation(tmp_path, field, value):
    import json
    path = source(tmp_path/'report.xlsx', [row()])
    workbook = load_workbook(path)
    metadata = json.loads(workbook[active.SHEET].cell(2, 2).value)
    metadata[field] = value
    workbook[active.SHEET].cell(2, 2, json.dumps(metadata))
    workbook.save(path); workbook.close()
    with pytest.raises(active.ActiveSourceError, match="核验信息无效"):
        active.require_matching_reports(path, path, store_name="shop")
