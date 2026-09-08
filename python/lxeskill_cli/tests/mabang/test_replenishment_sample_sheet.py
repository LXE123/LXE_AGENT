from copy import deepcopy
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from services.mabang.amazon.fba import store_msku_replenishment as rep
from services.mabang.amazon.fba import replenishment_template as templates
from services.mabang.amazon.fba.replenishment_formula_sheet import SOURCE_VALUE_COLUMNS
from services.mabang.amazon.fba.store_msku_sales_analysis import compute_sales_metrics
from test_mabang_store_msku_replenishment import _inventory_input_row, _write_sales_report


def sample_row(sales=(2, 4, 9, 81), *, actual=-86, template=None, trend=None):
    values = dict.fromkeys(SOURCE_VALUE_COLUMNS, 0.0)
    values.update(dict(zip(SOURCE_VALUE_COLUMNS[:4], sales)))
    values['可售'] = 28
    inv = replace(_inventory_input_row('LOW', fba_total_inventory=28), actual_inventory=actual)
    metrics = compute_sales_metrics(*sales[:3])
    detail = rep.SalesDetail(trend or metrics.trend, None, *sales[:3], source_values=values)
    return rep.calculate_replenishment_rows(
        [inv], {(inv.msku, inv.parent_asin, inv.asin, inv.local_sku): detail},
        template or templates.load_default_template(), {'LOW': 3},
    )[0]


@pytest.mark.parametrize('sales', [(0, 0, 0, 3), (1, 1, 1, 20), (2, 4, 9, 81), (2, 4, 10, 81)])
def test_low_sales_boundary_and_snapshot_columns(tmp_path, sales):
    row = sample_row(sales)
    assert (row.sheet_name == rep.SAMPLE_INSUFFICIENT_SHEET) == (sales[2] < 10)
    path = rep.write_replenishment_report([row], tmp_path / 'report.xlsx')
    book = load_workbook(path)
    try:
        sheet = book[rep.NON_SHIPPING_SHEET]
        assert [c.value for c in sheet[1]] == list(rep.NON_SHIPPING_COLUMNS)
        assert sheet.max_column == 21
        assert sheet.freeze_panes == 'A2'
        assert sheet.max_row == 2
        if sales[2] < 10:
            assert sheet.auto_filter.ref == 'A1:U2'
            record = dict(zip(rep.NON_SHIPPING_COLUMNS, (c.value for c in sheet[2])))
            assert [record[k] for k in SOURCE_VALUE_COLUMNS[:4]] == list(sales)
            assert record['加权日销'] == round(row.weighted_daily_sales, 2)
            assert record[rep.MABANG_FBA_TOTAL_COLUMN] == 28
            assert record[rep.ACTUAL_INVENTORY_QUANTITY_COLUMN] == -86
            assert record['未关联数量'] == 3
            assert record['具体原因'] == '近30天销量不足10件，本轮不计算备货'
            assert all(c.data_type != 'f' for c in sheet[2])
            assert all(c.data_type == 'n' and c.number_format == '0' for c in sheet[2][10:14])
            assert sheet['O2'].number_format == '0.00'
        assert all(d.height == 15 for d in sheet.row_dimensions.values())
        assert all(d.width == 15 for d in sheet.column_dimensions.values())
    finally:
        book.close()


def test_custom_skip_reason_and_effective_weights(tmp_path):
    template = templates.load_default_template()
    params = deepcopy(template.params)
    for group in ('growth', 'decline', 'stable'):
        params['trend_groups'][group] = [t for t in params['trend_groups'][group] if t != '增长']
    params['trend_groups']['skip'].append('增长')
    params['weighted_sales'].update({'7d_weight': .2, '14d_weight': .3, '30d_weight': .5})
    template = replace(template, name='自定义跳过', params=params)
    row = sample_row((14, 21, 30, 180), actual=None, template=template, trend='增长')
    payload = rep._non_shipping_payload(row)
    assert payload['具体原因'] == '参数方案「自定义跳过」将销量趋势「增长」设置为跳过，本轮不计算备货'
    assert payload[rep.ACTUAL_INVENTORY_QUANTITY_COLUMN] == ''
    assert payload['加权日销'] == 1.35
    assert row.sheet_name == rep.SAMPLE_INSUFFICIENT_SHEET


@pytest.mark.parametrize('bad', [None, '', 'invalid', float('nan'), float('inf'), -1, True, 'missing_inputs'])
@pytest.mark.parametrize('existing', [False, True])
def test_invalid_sample_sales_never_publish(tmp_path, bad, existing):
    target = tmp_path / 'report.xlsx'
    previous = b'previous valid report sentinel'
    if existing:
        target.write_bytes(previous)
    row = sample_row()
    if bad == 'missing_inputs':
        row = replace(row, formula_inputs=None)
    else:
        values = {**row.formula_inputs.values, '90天销量': bad}
        row = replace(row, formula_inputs=replace(row.formula_inputs, values=values))
    with pytest.raises((ValueError, rep.StoreMskuReplenishmentError), match='LOW'):
        rep.write_replenishment_report([row], target)
    assert target.read_bytes() == previous if existing else not target.exists()
    assert list(tmp_path.iterdir()) == ([target] if existing else [])


def test_numeric_strings_come_from_matching_sales_input(tmp_path):
    path = _write_sales_report(tmp_path / 'sales.xlsx')
    book = load_workbook(path)
    try:
        ws = book['MSKU明细']
        headers = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == 'SAMPLE-1':
                for field, value in zip(SOURCE_VALUE_COLUMNS[:4], ('1', '2', '9', '1,200')):
                    row[headers.index(field)].value = value
        book.save(path)
    finally:
        book.close()
    details = rep.load_sales_details(path)
    key = ('SAMPLE-1', 'PARENT-SAMPLE', 'ASIN-S', 'SKU-S')
    inv = replace(_inventory_input_row('SAMPLE-1', fba_total_inventory=80),
                  parent_asin=key[1], asin=key[2], local_sku=key[3])
    # Same MSKU with a different complete key must keep its own source values.
    other = replace(inv, asin='ASIN-OTHER')
    other_key = (other.msku, other.parent_asin, other.asin, other.local_sku)
    details[other_key] = replace(details[key], sales_7d=0, sales_14d=0, sales_30d=0,
                                source_values={**details[key].source_values, **dict.fromkeys(SOURCE_VALUE_COLUMNS[:4], 0)})
    rows = rep.calculate_replenishment_rows([inv, other], details)
    payloads = [rep._non_shipping_payload(row) for row in rows]
    assert [payloads[0][k] for k in SOURCE_VALUE_COLUMNS[:4]] == [1, 2, 9, 1200]
    assert [payloads[1][k] for k in SOURCE_VALUE_COLUMNS[:4]] == [0, 0, 0, 0]
