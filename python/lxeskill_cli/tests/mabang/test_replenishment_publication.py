from __future__ import annotations

import os
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from services.mabang.amazon.fba import report_staging, store_msku_replenishment as rep
from services.mabang.amazon.fba.replenishment_formula_sheet import cache_formula_values
from test_replenishment_formula_sheet import case_row


def existing_report(path):
    rep.write_replenishment_report([case_row("OLD")], path)
    return path.read_bytes()


@pytest.mark.parametrize("existing", [False, True])
@pytest.mark.parametrize("stage", ["save", "metadata", "cache", "integrity", "publish"])
def test_failure_never_publishes_or_changes_previous_report(tmp_path, monkeypatch, existing, stage):
    target = tmp_path / "202609071712-Amazon-Test_备货建议.xlsx"
    previous = existing_report(target) if existing else None
    failure = OSError(f"actual failure at {stage}")

    def fail(*args, **kwargs):
        raise failure

    if stage == "save":
        monkeypatch.setattr(Workbook, "save", fail)
    elif stage == "metadata":
        monkeypatch.setattr(rep, "stamp_report", fail)
    elif stage == "cache":
        monkeypatch.setattr(rep, "cache_formula_values", fail)
    elif stage == "integrity":
        monkeypatch.setattr(ZipFile, "testzip", fail)
    else:
        original_replace = os.replace

        def fail_publication(source, destination):
            if Path(destination) == target:
                raise failure
            return original_replace(source, destination)

        monkeypatch.setattr(os, "replace", fail_publication)
    with pytest.raises(OSError) as caught:
        rep.write_replenishment_report([case_row("NEW")], target, active_metadata={"fixture": True})
    assert caught.value is failure
    assert target.read_bytes() == previous if existing else not target.exists()
    assert list(tmp_path.iterdir()) == ([target] if existing else [])


def test_publish_happens_once_after_handles_close_and_complete_validation(tmp_path, monkeypatch):
    target = tmp_path / "result.xlsx"
    previous = existing_report(target)
    archives = []
    original_init = ZipFile.__init__
    original_replace = os.replace
    publications = []

    def track(self, *args, **kwargs):
        original_init(self, *args, **kwargs)
        archives.append(self)

    def replace(source, destination):
        assert all(archive.fp is None for archive in archives)
        assert target.read_bytes() == previous
        if Path(destination) == target:
            publications.append((source, destination))
            assert Path(source).parent.parent == target.parent
            book = load_workbook(source, data_only=True)
            try:
                assert book.active["A3"].value == "NEW"
                assert book.active["W3"].value is not None
                assert "Active核验信息" in book.sheetnames
            finally:
                book.close()
        return original_replace(source, destination)

    monkeypatch.setattr(ZipFile, "__init__", track)
    monkeypatch.setattr(os, "replace", replace)
    assert rep.write_replenishment_report([case_row("NEW")], target, active_metadata={"fixture": True}) == target
    assert len(publications) == 1
    assert target.read_bytes() != previous
    assert list(tmp_path.iterdir()) == [target]


def test_corrupt_zip_is_not_published(tmp_path, monkeypatch):
    target = tmp_path / "result.xlsx"
    monkeypatch.setattr(ZipFile, "testzip", lambda self: "xl/worksheets/sheet1.xml")
    with pytest.raises(BadZipFile, match="xl/worksheets/sheet1.xml"):
        rep.write_replenishment_report([], target)
    assert not list(tmp_path.iterdir())


def test_cache_validation_failure_preserves_input(tmp_path):
    path = tmp_path / "source.xlsx"
    book = Workbook()
    book.active["A1"] = 1
    book.save(path)
    book.close()
    original = path.read_bytes()
    with pytest.raises(ValueError, match="不是公式"):
        cache_formula_values(path, {"Sheet": {"A1": 2}})
    assert path.read_bytes() == original
    assert list(tmp_path.iterdir()) == [path]


def test_cleanup_failure_does_not_hide_original_failure(tmp_path, monkeypatch):
    failure = OSError("original save failed")
    clean_failure = PermissionError("cleanup denied")
    warnings = []

    def fail_cleanup(path):
        raise clean_failure

    monkeypatch.setattr(report_staging.shutil, "rmtree", fail_cleanup)
    monkeypatch.setattr(report_staging.logger, "warning", lambda *args: warnings.append(args))
    with pytest.raises(OSError) as caught:
        with report_staging.staged_report_path(tmp_path / "result.xlsx") as staged:
            staged.write_bytes(b"partial")
            raise failure
    assert caught.value is failure
    assert warnings[0][-1] is clean_failure
    assert not list(tmp_path.glob("*-*_备货建议.xlsx"))


@pytest.mark.skipif(os.name != "nt", reason="Requires Windows file sharing semantics")
def test_windows_locked_target_is_preserved_and_retry_succeeds(tmp_path):
    import ctypes
    from ctypes import wintypes

    target = tmp_path / "result.xlsx"
    previous = existing_report(target)
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                  wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    kernel.CreateFileW.restype = wintypes.HANDLE
    kernel.CloseHandle.argtypes = [wintypes.HANDLE]
    # Permit reads/writes, but deliberately omit FILE_SHARE_DELETE.
    handle = kernel.CreateFileW(str(target), 0x80000000, 3, None, 3, 0x80, None)
    assert handle != ctypes.c_void_p(-1).value, ctypes.get_last_error()
    try:
        with pytest.raises(PermissionError) as caught:
            rep.write_replenishment_report([case_row("NEW")], target)
        assert caught.value.winerror in (5, 32)
        assert target.read_bytes() == previous
        assert list(tmp_path.iterdir()) == [target]
    finally:
        kernel.CloseHandle(handle)
    rep.write_replenishment_report([case_row("NEW")], target)
    book = load_workbook(target, data_only=True)
    try:
        assert book.active["A3"].value == "NEW"
    finally:
        book.close()
