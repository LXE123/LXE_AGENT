"""User-confirmed A rules; no ERP requests and no alternate calculation engine."""
from __future__ import annotations

from decimal import Decimal
import os
from pathlib import Path
from shutil import copyfile
from zipfile import BadZipFile

import pytest
from openpyxl import Workbook, load_workbook

from services.mabang.amazon.fba import replenishment_v2_settings as settings
from services.mabang.amazon.fba import replenishment_v2_workbook as service
from services.mabang.amazon.fba import source_verification as verification
from services.mabang.amazon.fba import store_msku_actual_inventory as inventory
from services.mabang.amazon.fba.replenishment_v2_formulas import replenishment_v2_formulas
from services.mabang.amazon.fba.sku_catalog import LocalSkuDefinition
from test_replenishment_v2_workbook import STORE, _source_rows, _write_source, _write_confirmed_empty_snapshot


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))


def inputs(tmp_path, *, weight=50, stock=0, report_sku="SKU-SEA", sheet=None):
    rows = _source_rows()[:1]
    rows[0].update({"备注": "热销", "单品重量(g)(cm)": weight, "可售": -7})
    source = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx", rows,
                           local_skus=(LocalSkuDefinition("SKU-SEA", None),))
    report = tmp_path / "inventory.xlsx"
    book = Workbook()
    book.active.title = sheet or inventory.STOCK_ACTUAL_INVENTORY_SHEET
    book.active.append(["MSKU", inventory.ACTUAL_INVENTORY_QUANTITY_COLUMN, "本地SKU", "子SKU"])
    book.active.append(["SEA-1", stock, report_sku, ""])
    book.save(report)
    book.close()
    metadata = verification.load_verified_source(source, store_name=STORE).metadata
    verification.stamp_report(report, metadata)
    return dict(msku_source_xlsx_path=source, shenzhen_inventory_report_xlsx_path=report,
                unlinked_shipments_snapshot_path=_write_confirmed_empty_snapshot(tmp_path / "snapshot"))


def generate(tmp_path, monkeypatch, **kwargs):
    monkeypatch.setattr(service, "_recalculate", copyfile)
    monkeypatch.setattr(service, "_inspect_recalculated", lambda *_a, **_k: (0, ()))
    return service.generate_replenishment_v2(STORE, output_dir=tmp_path / "out", **kwargs)


def test_unverified_rows_keep_all_fourteen_formulas_and_negative_source_stock(tmp_path, monkeypatch):
    files = inputs(tmp_path)
    def unexpected_sales_analysis(*_args, **_kwargs):
        pytest.fail("unverified V2 rows must not trigger ordinary sales analysis")
    monkeypatch.setattr("services.mabang.amazon.fba.store_msku_sales_analysis.analyze_store_msku_sales", unexpected_sales_analysis)
    assert verification.load_verified_source(files["msku_source_xlsx_path"], store_name=STORE).records[0]["是否通过绑定核验"] is False
    result = generate(tmp_path, monkeypatch, **files)
    book = load_workbook(result.output_xlsx_path, data_only=False)
    try:
        sheet = book.active
        assert sheet.max_column == 35 and sheet["AI1"].value == "可供天数"
        assert len(replenishment_v2_formulas(2)) == 14
        assert sheet["K2"].value == -7
        assert sheet.auto_filter.ref == "A1:AI2"
        assert sheet["AI2"].number_format == "0.00"
        assert all(sheet[f"{col}2"].data_type == "f" for col in replenishment_v2_formulas(2))
    finally:
        book.close()


@pytest.mark.parametrize("stock,sheet", [
    (None, inventory.NO_INVENTORY_DATA_SHEET), (0, inventory.NO_INVENTORY_DATA_SHEET),
    (None, inventory.NO_LOCAL_SKU_SHEET), (0, inventory.NO_LOCAL_SKU_SHEET),
])
def test_unknown_shenzhen_placeholders_map_to_blank_and_keep_all_formulas(tmp_path, monkeypatch, stock, sheet):
    files = inputs(tmp_path, stock=stock, sheet=sheet)
    result = generate(tmp_path, monkeypatch, **files)
    assert result.success is False
    assert result.unknown_inventory_count == 1
    assert result.output_xlsx_path.endswith("-库存待核对.xlsx")
    assert "MSKU=SEA-1" in result.unknown_inventory_examples[0]
    assert sheet in result.unknown_inventory_examples[0]
    assert "原因=" in result.unknown_inventory_examples[0]
    book = load_workbook(result.output_xlsx_path)
    try:
        assert book.active["AG2"].value is None
        assert all(book.active[f"{col}2"].data_type == "f" for col in replenishment_v2_formulas(2))
    finally:
        book.close()


@pytest.mark.parametrize("stock,expected", [(0, 0), (-26, -26), (15, 15), (" 12.5 ", 12.5), ("-1,234.5", -1234.5)])
def test_ag_uses_finite_signed_shenzhen_values(tmp_path, monkeypatch, stock, expected):
    result = generate(tmp_path, monkeypatch, **inputs(tmp_path, stock=stock))
    book = load_workbook(result.output_xlsx_path)
    try:
        assert book.active["AG2"].value == expected
    finally:
        book.close()


def test_msku_source_local_inventory_is_not_required():
    rows = _source_rows()[:1]
    assert "本地库存" not in rows[0]
    assert service._unique_source_rows(rows)[1]["SEA-1"] is rows[0]


def test_shenzhen_quantity_maps_by_each_msku(tmp_path):
    rows = _source_rows()
    rows[-1]["本地库存"] = None
    source = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx", rows)
    verified = verification.load_verified_source(source, store_name=STORE)
    report = tmp_path / "inventory.xlsx"
    book = Workbook()
    book.active.title = inventory.STOCK_ACTUAL_INVENTORY_SHEET
    book.active.append(["MSKU", inventory.ACTUAL_INVENTORY_QUANTITY_COLUMN, "本地SKU", "子SKU"])
    for row in rows:
        book.active.append([row["MSKU"], 9, row["本地SKU"], ""])
    book.save(report)
    book.close()
    verification.stamp_report(report, verified.metadata)
    _, source_by_key = service._unique_source_rows(verified.records)
    inventory_by_key = service._load_shenzhen_rows(report, expected_keys=set(source_by_key))
    mapped = service._rows_for_workbook(
        verified.records, store_name=STORE, inventory_by_key=inventory_by_key, shipment_by_key={},
    )
    assert [row["本地库存"] for row in mapped] == [9, 9, 9, 9]


@pytest.mark.parametrize("failure", ["missing", "corrupt", "schema"])
def test_legal_placeholders_do_not_bypass_report_file_validation(tmp_path, monkeypatch, failure):
    files = inputs(tmp_path, stock=None, sheet=inventory.NO_INVENTORY_DATA_SHEET)
    report = files["shenzhen_inventory_report_xlsx_path"]
    if failure == "missing":
        report.unlink()
        expected, message = FileNotFoundError, "深圳库存报告不存在"
    elif failure == "corrupt":
        report.write_bytes(b"not an xlsx")
        expected, message = BadZipFile, "File is not a zip file"
    else:
        book = load_workbook(report)
        book.active["B1"] = "wrong-column"
        book.save(report)
        book.close()
        expected, message = service.ReplenishmentV2WorkbookError, "缺少列"
    with pytest.raises(expected, match=message):
        generate(tmp_path, monkeypatch, **files)
    assert not list((tmp_path / "out").glob("*.xlsx"))


@pytest.mark.parametrize("stock", [0, -9903])
def test_real_zero_and_signed_stock_are_not_missing(tmp_path, monkeypatch, stock):
    result = generate(tmp_path, monkeypatch, **inputs(tmp_path, stock=stock))
    assert result.success is True


def test_ah_prefers_report_without_weakening_source_fingerprint(tmp_path, monkeypatch):
    files = inputs(tmp_path, report_sku="REPORT-CODE")
    result = generate(tmp_path, monkeypatch, **files)
    book = load_workbook(result.output_xlsx_path)
    assert book.active["AH2"].value == "REPORT-CODE"
    book.close()
    source = load_workbook(files["msku_source_xlsx_path"])
    source.active["B2"] = "tampered"
    source.save(files["msku_source_xlsx_path"])
    source.close()
    with pytest.raises(verification.SourceVerificationError):
        generate(tmp_path, monkeypatch, **files)


@pytest.mark.parametrize("weight,expected", [(None, 0), ("", 0), (0, 0), ("0", 0), ("12.5", 12.5), ("12.5g 10*10", 12.5)])
def test_a_weight_blanks_and_valid_numeric_text(tmp_path, monkeypatch, weight, expected):
    result = generate(tmp_path, monkeypatch, **inputs(tmp_path, weight=weight))
    book = load_workbook(result.output_xlsx_path)
    assert book.active["AD2"].value == expected
    book.close()


def test_bad_weight_identifies_field_msku_stage_and_actual_reason(tmp_path, monkeypatch):
    with pytest.raises(service.ReplenishmentV2WorkbookError) as caught:
        generate(tmp_path, monkeypatch, **inputs(tmp_path, weight="invalid-weight"))
    message = str(caught.value)
    assert all(part in message for part in ("SEA-1", "单品重量(g)(cm)", "AD", "解析", "invalid-weight"))
    assert "下载失败" not in message


@pytest.mark.parametrize("saved,requested", [(True, False), (False, True), (True, True)])
def test_chat_request_controls_boost_without_rewriting_saved_preference(tmp_path, monkeypatch, saved, requested):
    settings.save_parameters({"peak_boost_enabled": saved, "peak_boost_days": 20})
    result = generate(tmp_path, monkeypatch, **inputs(tmp_path), peak_boost_enabled=requested)
    book = load_workbook(result.output_xlsx_path)
    try:
        assert book.active["V2"].value == replenishment_v2_formulas(2, peak_boost_enabled=requested, peak_boost_days=20)["V"]
    finally:
        book.close()
    assert settings.read_state() == {"peak_boost_enabled": saved, "peak_boost_days": 20}


@pytest.mark.parametrize("enabled", [False, True])
def test_omitted_chat_override_uses_the_saved_checkbox(tmp_path, monkeypatch, enabled):
    settings.save_parameters({"peak_boost_enabled": enabled, "peak_boost_days": 20})
    result = generate(tmp_path, monkeypatch, **inputs(tmp_path))
    book = load_workbook(result.output_xlsx_path)
    try:
        assert book.active["V2"].value == replenishment_v2_formulas(2, peak_boost_enabled=enabled, peak_boost_days=20)["V"]
    finally:
        book.close()


@pytest.mark.parametrize("saved_days", [None, 20])
def test_chat_days_enable_boost_and_override_saved_days_without_saving(tmp_path, monkeypatch, saved_days):
    saved = {"peak_boost_enabled": False, "peak_boost_days": saved_days}
    settings.save_parameters(saved)
    result = generate(tmp_path, monkeypatch, **inputs(tmp_path), peak_boost_days=15)
    book = load_workbook(result.output_xlsx_path)
    try:
        assert book.active["V2"].value == replenishment_v2_formulas(2, peak_boost_enabled=True, peak_boost_days=15)["V"]
    finally:
        book.close()
    assert settings.read_state() == saved


def test_enabled_request_without_days_stops_before_reading_inputs(tmp_path, monkeypatch):
    def unexpected(*_args, **_kwargs):
        pytest.fail("missing boost days must stop before reading inputs")
    monkeypatch.setattr(service, "_path", unexpected)
    with pytest.raises(ValueError, match="加成天数"):
        service.generate_replenishment_v2(STORE, msku_source_xlsx_path="source.xlsx",
            shenzhen_inventory_report_xlsx_path="inventory.xlsx", unlinked_shipments_snapshot_path="snapshot.xlsx",
            peak_boost_enabled=True)


def test_ai_recalculates_empty_zero_and_suggested_quantities(tmp_path):
    if not os.environ.get("LXE_OFFICE_NODE"):
        pytest.skip("NOT VERIFIED: Office Kit is unavailable")
    cases = [(None, 100, 20, 30, 40, None), (0, 100, 20, 30, 40, None),
             (2, 100, 20, None, None, 60), (2, 100, 20, 30, 40, 95),
             (3, 1, 0, None, None, 1 / 3)]
    book = Workbook()
    for row, (daily, stock, shipments, air, sea, _) in enumerate(cases, 2):
        for col, value in (("J", daily), ("R", stock), ("T", shipments), ("Y", air), ("AB", sea)):
            book.active[f"{col}{row}"] = value
        book.active[f"AI{row}"] = replenishment_v2_formulas(row)["AI"]
    source, output = tmp_path / "supply.xlsx", tmp_path / "recalculated.xlsx"
    book.save(source)
    book.close()
    service._recalculate(source, output)
    values = load_workbook(output, data_only=True)
    try:
        for row, case in enumerate(cases, 2):
            actual = values.active[f"AI{row}"].value
            assert actual is None if case[-1] is None else actual == pytest.approx(case[-1])
    finally:
        values.close()
