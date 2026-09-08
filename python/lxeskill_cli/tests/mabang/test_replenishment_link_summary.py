from dataclasses import replace

import pytest
from openpyxl import Workbook, load_workbook

from services.mabang.amazon.fba import replenishment_formula_sheet as formulas
from services.mabang.amazon.fba import store_msku_replenishment as rep
from test_replenishment_formula_sheet import case_row


def test_mixed_parent_ties_out_to_final_shipping_sheet(tmp_path):
    rows = [
        case_row("URGENT", daily=3, fba=90, unlinked=0, template_name="默认"),
        case_row("AIR", fba=300, unlinked=20, template_name="默认"),
        case_row("COMPANION"),
        case_row("SEA", fba=480, unlinked=0, template_name="默认"),
    ]
    rows += [replace(rows[0], msku="NO-SHIP", non_shipping_reason_code="sea_daily_threshold", non_shipping_reason_detail="fixture: sea daily threshold", sheet_name=rep.NO_SHIP_SHEET, replenish_quantity=999),
             replace(rows[0], msku="SAMPLE", non_shipping_reason_code="parameter_skip", non_shipping_reason_detail="fixture: skipped trend", sheet_name=rep.SAMPLE_INSUFFICIENT_SHEET, replenish_quantity=888)]
    rows = [replace(row, parent_asin="MIXED") for row in rows]
    rows += [replace(rows[0], msku="ZERO", parent_asin="ZERO-PARENT", replenish_quantity=0),
             replace(rows[0], msku="CLEAR", parent_asin="CLEAR-PARENT", non_shipping_reason_code="clearance", non_shipping_reason_detail="商品备注含‘清货’，本轮不安排备货", sheet_name=rep.CLEARANCE_SHEET)]
    path = rep.write_replenishment_report(rows, tmp_path / "summary.xlsx")
    book = load_workbook(path, data_only=True)
    try:
        summary_sheet = book[rep.SUMMARY_SHEET]
        records = list(summary_sheet.values)
        summaries = [dict(zip(records[0], record)) for record in records[1:]]
        assert [row["父ASIN"] for row in summaries] == ["MIXED", "ZERO-PARENT"]
        mixed, zero = summaries
        assert mixed["MSKU数"] == 6
        assert mixed["链接真实库存（深圳仓库）汇总"] == sum(r.actual_inventory for r in rows[:6])
        assert mixed["链接未关联数量汇总"] == sum(r.unlinked_quantity for r in rows[:6])
        assert (mixed["空运（急发）补货量"], mixed["空运补货量"], mixed["海运建议量"], mixed["总补货量"]) == (90, 140, 270, 500)
        assert all(zero[key] == 0 for key in ("总补货量", "空运（急发）补货量", "空运补货量", "海运建议量"))
        for summary in summaries:
            assert summary["总补货量"] == sum(summary[k] for k in ("空运（急发）补货量", "空运补货量", "海运建议量"))
        shipping = list(book[formulas.FINAL_SHIPPING_SHEET].iter_rows(min_row=3, values_only=True))
        assert {r[0] for r in shipping} == {"URGENT", "AIR", "COMPANION", "SEA"}
        assert sum(r[28] for r in shipping) == mixed["空运（急发）补货量"] + mixed["空运补货量"]
        assert sum(r[29] for r in shipping) == mixed["海运建议量"]
        assert sum(r[28] + r[29] for r in shipping) == sum(r["总补货量"] for r in summaries)
    finally:
        book.close()


def test_companion_air_is_reported_as_air_transport():
    summary = rep.summarize_links([case_row()])[0]
    assert summary["涉及运输方式"] == "空运、海运"
    assert (summary["空运补货量"], summary["海运建议量"], summary["总补货量"]) == (10, 210, 220)


def test_rejected_sea_minimum_keeps_group_but_contributes_no_shipping():
    row = case_row(fba=600, unlinked=31)
    assert row.sheet_name == rep.NO_SHIP_SHEET and row.replenish_quantity == 29
    summary = rep.summarize_links([row])[0]
    assert summary["MSKU数"] == 1
    assert summary["总补货量"] == summary["空运补货量"] == summary["海运建议量"] == 0


@pytest.mark.parametrize("changes", [
    {"sea_quantity": None}, {"sea_quantity": -1},
    {"companion_air_quantity": -1}, {"sea_quantity": 211},
])
def test_inconsistent_shipping_split_fails_both_report_views(changes):
    row = replace(case_row("BAD-SPLIT"), **changes)
    with pytest.raises(ValueError, match="最终发货数量拆分不一致: MSKU=BAD-SPLIT"):
        rep.summarize_links([row])
    book = Workbook()
    try:
        with pytest.raises(ValueError, match="最终发货数量拆分不一致: MSKU=BAD-SPLIT"):
            formulas.write_formula_sheet(book, [row], missing_snapshot=False)
    finally:
        book.close()
