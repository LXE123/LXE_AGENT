from __future__ import annotations

import math
import shutil
import subprocess
from dataclasses import replace
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook

from services.mabang.amazon.fba import replenishment_formula_sheet as formulas
from services.mabang.amazon.fba import replenishment_template as templates
from services.mabang.amazon.fba import store_msku_replenishment as rep
from test_mabang_store_msku_replenishment import _inventory_input_row, _write_sales_report


def case_row(msku="ITEM", *, daily=6, fba=430, unlinked=10, weight=100, template_name="test-companion", source_fba=None):
    values = {key: 0.0 for key in formulas.SOURCE_VALUE_COLUMNS}
    values.update({"7天销量": daily * 7, "14天销量": daily * 14, "30天销量": daily * 30,
                   "90天销量": daily * 90, "可售": fba if source_fba is None else source_fba, "计划入库": 1000})
    inventory = replace(_inventory_input_row(msku, fba_total_inventory=fba), actual_inventory=100,
                        remark="整箱包装", product_name="产品", local_sku_name="本地名")
    detail = rep.SalesDetail("平稳", weight, daily * 7, daily * 14, daily * 30, source_values=values)
    template = templates.get_template("2组-US站点-林美淇" if template_name == "test-companion" else template_name)
    if template_name == "test-companion":
        params = deepcopy(template.params)
        params["shipping"]["air_sales_days_lte"] = 70
        template = replace(template, name="test-companion", params=params)
    return rep.calculate_replenishment_rows(
        [inventory], {(msku, inventory.parent_asin, inventory.asin, inventory.local_sku): detail},
        template, {msku: unlinked},
    )[0]


def test_formulas_cached_initial_values_and_metadata(tmp_path):
    from mabang_test_helpers import _stamp_active_test_report
    from services.mabang.amazon.fba.source_verification import _read_info

    rows = [case_row("SEA"), case_row("AIR", fba=300, unlinked=20, template_name="默认"),
            case_row("URGENT", daily=3, fba=90, unlinked=0, weight=None, template_name="默认"),
            case_row("SEA-ONLY", fba=480, unlinked=0, weight=1, template_name="默认")]
    rows.append(replace(rows[0], msku="NO-SHIP", non_shipping_reason_code="sea_daily_threshold", non_shipping_reason_detail="fixture: sea daily threshold", sheet_name=rep.NO_SHIP_SHEET))
    rows.append(replace(rows[0], msku="CLEARANCE", non_shipping_reason_code="clearance", non_shipping_reason_detail="商品备注含‘清货’，本轮不安排备货", sheet_name=rep.CLEARANCE_SHEET))
    rows.append(replace(rows[0], msku="SAMPLE", non_shipping_reason_code="parameter_skip", non_shipping_reason_detail="fixture: skipped trend", sheet_name=rep.SAMPLE_INSUFFICIENT_SHEET, replenish_quantity=None))
    source = _write_sales_report(tmp_path / "202605251530-Amazon-Test_销量分析.xlsx")
    wb = load_workbook(source)
    metadata, _ = _read_info(wb)
    wb.close()
    path = rep.write_replenishment_report(rows, tmp_path / "result.xlsx", source_metadata=metadata, missing_unlinked_snapshot=True)
    book = load_workbook(path, data_only=True)
    raw = load_workbook(path, data_only=False)
    try:
        sheet = book[formulas.FINAL_SHIPPING_SHEET]
        headers = [cell.value for cell in sheet[2]]
        records = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=3, values_only=True)]
        by_msku = {record["MSKU"]: record for record in records}
        assert [record["MSKU"] for record in records] == ["URGENT", "AIR", "SEA", "SEA-ONLY"]
        assert [(record["空运建议量"], record["海运建议量"]) for record in records] == [(90, 0), (130, 0), (10, 210), (0, 60)]
        assert [record["预计总重量(kg)"] for record in records] == [None, 13, 22, 0.06]
        assert all(record["计划入库"] == 1000 for record in records)  # Planned inbound never enters FBA deduction.
        assert by_msku["SEA"]["FBA总库存（马帮）"] == 430
        assert by_msku["SEA"]["未发出货件总计"] == 10
        assert by_msku["SEA"]["深圳库存缺口"] == 120
        assert [record["可销售天数"] for record in records] == pytest.approx([30, 50, 430 / 6, 80])
        assert [record["建议运输方式（生成时）"] for record in records] == ["空运（急发）", "空运", "空运＋海运", "海运"]
        assert [record["空运补货天数"] for record in records] == [60, 75, 75, 0]
        assert [record["海运补货天数"] for record in records] == [0, 0, 110, 90]
        assert "备货公式参数" not in book.sheetnames
        assert book.sheetnames == ["最终备货意见", "本轮不备货", "链接备货汇总", "源数据核验信息"]
        assert all(record["备注"] == "整箱包装" for record in records)
        assert _read_info(book)[0] == metadata
        assert raw.calculation.calcMode == "auto"
        assert raw.active.title == formulas.FINAL_SHIPPING_SHEET
        assert raw.active.freeze_panes == "E3"
        assert not raw.active.tables
        assert raw.active.auto_filter.ref == "A2:AK6"
        assert not any(d.hidden for d in raw.active.column_dimensions.values())
        assert [c.value for c in raw.active[2]] == list(formulas.FINAL_SHIPPING_COLUMNS)
        assert len(formulas.FINAL_SHIPPING_COLUMNS) == 37
        assert [raw.active.cell(2, col).value for col in (19, 20, 21, 22, 23, 36, 37)] == [
            "深圳可用库存", "可销售天数", "建议运输方式（生成时）", "空运补货天数", "海运补货天数", "输入检查", "试算提示",
        ]
        for address in ("I3", "Q3", "T3", "X3", "Y3", "Z3", "AA3", "AB3", "AC3", "AD3", "AE3", "AG3", "AI3", "AJ3", "AK3"):
            assert raw.active[address].data_type == "f"
        assert raw.active["T3"].value == '=IF(AJ3="正常",IF(I3>0,Q3/I3,""),"")'
        assert raw.active["T3"].number_format == "#,##0.00"
        assert raw.active["AD3"].number_format == "#,##0.##"  # Air quantity, immediately before sea quantity.
        assert raw.active["AE3"].number_format == "#,##0.##"  # Sea quantity.
        assert raw.active["AF3"].number_format == "#,##0.00"  # Single-item grams, immediately after quantity.
        assert raw.active["AG3"].number_format == "#,##0.00"  # Total kilograms retain their existing format.
        assert raw.active["E3"].fill != raw.active["X3"].fill
        assert raw.active["T3"].fill.fgColor.rgb == raw.active["X3"].fill.fgColor.rgb == "00EAF2F8"
        assert raw.active.column_dimensions["T"].width == 15
        assert all(raw.active[address].fill.fgColor.rgb == "00FFF2CC" for address in ("S3", "V3", "W3", "AF3", "AH3"))
        assert raw.active["U3"].fill.patternType is None
        assert "COUNT(E3:R3,V3:W3)=16" in raw.active["AJ3"].value
        assert "T3" not in raw.active["AJ3"].value
        validation_ranges = {str(validation.sqref) for validation in raw.active.data_validations.dataValidation}
        assert "V3:W6" in validation_ranges
        decimal_validation = next(v for v in raw.active.data_validations.dataValidation if v.type == "decimal")
        assert "AF3:AF6" in str(decimal_validation.sqref)
        for ws in raw:
            for cells in ws:
                for cell in cells:
                    if cell.data_type == "f":
                        assert len(cell.value) < 600
                        assert "!" not in cell.value and "INDEX(" not in cell.value and "MATCH(" not in cell.value
        assert len(raw.active.data_validations.dataValidation) == 2
    finally:
        book.close()
        raw.close()
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        worksheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        assert worksheet.find("s:autoFilter", ns).attrib["ref"] == "A2:AK6"
        assert worksheet.find("s:tableParts", ns) is None



def test_empty_shipping_sheet_retains_instruction_and_headers(tmp_path):
    path = rep.write_replenishment_report([], tmp_path / "empty.xlsx")
    book = load_workbook(path)
    try:
        assert book.active.max_row == 2
        assert not book.active.tables
        assert book.active.auto_filter.ref == "A2:AK2"
    finally:
        book.close()


def test_weighted_sales_integer_boundary_uses_formal_evaluation_order(tmp_path):
    # YYH-US regression: sum() produced 184.00000000000003 before ceil,
    # while the formal left-associative calculation produced exactly 184.
    inventory = _inventory_input_row("CEIL-BOUNDARY", fba_total_inventory=82)
    values = dict.fromkeys(formulas.SOURCE_VALUE_COLUMNS, 0.0)
    values.update({"7天销量": 18.0, "14天销量": 33.0, "30天销量": 61.0,
                   "90天销量": 181.0, "可售": 31.0, "预留": 2.0,
                   "在途": 45.0, "待调仓": 2.0, "调仓中": 2.0})
    detail = rep.SalesDetail("增长", 100, 18, 33, 61, source_values=values)
    key = (inventory.msku, inventory.parent_asin, inventory.asin, inventory.local_sku)
    row = rep.calculate_replenishment_rows([inventory], {key: detail}, templates.load_default_template())[0]
    assert row.replenish_days == 75
    assert row.replenish_quantity == 102
    path = rep.write_replenishment_report([row], tmp_path / "boundary.xlsx")
    book = load_workbook(path, data_only=True)
    try:
        assert book[formulas.FINAL_SHIPPING_SHEET]["I3"].value == row.weighted_daily_sales
        assert book[formulas.FINAL_SHIPPING_SHEET]["AD3"].value == 102
        assert book[formulas.FINAL_SHIPPING_SHEET]["AE3"].value == 0
        assert book[formulas.FINAL_SHIPPING_SHEET]["X3"].value == 184
        assert not {rep.AIR_URGENT_SHEET, rep.AIR_SHEET, rep.SEA_SHEET} & set(book.sheetnames)
    finally:
        book.close()


@pytest.mark.parametrize("value", [None, "", "unknown", "NaN", float("inf"), -1, True])
def test_source_values_are_not_silently_zeroed(tmp_path, value):
    path = _write_sales_report(tmp_path / "source.xlsx")
    book = load_workbook(path)
    sheet = book["MSKU明细"]
    column = [cell.value for cell in sheet[1]].index("待调仓") + 1
    sheet.cell(2, column, value)
    if value is None:
        sheet.cell(2, column).value = None
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="备货公式输入无效"):
        rep.load_sales_details(path)


def test_load_source_columns_and_detect_missing_columns(tmp_path):
    path = _write_sales_report(tmp_path / "source.xlsx")
    details = rep.load_sales_details(path)
    first = next(iter(details.values()))
    assert first.source_values["90天销量"] == 540
    assert first.source_values["待调仓"] == 0.0
    book = load_workbook(path)
    sheet = book["MSKU明细"]
    sheet.delete_cols([c.value for c in sheet[1]].index("90天销量") + 1)
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="90天销量"):
        rep.load_sales_details(path)


def test_source_stock_total_must_match_inventory_report(tmp_path):
    path = _write_sales_report(tmp_path / "source.xlsx")
    details = rep.load_sales_details(path)
    row = replace(_inventory_input_row("SEA-1", fba_total_inventory=999), parent_asin="PARENT-SEA", asin="ASIN-SEA", local_sku="COMBO-SEA")
    with pytest.raises(ValueError, match="库存分项与真实库存报表不一致"):
        rep.calculate_replenishment_rows([row], details)


@pytest.mark.parametrize("weight", [None, 0.01, 1, 100000])
def test_sea_decision_and_minimum_quantity_independent_of_weight(weight):
    row = case_row(weight=weight, fba=600, unlinked=0)
    assert row.sheet_name == rep.SEA_SHEET
    assert row.sea_quantity == 60
    below = case_row(weight=weight, fba=600, unlinked=31)
    assert below.sheet_name == rep.NO_SHIP_SHEET


def test_legacy_weight_parameters_are_ignored_on_xlsx_import(tmp_path):
    path = templates.export_template_xlsx("默认", output_dir=tmp_path)
    book = load_workbook(path)
    sea = book["海运进入条件"]
    sea.insert_rows(5)
    sea.cell(5, 1, "海运最低重量kg")
    sea.cell(5, 2, 100000)
    special = book["特殊MSKU规则"]
    special.insert_cols(9)
    special.cell(1, 9, "海运最低重量kg")
    book.save(path)
    book.close()
    parsed = templates.validate_template_xlsx(path)
    assert "min_weight_kg" not in parsed.template.params["sea"]
    assert templates.validate_template(parsed.template)


def test_fractional_deductions_keep_each_rounding_stage(tmp_path):
    row = case_row(fba=300.2, unlinked=0.4, template_name="默认")
    assert row.replenish_quantity == 150
    path = rep.write_replenishment_report([row], tmp_path / "rounding.xlsx")
    book = load_workbook(path, data_only=True)
    try:
        assert book.active["AD3"].value == 150
    finally:
        book.close()


def test_zero_sea_component_is_not_replaced_with_total_quantity():
    row = replace(case_row(), original_replenish_quantity=450, replenish_quantity=450,
                  companion_air_quantity=450, sea_quantity=0, sea_net_quantity=0,
                  fba_total_inventory=430, unlinked_quantity=0)
    result = rep._with_inventory_deductions(row, min_sea_quantity=30)
    assert result.sheet_name == rep.NO_SHIP_SHEET
    assert result.sea_quantity == 0
    assert result.companion_air_quantity == 20


def test_formula_uses_effective_special_weights_with_nonuniform_sales(tmp_path):
    template = templates.get_template("默认")
    params = deepcopy(template.params)
    params["special_rules"] = [{"rule_name": "实际权重", "msku_list": ["SPECIAL"],
                               "overrides": {"weighted_sales": {"7d_weight": 0.1, "14d_weight": 0.2, "30d_weight": 0.7}}}]
    inventory = _inventory_input_row("SPECIAL", fba_total_inventory=300)
    values = dict.fromkeys(formulas.SOURCE_VALUE_COLUMNS, 0.0)
    values.update({"7天销量": 42, "14天销量": 112, "30天销量": 300, "90天销量": 999, "可售": 300})
    detail = rep.SalesDetail("平稳", None, 42, 112, 300, source_values=values)
    row = rep.calculate_replenishment_rows([inventory], {rep._row_key(inventory): detail}, replace(template, params=params))[0]
    assert row.formula_inputs.weights == (0.1, 0.2, 0.7)
    assert row.weighted_daily_sales == pytest.approx(9.2)
    assert row.sales_days == pytest.approx(300 / 9.2)
    path = rep.write_replenishment_report([row], tmp_path / "special.xlsx")
    book = load_workbook(path, data_only=True)
    raw = load_workbook(path, data_only=False)
    try:
        assert book.active["I3"].value == pytest.approx(9.2)
        assert book.active["Q3"].value == 300
        assert book.active["T3"].value == pytest.approx(row.sales_days)
        assert book.active["AD3"].value == row.replenish_quantity
        assert book.active["H3"].value == 999
        assert raw.active["T3"].value == '=IF(AJ3="正常",IF(I3>0,Q3/I3,""),"")'
        assert raw.active["T3"].number_format == "#,##0.00"
    finally:
        book.close()
        raw.close()


def test_sales_days_formula_uses_final_70_20_10_weights(tmp_path):
    template = templates.get_template("默认")
    params = deepcopy(template.params)
    params["weighted_sales"].update({"7d_weight": 0.7, "14d_weight": 0.2, "30d_weight": 0.1})
    inventory = _inventory_input_row("WEIGHTED-70-20-10", fba_total_inventory=10)
    values = dict.fromkeys(formulas.SOURCE_VALUE_COLUMNS, 0.0)
    values.update({"7天销量": 70, "14天销量": 70, "30天销量": 30, "90天销量": 999, "可售": 10})
    detail = rep.SalesDetail("平稳", None, 70, 70, 30, source_values=values)
    row = rep.calculate_replenishment_rows(
        [inventory], {rep._row_key(inventory): detail}, replace(template, name="70/20/10", params=params),
    )[0]
    assert row.formula_inputs.weights == (0.7, 0.2, 0.1)
    assert row.weighted_daily_sales == pytest.approx(8.1)
    path = rep.write_replenishment_report([row], tmp_path / "weighted-70-20-10.xlsx")
    book = load_workbook(path, data_only=True)
    raw = load_workbook(path, data_only=False)
    try:
        assert book.active["I3"].value == pytest.approx(8.1)
        assert book.active["Q3"].value == 10
        assert row.sales_days == pytest.approx(10 / 8.1)
        assert book.active["T3"].value == pytest.approx(10 / 8.1)
        assert raw.active["I3"].value == "=E3/7*0.7+F3/14*0.2+G3/30*0.1"
        assert raw.active["T3"].value == '=IF(AJ3="正常",IF(I3>0,Q3/I3,""),"")'
        assert raw.active["T3"].number_format == "#,##0.00"
    finally:
        book.close()
        raw.close()


@pytest.mark.parametrize("source_fba", [1, 1.0000008])
def test_zero_daily_sales_keeps_blank_days_on_non_shipping_sheet(tmp_path, source_fba):
    inventory = _inventory_input_row("ZERO-DAILY", fba_total_inventory=1)
    values = dict.fromkeys(formulas.SOURCE_VALUE_COLUMNS, 0.0)
    values["可售"] = source_fba
    detail = rep.SalesDetail("平稳", None, 0, 0, 0, source_values=values)
    row = rep.calculate_replenishment_rows([inventory], {rep._row_key(inventory): detail})[0]
    assert formulas._initial_values(row, False)["T"] is None
    with pytest.raises(ValueError, match="试算可销售天数与正式计算不一致"):
        formulas._initial_values(replace(row, sales_days=0), False)
    path = rep.write_replenishment_report([row], tmp_path / "zero-daily.xlsx")
    book = load_workbook(path, data_only=True)
    try:
        assert row.sales_days is None
        assert book[formulas.FINAL_SHIPPING_SHEET].max_row == 2
        headers = [cell.value for cell in book[rep.NON_SHIPPING_SHEET][1]]
        record = dict(zip(headers, next(book[rep.NON_SHIPPING_SHEET].iter_rows(min_row=2, values_only=True))))
        assert record["MSKU"] == "ZERO-DAILY"
        assert record["可销售天数"] in (None, "")
    finally:
        book.close()


def test_sales_days_less_than_one_keeps_fractional_cache(tmp_path):
    row = case_row("UNDER-DAY", daily=20, fba=17, unlinked=0, template_name="默认")
    assert row.sales_days == pytest.approx(0.85)
    path = rep.write_replenishment_report([row], tmp_path / "under-day.xlsx")
    book = load_workbook(path, data_only=True)
    raw = load_workbook(path, data_only=False)
    try:
        assert book.active["T3"].value == pytest.approx(0.85)
        assert raw.active["T3"].value == '=IF(AJ3="正常",IF(I3>0,Q3/I3,""),"")'
        assert raw.active["T3"].number_format == "#,##0.00"
    finally:
        book.close()
        raw.close()


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
@pytest.mark.parametrize("daily,fba,source_fba", [(9.2, 300, 300), (0.5, 0.25, 0.2500008)])
def test_sales_days_libreoffice_recalculation_preserves_precision_and_matches_initial_cache(tmp_path, daily, fba, source_fba):
    row = case_row("RECALC", daily=daily, fba=fba, source_fba=source_fba, unlinked=0, template_name="默认")
    path = rep.write_replenishment_report([row], tmp_path / "recalc-source.xlsx")
    source_values = load_workbook(path, data_only=True)
    try:
        initial_cache = source_values.active["T3"].value
        initial_values = {cell.coordinate: cell.value for cell in source_values.active[3]}
        assert row.sales_days == pytest.approx(fba / daily, rel=0, abs=1e-12)
        assert initial_cache == pytest.approx(source_fba / daily, rel=0, abs=1e-12)
    finally:
        source_values.close()

    input_dir = tmp_path / "office-inputs"
    output_dir = tmp_path / "office-recalculated"
    input_dir.mkdir()
    output_dir.mkdir()
    edits = {
        "initial": {},
        "sales-edited": {"E3": 70},
        "stock-edited": {"J3": 400},
        "zero-sales": {"E3": 0, "F3": 0, "G3": 0},
    }
    for name, values in edits.items():
        target = input_dir / f"{name}.xlsx"
        shutil.copy2(path, target)
        if values:
            workbook = load_workbook(target)
            sheet = workbook[formulas.FINAL_SHIPPING_SHEET]
            for address, value in values.items():
                sheet[address] = value
            workbook.save(target)
            workbook.close()

    command = [
        shutil.which("soffice"),
        f"-env:UserInstallation={(tmp_path / 'libreoffice-profile').as_uri()}",
        "--headless",
        "--convert-to",
        "xlsx",
        "--outdir",
        str(output_dir),
        *(str(input_dir / f"{name}.xlsx") for name in edits),
    ]
    result = subprocess.run(command, capture_output=True, text=True, timeout=60, check=False)
    assert result.returncode == 0, result.stdout + result.stderr

    calculated = {}
    for name in edits:
        workbook = load_workbook(output_dir / f"{name}.xlsx", data_only=True)
        try:
            sheet = workbook[formulas.FINAL_SHIPPING_SHEET]
            calculated[name] = sheet["T3"].value
            assert not any(cell.data_type == "e" for cells in sheet for cell in cells)
            if name == "initial":
                for address, expected in initial_values.items():
                    actual = sheet[address].value
                    assert actual == (pytest.approx(expected, rel=0, abs=1e-10) if isinstance(expected, (int, float)) else expected)
        finally:
            workbook.close()
    assert calculated["initial"] == pytest.approx(initial_cache, rel=0, abs=1e-12)
    assert calculated["sales-edited"] == pytest.approx(source_fba / (6 + daily * 0.3 + daily * 0.1), rel=0, abs=1e-10)
    assert calculated["stock-edited"] == pytest.approx(400 / daily, rel=0, abs=1e-10)
    assert calculated["zero-sales"] in (None, "")


@pytest.mark.parametrize("daily", [0.5, 6])
@pytest.mark.parametrize("source_fba", [0.25, 0.2500008, 0.2499992, 0.25 + 1e-6])
def test_sales_days_accepts_existing_stock_tolerance(tmp_path, daily, source_fba):
    assert abs(source_fba - 0.25) <= 1e-6
    row = case_row("TOLERANCE", daily=daily, fba=0.25, source_fba=source_fba, unlinked=0, template_name="默认")
    original = deepcopy(row)
    path = rep.write_replenishment_report([row], tmp_path / "tolerance.xlsx")
    book = load_workbook(path, data_only=True)
    try:
        sheet = book[formulas.FINAL_SHIPPING_SHEET]
        assert sheet["T3"].value == pytest.approx(source_fba / daily, rel=0, abs=1e-12)
        assert sheet["T3"].value - row.sales_days == pytest.approx((source_fba - 0.25) / daily, rel=0, abs=1e-12)
        assert sheet["I3"].value == row.weighted_daily_sales
        assert sheet["Q3"].value == source_fba
        assert sheet["U3"].value == row.sheet_name
        assert sheet["AD3"].value == row.replenish_quantity
        assert sheet["AE3"].value == 0
    finally:
        book.close()
    assert row == original


@pytest.mark.parametrize("daily", [0.5, 6])
@pytest.mark.parametrize("source_fba", [math.nextafter(0.25 + 1e-6, math.inf), 0.25 - 1.000001e-6])
def test_sales_days_rejects_stock_just_outside_existing_tolerance(tmp_path, daily, source_fba):
    assert abs(source_fba - 0.25) > 1e-6
    with pytest.raises(ValueError, match="库存分项与真实库存报表不一致"):
        case_row(daily=daily, fba=0.25, source_fba=source_fba, unlinked=0, template_name="默认")
    row = case_row(daily=daily, fba=0.25, unlinked=0, template_name="默认")
    row = replace(row, formula_inputs=replace(row.formula_inputs, values={**row.formula_inputs.values, "可售": source_fba}))
    target = tmp_path / "report.xlsx"
    target.write_bytes(b"previous")
    with pytest.raises(ValueError, match="试算FBA总库存与正式计算不一致"):
        rep.write_replenishment_report([row], target)
    assert target.read_bytes() == b"previous"
    assert list(tmp_path.iterdir()) == [target]


@pytest.mark.parametrize("changes,error", [
    ({"weighted_daily_sales": math.nextafter(0.5 + 1e-9, math.inf)}, "试算加权日销"),
    ({"sales_days": 0.5000016}, "试算可销售天数"),  # Matches Q/I, but not the formal inputs.
    ({"sales_days": None}, "试算可销售天数"),
    ({"sales_days": float("nan")}, "试算可销售天数"),
    ({"sales_days": float("inf")}, "试算可销售天数"),
])
def test_sales_days_tolerance_does_not_hide_inconsistent_formal_values(tmp_path, changes, error):
    row = case_row(daily=0.5, fba=0.25, source_fba=0.2500008, unlinked=0, template_name="默认")
    target = tmp_path / "report.xlsx"
    target.write_bytes(b"previous")
    with pytest.raises(ValueError, match=error):
        rep.write_replenishment_report([replace(row, **changes)], target)
    assert target.read_bytes() == b"previous"
    assert list(tmp_path.iterdir()) == [target]


def test_sales_days_respects_existing_daily_tolerance():
    row = case_row(daily=0.5, fba=1, source_fba=1.0000008, unlinked=0, template_name="默认")
    daily = row.weighted_daily_sales + 5e-10
    row = replace(row, weighted_daily_sales=daily, sales_days=row.fba_total_inventory / daily)
    assert formulas._initial_values(row, False)["T"] == pytest.approx(1.0000008 / 0.5, rel=0, abs=1e-12)


@pytest.mark.parametrize("daily,days", [(1e-10, None), (1e-10, 1e10), (0, 0)])
def test_sales_days_tolerance_preserves_zero_daily_state(daily, days):
    row = case_row(daily=0, fba=1, unlinked=0, template_name="默认")
    with pytest.raises(ValueError, match="试算可销售天数与正式计算不一致"):
        formulas._initial_values(replace(row, weighted_daily_sales=daily, sales_days=days), False)


def test_sales_days_source_mismatch_blocks_report_publication(tmp_path):
    row = replace(case_row("MISMATCH"), sales_days=999)
    target = tmp_path / "report.xlsx"
    target.write_bytes(b"previous")
    with pytest.raises(ValueError, match="试算可销售天数与正式计算不一致"):
        rep.write_replenishment_report([row], target)
    assert target.read_bytes() == b"previous"
    assert list(tmp_path.iterdir()) == [target]


@pytest.mark.parametrize('stock,expected', [(None,None), (-5,135), (0,130), (10.5,119.5), (130,0), (1000,0)])
def test_stock_gap_column_cached_and_sort_keys_preserved(tmp_path, stock, expected):
    row=replace(case_row('GAP', fba=300, unlinked=20, template_name='默认'),actual_inventory=stock)
    path=rep.write_replenishment_report([row],tmp_path/'gap.xlsx')
    book=load_workbook(path,data_only=True);raw=load_workbook(path)
    try:
        assert book.active['AI2'].value=='深圳库存缺口'
        assert book.active['AI3'].value==expected
        assert book.active['AD3'].value==130 and book.active['AE3'].value==0
        assert '记录标识' not in [c.value for c in book.active[2]]
        assert raw.active['AI3'].data_type=='f'
        assert 'COUNT(S3,AD3:AE3)=3' in raw.active['AI3'].value
        assert not any(d.hidden for d in raw.active.column_dimensions.values())
        assert raw.active['AI3'].fill.fgColor.rgb==raw.active['AD3'].fill.fgColor.rgb=='00EAF2F8'
        assert raw.active.auto_filter.ref=='A2:AK3'
        assert raw.active['AD3'].value=='=IF(AJ3="正常",AB3,"")'
    finally:book.close();raw.close()


def test_visible_deduction_stages_use_only_same_row_cells(tmp_path):
    row = case_row()
    path = rep.write_replenishment_report([row], tmp_path/'stages.xlsx')
    raw = load_workbook(path);cached = load_workbook(path,data_only=True)
    try:
        assert [cached.active.cell(3,c).value for c in range(24,32)] == [450,210,20,210,10,210,10,210]
        assert raw.active['X3'].value == '=IF(AJ3="正常",ROUNDUP(I3*V3,0),"")'
        assert raw.active['AB3'].value == '=IF(AJ3="正常",MAX(0,ROUNDUP(Z3-R3,0)),"")'
        assert raw.active['AC3'].value == '=IF(AJ3="正常",MAX(0,ROUNDUP(AA3-MAX(0,R3-Z3),0)),"")'
        assert 'AC3>=30.0' in raw.active['AE3'].value
        assert raw.active['AH3'].data_type == 's'
    finally:raw.close();cached.close()


def test_effective_weights_are_literal_and_row_specific(tmp_path):
    row = case_row(fba=300, unlinked=20, template_name='默认')
    path = rep.write_replenishment_report([row], tmp_path/'weights.xlsx')
    book = load_workbook(path)
    try:
        a,b,c=row.formula_inputs.weights
        assert book.active['I3'].value == f'=E3/7*{float(a)!r}+F3/14*{float(b)!r}+G3/30*{float(c)!r}'
        assert len(book.active['AD3'].value) < 80
    finally:book.close()
