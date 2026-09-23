"""Offline regressions derived from the two September 23 customs failures.

SKU, supplier, contract, source and SP identifiers are anonymized. Row numbers,
model names, prices, quantities and the overlapping source structure are fixed.
"""
from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.agent_cli.mabang import customs_erp as cli
from services.agent_cli.mabang.customs_prices import PriceMatchError, match_price, read_prices
from test_customs_erp import setup  # Reuse the offline ERP/download fixture.
from test_mabang_customs_declaration_fill_cli import _write_template


def _prices(path, specs):
    """spec = (Excel row, representative, original price, sale price, SKU quantities, kind)."""
    wb = Workbook()
    wb.active.title = "备货单"
    summary = wb.create_sheet("汇总表")
    headers = ["日期", "库存sku（第一行）", "型号", "发货量", "原价", "售价", "库存sku"]
    wb["备货单"].append(headers)
    summary.append(headers[:-1])
    for index, representative, cost, price, skus, kind in specs:
        # Deliberately untrusted quantity/old total; ERP owns shipped quantities.
        values = ["走库存" if kind == "carryover" else "2026-09-01", representative, None, 99999, cost]
        for sheet in (wb["备货单"], summary):
            for col, value in enumerate(values, 1):
                sheet.cell(index, col, value)
        wb["备货单"].cell(index, 6, "#N/A")
        wb["备货单"].cell(index, 7, "\n".join(f"{sku} × {quantity}" for sku, quantity in skus.items()))
        summary.cell(index, 6, price)
    wb.save(path)
    wb.close()
    return path


def _source(source_id, contract, quantity, cost, *, kind="carryover"):
    return {"source_id": source_id, "source_kind": kind, "source_contract_no": contract,
            "source_sp_no": f"SOURCE-{contract}", "planned_quantity": quantity,
            "actual_quantity": quantity, "purchase_price": cost}


def _line(sku, model, sources, *, product="尼龙表带", commodity="尼龙表带", planned=None):
    actual = sum(source["actual_quantity"] for source in sources)
    return {"stock_sku": sku, "model": model, "supplier_name": "供应商A", "product_name": product,
            "commodity_name": commodity, "unit": "个", "planned_quantity": actual if planned is None else planned,
            "actual_quantity": actual, "sources": sources}


def _case(name):
    if name == "us":
        specs = [(33, "SKU-A", "9", "13", {"SKU-A": 35}, "carryover"),
                 (34, "SKU-A", "9", "13.00", {"SKU-A": 5, "SKU-B": 30}, "carryover")]
        lines = [_line("SKU-A", "ZYZ-31", [_source("lot-1", "OLD-A", 35, "9"),
                                            _source("lot-2", "OLD-B", 5, "9")]),
                 _line("SKU-B", "ZYZ-31", [_source("lot-3", "OLD-B", 30, "9")])]
        return specs, lines, "ZYZ-31", 70, Decimal("13"), Decimal("910")
    specs = [(16, "SKU-A", "2.94", "3.2", {"SKU-A": 20, "SKU-B": 10}, "carryover"),
             (17, "SKU-B", "2.94", "3.2", {"SKU-B": 10, "SKU-C": 20, "SKU-D": 20}, "carryover"),
             (2, "SKU-H", "3.84", "4", {"SKU-H": 885}, "current_purchase")]
    lines = [_line("SKU-A", "JZ-23", [_source("lot-1", "OLD-A", 20, "2.94")], product="PC表壳", commodity="手表壳"),
             _line("SKU-B", "JZ-23", [_source("lot-2", "OLD-A", 10, "2.94"), _source("lot-3", "OLD-B", 10, "2.94")],
                   product="PC表壳", commodity="手表壳"),
             _line("SKU-C", "JZ-23", [_source("lot-4", "OLD-B", 5, "2.94"), _source("lot-5", "OLD-B", 15, "2.94")],
                   product="PC表壳", commodity="手表壳"),
             _line("SKU-D", "JZ-23", [_source("lot-6", "OLD-B", 15, "2.94"), _source("lot-7", "OLD-B", 5, "2.94")],
                   product="PC表壳", commodity="手表壳"),
             _line("SKU-H", "HS-6", [_source("lot-8", "NEW", 860, "3.84", kind="current_purchase")], planned=885)]
    lines[-1]["sources"][0]["planned_quantity"] = 885
    return specs, lines, "JZ-23", 80, Decimal("3.2"), Decimal("256")


@pytest.mark.parametrize("case", ["us", "uk"])
def test_real_failure_shapes_preview_fill_and_provenance(setup, case):
    args, erp, _, _, _ = setup
    specs, lines, model, quantity, price, amount = _case(case)
    _prices(Path(args["input_xlsx"][0]), specs)
    erp["shipments"][0]["lines"] = lines
    # US had two spreadsheet rows but only needs one template block after merge.
    _write_template(Path(args["template_xlsx"]), customs_detail_blocks=1 if case == "us" else 2)
    preview = cli.preview(args)
    assert preview["can_generate"], preview["issues"]
    assert preview["row_count"] == (1 if case == "us" else 2)
    data = json.loads(Path(preview["preview_path"]).read_text(encoding="utf-8"))
    rows = [row for row in data["bundles"][0]["rows"] if row["model"] == model]
    assert len(rows) == 1
    assert (Decimal(rows[0]["quantity"]), Decimal(rows[0]["sale_price"]), Decimal(rows[0]["total_price"])) == (quantity, price, amount)
    summary = next(row for row in preview["model_summary"] if row["model"] == model)
    assert summary["planned_quantity"] == summary["actual_quantity"] == str(quantity)
    matches = [row for row in data["price_diagnostics"] if row["model"] == model]
    assert sum(Decimal(row["actual_quantity"]) for row in matches) == quantity
    assert len(matches) == sum(len(line["sources"]) for line in lines if line["model"] == model)
    assert {row["source_contract_no"] for row in matches} == {"OLD-A", "OLD-B"}
    assert all(row["source_id"] and row["candidates"] for row in matches)
    assert len({row["declaration_row_number"] for row in matches}) == 1
    assert all(Decimal(row["sale_price"]) == price for row in matches)
    overlap = next(row for row in matches if len(row["candidates"]) == 2)
    assert overlap["matched_summary_rows"] == ([33, 34] if case == "us" else [16, 17])
    if case == "us":
        assert all([item["row_number"] for item in candidate["restock_candidates"]] == [33, 34]
                   for candidate in overlap["candidates"])
    else:
        shortage = next(row for row in preview["model_summary"] if row["model"] == "HS-6")
        assert (shortage["planned_quantity"], shortage["actual_quantity"], shortage["shortage_quantity"]) == ("885", "860", "25")
    report = load_workbook(preview["validation_report_xlsx"], data_only=True)
    assert report["售价匹配"].max_row == len(data["price_diagnostics"]) + 1
    assert "汇总表" in report["售价匹配"].cell(2, 13).value
    report.close()
    result = cli.fill(preview["preview_path"])
    assert result["total_amount"] == (910 if case == "us" else 3696)
    assert result["row_count"] == preview["row_count"]
    wb = load_workbook(result["output_xlsx"], data_only=False)
    sheet = wb["申报要素"]
    assert sum(cell.value == model for row in sheet for cell in row) == 1
    wb.close()


@pytest.mark.parametrize("case", ["us", "uk"])
def test_different_prices_without_source_mapping_block_with_actual_diagnostics(setup, case):
    args, erp, _, _, _ = setup
    specs, lines, model, *_ = _case(case)
    spec = list(specs[1])
    spec[3] = "15" if case == "us" else "3.4"
    specs[1] = tuple(spec)
    _prices(Path(args["input_xlsx"][0]), specs)
    erp["shipments"][0]["lines"] = lines
    result = cli.preview(args)
    assert result["status"] == "blocked"
    message = " ".join(result["issues"])
    for token in (erp["shipments"][0]["sp_no"], model, "SKU=", "ERP来源采购价=", "售价来源不唯一",
                  f"汇总表第{specs[0][0]}行", f"汇总表第{specs[1][0]}行", "原价=", "售价=", "备货单候选行="):
        assert token in message
    data = json.loads(Path(result["preview_path"]).read_text(encoding="utf-8"))
    failed = next(item for item in data["price_diagnostics"] if item["status"] == "blocked")
    assert len(failed["candidates"]) == 2
    assert failed["error"] in result["issues"]
    assert "sale_price" not in failed
    report = load_workbook(result["validation_report_xlsx"], data_only=True)
    failed_rows = [row for row in report["售价匹配"].iter_rows(min_row=2, values_only=True) if row[15] == "blocked"]
    assert len(failed_rows) == 1
    assert failed_rows[0][10] is None and failed_rows[0][16] == failed["error"]
    report.close()
    with pytest.raises(ValueError, match="不能生成"):
        cli.fill(result["preview_path"])


def test_same_price_across_costs_and_source_kinds_merges_and_rounds_once(tmp_path):
    specs = [(2, "A", "2", "0.0005", {"A": 1}, "carryover"),
             (3, "A", "3", "0.00050", {"A": 1}, "current_purchase")]
    prices = read_prices(_prices(tmp_path / "prices.xlsx", specs))
    sources = [_source("old", "OLD", 1, "2"), _source("new", "NEW", 1, "3", kind="current_purchase")]
    rows = cli.declaration_rows(prices, {"sp_no": "SP1", "lines": [_line("A", "M1", sources)]})
    assert len(rows) == 1
    # Existing customs amounts use three decimals: rounding each source first
    # would incorrectly produce 0.002 instead of the merged amount 0.001.
    assert rows[0]["quantity"] == 2 and rows[0]["total_price"] == Decimal("0.001")
    assert rows[0]["source_kind"] == "mixed" and rows[0]["purchase_price"] is None


def test_different_prices_resolve_by_sku_kind_or_cost(tmp_path):
    specs = [(2, "A", "2", "5", {"A": 1}, "carryover"),
             (3, "A", "3", "6", {"A": 1}, "carryover"),
             (4, "A", "3", "7", {"A": 1}, "current_purchase"),
             (5, "B", "3", "8", {"B": 1}, "carryover")]
    prices = read_prices(_prices(tmp_path / "prices.xlsx", specs))
    sources = [_source("o1", "O1", 2, "2"), _source("o2", "O2", 3, "3"),
               _source("n", "N", 4, "3", kind="current_purchase")]
    diagnostic = []
    rows = cli.declaration_rows(prices, {"sp_no": "SP1", "lines": [
        _line("A", "M1", sources), _line("B", "M1", [_source("b", "B", 5, "3")])]}, price_diagnostics=diagnostic)
    assert [(row["quantity"], row["sale_price"]) for row in rows] == [(2, 5), (3, 6), (4, 7), (5, 8)]
    assert diagnostic[0]["method"] == "purchase_price"
    assert len(diagnostic[0]["candidates"]) == 2  # Preserve candidates excluded by cost, too.
    assert diagnostic[0]["matched_summary_rows"] == [2]
    with pytest.raises(PriceMatchError, match="不唯一"):
        match_price(prices, "A", _source("x", "X", 1, "999"))


def test_same_price_does_not_require_equal_cost_and_does_not_round_conflicts(tmp_path):
    specs = [(2, "A", "2", "13", {"A": 1}, "carryover"),
             (3, "B", "3", "13.00", {"A": 1, "B": 1}, "carryover")]
    prices = read_prices(_prices(tmp_path / "prices.xlsx", specs))
    resolved = match_price(prices, "A", _source("a", "A", 1, "2"))
    assert len(resolved.matched_rows) == 2 and resolved.sale_price == 13
    conflicting = [prices[0], replace(prices[1], purchase_price=Decimal(2), sale_price=Decimal("13.0001"))]
    with pytest.raises(PriceMatchError):
        match_price(conflicting, "A", _source("a", "A", 1, "2"))


@pytest.mark.parametrize("difference", ["supplier_name", "model", "commodity_name", "unit", "classification", "hs_code"])
def test_different_declaration_dimensions_do_not_merge(tmp_path, difference):
    specs = [(2, "A", "2", "5", {"A": 1, "B": 1}, "carryover")]
    prices = read_prices(_prices(tmp_path / "prices.xlsx", specs))
    a = _line("A", "M1", [_source("a", "O", 1, "2")], product="尼龙表带", commodity="表带")
    b = deepcopy(a)
    b.update(stock_sku="B", sources=[_source("b", "O", 1, "2")])
    if difference == "classification":
        b["product_name"] = "硅胶表带"  # Same HS code, different material.
    elif difference == "hs_code":
        b["product_name"] = "米兰尼斯表带"
    else:
        b[difference] = "different"
    rows = cli.declaration_rows(prices, {"sp_no": "SP1", "lines": [a, b]})
    assert len(rows) == 2


def test_quantity_mismatch_is_not_hidden_by_merging(tmp_path):
    specs, lines, *_ = _case("us")
    prices = read_prices(_prices(tmp_path / "prices.xlsx", specs))
    lines[0]["actual_quantity"] += 1
    with pytest.raises(ValueError, match="不守恒"):
        cli.declaration_rows(prices, {"sp_no": "SP1", "lines": lines})


def test_zero_quantity_source_needs_no_price_and_adds_no_match(tmp_path):
    specs = [(2, "A", "2", "5", {"A": 1}, "carryover")]
    prices = read_prices(_prices(tmp_path / "prices.xlsx", specs))
    diagnostics = []
    rows = cli.declaration_rows(prices, {"sp_no": "SP1", "lines": [
        _line("A", "M1", [_source("a", "A", 1, "2")]),
        _line("UNPRICED", "M2", [_source("b", "B", 0, "3")])]}, price_diagnostics=diagnostics)
    assert len(rows) == len(diagnostics) == 1


def test_multi_sp_same_model_same_price_stays_separate(setup, monkeypatch):
    args, erp, _, _, _ = setup
    specs, lines, *_ = _case("us")
    first_path = Path(args["input_xlsx"][0])
    _prices(first_path, specs)
    other_path = _prices(first_path.with_name("SP260831009-美国.xlsx"), specs)
    args["input_xlsx"].append(str(other_path))
    weight = Path(args.pop("consignment_excel"))
    monkeypatch.setattr(cli.template, "_resolve_consignment_excel_path", lambda *a: weight)
    erp["shipments"][0]["lines"] = lines
    other = deepcopy(erp["shipments"][0])
    other.update(sp_no="SP260831009", packing_sp_no="SP260831009")
    erp["shipments"].append(other)
    result = cli.preview(args)
    assert result["row_count"] == 2
    final = cli.fill(result["preview_path"])
    assert final["row_count"] == 2 and final["total_amount"] == 1820 and final["box_count"] == 2


def test_legacy_confirmed_preview_keeps_frozen_split_rows(setup):
    args, _, _, _, _ = setup
    result = cli.preview(args)
    path = Path(result["preview_path"])
    data = json.loads(path.read_text(encoding="utf-8"))
    data.pop("price_diagnostics")
    # A pre-upgrade preview may have separate rows whose prices are identical.
    row = data["bundles"][0]["rows"][0]
    duplicate = deepcopy(row)
    original_quantity = Decimal(row["quantity"])
    row["quantity"] = duplicate["quantity"] = str(original_quantity / 2)
    row["total_price"] = duplicate["total_price"] = str(Decimal(row["total_price"]) / 2)
    data["bundles"][0]["rows"].insert(1, duplicate)
    # Expand only this synthetic old snapshot's template and update its digest.
    _write_template(path.parent / data["template_name"], customs_detail_blocks=3)
    next(item for item in data["files"] if item["name"] == data["template_name"])["sha256"] = cli._sha(path.parent / data["template_name"])
    serialized = cli._json(data)
    legacy_path = path.with_name(hashlib.sha256(serialized.encode()).hexdigest() + ".json")
    legacy_path.write_text(serialized, encoding="utf-8")
    final = cli.fill(legacy_path)
    assert final["row_count"] == 3 and final["total_amount"] == 48
