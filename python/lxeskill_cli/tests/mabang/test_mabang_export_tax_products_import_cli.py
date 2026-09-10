from __future__ import annotations

import asyncio
import json
from collections import OrderedDict
from pathlib import Path
from types import SimpleNamespace

import pytest

from lxeskill import business, cli as lxeskill
from services.agent_cli.mabang import import_export_tax_products as cli
from services.agent_cli.mabang import summarize_fba_delivery_tax_sku as summary_cli
from shared import input_assets


def _write_products(
    path: Path,
    rows: list[dict[str, str]],
    *,
    columns: list[str] | None = None,
    sheet_name: str = "Sheet1",
) -> None:
    from openpyxl import Workbook

    if columns is None:
        columns = ["sku", "产品名称"]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(columns)
    for row in rows:
        worksheet.append([row.get(column, "") for column in columns])
    workbook.save(path)


def _read_rows(path: Path) -> list[tuple]:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    worksheet = workbook["Sheet1"]
    return list(worksheet.iter_rows(values_only=True))


def test_missing_sku_returns_failure_json(monkeypatch, capsys):

    payload = cli.run({})
    assert payload == {"success": False, "exception": "sku 不能为空"}


def test_normalize_input_skus_dedupes_and_ignores_empty_values():
    assert cli.normalize_input_skus([" SKU-A ", "", "SKU-A", "SKU-B\nSKU-C，SKU-D"]) == [
        "SKU-A",
        "SKU-B",
        "SKU-C",
        "SKU-D",
    ]


def test_import_skips_existing_and_not_found_then_appends_found(monkeypatch, tmp_path):
    products_path = tmp_path / "products.xlsx"
    backup_dir = tmp_path / "backup"
    output_dir = tmp_path / "artifacts"
    _write_products(
        products_path,
        [
            {"sku": "SKU-EXIST", "产品名称": "已有产品"},
        ],
    )

    async def fake_export_stock_sku_names(skus, **kwargs):
        assert skus == ["SKU-NEW", "SKU-MISSING"]
        return SimpleNamespace(
            names_by_key=OrderedDict([("SKU-NEW", "新产品")]),
            xlsx_paths=[],
        )

    monkeypatch.setattr(cli, "export_stock_sku_names", fake_export_stock_sku_names)
    monkeypatch.setattr(cli, "_timestamp", lambda: "20260512_120000")

    payload = asyncio.run(
        cli.import_export_tax_products(
            ["SKU-EXIST", "SKU-NEW", "SKU-MISSING"],
            products_path=products_path,
            backup_dir=backup_dir,
            output_dir=output_dir,
        )
    )

    assert payload["success"] is True
    assert payload["requested_sku_count"] == 3
    assert payload["imported_count"] == 1
    assert payload["skipped_duplicate_count"] == 1
    assert payload["skipped_not_found_count"] == 1
    assert payload["imported_skus"] == ["SKU-NEW"]
    assert payload["skipped_duplicate_skus"] == ["SKU-EXIST"]
    assert payload["skipped_not_found_skus"] == ["SKU-MISSING"]
    assert payload["backup_path"] == str(backup_dir / "products_20260512_120000.xlsx")
    assert Path(payload["backup_path"]).is_file()
    assert _read_rows(products_path) == [
        ("sku", "产品名称"),
        ("SKU-EXIST", "已有产品"),
    ]
    assert payload["input_products_path"] == str(products_path)
    assert payload["output_xlsx"] == str(output_dir / "products_updated_20260512_120000.xlsx")
    assert _read_rows(Path(payload["output_xlsx"])) == [
        ("sku", "产品名称"),
        ("SKU-EXIST", "已有产品"),
        ("SKU-NEW", "新产品"),
    ]
    assert _read_rows(Path(payload["backup_path"])) == [
        ("sku", "产品名称"),
        ("SKU-EXIST", "已有产品"),
    ]


def test_import_does_not_write_or_call_api_when_all_duplicate(monkeypatch, tmp_path):
    products_path = tmp_path / "products.xlsx"
    backup_dir = tmp_path / "backup"
    _write_products(products_path, [{"sku": "SKU-EXIST", "产品名称": "已有产品"}])

    async def fail_export_stock_sku_names(*args, **kwargs):
        raise AssertionError("stock export should not be called")

    monkeypatch.setattr(cli, "export_stock_sku_names", fail_export_stock_sku_names)

    payload = asyncio.run(
        cli.import_export_tax_products(
            ["SKU-EXIST"],
            products_path=products_path,
            backup_dir=backup_dir,
        )
    )

    assert payload["imported_count"] == 0
    assert payload["skipped_duplicate_skus"] == ["SKU-EXIST"]
    assert payload["backup_path"] == ""
    assert payload["output_xlsx"] == ""
    assert not backup_dir.exists()
    assert _read_rows(products_path) == [
        ("sku", "产品名称"),
        ("SKU-EXIST", "已有产品"),
    ]


def test_import_does_not_write_when_api_finds_no_skus(monkeypatch, tmp_path):
    products_path = tmp_path / "products.xlsx"
    backup_dir = tmp_path / "backup"
    _write_products(products_path, [{"sku": "SKU-EXIST", "产品名称": "已有产品"}])

    async def fake_export_stock_sku_names(skus, **kwargs):
        assert skus == ["SKU-MISSING"]
        return SimpleNamespace(names_by_key=OrderedDict(), xlsx_paths=[])

    monkeypatch.setattr(cli, "export_stock_sku_names", fake_export_stock_sku_names)

    payload = asyncio.run(
        cli.import_export_tax_products(
            ["SKU-MISSING"],
            products_path=products_path,
            backup_dir=backup_dir,
        )
    )

    assert payload["imported_count"] == 0
    assert payload["skipped_not_found_skus"] == ["SKU-MISSING"]
    assert payload["backup_path"] == ""
    assert payload["output_xlsx"] == ""
    assert not backup_dir.exists()


def test_main_success_outputs_json(monkeypatch, tmp_path, capsys):
    products_path = tmp_path / "products.xlsx"
    backup_dir = tmp_path / "backup"
    _write_products(products_path, [{"sku": "SKU-EXIST", "产品名称": "已有产品"}])

    async def fake_export_stock_sku_names(skus, **kwargs):
        return SimpleNamespace(names_by_key=OrderedDict([("SKU-NEW", "新产品")]), xlsx_paths=[])

    monkeypatch.setattr(cli, "export_stock_sku_names", fake_export_stock_sku_names)
    monkeypatch.setattr(cli, "_timestamp", lambda: "20260512_120000")
    monkeypatch.setattr(cli, "DEFAULT_BACKUP_DIR", backup_dir)
    monkeypatch.setattr(cli, "DEFAULT_OUTPUT_DIR", tmp_path / "artifacts")

    payload = cli.run({"sku": "SKU-NEW", "products_path": str(products_path)})
    assert payload["success"] is True
    assert payload["input_products_path"] == str(products_path)
    assert payload["products_path"] == str(tmp_path / "artifacts" / "products_updated_20260512_120000.xlsx")
    assert payload["output_xlsx"] == payload["products_path"]
    assert payload["backup_path"] == str(backup_dir / "products_20260512_120000.xlsx")
    assert payload["source"] == "export_tax_products_import"


def test_main_validation_failure_returns_failure_json(monkeypatch, tmp_path, capsys):
    products_path = tmp_path / "products.xlsx"
    _write_products(products_path, [{"sku": "SKU-A"}], columns=["sku"])

    payload = cli.run({"sku": "SKU-NEW", "products_path": str(products_path)})
    assert payload["success"] is False
    assert "缺少列: 产品名称" in payload["exception"]


@pytest.fixture
def stored_products(monkeypatch, tmp_path):
    artifacts = tmp_path / "artifacts"
    monkeypatch.setattr(input_assets, "input_root", lambda: tmp_path / "inputs")
    monkeypatch.setattr(business, "artifact_root", lambda: artifacts)
    monkeypatch.setattr(lxeskill, "activate_project_workspace", lambda: tmp_path)
    monkeypatch.setattr(cli, "DEFAULT_BACKUP_DIR", artifacts / "backup")
    monkeypatch.setattr(cli, "DEFAULT_OUTPUT_DIR", artifacts / "products")
    original = tmp_path / "products.xlsx"
    _write_products(original, [{"sku": "SKU-EXIST", "产品名称": "已有产品"}])
    input_assets.promote_asset("export_tax_products", original)

    async def fake_export(skus, **kwargs):
        assert skus == ["SKU-NEW"]
        return SimpleNamespace(names_by_key={"SKU-NEW": "新产品"}, xlsx_paths=[])

    monkeypatch.setattr(cli, "export_stock_sku_names", fake_export)
    return original


@pytest.mark.parametrize("use_upload", [False, True])
def test_import_updates_default_asset_for_next_import_and_summary(
    stored_products, use_upload, monkeypatch, capsys
):
    original_bytes = stored_products.read_bytes()
    arguments = ["fba", "export-tax", "products-import", "--sku", "SKU-NEW"]
    if use_upload:
        arguments.extend(["--products-path", str(stored_products)])

    assert lxeskill.main(arguments) == 0
    record = json.loads(capsys.readouterr().out)
    updated = input_assets.current_asset("export_tax_products")
    previous = input_assets.previous_asset("export_tax_products")
    assert updated.path.read_bytes() == Path(record["data"]["output_xlsx"]).read_bytes()
    assert _read_rows(updated.path) == [
        ("sku", "产品名称"), ("SKU-EXIST", "已有产品"), ("SKU-NEW", "新产品"),
    ]
    assert previous.path.read_bytes() == original_bytes
    assert stored_products.read_bytes() == original_bytes
    assert record["files"] == [record["data"]["output_xlsx"]]
    assert record["data"]["asset_sources"]["products_path"] == {
        "slot": "export_tax_products", "from": "generated",
        "file_name": updated.file_name, "updated_at": updated.updated_at,
    }

    async def unexpected_export(*args, **kwargs):
        pytest.fail("the next import must recognize the SKU in the stored updated file")

    monkeypatch.setattr(cli, "export_stock_sku_names", unexpected_export)
    assert lxeskill.main(["fba", "export-tax", "products-import", "--sku", "SKU-NEW"]) == 0
    duplicate = json.loads(capsys.readouterr().out)
    assert duplicate["data"]["skipped_duplicate_count"] == 1
    assert duplicate["data"]["output_xlsx"] == ""
    assert input_assets.current_asset("export_tax_products") == updated
    assert input_assets.previous_asset("export_tax_products") == previous

    def read_summary_input(arguments):
        assert Path(arguments["products_path"]) == updated.path
        assert ("SKU-NEW", "新产品") in _read_rows(Path(arguments["products_path"]))
        return {"success": True}

    monkeypatch.setattr(summary_cli, "run", read_summary_input)
    assert lxeskill.main(["fba", "export-tax", "delivery-summary", "--delivery-no", "SP123"]) == 0


def test_failed_output_validation_keeps_current_asset(stored_products, monkeypatch, capsys):
    current = input_assets.current_asset("export_tax_products")
    original_bytes = current.path.read_bytes()
    validate = cli.validate_export_tax_products

    def reject_updated(path):
        if "_updated_" in Path(path).name:
            raise ValueError("updated workbook validation failed")
        return validate(path)

    monkeypatch.setattr(cli, "validate_export_tax_products", reject_updated)
    assert lxeskill.main(["fba", "export-tax", "products-import", "--sku", "SKU-NEW"]) == lxeskill.EXIT_BUSINESS
    record = json.loads(capsys.readouterr().out)
    assert record["ok"] is False
    assert "updated workbook validation failed" in record["error"]["message"]
    assert input_assets.current_asset("export_tax_products") == current
    assert current.path.read_bytes() == original_bytes
    assert input_assets.previous_asset("export_tax_products") is None


@pytest.mark.parametrize("error_type", [input_assets.InputAssetError, PermissionError])
def test_asset_save_failure_reports_error_and_keeps_generated_file(
    stored_products, monkeypatch, capsys, error_type
):
    def fail_promotion(*args, **kwargs):
        raise error_type("cannot save updated whitelist")

    monkeypatch.setattr(lxeskill, "promote_asset", fail_promotion)
    assert lxeskill.main(["fba", "export-tax", "products-import", "--sku", "SKU-NEW"]) == lxeskill.EXIT_BUSINESS
    record = json.loads(capsys.readouterr().out)
    assert record["ok"] is False
    assert record["data"]["success"] is False
    assert record["error"]["code"] == "input_asset_update_failed"
    assert f"{error_type.__name__}: cannot save updated whitelist" in record["error"]["message"]
    assert record["files"] == [record["data"]["output_xlsx"]]
    assert ("SKU-NEW", "新产品") in _read_rows(Path(record["files"][0]))
