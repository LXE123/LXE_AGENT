from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from services.agent_cli.mabang import download_store_unlinked_shipments as cli
from services.mabang.amazon.fba import unlinked_shipments as ship
from services.mabang.amazon.fba import store_msku_replenishment as rep
from test_mabang_store_msku_replenishment import _write_inventory_report, _write_sales_report


def query(*, store="Amazon-Test", totals=(0, 0, 0), raw=""):
    return ship.StoreUnlinkedShipmentDownloadResult(
        store_name=store, store_id=401403, download_time="202605251700",
        status_results=[ship.UnlinkedShipmentStatusResult(s.status_name, t, raw_file_path=raw if t else "")
                        for s, t in zip(ship.UNLINKED_SHIPMENT_STATUS_SPECS, totals)],
    )


def build(tmp_path, result=None, time="202605251800"):
    result = result or query()
    paths = [r.raw_file_path for r in result.status_results if r.raw_file_path]
    return ship.build_store_unlinked_shipments_snapshot(paths, store_name=result.store_name,
        query_result=result, snapshot_time=time, output_dir=tmp_path)


def test_confirmed_empty_contains_evidence_and_no_fake_rows(tmp_path):
    result = build(tmp_path)
    assert result.confirmed_empty and result.to_payload()["confirmed_empty"]
    assert (result.raw_file_count, result.detail_count, result.msku_count, result.total_unlinked_quantity) == (0, 0, 0, 0)
    assert ship.load_unlinked_shipment_quantities(result.snapshot_xlsx_path, store_name="Amazon-Test") == {}
    with ZipFile(result.snapshot_xlsx_path) as archive:
        assert archive.testzip() is None
    book = load_workbook(result.snapshot_xlsx_path)
    try:
        assert book[ship.SNAPSHOT_SUMMARY_SHEET].max_row == book[ship.SNAPSHOT_DETAIL_SHEET].max_row == 1
        assert book[ship.SNAPSHOT_VERIFICATION_SHEET].sheet_state == "hidden"
        evidence = json.loads(book[ship.SNAPSHOT_VERIFICATION_SHEET]["B2"].value)
        assert evidence["store_id"] == 401403
        assert evidence["status_totals"] == dict.fromkeys((s.status_name for s in ship.UNLINKED_SHIPMENT_STATUS_SPECS), 0)
    finally:
        book.close()


@pytest.mark.parametrize("totals", [(0,), (0, 0), (0, -1, 0), (0, None, 0), (0, False, 0), (0, "0", 0), (0, 0.5, 0), (0, 1, 0)])
def test_incomplete_or_invalid_query_never_publishes(tmp_path, totals):
    with pytest.raises(ship.UnlinkedShipmentError):
        build(tmp_path, query(totals=totals))
    assert not list(tmp_path.glob("*.xlsx"))


def test_duplicate_statuses_fail(tmp_path):
    result = query()
    result.status_results[-1] = result.status_results[0]
    with pytest.raises(ship.UnlinkedShipmentError, match="重复"):
        build(tmp_path, result)


@pytest.mark.parametrize("total", [None, False, True, -1, 0.1, "0.1", "NaN", "", "-1"])
def test_status_total_does_not_turn_invalid_values_into_zero(total):
    with pytest.raises(ship.UnlinkedShipmentError):
        ship._list_total({"data": {"total": total}})


def test_mixed_query_persists_real_quantities(tmp_path):
    raw = tmp_path / "WMS待装箱.csv"
    raw.write_text("店铺,MSKU,MSKU发货量\nAmazon-Test,ITEM,17\n", encoding="utf-8-sig")
    result = build(tmp_path, query(totals=(0, 1, 0), raw=str(raw)))
    assert not result.confirmed_empty
    assert ship.load_unlinked_shipment_quantities(result.snapshot_xlsx_path, store_name="Amazon-Test") == {"ITEM": 17}


@pytest.mark.parametrize("key,value", [("version", 2), ("store_id", True), ("snapshot_time", "202605261800"),
                                     ("store_name", "Other"), ("confirmed_empty", False), ("status_totals", {})])
def test_corrupt_confirmation_is_rejected(tmp_path, key, value):
    path = Path(build(tmp_path).snapshot_xlsx_path)
    book = load_workbook(path)
    try:
        cell = book[ship.SNAPSHOT_VERIFICATION_SHEET]["B2"]
        data = json.loads(cell.value); data[key] = value; cell.value = json.dumps(data)
        book.save(path)
    finally:
        book.close()
    with pytest.raises((ship.UnlinkedShipmentError, RuntimeError)):
        ship.load_unlinked_shipment_quantities(path, store_name="Amazon-Test")


def test_empty_snapshot_cannot_be_renamed_or_used_for_another_store_or_day(tmp_path):
    path = Path(build(tmp_path).snapshot_xlsx_path)
    with pytest.raises(RuntimeError, match="店铺不一致"):
        ship.load_unlinked_shipment_quantities(path, store_name="Other")
    with pytest.raises(rep.StoreMskuReplenishmentError, match="日期不一致"):
        rep.validate_unlinked_shipments_snapshot_same_day(path, "202605261800")
    other = tmp_path / path.name.replace("20260525", "20260526")
    other.write_bytes(path.read_bytes())
    with pytest.raises(RuntimeError, match="时间.*不一致"):
        ship.load_unlinked_shipment_quantities(other, store_name="Amazon-Test")


@pytest.mark.parametrize("detail", [False, True])
def test_confirmed_empty_rejects_injected_business_rows(tmp_path, detail):
    path = Path(build(tmp_path).snapshot_xlsx_path)
    book = load_workbook(path)
    try:
        book[ship.SNAPSHOT_DETAIL_SHEET if detail else ship.SNAPSHOT_SUMMARY_SHEET].append(["Amazon-Test", "ITEM", 1])
        book.save(path)
    finally:
        book.close()
    with pytest.raises((RuntimeError, ship.UnlinkedShipmentError)):
        ship.load_unlinked_shipment_quantities(path)


def test_legacy_nonempty_works_but_unverified_empty_is_not_zero(tmp_path):
    path = tmp_path / "legacy.xlsx"
    book = Workbook();book.active.append(ship.SNAPSHOT_SUMMARY_COLUMNS);book.active.append(["Amazon-Test", "ITEM", 7]);book.save(path);book.close()
    assert ship.load_unlinked_shipment_quantities(path, store_name="Amazon-Test") == {"ITEM": 7}
    book = load_workbook(path);book.active.delete_rows(2);book.save(path);book.close()
    with pytest.raises(ship.UnlinkedShipmentError, match="重新查询"):
        ship.load_unlinked_shipment_quantities(path)


@pytest.mark.parametrize("existing", [False, True])
@pytest.mark.parametrize("stage", ["save", "integrity", "validate", "publish"])
def test_snapshot_failure_preserves_previous_file(tmp_path, monkeypatch, existing, stage):
    path = tmp_path / "202605251800-Amazon-Test_未关联货件快照.xlsx"
    previous = None
    if existing:
        build(tmp_path);previous = path.read_bytes()
    error = OSError(f"actual snapshot {stage} failure")
    def fail(*args, **kwargs): raise error
    if stage == "save": monkeypatch.setattr(Workbook, "save", fail)
    elif stage == "integrity": monkeypatch.setattr(ZipFile, "testzip", fail)
    elif stage == "validate": monkeypatch.setattr(ship, "_snapshot_records", fail)
    else: monkeypatch.setattr(ship.os, "replace", fail)
    with pytest.raises(OSError) as caught: build(tmp_path)
    assert caught.value is error
    assert path.read_bytes() == previous if existing else not path.exists()
    assert not list(tmp_path.glob(".replenishment-*"))


def test_later_confirmed_empty_supersedes_old_nonzero_in_calculation(tmp_path):
    sales, inventory, snapshots = (tmp_path / name for name in ("sales", "inventory", "snapshots"))
    _write_sales_report(sales / "202605251530-Amazon-Test_销量分析.xlsx")
    _write_inventory_report(inventory / "202605251530-Amazon-Test_真实库存（深圳仓库）.xlsx")
    snapshots.mkdir()
    ship.write_unlinked_shipments_snapshot([{"店铺": "Amazon-Test", "MSKU": "AIR-1", "未关联数量": 20}], [],
                                          snapshots / "202605251200-Amazon-Test_未关联货件快照.xlsx")
    zero = build(snapshots)
    result = rep.calculate_store_msku_replenishment("Amazon-Test", sales_analysis_dir=sales,
        actual_inventory_dir=inventory, output_dir=tmp_path / "reports", unlinked_shipments_snapshot_dir=snapshots)
    assert result.unlinked_shipments_snapshot_path == zero.snapshot_xlsx_path
    assert not result.unlinked_shipments_snapshot_warning
    book = load_workbook(result.report_xlsx_path, data_only=True)
    try:
        rows = list(book["最终备货意见"].iter_rows(min_row=2, values_only=True)); cols = {v:i for i,v in enumerate(rows[0])}
        row = next(row for row in rows[1:] if row[cols["MSKU"]] == "AIR-1")
        assert row[cols["未发出货件总计"]] == 0
        assert row[cols["空运建议量"]] == 60
    finally: book.close()


def test_cli_query_failure_does_not_build_snapshot(monkeypatch):
    async def fail(*args, **kwargs): raise RuntimeError("actual upstream failure 401403")
    def forbidden(*args, **kwargs): pytest.fail("must not build a snapshot after query failure")
    monkeypatch.setattr(cli, "download_store_unlinked_shipments", fail)
    monkeypatch.setattr(cli, "build_store_unlinked_shipments_snapshot", forbidden)
    result = cli.run({"store_name": "Amazon-Test"})
    assert result["success"] is False and result["exception"] == "actual upstream failure 401403"
