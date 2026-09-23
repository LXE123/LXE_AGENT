from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.agent_cli.mabang import customs_erp as cli
from services.agent_cli.mabang.customs_prices import read_prices, match_price
from test_mabang_customs_declaration_fill_cli import _write_template, _write_consignment_excel


def prices_file(path, *, bad=False):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "备货单"
    headers = ["日期", "库存sku（第一行）", "型号", "发货量", "原价", "售价", "库存sku"]
    sheet.append(headers)
    sheet.append(["2026-09-01", "SKU-A", None, 99999, 3, "#N/A", "SKU-A × 70\nSKU-B × 110"])
    sheet.append(["走库存", "SKU-A", None, 88888, 2, 5, "SKU-A × 30"])
    summary = wb.create_sheet("汇总表")
    summary.append(headers[:-1])
    summary.append(["2026-09-01", "SKU-A", None, 99999, 3, "#N/A" if bad else 12])
    summary.append(["走库存", "SKU-A", None, 88888, 2, 6])
    summary.append(["合计", "", None, 188887])
    wb.save(path)
    wb.close()
    return path


def response(actual=6):
    return {"response_schema": cli.API_SCHEMA, "calculation_fingerprint": "original", "status": "ready",
            "inventory_changes_committed": False, "snapshot_changes_committed": False,
            "shipments": [{"sp_no": "SP260831008", "packing_sp_no": "SP260831008", "country": "美国",
                "batch_no": "B-1", "revision_id": "r1", "version_no": 1, "issues": [], "excluded_lines": [],
                "lines": [{"stock_sku": "SKU-A", "model": "M-1", "supplier_name": "工厂",
                           "product_name": "硅胶表带", "commodity_name": "硅胶表带", "unit": "条",
                           "planned_quantity": 10, "actual_quantity": actual,
                           "sources": [
                               {"source_id": "old", "source_kind": "carryover", "purchase_price": 2,
                                "planned_quantity": 4, "actual_quantity": min(4, actual), "source_contract_no": "OLD"},
                               {"source_id": "new", "source_kind": "current_purchase", "purchase_price": 3,
                                "planned_quantity": 6, "actual_quantity": max(actual - 4, 0), "source_contract_no": "NEW"}]}]}]}


@pytest.fixture
def setup(monkeypatch, tmp_path):
    prices = prices_file(tmp_path / "9.1-SP260831008-美国.xlsx")
    csv = tmp_path / "fresh.csv"
    csv.write_text("MSKU,MSKU发货量\nMSKU-A,3\n", encoding="utf-8-sig")
    template_path = tmp_path / "template.xlsx"
    _write_template(template_path, customs_detail_blocks=2)
    weight = _write_consignment_excel(tmp_path / "weight.xlsx", [{"箱序号": 1, "毛重": 10}])
    downloads, calls = [], []
    erp = response()
    def download(sp):
        downloads.append(sp)
        return csv
    def request(method, path, **kwargs):
        calls.append((method, path, deepcopy(kwargs["json_payload"])))
        return 200, deepcopy(erp)
    monkeypatch.setattr(cli, "_download", download)
    monkeypatch.setattr(cli, "request_json", request)
    arguments = {"input_xlsx": [str(prices)], "template_xlsx": str(template_path),
                 "consignment_excel": str(weight), "output_dir": str(tmp_path / "output")}
    return arguments, erp, downloads, calls, csv


def test_summary_prices_and_sku_coverage(tmp_path):
    rows = read_prices(prices_file(tmp_path / "prices.xlsx"))
    assert rows[0].sale_price == 12
    assert rows[1].sale_price == 6
    assert rows[0].skus == {"SKU-A", "SKU-B"}
    assert match_price(rows, "sku-b", {"source_kind": "current_purchase", "purchase_price": 3}).matched_rows == (rows[0],)
    assert match_price(rows, "SKU-A", {"source_kind": "carryover", "purchase_price": 2}).matched_rows == (rows[1],)


def test_price_conflict_is_not_first_match(tmp_path):
    rows = read_prices(prices_file(tmp_path / "prices.xlsx"))
    with pytest.raises(ValueError, match="不唯一"):
        match_price(rows + [replace(rows[0], row_number=4, sale_price=Decimal(13))],
                    "SKU-A", {"source_kind": "current_purchase", "purchase_price": 3})


def test_plain_sku_ending_in_x_digits_is_not_a_quantity(tmp_path):
    path = prices_file(tmp_path / "prices.xlsx")
    workbook = load_workbook(path)
    workbook["备货单"]["G2"] = "SKU-A × 70\nSKU-X123\nSKU-BX456 × 110"
    workbook.save(path)
    workbook.close()
    row = read_prices(path)[0]
    assert row.skus == {"SKU-A", "SKU-X123", "SKU-BX456"}


def test_preview_then_fill_uses_erp_quantities_and_frozen_prices(setup):
    args, erp, downloads, calls, csv = setup
    preview = cli.preview(args)
    assert preview["status"] == "confirmation_required"
    assert preview["model_summary"][0]["shortage_quantity"] == "4"
    assert not list(Path(preview["preview_path"]).parent.glob("*documents.xlsx"))
    # Original downloads/uploads may change after review; captured files remain authoritative.
    csv.write_text("MSKU,MSKU发货量\nMSKU-A,99\n")
    Path(args["input_xlsx"][0]).write_bytes(b"replaced upload")
    result = cli.fill(preview["preview_path"])
    assert result["row_count"] == 2
    assert result["total_amount"] == 48  # 4 old * 6 + 2 new * 12; never spreadsheet quantity
    assert result["box_count"] == 1
    assert downloads == ["SP260831008"]
    assert len(calls) == 2 and all(c[1] == cli.API_PATH for c in calls)
    assert calls[0][2] == calls[1][2]
    wb = load_workbook(result["output_xlsx"])
    assert {"申报要素", "报关单", "发票", "箱单", "合同"}.issubset(wb.sheetnames)
    wb.close()


@pytest.mark.parametrize("damage", ["manifest", "prices", "erp"])
def test_changed_review_cannot_generate(setup, damage):
    args, erp, _, _, _ = setup
    preview = cli.preview(args)
    path = Path(preview["preview_path"])
    if damage == "manifest":
        path.write_text(path.read_text() + " ")
    elif damage == "prices":
        (path.parent / "prices-0.xlsx").write_bytes(b"changed")
    else:
        erp["calculation_fingerprint"] = "changed"
    with pytest.raises(ValueError, match="重新预览"):
        cli.fill(path)
    assert not list(path.parent.glob("*documents.xlsx"))


def test_zero_shipment_is_previewable_without_empty_declaration(setup):
    args, erp, _, _, csv = setup
    erp.update(response(actual=0))
    csv.write_text("MSKU,MSKU发货量\nMSKU-A,0\n")
    preview = cli.preview(args)
    assert preview["status"] == "no_shipment"
    assert preview["model_summary"][0]["shortage_quantity"] == "10"
    with pytest.raises(ValueError, match="不能生成"):
        cli.fill(preview["preview_path"])


@pytest.mark.parametrize("reason", ["unknown", "price", "country", "no_price"])
def test_unresolved_inputs_are_blocked(setup, reason):
    args, erp, _, _, _ = setup
    shipment = erp["shipments"][0]
    if reason == "unknown":
        shipment["issues"] = [{"status": "incomplete", "message": "Unknown MSKU: X"}]
        erp["status"] = "incomplete"
    elif reason == "price":
        prices_file(Path(args["input_xlsx"][0]), bad=True)
    elif reason == "country":
        shipment["country"] = "德国"
    else:
        shipment["lines"][0]["stock_sku"] = "UNPRICED"
    preview = cli.preview(args)
    assert preview["status"] == "blocked"
    assert preview["issues"]
    assert Path(preview["validation_report_xlsx"]).is_file()


@pytest.mark.parametrize("text", ["M,", "M,-1", "M,NaN", "M,1.5", ",2", "M,1\nM,"])
def test_csv_rejects_partial_and_invalid_quantities(tmp_path, text):
    path = tmp_path / "delivery.csv"
    path.write_text("MSKU,MSKU发货量\n" + text)
    with pytest.raises(ValueError):
        cli.read_actual_lines(path)


def test_csv_aggregates_case_and_keeps_full_data(tmp_path):
    path = tmp_path / "delivery.csv"
    path.write_text("MSKU,MSKU发货量\n" + "\n".join(f"M{i},1" for i in range(250)) + "\nm0,2\nZERO,0\n")
    lines = cli.read_actual_lines(path)
    assert len(lines) == 250
    assert lines[0] == {"msku": "M0", "actual_quantity": "3"}


def test_model_summary_does_not_net_shortages_against_overages():
    result = response()
    other = deepcopy(result["shipments"][0]["lines"][0])
    other.update(stock_sku="SKU-B", actual_quantity=14)
    result["shipments"][0]["lines"].append(other)
    row = cli.model_summary(result)[0]
    assert row["planned_quantity"] == row["actual_quantity"] == "20"
    assert row["shortage_quantity"] == row["overage_quantity"] == "4"


def test_old_fill_entry_requires_preview(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("legacy quantity algorithm must not run")
    monkeypatch.setattr(cli.template, "fill_customs_declaration", forbidden)
    result = cli.template.run({"input_xlsx": ["old.xlsx"]})
    assert not result["success"]
    assert "先运行" in result["exception"]


def test_template_capacity_blocks_before_confirmation(setup):
    args, _, _, _, _ = setup
    _write_template(Path(args["template_xlsx"]), customs_detail_blocks=1)
    result = cli.preview(args)
    assert result["status"] == "blocked"
    assert "容量不足" in " ".join(result["issues"])


def test_multi_sp_keeps_prices_and_weights_separate(setup, monkeypatch):
    args, erp, _, _, _ = setup
    path = prices_file(Path(args["input_xlsx"][0]).with_name("SP260831009-美国.xlsx"))
    args["input_xlsx"].append(str(path))
    weight_path = args.pop("consignment_excel")
    monkeypatch.setattr(cli.template, "_resolve_consignment_excel_path", lambda *a: Path(weight_path))
    other = deepcopy(erp["shipments"][0])
    other.update(sp_no="SP260831009", packing_sp_no="SP260831009")
    erp["shipments"].append(other)
    _write_template(Path(args["template_xlsx"]), customs_detail_blocks=4)
    result = cli.preview(args)
    assert result["row_count"] == 4
    final = cli.fill(result["preview_path"])
    assert final["total_amount"] == 96
    assert final["box_count"] == 2


def test_complete_erp_detail_is_not_capped_at_200(setup):
    args, erp, _, _, _ = setup
    line = erp["shipments"][0]["lines"][0]
    # Repeated source records must all participate, not just the display prefix.
    erp["shipments"][0]["lines"] = [deepcopy(line) for _ in range(251)]
    result = cli.preview(args)
    assert result["model_summary"][0]["actual_quantity"] == str(6 * 251)
    assert cli.fill(result["preview_path"])["total_amount"] == 48 * 251


def test_duplicate_sp_rejected_before_download(setup):
    args, _, downloads, _, _ = setup
    args["input_xlsx"] *= 2
    with pytest.raises(ValueError, match="重复 SP"):
        cli.preview(args)
    assert not downloads


def test_alias_retry_downloads_actual_packing_sp(setup, monkeypatch):
    from services.agent_cli.mabang.erp_http import ErpHttpError
    args, erp, downloads, calls, _ = setup
    erp["shipments"][0]["packing_sp_no"] = "SP999"
    existing = cli.request_json
    first = True
    def request(*a, **kw):
        nonlocal first
        if first:
            first = False
            raise ErpHttpError("packing_sp_alias_required", "use alias", detail={
                "sp_no": "SP260831008", "required_packing_sp_no": "SP999"})
        return existing(*a, **kw)
    monkeypatch.setattr(cli, "request_json", request)
    result = cli.preview(args)
    assert result["can_generate"]
    assert downloads == ["SP260831008", "SP999"]
    assert calls[0][2]["shipments"][0]["sp_no"] == "SP999"


@pytest.mark.parametrize("count,allowed", [(50, True), (51, False)])
def test_customs_product_limit(setup, count, allowed):
    args, erp, _, _, _ = setup
    workbook = Workbook()
    workbook.active.title = "备货单"
    summary = workbook.create_sheet("汇总表")
    for sheet in workbook:
        sheet.append(["日期", "库存sku（第一行）", "售价", "库存sku"])
        for i in range(count):
            sheet.append(["2026-09-01", f"SKU-{i}", 12, f"SKU-{i} × 9999"])
    workbook.save(args["input_xlsx"][0])
    workbook.close()
    base = erp["shipments"][0]["lines"][0]
    lines = []
    for i in range(count):
        line = deepcopy(base)
        line.update(stock_sku=f"SKU-{i}", model=f"M-{i}")
        line["sources"] = [{**base["sources"][1], "actual_quantity": 6}]
        lines.append(line)
    erp["shipments"][0]["lines"] = lines
    _write_template(Path(args["template_xlsx"]), customs_detail_blocks=count)
    result = cli.preview(args)
    assert result["can_generate"] is allowed
    assert result["row_count"] == count
