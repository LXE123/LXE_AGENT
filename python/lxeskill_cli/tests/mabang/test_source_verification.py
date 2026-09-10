from __future__ import annotations

import asyncio
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.mabang.amazon.fba import source_verification as active
from services.mabang.amazon.fba import combo_sku as combo
from services.mabang.amazon.fba.sku_catalog import LocalSkuDefinition, SkuCatalogSnapshot, NOT_FOUND_REASON
from services.mabang.amazon.fba import store_msku as download
from services.mabang.amazon.fba import store_msku_actual_inventory as inventory
from services.mabang.amazon.fba import store_msku_sales_analysis as sales
from services.mabang.amazon.fba import store_msku_replenishment as replenishment
from services.mabang.amazon.fba.store_resolver import FbaStore
from services.mabang.official_api import OfficialApiError


def row(msku='M', asin='A', sku='S', **extra):
    return {'MSKU': msku, 'ASIN': asin, '本地SKU': sku, **extra}


def binding(msku='M', asin='A', sku='S', kind=1):
    return LocalSkuDefinition(sku, kind, (combo.ComboComponent('S', Decimal(2)),) if kind == 2 else ())


def source(path, records, bindings=(binding(),)):
    headers = list(dict.fromkeys([*inventory.SOURCE_COLUMNS, *sales.REQUIRED_COLUMNS, '本地SKU']))
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.active.append(headers)
    for record in records:
        workbook.active.append([record.get(header) for header in headers])
    workbook.save(path)
    workbook.close()
    if bindings is not None:
        active.annotate_source(path, SkuCatalogSnapshot('shop', '10', 'us', bindings), requested_store_name='shop')
    return path


def test_classification_uses_source_sku_for_every_msku():
    rows = [row(), row('Amazon.Found.A'), row('inactive', sku='U'), row('unpaired', sku='')]
    skus = (binding(), binding(sku='U', kind=None))
    assert active.classify(rows, skus) == [('已核验', True, ''), ('已核验', True, ''), ('未匹配', False, NOT_FOUND_REASON), ('缺少本地SKU', False, '源表无本地SKU，不参与备货计算')]
    assert rows[3]['本地SKU'] == ''
    assert active.classify([row(' M ', sku=' S ')], (binding(),))[0][1] is True
    assert active.classify([row('m', asin='')], (binding(),))[0][1] is True


@pytest.mark.parametrize('skus', [(), (binding(sku='changed'),), (binding(), binding(kind=2)), (binding(), binding())])
def test_incomplete_or_duplicate_catalog_snapshot_is_not_an_exclusion(skus):
    with pytest.raises(OfficialApiError):
        active.classify([row()], skus)


def test_sku_normalization_preserves_case():
    assert active.classify([row(sku='s')], (binding(sku='s', kind=None),))[0][1] is False


@pytest.mark.parametrize('mutation', ['sales', 'flag', 'binding', 'store'])
def test_snapshot_rejects_tampering(tmp_path, mutation):
    path = source(tmp_path/'source.xlsx', [row(**{'7天销量': 7})])
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    columns = {cell.value: cell.column for cell in sheet[1]}
    if mutation == 'sales':
        sheet.cell(2, columns['7天销量'], 999)
    elif mutation == 'flag':
        sheet.cell(2, columns['是否通过绑定核验'], False)
    elif mutation == 'binding':
        workbook[active.SHEET].cell(3, 2, workbook[active.SHEET].cell(3, 2).value.replace('"S"', '"changed"'))
    workbook.save(path); workbook.close()
    with pytest.raises(active.SourceVerificationError):
        active.load_verified_source(path, store_name='other' if mutation == 'store' else 'shop')


def test_legacy_zero_active_and_report_mix_are_blocked(tmp_path):
    path = source(tmp_path/'legacy.xlsx', [row()], None)
    with pytest.raises(active.SourceVerificationError, match='重新下载'):
        active.load_verified_source(path, store_name='shop')
    path = source(tmp_path/'zero.xlsx', [row()], (binding(kind=None),))
    assert len(active.load_verified_source(path, store_name='shop').records) == 1
    with pytest.raises(active.SourceVerificationError, match='无可计算记录'):
        active.require_matching_reports(path, path, store_name='shop')
    a = source(tmp_path/'a.xlsx', [row()]); b = source(tmp_path/'b.xlsx', [row(**{'7天销量': 8})])
    with pytest.raises(active.SourceVerificationError, match='指纹不一致'):
        active.require_matching_reports(a, b, store_name='shop')
    with pytest.raises(active.SourceVerificationError, match='重新下载'):
        active.require_matching_reports(a, tmp_path/'legacy.xlsx', store_name='shop')


def test_full_source_sales_and_verified_inventory_reuse(monkeypatch, tmp_path):
    common = {'父ASIN': 'P', '商品链接': 'https://example.test/A', '7天销量': 7, '14天销量': 14, '30天销量': 30, '90天销量': 90, '可售': 2}
    records = [row(**common), row('COMBO', sku='C', **common), row('no-local', sku='', **common), row('inactive', sku='U', **{**common, '7天销量': 9000}), row('Amazon.Found.A', sku='U', **common)]
    bindings = (binding(), binding('COMBO', sku='C', kind=2), binding(sku='U', kind=None))
    path = source(tmp_path/'source'/'202609071200-shop_店铺MSKU数据.xlsx', records, bindings)
    snapshot = active.load_verified_source(path, store_name='shop')
    assert snapshot.counts == {'original_row_count': 5, 'binding_verified_row_count': 2, 'binding_unverified_row_count': 3}
    queries = []
    async def post(endpoint, body, **kwargs):
        queries.append((endpoint, body))
        assert endpoint == 'combo-skus/search'
        assert body['comboSku'] == 'C'
        return {'code': 200, 'data': {'page': 1, 'rowsPerPage': 20, 'total': 1, 'data': [{'comboSku': 'C', 'comboProductDetail': [{'stockSku': 'S', 'quantity': 2}]}]}}
    async def search(skus):
        assert skus == ['S']
    async def warehouse(**kwargs):
        return tmp_path/'stock.xlsx'
    monkeypatch.setattr(combo, 'post_json', post)
    monkeypatch.setattr(inventory, 'search_warehouse_stock', search)
    monkeypatch.setattr(inventory, 'download_warehouse_stock_xlsx', warehouse)
    monkeypatch.setattr(inventory, 'parse_stock_inventory_xlsx', lambda path: {'S': Decimal(100)})
    analysis = sales.analyze_store_msku_sales('shop', input_dir=tmp_path/'source', output_dir=tmp_path/'sales')
    result = asyncio.run(inventory.export_store_msku_actual_inventory('shop', input_dir=tmp_path/'source', output_dir=tmp_path/'inventory'))
    assert analysis.msku_count == 5
    assert result.matched_warehouse_inventory_msku_row_count == 2
    assert result.missing_local_sku_msku_row_count == 1
    assert result.missing_warehouse_inventory_msku_row_count == 2
    assert queries == [], 'Inventory must reuse catalog snapshot, without any official API query'
    workbook = load_workbook(analysis.report_xlsx_path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            if sheet.title == active.SHEET:
                continue
            values = sheet.values; headers = next(values); rows = [dict(zip(headers, values)) for values in values]
            if rows and '7天销量' in headers:
                assert sum(float(row.get('7天销量') or 0) for row in rows) == 9028
    finally:
        workbook.close()
    inputs = replenishment.load_inventory_rows(result.shenzhen_warehouse_inventory_report_xlsx_path)
    assert {item.msku for item in inputs} == {'M', 'COMBO'}
    active.require_matching_reports(Path(analysis.report_xlsx_path), Path(result.shenzhen_warehouse_inventory_report_xlsx_path), store_name='shop')
    # Replaying fixed source/warehouse yields identical rows and no extra Listing request.
    again = asyncio.run(inventory.export_store_msku_actual_inventory('shop', input_dir=tmp_path/'source', output_dir=tmp_path/'inventory2'))
    assert replenishment.load_inventory_rows(again.shenzhen_warehouse_inventory_report_xlsx_path) == inputs


@pytest.mark.parametrize('group', [False, True])
def test_download_publication_and_group_scope(monkeypatch, tmp_path, group):
    async def stores():
        return [FbaStore('shop', '1'), *([FbaStore('child', '2', 'shopId', parent_store_id='1', parent_id_type='fbaWarehouseIds[]')] if group else [])]
    async def pipeline(spec):
        staged = spec.download_file.keywords['output_dir']
        path = source(staged/'202609071200-shop_店铺MSKU数据.xlsx', [row()], None)
        return download.StoreMskuExcelResult('shop', '1', 'fbaWarehouseIds[]', 1, str(path), False, False)
    async def failing_snapshot(name, skus):
        assert not list(tmp_path.glob('*.xlsx')), 'unverified source was published'
        raise OfficialApiError('Listing shop', 'HTTP 401 upstream-denied')
    monkeypatch.setattr(download, 'fetch_fba_stores', stores)
    monkeypatch.setattr(download, 'run_export_pipeline', pipeline)
    monkeypatch.setattr(download, 'fetch_sku_catalog_snapshot', failing_snapshot)
    if group:
        async def unexpected(spec):
            pytest.fail('Group must be blocked before the export pipeline')
        monkeypatch.setattr(download, 'run_export_pipeline', unexpected)
        with pytest.raises(download.StoreMskuGroupNotSupportedError) as exc:
            asyncio.run(download.download_store_msku_excel('1', 'fbaWarehouseIds[]', store_name='shop', output_dir=tmp_path/'not-created'))
        assert exc.value.context['candidates'] == [{'store_name': 'child', 'store_id': '2', 'id_type': 'shopId'}]
        assert not (tmp_path/'not-created').exists()
    else:
        with pytest.raises(OfficialApiError, match='upstream-denied'):
            asyncio.run(download.download_store_msku_excel('1', 'fbaWarehouseIds[]', store_name='shop', output_dir=tmp_path))
        assert not list(tmp_path.glob('*.xlsx'))
    assert not list(tmp_path.glob('.verified-msku-*'))


@pytest.mark.parametrize('status', ['Active', '在售', 'Inactive', '停售', 'delete', 'Incomplete', None])
def test_listing_accepts_all_statuses_without_filter(monkeypatch, status):
    async def post(endpoint, body, **kwargs):
        if endpoint == 'shops/list':
            return {'code': 200, 'data': {'profile': {'name': 'shop', 'sid': 10, 'amazonsite': 'us'}}}
        assert 'pStatus' not in body
        return {'code': 200, 'data': {'list': [{'platformSku': 'M', 'asin': 'A', 'stockSku': 'S', 'stockType': 1, 'shopIds': '10', 'amazonsite': 'us', 'pStatus': status}], 'total': 1, 'totalPage': 1, 'nowPage': 1}}
    monkeypatch.setattr(combo, 'post_json', post)
    assert asyncio.run(combo.fetch_listing_snapshot('shop')).bindings == (combo.ListingSkuBinding('M', 'A', 'S', 1),)


def test_source_shop_is_not_inferred_from_matching_sku(tmp_path):
    extra = {'店铺名称': 'other', '站点': '欧洲站'}
    path = tmp_path/'source.xlsx'
    workbook = Workbook(); workbook.active.append(['MSKU', 'ASIN', '本地SKU', *extra])
    workbook.active.append(['M', 'A', 'S', *extra.values()]); workbook.save(path); workbook.close()
    with pytest.raises(OfficialApiError, match='源表店铺名称与所选店铺不一致'):
        active.annotate_source(path, SkuCatalogSnapshot('shop', '10', 'us', (binding(),)), requested_store_name='shop')


@pytest.mark.parametrize('site,label', [('de', '欧洲站'), ('fr', '欧洲站'), ('gb', '欧洲站'), ('us', '美国站'), ('de', None)])
def test_source_region_label_is_preserved_without_country_comparison(tmp_path, site, label):
    path = tmp_path/'source.xlsx'
    original = {'店铺名称': 'shop', '站点': label, 'MSKU': 'Amazon.Found.B0BN5NRBP1',
                'ASIN': 'B0BN5NRBP1', '本地SKU': 'S', '7天销量': 7, '可售': 5}
    workbook = Workbook()
    workbook.active.append(list(original))
    workbook.active.append(list(original.values()))
    workbook.save(path)
    workbook.close()
    active.annotate_source(path, SkuCatalogSnapshot('shop', '10', site, (binding(),)), requested_store_name='shop')
    verified = active.load_verified_source(path, store_name='shop')
    assert verified.metadata['site'] == site
    assert verified.counts['binding_verified_row_count'] == 1
    assert {key: verified.records[0][key] for key in original} == original
    # A display-only field is still part of the immutable source fingerprint.
    workbook = load_workbook(path)
    workbook.active.cell(2, 2, 'changed-label')
    workbook.save(path)
    workbook.close()
    with pytest.raises(active.SourceVerificationError, match='发生变化'):
        active.load_verified_source(path, store_name='shop')


def test_download_listing_deadline_cancels_without_publication(monkeypatch, tmp_path):
    cancelled = []
    async def stores():
        return [FbaStore('shop', '1')]
    async def pipeline(spec):
        path = source(spec.download_file.keywords['output_dir']/'202609071200-shop_店铺MSKU数据.xlsx', [row()], None)
        return download.StoreMskuExcelResult('shop', '1', 'fbaWarehouseIds[]', 1, str(path), False, False)
    async def forever(name, skus):
        try:
            await asyncio.Event().wait()
        finally:
            cancelled.append(True)
    monkeypatch.setattr(download, 'fetch_fba_stores', stores)
    monkeypatch.setattr(download, 'run_export_pipeline', pipeline)
    monkeypatch.setattr(download, 'fetch_sku_catalog_snapshot', forever)
    monkeypatch.setattr(download, 'OFFICIAL_LOOKUP_TIMEOUT_SECONDS', .01)
    with pytest.raises(OfficialApiError, match='TimeoutError'):
        asyncio.run(download.download_store_msku_excel('1', 'fbaWarehouseIds[]', store_name='shop', output_dir=tmp_path))
    assert cancelled == [True]
    assert list(tmp_path.iterdir()) == []


def test_download_cli_business_error_does_not_request_cookie_refresh(monkeypatch):
    from services.agent_cli.mabang import download_store_msku_excel as cli
    async def failure(*args, **kwargs):
        raise OfficialApiError('sid=401403', 'HTTP 401 upstream actual-denial')
    monkeypatch.setattr(cli, 'download_store_msku_excel', failure)
    payload = cli.run({'store_name': 'shop', 'store_id': '1', 'id_type': 'shopId'})
    assert payload['auth_refresh_required'] is False
    assert 'actual-denial' in payload['exception']


@pytest.mark.parametrize("field,value", [("version", 1), ("version", 2), ("verification_method", "listing"), ("binding_verified_row_count", -1), ("original_row_count", None), ("snapshot_id", ""), ("site", None)])
def test_invalid_metadata_is_rejected_before_calculation(tmp_path, field, value):
    import json
    path = source(tmp_path/'report.xlsx', [row()])
    workbook = load_workbook(path)
    metadata = json.loads(workbook[active.SHEET].cell(2, 2).value)
    metadata[field] = value
    workbook[active.SHEET].cell(2, 2, json.dumps(metadata))
    workbook.save(path); workbook.close()
    with pytest.raises(active.SourceVerificationError, match="核验信息无效"):
        active.require_matching_reports(path, path, store_name="shop")


def test_unverified_marker_blocks_positive_inventory_and_preserves_full_key(tmp_path):
    from dataclasses import replace
    records = [row('unknown', sku='U', **{'父ASIN': 'P'}), row('known', **{'父ASIN': 'P'})]
    path = source(tmp_path/'source.xlsx', records, (binding('known'), binding(sku='U', kind=None)))
    verified = active.load_verified_source(path, store_name='shop')
    inputs = inventory.load_store_msku_rows(path, records=verified.records)
    rows, _ = inventory.calculate_inventory_rows(inputs, combo_map={}, stock_quantities={'S': Decimal(100)}, unverified_rows=active.unverified_reasons(verified.metadata))
    assert rows[0].binding_verified is False
    assert rows[1].binding_verified is True
    # Even a positive quantity cannot promote an unverified record into stock rows.
    changed = replace(rows[0], actual_inventory=Decimal(100))
    assert inventory.split_inventory_rows([changed]).no_inventory_rows == [changed]
    # The final loader checks the explicit metadata even if a row is moved to a business sheet.
    report = inventory.write_actual_inventory_xlsx([replace(changed, binding_verified=True), rows[1]], tmp_path/'stock.xlsx')
    active.stamp_report(report, verified.metadata)
    assert [item.msku for item in replenishment.load_inventory_rows(report)] == ['known']


def test_old_active_metadata_and_wrong_scope_are_rejected(tmp_path):
    import json
    path = source(tmp_path/'source.xlsx', [row()])
    workbook = load_workbook(path)
    workbook[active.SHEET].title = 'Active核验信息'
    workbook.save(path); workbook.close()
    with pytest.raises(active.SourceVerificationError, match='重新下载'):
        active.load_verified_source(path, store_name='shop')
    path = source(tmp_path/'new.xlsx', [row()])
    workbook = load_workbook(path)
    metadata = json.loads(workbook[active.SHEET].cell(2,2).value)
    metadata['scope'] = 'active_only'
    workbook[active.SHEET].cell(2,2,json.dumps(metadata))
    workbook.save(path); workbook.close()
    with pytest.raises(active.SourceVerificationError, match='不是全量'):
        active.read_report_metadata(path)
