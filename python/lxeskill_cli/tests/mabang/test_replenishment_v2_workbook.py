"""Workbook-owned mapping, provenance and publication; formula oracles live separately."""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from shutil import copyfile
import subprocess
from unittest.mock import Mock

import pytest
from openpyxl import Workbook, load_workbook

from services.mabang.amazon.fba import replenishment_v2_workbook as v2_workbook
from services.mabang.amazon.fba import replenishment_v2_settings as settings
from services.mabang.amazon.fba import source_verification
from services.mabang.amazon.fba import store_msku_actual_inventory as inventory
from services.mabang.amazon.fba import unlinked_shipments as shipments
from services.mabang.amazon.fba.replenishment_v2_formulas import replenishment_v2_formulas
from services.mabang.amazon.fba.sku_catalog import LocalSkuDefinition, SkuCatalogSnapshot
from services.mabang.amazon.fba.store_msku_replenishment import StoreMskuReplenishmentError
from services.mabang.amazon.fba.replenishment_v2_workbook import ReplenishmentV2WorkbookError, generate_replenishment_v2

STORE = "Amazon-Test"
SOURCE_HEADERS = (
    "店铺名称", "MSKU", "父ASIN", "ASIN", "本地SKU", "本地SKU名称", "商品名称", "备注",
    "7天销量", "14天销量", "30天销量", "90天销量", "可售", "待入库", "预留",
    "计划入库", "在途", "待调仓", "调仓中", "单品重量(g)(cm)",
)



@pytest.fixture(autouse=True)
def isolated_settings(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "app"))


def _source_rows() -> list[dict[str, object]]:
    return [
        {
            "MSKU": "SEA-1", "父ASIN": "PARENT-SEA", "ASIN": "ASIN-SEA", "本地SKU": "SKU-SEA",
            "本地SKU名称": "Ocean item", "备注": "", "7天销量": 77, "14天销量": 154,
            "30天销量": 330, "90天销量": 990, "可售": 880, "待入库": 0, "预留": 0,
            "计划入库": 17, "在途": 0, "待调仓": 0, "调仓中": 0, "单品重量(g)(cm)": "50g 10*10*10",
        },
        {
            "MSKU": "CLEAR-1", "父ASIN": "PARENT-CLEAR", "ASIN": "ASIN-CLEAR", "本地SKU": "SKU-CLEAR",
            "本地SKU名称": "Clearance item", "备注": "清货", "7天销量": 7, "14天销量": 14,
            "30天销量": 30, "90天销量": 90, "可售": 0, "待入库": 0, "预留": 0,
            "计划入库": 0, "在途": 0, "待调仓": 0, "调仓中": 0, "单品重量(g)(cm)": 50,
        },
        {
            "MSKU": "COMBO-1", "父ASIN": "PARENT-COMBO", "ASIN": "ASIN-COMBO", "本地SKU": "SKU-COMBO",
            "本地SKU名称": "Combo item", "备注": "", "7天销量": 7, "14天销量": 14,
            "30天销量": 30, "90天销量": 90, "可售": 0, "待入库": 0, "预留": 0,
            "计划入库": 0, "在途": 0, "待调仓": 0, "调仓中": 0, "单品重量(g)(cm)": 25,
        },
        {
            "MSKU": "UNBOUND-1", "父ASIN": "PARENT-UNBOUND", "ASIN": "ASIN-UNBOUND", "本地SKU": "UNVERIFIED-LOCAL",
            "本地SKU名称": "Unbound item", "备注": "", "7天销量": 7, "14天销量": 14,
            "30天销量": 30, "90天销量": 90, "可售": 0, "待入库": 0, "预留": 0,
            "计划入库": 0, "在途": 0, "待调仓": 0, "调仓中": 0, "单品重量(g)(cm)": 50,
        },
    ]


def _write_source(
    path: Path, rows: list[dict[str, object]] | None = None,
    *, headers: tuple[str, ...] = SOURCE_HEADERS,
    local_skus: tuple[LocalSkuDefinition, ...] | None = None,
) -> Path:
    rows = _source_rows() if rows is None else rows
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
        # Model an authoritative source whose strings remain literal text.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    workbook.save(path)
    workbook.close()
    if local_skus is None:
        local_skus = tuple(
            LocalSkuDefinition(str(row["本地SKU"]), None if row["MSKU"] == "UNBOUND-1" else 1)
            for row in rows if str(row.get("本地SKU") or "").strip()
        )
    source_verification.annotate_source(
        path,
        SkuCatalogSnapshot(STORE, "401403", "shopId", local_skus),
        requested_store_name=STORE,
    )
    return path


def _write_inventory_report(path: Path, source_path: Path, rows: list[dict[str, object]] | None = None) -> Path:
    source = source_verification.load_verified_source(source_path, store_name=STORE)
    rows = _source_rows() if rows is None else rows
    inventory_rows = []
    verdicts = {record["MSKU"]: record["是否通过绑定核验"] for record in source.records}
    for row in rows:
        child_skus = "SKU-CHILD * 2" if row["MSKU"] == "COMBO-1" else ""
        inventory_rows.append(
            inventory.ActualInventoryRow(
                msku=str(row["MSKU"]),
                parent_asin=str(row["父ASIN"]),
                asin=str(row["ASIN"]),
                local_sku=str(row["本地SKU"]),
                product_link="",
                actual_inventory=Decimal("12") if row["MSKU"] == "SEA-1" else Decimal(0),
                child_skus=child_skus,
                binding_verified=verdicts[str(row["MSKU"])],
            )
        )
    inventory.write_actual_inventory_xlsx(inventory_rows, path)
    source_verification.stamp_report(path, source.metadata)
    return path


def _write_confirmed_empty_snapshot(directory: Path, *, snapshot_time: str = "202610131100") -> Path:
    query = shipments.StoreUnlinkedShipmentDownloadResult(
        store_name=STORE,
        store_id=401403,
        download_time="202610131000",
        status_results=[
            shipments.UnlinkedShipmentStatusResult(spec.status_name, 0)
            for spec in shipments.UNLINKED_SHIPMENT_STATUS_SPECS
        ],
    )
    return Path(
        shipments.build_store_unlinked_shipments_snapshot(
            [], store_name=STORE, query_result=query, snapshot_time=snapshot_time, output_dir=directory,
        ).snapshot_xlsx_path
    )


def _generate(tmp_path: Path, *, source_path: Path | None = None, inventory_path: Path | None = None,
              snapshot_path: Path | None = None, **generation_settings):
    source_path = source_path or _write_source(
        tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx"
    )
    inventory_path = inventory_path or _write_inventory_report(
        tmp_path / "inventory.xlsx", source_path
    )
    snapshot_path = snapshot_path or _write_confirmed_empty_snapshot(tmp_path / "snapshots")
    return generate_replenishment_v2(
        STORE,
        msku_source_xlsx_path=source_path,
        shenzhen_inventory_report_xlsx_path=inventory_path,
        unlinked_shipments_snapshot_path=snapshot_path,
        output_dir=tmp_path / "out",
        **generation_settings,
    )


def test_workbook_layout_mapping_and_pr1_formulas_before_office(tmp_path: Path) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    report_path = _write_inventory_report(tmp_path / "inventory.xlsx", source_path)
    snapshot_path = _write_confirmed_empty_snapshot(tmp_path / "snapshots")
    verified = source_verification.load_verified_source(source_path, store_name=STORE)
    _, source_by_key = v2_workbook._unique_source_rows(verified.records)
    inventory_by_key = v2_workbook._load_shenzhen_rows(report_path, expected_keys=set(source_by_key))
    shipment_by_key = v2_workbook._shipment_quantities(
        snapshot_path, store_name=STORE, source_data_time="202610131000"
    )
    rows = v2_workbook._rows_for_workbook(
        verified.records,
        store_name=STORE,
        inventory_by_key=inventory_by_key,
        shipment_by_key=shipment_by_key,
    )
    rows[0]["本地SKU名称"] = "=2+2"
    workbook_path = tmp_path / "formulas.xlsx"
    v2_workbook._write_workbook(workbook_path, rows)

    workbook = load_workbook(workbook_path, data_only=False)
    try:
        assert workbook.sheetnames == ["mskulist"]
        sheet = workbook["mskulist"]
        assert sheet.max_column == 35
        assert sheet.max_row == len(verified.records) + 1
        assert tuple(sheet.cell(1, column).value for column in range(1, 36)) == (
            "店铺名称", "MSKU", "父ASIN", "ASIN", "7天销量", "14天销量", "30天销量", "90天销量",
            "销量趋势", "加权日销", "可售", "待入库", "预留", "计划入库", "在途", "运营中心转运",
            "运营中心处理中", "FBA总库存（马帮）", "可销售天数", "未发出货件总计", "建议运输方式",
            "空运补货天数", "海运补货天数", "空运计划", "空运实际", "空运总重", "海运计划",
            "海运实际", "海运总重", "单件重量(g)", "本地SKU名称", "备注", "本地库存", "本地SKU", "可供天数",
        )
        assert sheet["AE2"].value == "=2+2" and sheet["AE2"].data_type == "s"
        assert sheet["AH4"].value == "SKU-COMBO"
        assert sheet["AH3"].value == "SKU-CLEAR"
        assert sheet["AG3"].value == 0
        assert sheet["AG2"].value == 12
        assert sheet["N2"].value == 17
        assert sheet["T2"].value == 0
        assert sheet["AD2"].value == 50
        for row_number in range(2, len(rows) + 2):
            expected = replenishment_v2_formulas(row_number)
            actual = {
                column: sheet[f"{column}{row_number}"].value
                for column in expected
            }
            assert actual == expected
        assert all(sheet[f"{column}5"].data_type == "f" for column in ("U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC", "AI"))
        assert [sheet.cell(row, 2).value for row in range(2, 6)] == ["SEA-1", "CLEAR-1", "COMBO-1", "UNBOUND-1"]
    finally:
        workbook.close()


@pytest.mark.parametrize("mutation", ["store", "fingerprint"])
def test_source_store_and_fingerprint_are_verified(tmp_path: Path, mutation: str) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    inventory_path = _write_inventory_report(tmp_path / "inventory.xlsx", source_path)
    snapshot_path = _write_confirmed_empty_snapshot(tmp_path / "snapshots")
    if mutation == "fingerprint":
        workbook = load_workbook(source_path)
        workbook.active["I2"] = 78  # Change raw 7-day sales after source verification.
        workbook.save(source_path)
        workbook.close()
    diagnostic = "店铺不一致" if mutation == "store" else "发生变化"
    with pytest.raises(source_verification.SourceVerificationError, match=diagnostic):
        generate_replenishment_v2(
            "Wrong-Store" if mutation == "store" else STORE,
            msku_source_xlsx_path=source_path,
            shenzhen_inventory_report_xlsx_path=inventory_path,
            unlinked_shipments_snapshot_path=snapshot_path,
            output_dir=tmp_path / "out",
        )


@pytest.mark.parametrize("different_content", [True, False])
def test_inventory_report_must_have_the_same_source_batch(tmp_path: Path, different_content: bool) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    different_rows = _source_rows()
    if different_content:
        different_rows[0]["7天销量"] = 78
    other_source = _write_source(tmp_path / "other.xlsx", different_rows)
    source_metadata = source_verification.load_verified_source(source_path, store_name=STORE).metadata
    other_metadata = source_verification.load_verified_source(other_source, store_name=STORE).metadata
    assert source_metadata["snapshot_id"] != other_metadata["snapshot_id"]
    assert (source_metadata["source_fingerprint"] != other_metadata["source_fingerprint"]) is different_content
    report = _write_inventory_report(tmp_path / "inventory.xlsx", other_source)
    snapshot = _write_confirmed_empty_snapshot(tmp_path / "snapshots")

    with pytest.raises(ReplenishmentV2WorkbookError, match="批次"):
        _generate(
            tmp_path,
            source_path=source_path,
            inventory_path=report,
            snapshot_path=snapshot,
        )


def test_inventory_report_must_cover_every_source_msku(tmp_path: Path) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    report = _write_inventory_report(tmp_path / "inventory.xlsx", source_path, _source_rows()[:2])
    snapshot = _write_confirmed_empty_snapshot(tmp_path / "snapshots")

    with pytest.raises(ReplenishmentV2WorkbookError, match="覆盖与源表不一致"):
        _generate(tmp_path, source_path=source_path, inventory_path=report, snapshot_path=snapshot)


def test_inventory_report_rejects_duplicate_msku_across_sheets(tmp_path: Path, monkeypatch) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    report = _write_inventory_report(tmp_path / "inventory.xlsx", source_path)
    snapshot = _write_confirmed_empty_snapshot(tmp_path / "snapshots")
    workbook = load_workbook(report)
    try:
        stock_sheet = workbook[inventory.STOCK_ACTUAL_INVENTORY_SHEET]
        combo_sheet = workbook[inventory.COMBO_ACTUAL_INVENTORY_SHEET]
        headers = [cell.value for cell in stock_sheet[1]]
        assert [cell.value for cell in combo_sheet[1]] == headers
        msku_index = headers.index("MSKU")
        duplicate = next(list(row) for row in stock_sheet.iter_rows(min_row=2, values_only=True)
                         if row[msku_index] == "SEA-1")
        duplicate[headers.index(inventory.ACTUAL_INVENTORY_QUANTITY_COLUMN)] = 999
        combo_sheet.append(duplicate)
        assert stock_sheet.title != combo_sheet.title
        assert sum(row[msku_index] == "SEA-1" for sheet in (stock_sheet, combo_sheet)
                   for row in sheet.iter_rows(min_row=2, values_only=True)) == 2
        workbook.save(report)
    finally:
        workbook.close()

    write = Mock(side_effect=AssertionError("重复库存报告不得生成工作簿"))
    recalculate = Mock(side_effect=AssertionError("重复库存报告不得进入Office重算"))
    monkeypatch.setattr(v2_workbook, "_write_workbook", write)
    monkeypatch.setattr(v2_workbook, "_recalculate", recalculate)
    with pytest.raises(ReplenishmentV2WorkbookError, match="深圳库存报告存在重复 MSKU: 'SEA-1'"):
        _generate(tmp_path, source_path=source_path, inventory_path=report, snapshot_path=snapshot)
    write.assert_not_called()
    recalculate.assert_not_called()
    assert not list((tmp_path / "out").glob("*.xlsx"))


def test_snapshot_from_a_different_day_is_rejected(tmp_path: Path) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    report = _write_inventory_report(tmp_path / "inventory.xlsx", source_path)
    snapshot = _write_confirmed_empty_snapshot(tmp_path / "snapshots", snapshot_time="202610121100")

    with pytest.raises(StoreMskuReplenishmentError, match="日期与备货数据日期不一致"):
        _generate(tmp_path, source_path=source_path, inventory_path=report, snapshot_path=snapshot)


def test_office_failure_preserves_real_diagnostic(monkeypatch, tmp_path: Path) -> None:
    def fail(arguments, *, timeout_seconds):
        assert arguments[0] == "recalculate"
        assert timeout_seconds == v2_workbook.OFFICE_TIMEOUT_SECONDS
        return subprocess.CompletedProcess(arguments, 1, "", "LXE_OFFICE_NODE is not configured")

    monkeypatch.setattr(v2_workbook, "run_kit_bounded", fail)
    with pytest.raises(ReplenishmentV2WorkbookError, match="LXE_OFFICE_NODE is not configured"):
        v2_workbook._recalculate(tmp_path / "in.xlsx", tmp_path / "out.xlsx")


def test_formula_error_does_not_replace_same_day_successful_output(monkeypatch, tmp_path: Path) -> None:
    source_rows = _source_rows()[:3]
    source = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx", source_rows)
    report = _write_inventory_report(tmp_path / "inventory.xlsx", source, source_rows)
    snapshot = _write_confirmed_empty_snapshot(tmp_path / "snapshots")
    recalculation_count = 0

    def recalculate_with_distinct_valid_xlsx(input_path: Path, output_path: Path) -> None:
        nonlocal recalculation_count
        recalculation_count += 1
        workbook = load_workbook(input_path)
        try:
            workbook.properties.description = f"publication-test-{recalculation_count}"
            workbook.save(output_path)
        finally:
            workbook.close()

    inspections = iter([(0, ()), (1, ("AC3=#VALUE!",))])
    monkeypatch.setattr(v2_workbook, "_recalculate", recalculate_with_distinct_valid_xlsx)
    monkeypatch.setattr(
        v2_workbook, "_inspect_recalculated",
        lambda path, *, rows: next(inspections),
    )
    first = _generate(tmp_path, source_path=source, inventory_path=report, snapshot_path=snapshot)
    official_path = Path(first.output_xlsx_path)
    official_bytes = official_path.read_bytes()
    assert first.success

    second = _generate(tmp_path, source_path=source, inventory_path=report, snapshot_path=snapshot)
    diagnostic_path = Path(second.output_xlsx_path)
    assert not second.success
    assert second.formula_error_examples == ("AC3=#VALUE!",)
    assert diagnostic_path.is_file() and diagnostic_path != official_path
    assert official_path.read_bytes() == official_bytes
    assert diagnostic_path.read_bytes() != official_bytes


def test_report_mapping_owns_ag_ah_ad_and_shipment_values():
    records = _source_rows()[:2]
    records[0].update({"本地库存": 999, "单品重量(g)(cm)": None})
    records[1]["单品重量(g)(cm)"] = "12.5g 10*10"
    report = {"SEA-1": v2_workbook.ShenzhenInventoryRow(-3, "REPORT-SKU", "库存", ""),
              "CLEAR-1": v2_workbook.ShenzhenInventoryRow(0, "", "库存", "")}
    rows = v2_workbook._rows_for_workbook(records, store_name=STORE, inventory_by_key=report, shipment_by_key={"SEA-1": 17})
    assert [(r["本地库存"], r["本地SKU"], r["单件重量(g)"], r["未发出货件总计"]) for r in rows] == [
        (-3, "REPORT-SKU", 0, 17), (0, "SKU-CLEAR", 12.5, 0)]
    with pytest.raises(ReplenishmentV2WorkbookError, match="源表没有的 MSKU"):
        v2_workbook._rows_for_workbook(records, store_name=STORE, inventory_by_key=report, shipment_by_key={"ORPHAN": 1})
    with pytest.raises(ReplenishmentV2WorkbookError, match="重复 MSKU"):
        v2_workbook._unique_source_rows([records[0], records[0]])
    with pytest.raises(ReplenishmentV2WorkbookError, match="可售"):
        v2_workbook._rows_for_workbook([{**records[0], "可售": None}], store_name=STORE,
                                     inventory_by_key=report, shipment_by_key={})
    records[0]["单品重量(g)(cm)"] = "invalid"
    with pytest.raises(ReplenishmentV2WorkbookError, match="AD.*MSKU=SEA-1"):
        v2_workbook._rows_for_workbook(records, store_name=STORE, inventory_by_key=report, shipment_by_key={})


def test_inventory_number_parser_preserves_unknown_and_signed_values():
    for value, expected in [(None, None), (0, 0), (-3, -3), (" 2.5 ", 2.5)]:
        assert v2_workbook._optional_number(value, context="深圳库存", allow_negative=True) == expected
    with pytest.raises(ReplenishmentV2WorkbookError):
        v2_workbook._optional_number("NaN", context="深圳库存", allow_negative=True)


@pytest.mark.parametrize("saved,override,enabled,days", [
    ((False, None), {}, False, 0),
    ((True, 20), {}, True, 20),
    ((False, 20), {"peak_boost_enabled": True}, True, 20),
    ((True, 20), {"peak_boost_enabled": False, "peak_boost_days": 40}, False, 40),
    ((False, None), {"peak_boost_days": 15}, True, 15),
])
def test_saved_settings_and_chat_override_change_only_this_run(tmp_path, monkeypatch, saved, override, enabled, days):
    parameters = dict(zip(("peak_boost_enabled", "peak_boost_days"), saved))
    settings.save_parameters(parameters)
    monkeypatch.setattr(v2_workbook, "_recalculate", copyfile)
    monkeypatch.setattr(v2_workbook, "_inspect_recalculated", lambda *_a, **_k: (0, ()))
    result = _generate(tmp_path, **override)
    book = load_workbook(result.output_xlsx_path)
    expected = replenishment_v2_formulas(2, peak_boost_enabled=enabled, peak_boost_days=days)
    try:
        assert [book.active[f"{col}2"].value for col in ("V", "W")] == [expected["V"], expected["W"]]
    finally:
        book.close()
    assert settings.read_state() == parameters


def test_running_generation_uses_its_settings_snapshot(tmp_path, monkeypatch):
    settings.save_parameters({"peak_boost_enabled": True, "peak_boost_days": 20})
    write_book = v2_workbook._write_workbook
    def change_setting_then_write(path, rows, **parameters):
        settings.save_parameters({"peak_boost_enabled": True, "peak_boost_days": 40})
        write_book(path, rows, **parameters)
    monkeypatch.setattr(v2_workbook, "_write_workbook", change_setting_then_write)
    monkeypatch.setattr(v2_workbook, "_recalculate", copyfile)
    monkeypatch.setattr(v2_workbook, "_inspect_recalculated", lambda *_a, **_k: (0, ()))
    result = _generate(tmp_path)
    book = load_workbook(result.output_xlsx_path)
    try:
        assert book.active["V2"].value == replenishment_v2_formulas(2, peak_boost_days=20)["V"]
    finally:
        book.close()
    assert settings.read_state()["peak_boost_days"] == 40


def test_enabled_run_requires_days_before_inputs_are_read():
    with pytest.raises(ValueError, match="加成天数"):
        generate_replenishment_v2(STORE, msku_source_xlsx_path="missing", shenzhen_inventory_report_xlsx_path="missing",
                                  unlinked_shipments_snapshot_path="missing", peak_boost_enabled=True)
