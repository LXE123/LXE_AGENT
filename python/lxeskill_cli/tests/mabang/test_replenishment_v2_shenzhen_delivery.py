"""Offline report-to-CLI acceptance with the project's real Office Kit engine."""
from __future__ import annotations

from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook

from lxeskill import cli as lxeskill
from services.mabang.amazon.fba import source_verification as verification
from services.mabang.amazon.fba import store_msku_actual_inventory as inventory
from services.mabang.amazon.fba import store_msku_replenishment as legacy_replenishment
from services.mabang.amazon.fba import replenishment_v2_workbook as workbook
from services.mabang.amazon.fba.sku_catalog import LocalSkuDefinition
from test_replenishment_v2_workbook import (
    SOURCE_HEADERS,
    STORE,
    _source_rows,
    _write_confirmed_empty_snapshot,
    _write_source,
)


@pytest.mark.parametrize("include_unknown", [False, True], ids=["formal-no-uq122", "inventory-review-old-value-ignored"])
def test_shenzhen_report_controls_ag_and_office_delivery_status(
    tmp_path: Path, monkeypatch, capsys, include_unknown: bool,
) -> None:
    if not os.environ.get("LXE_OFFICE_NODE") or not os.environ.get("LXE_OFFICE_CLI"):
        pytest.skip("Office Kit is not configured; real recalculation remains NOT VERIFIED")
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path))
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", "replenishment-workflow-map")
    expected_stock = {"POS-1": 12, "ZERO-1": 0, "NEG-1": -3, "COMBO-1": 7}
    rows = [
        {**_source_rows()[0], "MSKU": msku, "本地SKU": sku,
         **({"本地库存": 900 + index} if include_unknown else {}),
         "商品链接": f"https://example.test/{msku}"}
        for index, (msku, sku) in enumerate([
            ("POS-1", "SKU-POS"), ("ZERO-1", "SKU-ZERO"),
            ("NEG-1", "SKU-NEG"), ("COMBO-1", "SKU-COMBO"),
        ])
    ]
    if include_unknown:
        rows.extend([
            {**rows[0], "MSKU": "NO-LOCAL-1", "本地SKU": "", "本地库存": 904},
            {**rows[0], "MSKU": "NO-STOCK-1", "本地SKU": "SKU-MISSING", "本地库存": 905},
        ])
        expected_stock.update({"NO-LOCAL-1": None, "NO-STOCK-1": None})
    headers = (
        (*SOURCE_HEADERS, "商品链接") if not include_unknown
        else (*SOURCE_HEADERS, "本地库存", "商品链接")
    )
    source = _write_source(
        tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx", rows,
        headers=headers,
        local_skus=(
            LocalSkuDefinition("SKU-POS", 1), LocalSkuDefinition("SKU-ZERO", 1),
            LocalSkuDefinition("SKU-NEG", 1),
            LocalSkuDefinition("SKU-COMBO", 2, (
                inventory.ComboComponent("CHILD-A", Decimal(2)),
                inventory.ComboComponent("CHILD-B", Decimal(3)),
            )),
        ) + ((LocalSkuDefinition("SKU-MISSING", 1),) if include_unknown else ()),
    )
    if not include_unknown:
        source_book = load_workbook(source, read_only=True, data_only=True)
        try:
            assert "本地库存" not in [cell.value for cell in source_book.active[1]]
        finally:
            source_book.close()
    verified = verification.load_verified_source(source, store_name=STORE)
    if not include_unknown:
        assert all("本地库存" not in record for record in verified.records)
    # Produce the report with the existing Shenzhen algorithm, without ERP access.
    report_rows, missing = inventory.calculate_inventory_rows(
        inventory.load_store_msku_rows(source, records=verified.records),
        combo_map=verified.combo_map,
        stock_quantities={
            "SKU-POS": Decimal(12), "SKU-ZERO": Decimal(0), "SKU-NEG": Decimal(-3),
            "CHILD-A": Decimal(20), "CHILD-B": Decimal(21),
        },
        unverified_rows=verification.unverified_reasons(verified.metadata),
    )
    assert {row.msku: row.actual_inventory for row in report_rows} == expected_stock
    assert missing == (["SKU-MISSING"] if include_unknown else [])
    report = inventory.write_actual_inventory_xlsx(report_rows, tmp_path / "inventory.xlsx")
    verification.stamp_report(report, verified.metadata)
    snapshot = _write_confirmed_empty_snapshot(tmp_path / "snapshots")
    inputs = [source, report, snapshot]
    input_hashes = [hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs]

    def unexpected_inventory_call(*_args, **_kwargs):
        pytest.fail("the workbook must consume the report without rerunning inventory calculation or ERP")

    monkeypatch.setattr(inventory, "calculate_inventory_rows", unexpected_inventory_call)
    monkeypatch.setattr(inventory, "search_warehouse_stock", unexpected_inventory_call)
    monkeypatch.setattr(inventory, "download_warehouse_stock_xlsx", unexpected_inventory_call)
    monkeypatch.setattr(legacy_replenishment, "calculate_store_msku_replenishment", unexpected_inventory_call)
    monkeypatch.setattr(legacy_replenishment, "calculate_replenishment_rows", unexpected_inventory_call)
    code = lxeskill.main([
        "replenish", "v2", "generate", "--store-name", STORE,
        "--msku-source-xlsx-path", str(source),
        "--shenzhen-inventory-report-xlsx-path", str(report),
        "--unlinked-shipments-snapshot-path", str(snapshot),
    ])
    terminal = json.loads(capsys.readouterr().out)
    assert code == (lxeskill.EXIT_BUSINESS if include_unknown else 0)
    assert terminal["ok"] is (not include_unknown)
    assert terminal["data"]["success"] is (not include_unknown)
    assert terminal["data"]["formula_error_count"] == 0
    assert terminal["data"]["unknown_inventory_count"] == (2 if include_unknown else 0)
    output = Path(terminal["data"]["output_xlsx_path"])
    assert terminal["files"] == [str(output)]
    assert ("-库存待核对" in output.name) is include_unknown
    if include_unknown:
        examples = terminal["data"]["unknown_inventory_examples"]
        assert any("NO-LOCAL-1" in example and "无本地SKU" in example for example in examples)
        assert any("NO-STOCK-1" in example and "无库存数据" in example for example in examples)
        assert "须运营确认后才能正式交付" in terminal["error"]["message"]
    formulas = load_workbook(output, read_only=True, data_only=False)
    values = load_workbook(output, read_only=True, data_only=True)
    try:
        sheet, cached = formulas["mskulist"], values["mskulist"]
        assert sheet.max_column == 35 and sheet.max_row == len(rows) + 1
        assert [sheet[f"B{index}"].value for index in range(2, len(rows) + 2)] == list(expected_stock)
        for index, row in enumerate(rows, 2):
            assert sheet[f"AG{index}"].value == expected_stock[row["MSKU"]]
            if include_unknown:
                assert sheet[f"AG{index}"].value != row["本地库存"]
            assert len([cell for cell in sheet[index] if cell.data_type == "f"]) == 14
            assert isinstance(cached[f"J{index}"].value, (int, float))
        assert not any(cell.data_type == "e" for row in cached for cell in row)
    finally:
        formulas.close()
        values.close()
    assert [hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs] == input_hashes

    if not include_unknown:
        # Reuse the recalculated file to prove the real inspector rejects an Excel error.
        diagnostic = tmp_path / "office-error.xlsx"
        with ZipFile(output) as original, ZipFile(diagnostic, "w") as changed:
            for entry in original.infolist():
                content = original.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    sheet_xml = ET.fromstring(content)
                    cell = sheet_xml.find(".//{*}c[@r='AC2']")
                    assert cell is not None
                    cell.set("t", "e")
                    cell.find("{*}v").text = "#VALUE!"
                    content = ET.tostring(sheet_xml)
                changed.writestr(entry, content)
        count, examples = workbook._inspect_recalculated(diagnostic, rows=rows)
        assert count == 1 and examples == ("AC2=#VALUE!",)
