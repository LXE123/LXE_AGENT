from copy import deepcopy
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from services.mabang.amazon.fba import store_msku_replenishment as rep
from services.mabang.amazon.fba import replenishment_template as templates
from test_mabang_store_msku_replenishment import _inventory_input_row
from test_replenishment_formula_sheet import case_row
from test_replenishment_sample_sheet import sample_row


def decide(*, daily=6, fba=480, sea_changes=None, trend='平稳', unlinked=0):
    template = templates.load_default_template()
    params = deepcopy(template.params)
    params['sea'].update(sea_changes or {})
    template = replace(template, name='生效方案', params=params)
    values = dict.fromkeys(rep.SOURCE_VALUE_COLUMNS, 0)
    values.update({'7天销量': daily*7, '14天销量': daily*14, '30天销量': daily*30,
                   '90天销量': daily*90, '可售': fba})
    inv = _inventory_input_row('CHECK', fba_total_inventory=fba)
    detail = rep.SalesDetail(trend, None, daily*7, daily*14, daily*30, source_values=values)
    return rep.calculate_replenishment_rows([inv], {(inv.msku,inv.parent_asin,inv.asin,inv.local_sku):detail},
                                            template, {'CHECK':unlinked})[0]


@pytest.mark.parametrize('options,code,pieces', [
    ({'daily':1,'fba':60}, 'fba_covered', ['理论需求45件','FBA库存60件']),
    ({'daily':1,'fba':30,'unlinked':20}, 'inventory_shipments_covered', ['理论需求45件','库存30件','货件20件']),
    ({'daily':0,'fba':1}, 'zero_daily_sales', ['加权日销为零']),
    ({'sea_changes':{'enabled':False}}, 'sea_disabled', ['80.00天','70天','生效方案']),
    ({'daily':4,'fba':400}, 'sea_daily_threshold', ['100.00天','4.00件','超过5件']),
    ({'daily':4,'fba':400,'sea_changes':{'min_daily_sales':7,'min_daily_sales_inclusive':True}},
     'sea_daily_threshold', ['达到7件']),
    ({'sea_changes':{'tiers':[]}}, 'sea_tier_missing', ['未匹配海运天数分档']),
    ({'sea_changes':{'companion_air_enabled':True,'companion_air_tiers':[]}},
     'sea_tier_missing', ['未匹配海运同时空运天数分档']),
])
def test_reason_from_actual_branch(options, code, pieces):
    row = decide(**options)
    assert row.sheet_name == rep.NO_SHIP_SHEET
    assert row.non_shipping_reason_code == code
    assert all(p in row.non_shipping_reason_detail for p in pieces)
    payload = rep._non_shipping_payload(row)
    assert payload['未备货原因分类'] == rep.NON_SHIPPING_REASONS[code]
    # Rendering must not parse the legacy diagnostic string.
    assert rep._non_shipping_payload(replace(row, decision_reason='unrelated 401/403 text')) == payload


def test_sea_minimum_and_zero_days_remain_actual_decisions():
    row = case_row(fba=600, unlinked=31)
    assert row.replenish_quantity == 29
    assert row.non_shipping_reason_code == 'sea_quantity_threshold'
    assert '海运数量29件' in row.non_shipping_reason_detail
    assert '最低30件' in row.non_shipping_reason_detail
    assert rep._non_shipping_payload(decide(daily=0,fba=1))['可销售天数'] == ''


def test_merge_preserves_records_and_sorts_on_daily_sales(tmp_path):
    rows = [replace(sample_row(), msku='A'), replace(decide(daily=4,fba=400), msku='HIGH'),
            replace(decide(daily=1,fba=60), msku='SAME',asin='B'),
            replace(decide(daily=1,fba=60), msku='SAME',asin='A')]
    path = rep.write_replenishment_report(rows, tmp_path/'merged.xlsx')
    book = load_workbook(path)
    try:
        assert rep.NO_SHIP_SHEET not in book.sheetnames and rep.SAMPLE_INSUFFICIENT_SHEET not in book.sheetnames
        assert book.sheetnames == ['最终备货意见','真实库存（深圳仓库）不足','清货','本轮不备货','链接备货汇总','备货公式参数']
        sheet = book[rep.NON_SHIPPING_SHEET]
        data = list(sheet.iter_rows(min_row=2, values_only=True))
        assert [(r[0],r[2]) for r in data] == [(r.msku,r.asin) for r in sorted(rows,key=rep._non_shipping_sort_key)]
        assert len(data)==4 and sheet.max_column==21
        assert sheet.auto_filter.ref == 'A1:U5'
        assert sheet.freeze_panes == 'A2'
        assert {r[19] for r in data} == {'样本不足','未达海运日销门槛','FBA库存已覆盖'}
        assert not any(c.data_type=='f' for cells in sheet for c in cells)
        assert sheet['P2'].number_format=='0.00'
    finally:book.close()


def test_empty_merged_sheet_has_only_header(tmp_path):
    book = load_workbook(rep.write_replenishment_report([], tmp_path/'empty.xlsx'))
    try:
        sheet=book[rep.NON_SHIPPING_SHEET]
        assert sheet.max_row==1
        assert [c.value for c in sheet[1]]==list(rep.NON_SHIPPING_COLUMNS)
    finally:book.close()


@pytest.mark.parametrize('code,detail', [('', 'reason'), ('unknown', 'reason'), ('fba_covered','')])
def test_missing_structured_reason_blocks_publication(tmp_path, code, detail):
    path=tmp_path/'report.xlsx';path.write_bytes(b'previous')
    row=replace(decide(daily=1,fba=60), non_shipping_reason_code=code, non_shipping_reason_detail=detail)
    with pytest.raises(rep.StoreMskuReplenishmentError, match='缺少有效原因'):
        rep.write_replenishment_report([row],path)
    assert path.read_bytes()==b'previous'
    assert list(tmp_path.iterdir())==[path]
