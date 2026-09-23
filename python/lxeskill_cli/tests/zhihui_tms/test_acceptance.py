from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from lxeskill import cli as lxeskill_cli
from services.agent_cli.zhihui import export_products
from services.zhihui_tms.client import ZhihuiTmsClient
from services.zhihui_tms.xlsx_delivery import ZhihuiTmsDeliveryError
from shared.workspace import activate_external_workspace, activate_project_workspace


ARGUMENTS = ["--platform", "zhihui_tms", "--warehouse", "PH", "--intent", "product_export"]


class _Response:
    def __init__(self, *, payload: dict[str, Any] | None = None, content: bytes = b"") -> None:
        self.status_code = 200
        self.payload = payload
        self.content = content
        self.text = json.dumps(payload, ensure_ascii=False) if payload is not None else ""
        self.headers = {"Content-Type": "application/octet-stream"} if content else {}

    def json(self) -> dict[str, Any]:
        assert self.payload is not None
        return self.payload

    def iter_content(self, *, chunk_size: int):
        yield self.content

    def close(self) -> None:
        pass


class _Session:
    def __init__(self, content: bytes) -> None:
        self.responses = [
            _Response(payload={"code": "200", "data": {"apiToken": "fixture-token"}}),
            _Response(payload={"code": "200", "datas": [{"id": 101}], "totalNum": 1}),
            _Response(payload={"code": "200", "pop": "https://tms-cos.mabangerp.com/fixture.xls"}),
            _Response(content=content),
        ]

    def request(self, method: str, url: str, **kwargs: Any) -> _Response:
        return self.responses.pop(0)

    def close(self) -> None:
        pass


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["商品ID", "库存"])
    worksheet.append([101, 5])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_execute_catalog_path_delivers_fixture_workbook(monkeypatch, tmp_path: Path, capsys) -> None:
    monkeypatch.setenv("ZHIHUI_TMS_PRODUCTION_ENABLED", "1")
    monkeypatch.setenv("ZHIHUI_TMS_ACCOUNT", "fixture-account")
    monkeypatch.setenv("ZHIHUI_TMS_PASSWORD", "fixture-secret")
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "desktop-data"))
    monkeypatch.setattr(
        export_products,
        "ZhihuiTmsClient",
        lambda *, api_token="": ZhihuiTmsClient(
            session=_Session(_workbook_bytes()), min_request_interval_seconds=0, api_token=api_token,
        ),
    )
    activate_external_workspace(tmp_path)
    try:
        exit_code = lxeskill_cli._main(["tms", "philippines", "products-export", "execute", *ARGUMENTS])
        records = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line]
    finally:
        monkeypatch.undo()
        activate_project_workspace()
    assert exit_code == 0
    progress = [record for record in records if record.get("type") == "progress"]
    assert progress
    assert all(record["command"] == "tms philippines products-export execute" for record in progress)
    assert all(record.get("stage") and record.get("message", "").startswith("智汇 TMS：") for record in progress)
    result = records[-1]
    assert result["ok"] is True
    assert result["data"] == {
        "platform": "zhihui_tms",
        "country": "PH",
        "business_type": "product_export",
        "row_count": 1,
    }
    assert len(result["files"]) == 1
    assert "artifacts" not in result["data"]
    assert list(load_workbook(result["files"][0], read_only=True).active.values) == [("商品ID", "库存"), (101, 5)]
    assert "fixture-secret" not in json.dumps(records, ensure_ascii=False)


def test_execute_catalog_path_reports_partial_delivery_as_structured_failure(monkeypatch, tmp_path: Path, capsys) -> None:
    monkeypatch.setenv("ZHIHUI_TMS_PRODUCTION_ENABLED", "1")
    monkeypatch.setenv("ZHIHUI_TMS_ACCOUNT", "fixture-account")
    monkeypatch.setenv("ZHIHUI_TMS_PASSWORD", "fixture-secret")
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "desktop-data"))
    monkeypatch.setattr(
        export_products,
        "ZhihuiTmsClient",
        lambda *, api_token="": ZhihuiTmsClient(
            session=_Session(_workbook_bytes()), min_request_interval_seconds=0, api_token=api_token,
        ),
    )

    def partial_delivery(_client, _result, *, output_dir: Path, date_label: str, on_event=None):
        assert callable(on_event)
        output_dir.mkdir(parents=True)
        partial = output_dir / f"智汇tms-商品-部分合并-{date_label}.xlsx"
        partial.write_bytes(_workbook_bytes())
        error = ZhihuiTmsDeliveryError("tms_download_failed", "fixture download failed")
        error.partial_artifacts = (type("Artifact", (), {"path": str(partial), "kind": "merged_partial", "page": None, "total_pages": 2})(),)
        error.partial_pages = 1
        error.partial_rows = 1
        raise error

    monkeypatch.setattr(export_products, "deliver_product_exports", partial_delivery)
    activate_external_workspace(tmp_path)
    try:
        exit_code = lxeskill_cli._main(["tms", "philippines", "products-export", "execute", *ARGUMENTS])
        records = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line]
    finally:
        monkeypatch.undo()
        activate_project_workspace()

    assert exit_code == lxeskill_cli.EXIT_BUSINESS
    result = records[-1]
    assert result["ok"] is False
    assert result["data"] == {
        "platform": "zhihui_tms",
        "country": "PH",
        "business_type": "product_export",
        "partial": True,
        "partial_pages": 1,
        "partial_rows": 1,
    }
    assert result["error"] == {"code": "tms_export_partial", "message": "ZhihuiTmsDeliveryError: fixture download failed"}
    assert len(result["files"]) == 1
    assert list(load_workbook(result["files"][0], read_only=True).active.values) == [("商品ID", "库存"), (101, 5)]
