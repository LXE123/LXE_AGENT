"""The five-sheet writer uses only the current export and explicit SKU inputs."""

from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula
import pytest

from services.vietnam_replenishment import workbook as writer
from services.vietnam_replenishment.asset_contract import SkuParameters
from services.vietnam_replenishment.yacang_sources import VietnamSources
from services.yacang.validation import (
    INVENTORY_LIST_HEADERS,
    INVENTORY_SALES_HEADERS,
    WAREHOUSE_PRODUCTS_HEADERS,
)

_LOAD_REAL_SKELETON = writer._load_skeleton


INVENTORY_HEADERS = (
    "条码", "SKU", "映射条码", "仓库", "规格", "库存数量", "占用数量",
    "在途数量", "冻结库存", "可用库存", "中文标题", "英文标题", "图片链接",
)


def _skeleton() -> Workbook:
    book = Workbook()
    main = book.active
    main.title = "越南备货清单"
    for name in ("雅仓库存", "雅仓动销", "数据更改", "库存商品信息"):
        book.create_sheet(name)
    for coordinate, value in {
        "B1": "判断热销", "E1": "SKU", "F1": "名称", "G1": "成本", "H1": "陆运",
        "AA1": "上架时间", "AE1": "跨境价", "AJ1": "折扣价", "AN1": "总在途",
        "AV1": "30天", "AW1": "15天", "AX1": "7天", "AY1": "汇率",
    }.items():
        main[coordinate] = value
    main["H2"] = "=ROUNDUP(T2*AC2-J2-AN2,-1)"
    main["J2"] = "=VLOOKUP(E2,雅仓库存!B:J,9,0)"
    main["K2"] = "=VLOOKUP(E2,雅仓动销!A:G,7,0)"
    main["S2"] = "=K2*数据更改!$E$2/(30+AV2)+L2*数据更改!$F$2/(15+AW2)+M2*数据更改!$G$2/(7+AX2)"
    main["Q2"] = '=IF(K2>0,"有动销","无动销")'
    main["AA2"] = '=IFERROR(VLOOKUP(E2,库存商品信息!B:K,10,0),"未收录")'
    main["AC2"] = ArrayFormula("AC2", text="=IF(B2=1,80,40)")
    for column, source in zip("VWXY", "ABCD"):
        # The actual parameter cells occupy AV:AY. Keep row 2 in every copy.
        main[f"A{column}2"] = f"=数据更改!{source}2"
    for index, value in enumerate(INVENTORY_HEADERS, 1):
        book["雅仓库存"].cell(1, index, value)
    for index, value in enumerate(INVENTORY_SALES_HEADERS, 1):
        book["雅仓动销"].cell(1, index, value)
    for index, value in enumerate(WAREHOUSE_PRODUCTS_HEADERS, 1):
        book["库存商品信息"].cell(1, index, value)
    for index, value in enumerate(("30天", "15天", "7天", "汇率", "30天销量权重", "15天销量权重", "7天销量权重"), 1):
        book["数据更改"].cell(1, index, value)
    return book


@pytest.fixture(autouse=True)
def use_synthetic_skeleton(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(writer, "_load_skeleton", _skeleton)


def _source_row(headers: tuple[str, ...], **values: object) -> dict[str, object]:
    return {header: values.get(header) for header in headers}


def _sources() -> VietnamSources:
    sales = {
        "VN-A": _source_row(INVENTORY_SALES_HEADERS, SKU="VN-A", 商品名="动销名称 A", 仓库="VN8806", **{"7天销量": 7, "15天销量": 15, "30天销量": 30}),
        "VN-B": _source_row(INVENTORY_SALES_HEADERS, SKU="VN-B", 商品名="动销名称 B", 仓库="VN8806", **{"7天销量": 0, "15天销量": 0, "30天销量": 0}),
    }
    inventory = {
        "VN-A": _source_row(INVENTORY_LIST_HEADERS, SKU="VN-A", 仓库="VN8806", 库存数量=12, 占用数量=2, 在途数量=3, 可用库存=10, 中文标题="库存名称 A"),
        "VN-B": _source_row(INVENTORY_LIST_HEADERS, SKU="VN-B", 仓库="VN8806", 库存数量=0, 占用数量=0, 在途数量=0, 可用库存=0),
    }
    products = {
        "VN-A": _source_row(WAREHOUSE_PRODUCTS_HEADERS, SKU="VN-A", 中文标题="产品名称 A", 创建时间="2026-9-2 1:2"),
        "VN-B": _source_row(WAREHOUSE_PRODUCTS_HEADERS, SKU="VN-B", 中文标题="产品名称 B", 创建时间="2024-01-02 03:04"),
        "GLOBAL-OTHER": _source_row(WAREHOUSE_PRODUCTS_HEADERS, SKU="GLOBAL-OTHER", 创建时间="2026-09-23 11:00"),
    }
    return VietnamSources(
        skus=("VN-A", "VN-B"), sales=sales, inventory=inventory, products=products,
        missing_sales=(), missing_inventory=(), missing_products=(),
        in_transit={"VN-A": Decimal("3"), "VN-B": Decimal("0")},
        missing_in_transit=(), in_transit_mismatch=(), artifacts={},
    )


def _parameters() -> dict[str, SkuParameters]:
    return {
        "VN-A": SkuParameters(cost=Decimal("0"), cross_border_price=Decimal("120.25"), discount_price=Decimal("99.5"), hot_flag=1),
        "VN-B": SkuParameters(cost=Decimal("3.5"), cross_border_price=Decimal("50"), discount_price=Decimal("40")),
        "GLOBAL-OTHER": SkuParameters(cost=Decimal("999")),
    }


def test_writer_projects_only_current_skus_and_translates_formulas(tmp_path: Path) -> None:
    output = tmp_path / "five-sheets.xlsx"
    writer.write_vietnam_workbook(output, _sources(), _parameters(), writer.RecommendationConfig())

    book = load_workbook(output, data_only=False)
    try:
        assert book.sheetnames == ["越南备货清单", "雅仓库存", "雅仓动销", "数据更改", "库存商品信息"]
        main = book["越南备货清单"]
        assert main.max_row == 3
        assert [main[f"E{row}"].value for row in (2, 3)] == ["VN-A", "VN-B"]
        assert main["F2"].value == "产品名称 A"
        assert main["B2"].value == 1
        assert main["B3"].value == 2
        assert main["G2"].value == 0
        assert main["AE2"].value == 120.25
        assert main["AJ2"].value == 99.5
        assert main["AN2"].value == 3
        assert main["AN3"].value == 0
        assert main["H3"].value == "=ROUNDUP(T3*AC3-J3-AN3,-1)"
        assert main["Q3"].value == '=IF(K3>0,"有动销","无动销")'
        assert main["S3"].value == "=K3*数据更改!$E$2/(30+AV3)+L3*数据更改!$F$2/(15+AW3)+M3*数据更改!$G$2/(7+AX3)"
        assert main["AA3"].value == '=IFERROR(VLOOKUP(E3,库存商品信息!B:K,10,0),"未收录")'
        assert isinstance(main["AC3"].value, ArrayFormula)
        assert main["AC3"].value.ref == "AC3"
        assert main["AC3"].value.text == "=IF(B3=1,80,40)"
        assert [main[f"{column}3"].value for column in ("AV", "AW", "AX", "AY")] == [
            "=数据更改!A2", "=数据更改!B2", "=数据更改!C2", "=数据更改!D2",
        ]
        for column in ("A", "C", "D", "AD", "AO", "AP", "AQ", "AR", "AS", "AT", "AU"):
            assert main[f"{column}2"].value is None
            assert main[f"{column}3"].value is None

        inventory = book["雅仓库存"]
        assert [inventory.cell(1, i).value for i in range(1, 14)] == list(INVENTORY_HEADERS)
        assert inventory["F2"].value == 12
        assert inventory["H2"].value == 3
        assert inventory["J2"].value == 10
        assert inventory["K2"].value == "库存名称 A"
        sales = book["雅仓动销"]
        assert sales["A2"].value == "VN-A"
        assert sales["E2"].value == 7
        products = book["库存商品信息"]
        assert [products[f"B{row}"].value for row in (2, 3)] == ["VN-A", "VN-B"]
        assert products["K2"].value == "2026-09-02 01:02"
        assert products["K2"].data_type == "s"
        assert book["数据更改"]["A2"].value == 0.8
        assert book["数据更改"]["D2"].value == 3900
    finally:
        book.close()


def test_config_overrides_seven_inputs(tmp_path: Path) -> None:
    config = writer.RecommendationConfig(
        day_adjustment_30d=Decimal("1.5"), day_adjustment_15d=Decimal("2.5"),
        day_adjustment_7d=Decimal("3.5"), exchange_rate=Decimal("4000"),
        sales_weight_30d=Decimal("0.2"), sales_weight_15d=Decimal("0.5"),
        sales_weight_7d=Decimal("0.3"),
    )
    output = tmp_path / "configured.xlsx"
    writer.write_vietnam_workbook(output, _sources(), _parameters(), config)
    book = load_workbook(output, read_only=True, data_only=False)
    try:
        assert [book["数据更改"].cell(2, i).value for i in range(1, 8)] == [1.5, 2.5, 3.5, 4000, 0.2, 0.5, 0.3]
        assert all(book["数据更改"].cell(2, i).data_type == "n" for i in range(1, 8))
    finally:
        book.close()


@pytest.mark.parametrize(
    ("field", "value", "error"),
    [
        ("day_adjustment_30d", Decimal("-1"), "有限非负数"),
        ("day_adjustment_15d", Decimal("NaN"), "有限非负数"),
        ("day_adjustment_7d", True, "布尔值"),
        ("exchange_rate", Decimal("0"), "有限正数"),
        ("exchange_rate", Decimal("Infinity"), "有限正数"),
        ("sales_weight_30d", Decimal("-0.1"), "有限非负数"),
        ("sales_weight_15d", Decimal("Infinity"), "有限非负数"),
        ("sales_weight_7d", Decimal("0.1234567890123456"), "Excel 精度"),
        ("sales_weight_7d", Decimal("1E-1000"), "Excel"),
        ("sales_weight_7d", Decimal("0.5"), "合计.*1"),
    ],
)
def test_invalid_configuration_stops_without_file(tmp_path: Path, field: str, value: object, error: str) -> None:
    output = tmp_path / "invalid-config.xlsx"
    config = replace(writer.RecommendationConfig(), **{field: value})
    with pytest.raises(writer.WorkbookInputError, match=error):
        writer.write_vietnam_workbook(output, _sources(), _parameters(), config)
    assert not output.exists()


def test_sales_weight_total_cannot_round_to_one() -> None:
    config = writer.RecommendationConfig(
        sales_weight_30d=Decimal("1"), sales_weight_15d=Decimal("1E-100"),
        sales_weight_7d=Decimal("0"),
    )
    with pytest.raises(writer.WorkbookInputError, match="合计.*1"):
        writer.validate_recommendation_config(config)


def test_configuration_accepts_decimal_boundary_and_signed_zero() -> None:
    config = writer.RecommendationConfig(
        day_adjustment_30d=Decimal("-0"), sales_weight_30d=Decimal("0.123456789012345"),
        sales_weight_15d=Decimal("0.376543210987655"), sales_weight_7d=Decimal("0.5"),
    )
    assert writer.validate_recommendation_config(config) == (
        Decimal("-0"), Decimal("0.8"), Decimal("0"), Decimal("3900"),
        Decimal("0.123456789012345"), Decimal("0.376543210987655"), Decimal("0.5"),
    )


@pytest.mark.parametrize(
    ("change", "expected"),
    [
        (lambda s: replace(s, sales={"VN-B": s.sales["VN-B"]}), "VN-A.*动销"),
        (lambda s: replace(s, inventory={"VN-B": s.inventory["VN-B"]}), "VN-A.*库存"),
        (lambda s: replace(s, products={"VN-B": s.products["VN-B"]}), "VN-A.*产品"),
        (lambda s: replace(s, in_transit={"VN-A": None, "VN-B": Decimal("0")}), "VN-A.*在途数量"),
        (lambda s: replace(s, products={**s.products, "VN-A": {**s.products["VN-A"], "创建时间": None}}), "VN-A.*创建时间"),
        (lambda s: replace(s, products={**s.products, "VN-A": {**s.products["VN-A"], "中文标题": None}}), "VN-A.*中文标题"),
        (lambda s: replace(s, sales={**s.sales, "VN-A": {**s.sales["VN-A"], "仓库": "OTHER"}}), "VN-A.*动销.*VN8806"),
        (lambda s: replace(s, sales={**s.sales, "VN-A": {**s.sales["VN-A"], "30天销量": None}}), "VN-A.*30天销量"),
        (lambda s: replace(s, inventory={**s.inventory, "VN-A": {**s.inventory["VN-A"], "可用库存": None}}), "VN-A.*可用库存"),
    ],
)
def test_missing_current_source_stops_without_file(tmp_path: Path, change, expected: str) -> None:
    output = tmp_path / "absent.xlsx"
    with pytest.raises(writer.WorkbookInputError, match=expected):
        writer.write_vietnam_workbook(output, change(_sources()), _parameters(), writer.RecommendationConfig())
    assert not output.exists()


@pytest.mark.parametrize("field", ("cost", "cross_border_price", "discount_price"))
def test_missing_explicit_price_stops_without_file(tmp_path: Path, field: str) -> None:
    values = _parameters()
    values["VN-A"] = replace(values["VN-A"], **{field: None})
    output = tmp_path / "absent.xlsx"
    with pytest.raises(writer.WorkbookInputError, match=f"VN-A.*{field}"):
        writer.write_vietnam_workbook(output, _sources(), values, writer.RecommendationConfig())
    assert not output.exists()


def test_missing_parameter_row_and_bad_hot_flag_fail(tmp_path: Path) -> None:
    values = _parameters()
    del values["VN-A"]
    with pytest.raises(writer.WorkbookInputError, match="VN-A.*映射"):
        writer.write_vietnam_workbook(tmp_path / "missing.xlsx", _sources(), values, writer.RecommendationConfig())
    values["VN-A"] = replace(_parameters()["VN-A"], hot_flag=3)
    with pytest.raises(writer.WorkbookInputError, match="VN-A.*热销标记"):
        writer.write_vietnam_workbook(tmp_path / "bad-hot.xlsx", _sources(), values, writer.RecommendationConfig())


def test_old_map_listing_time_conflict_fails_and_date_only_match_passes(tmp_path: Path) -> None:
    values = _parameters()
    values["VN-A"] = replace(values["VN-A"], listed_at="2026-09-03")
    with pytest.raises(writer.WorkbookInputError, match="VN-A.*上架时间.*创建时间"):
        writer.write_vietnam_workbook(tmp_path / "conflict.xlsx", _sources(), values, writer.RecommendationConfig())
    values["VN-A"] = replace(values["VN-A"], listed_at="2026-09-02")
    writer.write_vietnam_workbook(tmp_path / "same-date.xlsx", _sources(), values, writer.RecommendationConfig())


def test_price_over_excel_precision_limit_is_rejected(tmp_path: Path) -> None:
    values = _parameters()
    values["VN-A"] = replace(values["VN-A"], cost=Decimal("1.234567890123456"))
    with pytest.raises(writer.WorkbookInputError, match="VN-A.*成本.*Excel 精度"):
        writer.write_vietnam_workbook(tmp_path / "imprecise.xlsx", _sources(), values, writer.RecommendationConfig())


@pytest.mark.parametrize("cost", (Decimal("1E1000"), Decimal("1E-1000")))
def test_price_outside_excel_numeric_range_is_rejected(tmp_path: Path, cost: Decimal) -> None:
    values = _parameters()
    values["VN-A"] = replace(values["VN-A"], cost=cost)
    with pytest.raises(writer.WorkbookInputError, match="VN-A.*成本.*Excel"):
        writer.write_vietnam_workbook(tmp_path / "out-of-range.xlsx", _sources(), values, writer.RecommendationConfig())


def test_price_at_excel_precision_limit_round_trips(tmp_path: Path) -> None:
    values = _parameters()
    values["VN-A"] = replace(values["VN-A"], cost=Decimal("1.23456789012345"))
    output = tmp_path / "precise.xlsx"
    writer.write_vietnam_workbook(output, _sources(), values, writer.RecommendationConfig())
    book = load_workbook(output, read_only=True, data_only=False)
    try:
        assert Decimal(str(book["越南备货清单"]["G2"].value)) == Decimal("1.23456789012345")
        assert book["越南备货清单"]["G2"].data_type == "n"
    finally:
        book.close()


def test_packaged_skeleton_supports_real_formula_set(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(writer, "_load_skeleton", _LOAD_REAL_SKELETON)
    output = tmp_path / "packaged.xlsx"
    writer.write_vietnam_workbook(output, _sources(), _parameters(), writer.RecommendationConfig())
    book = load_workbook(output, data_only=False)
    try:
        main = book["越南备货清单"]
        assert main["H3"].data_type == "f"
        assert "T3*AC3-J3-AN3" in main["H3"].value
        assert "DATEVALUE(LEFT(AA3,10))" in main["U3"].value
        assert "E3" in main["AA3"].value
        assert isinstance(main["AC3"].value, ArrayFormula)
        assert main["AC3"].value.ref == "AC3"
        assert "B3=1" in main["AC3"].value.text
        assert main["AV3"].value == "=数据更改!A2"
        assert main["AY3"].value == "=数据更改!D2"
        assert main["S3"].value == "=K3*数据更改!$E$2/(30+AV3)+L3*数据更改!$F$2/(15+AW3)+M3*数据更改!$G$2/(7+AX3)"
    finally:
        book.close()


def test_existing_output_is_never_overwritten(tmp_path: Path) -> None:
    output = tmp_path / "five-sheets.xlsx"
    output.write_bytes(b"untouched")
    with pytest.raises(FileExistsError):
        writer.write_vietnam_workbook(output, _sources(), _parameters(), writer.RecommendationConfig())
    assert output.read_bytes() == b"untouched"
