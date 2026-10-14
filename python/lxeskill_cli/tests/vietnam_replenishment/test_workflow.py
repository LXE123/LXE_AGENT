"""The Vietnam workflow requires a nine-sheet template for every run."""

from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from importlib.resources import files
from pathlib import Path

from openpyxl import load_workbook
import pytest

from services.vietnam_replenishment import workflow
from services.vietnam_replenishment.full_workbook import REQUIRED_ORDER, UNCHANGED
from services.vietnam_replenishment.workbook import RecommendationConfig
from services.vietnam_replenishment.yacang_sources import VietnamSourceError, VietnamSources
from services.vietnam_replenishment import settings as store
from shared import workspace


@pytest.fixture(autouse=True)
def isolated_state(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path / "workspace"))
    original = {
        name: getattr(workspace, name)
        for name in ("_workspace_root", "_internal_root", "_artifact_root", "_input_root")
    }
    workspace.activate_project_workspace()
    yield
    for name, value in original.items():
        setattr(workspace, name, value)


def _template(path: Path) -> Path:
    resource = files("services.vietnam_replenishment").joinpath("resources/skeleton.xlsx")
    with resource.open("rb") as handle:
        book = load_workbook(handle)
    try:
        main = book["越南备货清单"]
        for cell, value in {"B2": 2, "E2": "VN-A", "G2": 11, "AE2": 33, "AJ2": 25}.items():
            main[cell] = value
        for index, name in enumerate(UNCHANGED, 1):
            book.create_sheet(name, index)["A1"] = name
        book.save(path)
    finally:
        book.close()
    return path


def _sources() -> VietnamSources:
    sku = "VN-A"
    return VietnamSources(
        skus=(sku,),
        sales={sku: {
            "SKU": sku, "仓库": "VN8806", "7天销量": 0,
            "15天销量": 0, "30天销量": 0, "在途": 99,
        }},
        inventory={sku: {
            "SKU": sku, "仓库": "VN8806", "库存数量": 7,
            "占用数量": 2, "在途数量": 3, "可用库存": 5,
        }},
        products={sku: {
            "SKU": sku, "中文标题": "合成产品 A", "创建时间": "2026-09-23 10:00",
        }},
        missing_sales=(), missing_inventory=(), missing_products=(),
        in_transit={sku: Decimal("3")}, missing_in_transit=(),
        in_transit_mismatch=(sku,), artifacts={},
    )


def test_missing_template_fails_before_loading_reports(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    (store.data_directory()).mkdir(parents=True)
    (store.data_directory() / "sku-map.xlsx").write_bytes(b"obsolete")
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: pytest.fail("Reports must not load"))
    with pytest.raises(workflow.VietnamWorkflowError, match="九表模板") as caught:
        workflow.generate_current_vietnam_recommendation(
            sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx",
        )
    assert caught.value.code == "template_missing"


def test_template_fields_reach_generator_without_old_map(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    template = _template(tmp_path / "template.xlsx")
    (store.data_directory()).mkdir(parents=True)
    (store.data_directory() / "sku-map.xlsx").write_bytes(b"obsolete")
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: _sources())
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: tmp_path / "artifacts" / parts[-1])

    def fake_generate(output: Path, *, sources, config, template, parameters) -> Path:
        assert template.read_bytes() == (tmp_path / "template.xlsx").read_bytes()
        assert sources.skus == ("VN-A",)
        assert parameters["VN-A"].hot_flag == 2
        assert parameters["VN-A"].cost == Decimal(11)
        assert parameters["VN-A"].cross_border_price == Decimal(33)
        assert parameters["VN-A"].discount_price == Decimal(25)
        assert config == RecommendationConfig()
        output.parent.mkdir(parents=True)
        output.write_bytes(b"synthetic final workbook")
        return output

    monkeypatch.setattr(workflow, "generate_vietnam_workbook", fake_generate)
    result = workflow.generate_current_vietnam_recommendation(
        sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx", template=str(template),
    )
    assert result.template_source == str(template)
    assert result.validation["sheets"] == list(REQUIRED_ORDER)
    assert result.validation["missing_mapping_count"] == 0
    assert result.output_xlsx.read_bytes() == b"synthetic final workbook"


def test_nondefault_config_reaches_generator(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    template = _template(tmp_path / "template.xlsx")
    config = replace(RecommendationConfig(), exchange_rate=Decimal(4200))
    store.save_parameters(store.config_json(config))
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: _sources())
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: tmp_path / "artifacts" / parts[-1])

    def fake_generate(output: Path, *, sources, config: RecommendationConfig, template, parameters) -> Path:
        assert config.exchange_rate == Decimal(4200)
        output.parent.mkdir(parents=True)
        output.write_bytes(b"synthetic final workbook")
        return output

    monkeypatch.setattr(workflow, "generate_vietnam_workbook", fake_generate)
    result = workflow.generate_current_vietnam_recommendation(
        sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx", template=str(template),
    )
    assert result.config_source == str(store.data_directory() / "parameters.json")
    assert result.config.exchange_rate == Decimal(4200)


def test_source_error_does_not_start_generation(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    template = _template(tmp_path / "template.xlsx")
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: (_ for _ in ()).throw(VietnamSourceError("库存列表失败")))
    monkeypatch.setattr(workflow, "generate_vietnam_workbook", lambda *a, **kw: pytest.fail("Generator called"))
    with pytest.raises(VietnamSourceError, match="库存列表失败"):
        workflow.generate_current_vietnam_recommendation(
            sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx", template=str(template),
        )


def test_empty_source_sku_set_does_not_start_generation(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    template = _template(tmp_path / "template.xlsx")
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: replace(_sources(), skus=()))
    monkeypatch.setattr(workflow, "generate_vietnam_workbook", lambda *a, **kw: pytest.fail("Generator called"))
    with pytest.raises(workflow.VietnamWorkflowError, match="没有 SKU") as caught:
        workflow.generate_current_vietnam_recommendation(
            sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx", template=str(template),
        )
    assert caught.value.code == "current_skus_empty"


def test_failed_generation_removes_partial_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    template = _template(tmp_path / "template.xlsx")
    output_root = tmp_path / "artifacts"
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: output_root / parts[-1])
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: _sources())

    def fail_generate(output: Path, **_kwargs) -> Path:
        output.parent.mkdir(parents=True)
        output.write_bytes(b"incomplete")
        raise RuntimeError("Office actual failure")

    monkeypatch.setattr(workflow, "generate_vietnam_workbook", fail_generate)
    with pytest.raises(RuntimeError, match="Office actual failure"):
        workflow.generate_current_vietnam_recommendation(
            sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx", template=str(template),
        )
    assert not list(output_root.rglob("*.xlsx"))


def test_generator_without_final_file_is_not_success(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    template = _template(tmp_path / "template.xlsx")
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: tmp_path / "artifacts" / parts[-1])
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: _sources())
    monkeypatch.setattr(workflow, "generate_vietnam_workbook", lambda output, **_kw: output)
    with pytest.raises(workflow.VietnamWorkflowError, match="没有生成最终 XLSX") as caught:
        workflow.generate_current_vietnam_recommendation(
            sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx", template=str(template),
        )
    assert caught.value.code == "output_missing"
