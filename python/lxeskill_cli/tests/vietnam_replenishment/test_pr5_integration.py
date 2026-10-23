"""Synthetic local report files through the CLI and real Office Kit."""

from __future__ import annotations

from dataclasses import replace
from contextlib import closing
from decimal import Decimal
from importlib.resources import files
from itertools import product
import json
import os
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900
import pytest

from lxeskill import cli as lxeskill
from services.vietnam_replenishment import settings
from services.vietnam_replenishment.workbook import RecommendationConfig
from services.vietnam_replenishment.full_workbook import REQUIRED_ORDER, UNCHANGED, _sheet_parts
from services.yacang.validation import INVENTORY_SALES_HEADERS, INVENTORY_LIST_HEADERS, WAREHOUSE_PRODUCTS_HEADERS
from services.vietnam_replenishment.yacang_sources import VietnamSources
from shared import workspace


def _sources() -> VietnamSources:
    """Use the same single-VN8806 source shape as the workflow's Office test."""
    sku = "VN-A"
    return VietnamSources(
        skus=(sku,),
        sales={sku: {
            "SKU": sku, "仓库": "VN8806", "7天销量": 0,
            "15天销量": 0, "30天销量": 0, "在途": 99,
        }},
        inventory={sku: {
            "SKU": sku, "仓库": "VN8806", "库存数量": 7,
            "占用数量": 2, "在途数量": 3, "可用库存": 5,
        }},
        products={sku: {
            "SKU": sku, "中文标题": "合成产品 A", "创建时间": "2026-09-23 10:00",
        }},
        missing_sales=(), missing_inventory=(), missing_products=(),
        in_transit={sku: Decimal("3")}, missing_in_transit=(),
        in_transit_mismatch=(sku,), artifacts={},
    )


def _three_sources() -> VietnamSources:
    base = _sources()
    sales = dict(base.sales)
    inventory = dict(base.inventory)
    products = dict(base.products)
    in_transit = dict(base.in_transit)
    for sku in ("VN-B", "VN-C"):
        sales[sku] = {
            **sales["VN-A"], "SKU": sku,
            "7天销量": 7, "15天销量": 15, "30天销量": 30,
        }
        inventory[sku] = {
            **inventory["VN-A"], "SKU": sku,
            "库存数量": 12, "占用数量": 2,
            "在途数量": 2, "可用库存": 10,
        }
        products[sku] = {
            **products["VN-A"], "SKU": sku,
            "中文标题": f"合成产品 {sku}", "创建时间": "2026-09-20 10:00",
        }
        in_transit[sku] = Decimal("2")
    return replace(
        base, skus=("VN-A", "VN-B", "VN-C"),
        sales=sales, inventory=inventory, products=products,
        in_transit=in_transit, in_transit_mismatch=(),
    )


def _template(path: Path, *, rows: list[tuple] | None = None) -> Path:
    resource = files("services.vietnam_replenishment").joinpath("resources/skeleton.xlsx")
    with resource.open("rb") as handle:
        book = load_workbook(handle)
    try:
        main = book["越南备货清单"]
        rows = rows if rows is not None else [("VN-A", 2, 12, 22, 16)]
        for row_number, (sku, hot, cost, cross, discount) in enumerate(rows, 2):
            for column, value in (("B", hot), ("E", sku), ("G", cost), ("AE", cross), ("AJ", discount)):
                main[f"{column}{row_number}"] = value
        for index, name in enumerate(UNCHANGED, 1):
            sheet = book.create_sheet(name, index)
            sheet["A1"] = f"保留 {name}"
            sheet["B2"] = "=越南备货清单!H2"
        book.save(path)
    finally:
        book.close()
    return path




def _report_arguments(directory: Path, sources: VietnamSources) -> list[str]:
    args = []
    for flag, headers, rows in (
        ("sales", INVENTORY_SALES_HEADERS, sources.sales),
        ("inventory", INVENTORY_LIST_HEADERS, sources.inventory),
        ("products", WAREHOUSE_PRODUCTS_HEADERS, sources.products),
    ):
        path = directory / f"{flag}.xlsx"
        book = Workbook()
        book.active.append(headers)
        for row in rows.values():
            book.active.append([row.get(header) for header in headers])
        book.save(path)
        book.close()
        args.extend((f"--{flag}-file", str(path)))
    return args


@pytest.fixture()
def isolated_state(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    state = tmp_path / "state"
    original = {
        name: getattr(workspace, name)
        for name in ("_workspace_root", "_internal_root", "_artifact_root", "_input_root")
    }
    monkeypatch.setenv("LXE_DATA_ROOT", str(state))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path / "workspace"))
    workspace.activate_project_workspace()
    yield state
    for name, value in original.items():
        setattr(workspace, name, value)


@pytest.mark.skipif(
    not (os.environ.get("LXE_OFFICE_NODE") and os.environ.get("LXE_OFFICE_CLI")),
    reason="Host Office Kit paths are not configured",
)
def test_upload_replace_then_generate_one_final_workbook(
    tmp_path: Path, isolated_state: Path,
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.delenv("LXESKILL_SKILL_SCOPE", raising=False)
    settings.save_parameters(settings.config_json(RecommendationConfig(Decimal("0.7"), Decimal("0.6"), Decimal("0.1"), Decimal("4000"))))
    report_args = _report_arguments(tmp_path, _sources())
    template = _template(tmp_path / "完整模板.xlsx")
    assert lxeskill.main(["vietnam", "replenishment", "calculate", *report_args, "--template-file", str(template)]) == 0
    records = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line.strip()]
    assert len(records) == 1
    result = records[0]
    assert result["type"] == "result" and result["ok"] is True
    assert result["data"]["config"] == {
        "day_adjustment_30d": "0.7", "day_adjustment_15d": "0.6",
        "day_adjustment_7d": "0.1", "exchange_rate": "4000",
        "sales_weight_7d": "0.6", "sales_weight_15d": "0.3", "sales_weight_30d": "0.1",
    }
    assert result["data"]["config_source"] == str(settings.data_directory() / "parameters.json")
    assert result["data"]["sku_count"] == 1
    output = Path(result["data"]["output_xlsx"])
    assert output.is_relative_to(tmp_path / "workspace" / ".lxeagent" / "artifacts")
    assert result["files"] == [str(output)]
    assert result["data"]["validation"]["status"] == "passed"
    assert len(result["data"]["source_files"]) == 3
    assert all(len(source["sha256"]) == 64 for source in result["data"]["source_files"])
    assert output.is_file() and output.stat().st_size > 0

    book = load_workbook(output, read_only=True, data_only=True)
    try:
        assert book.epoch == CALENDAR_WINDOWS_1900
        assert tuple(book.sheetnames) == REQUIRED_ORDER
        main = book["越南备货清单"]
        assert main["E2"].value == "VN-A"
        assert [main[cell].value for cell in ("G2", "AE2", "AJ2")] == [12, 22, 16]
        assert [book["数据更改"].cell(2, column).value for column in range(1, 5)] == [
            0.7, 0.6, 0.1, 4000,
        ]
        assert [main.cell(2, column).value for column in range(48, 52)] == [
            0.7, 0.6, 0.1, 4000,
        ]
    finally:
        book.close()
    with ZipFile(template) as original, ZipFile(output) as completed:
        names = _sheet_parts(original)
        assert all(original.read(names[name]) == completed.read(names[name]) for name in UNCHANGED)


@pytest.mark.skipif(
    not (os.environ.get("LXE_OFFICE_NODE") and os.environ.get("LXE_OFFICE_CLI")),
    reason="Host Office Kit paths are not configured",
)
@pytest.mark.parametrize("case", ["chat", "legacy_current", "after_office"])
def test_date_system_failures_never_deliver_a_successful_cli_file(
    tmp_path, isolated_state, monkeypatch, capsys, case,
):
    from services.vietnam_replenishment import recalculation
    from shared.input_assets import current_asset, promote_asset

    monkeypatch.delenv("LXESKILL_SKILL_SCOPE", raising=False)
    template = _template(tmp_path / "template.xlsx")
    settings.upload_template(template)
    arguments = ["vietnam", "replenishment", "calculate", *_report_arguments(tmp_path, _sources())]
    seen = []
    if case in ("chat", "legacy_current"):
        unsupported = tmp_path / "1904.xlsx"
        book = load_workbook(template)
        book.epoch = CALENDAR_MAC_1904
        book.save(unsupported)
        book.close()
        if case == "chat":
            arguments += ["--template-file", str(unsupported)]
        else:
            promote_asset("vietnam_full_workbook", unsupported)
    else:
        merge = recalculation.combine_full_workbook
        def mismatched_office_result(original, calculated, output):
            # Inject only the epoch flag after real Office calculation and its
            # numerical validation; the final merge must catch the discrepancy.
            staged = calculated.with_suffix(".tmp")
            with ZipFile(calculated) as source, ZipFile(staged, "w") as target:
                for item in source.infolist():
                    data = source.read(item.filename)
                    if item.filename == "xl/workbook.xml":
                        root = ET.fromstring(data)
                        root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}workbookPr").set("date1904", "1")
                        data = ET.tostring(root)
                    target.writestr(item, data)
            staged.replace(calculated)
            book = load_workbook(calculated, read_only=True)
            assert book.epoch == CALENDAR_MAC_1904
            book.close()
            seen.append("real Office result with 1904 flag")
            return merge(original, calculated, output)
        monkeypatch.setattr(recalculation, "combine_full_workbook", mismatched_office_result)
    saved = current_asset("vietnam_full_workbook")
    before = saved.path.read_bytes()
    assert lxeskill.main(arguments) == lxeskill.EXIT_BUSINESS
    records = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line.strip()]
    assert len(records) == 1
    result = records[0]
    assert result["ok"] is False and result["files"] == []
    assert result["data"]["status"] == "failed" and "validation" not in result["data"]
    assert "日期系统" in result["error"]["message"]
    assert not list((tmp_path / "workspace" / ".lxeagent" / "artifacts").rglob("*.xlsx"))
    assert saved.path.read_bytes() == before
    assert seen == (["real Office result with 1904 flag"] if case == "after_office" else [])


@pytest.mark.skipif(
    not (os.environ.get("LXE_OFFICE_NODE") and os.environ.get("LXE_OFFICE_CLI")),
    reason="Host Office Kit paths are not configured",
)
def test_sparse_template_keeps_all_yacang_skus(
    tmp_path: Path, isolated_state: Path,
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.delenv("LXESKILL_SKILL_SCOPE", raising=False)
    settings.save_parameters(settings.config_json(RecommendationConfig(Decimal("0.7"), Decimal("0.6"), Decimal("0.1"), Decimal("4000"))))
    report_args = _report_arguments(tmp_path, _three_sources())
    template = _template(tmp_path / "完整模板.xlsx", rows=[
        ("VN-A", 2, 10, 20, 15), ("VN-B", 2, 5, 25, 18),
    ])
    assert lxeskill.main(["vietnam", "replenishment", "calculate", *report_args, "--template-file", str(template)]) == 0
    records = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line.strip()]
    assert len(records) == 1
    result = records[0]
    assert result["type"] == "result" and result["ok"] is True
    assert result["data"]["validation"]["missing_mapping_count"] == 1
    assert result["data"]["sku_count"] == 3
    output = Path(result["data"]["output_xlsx"])
    assert output.is_relative_to(tmp_path / "workspace" / ".lxeagent" / "artifacts")
    assert result["files"] == [str(output)]
    assert result["data"]["validation"]["status"] == "passed"
    assert len(result["data"]["source_files"]) == 3
    assert all(len(source["sha256"]) == 64 for source in result["data"]["source_files"])
    assert output.is_file() and output.stat().st_size > 0
    assert not list(output.parent.glob(".vietnam-workbook-*"))

    book = load_workbook(output, data_only=True)
    try:
        assert tuple(book.sheetnames) == REQUIRED_ORDER
        main = book["越南备货清单"]
        assert [main[f"E{row}"].value for row in (2, 3, 4)] == ["VN-A", "VN-B", "VN-C"]
        for name, column in (("雅仓库存", "B"), ("雅仓动销", "A"), ("库存商品信息", "B")):
            assert [book[name][f"{column}{row}"].value for row in (2, 3, 4)] == ["VN-A", "VN-B", "VN-C"]

        assert main["B3"].value == 2
        assert main["AJ3"].value == 18
        assert all(main[f"{column}3"].value is not None for column in ("H", "AC", "AF", "AG", "AH", "AI"))
        assert all(main[f"{column}3"].value is not None for column in ("AK", "AL", "AM"))

        assert all(main[f"{column}4"].value in (None, "") for column in (
            "B", "G", "AE", "AJ", "H", "AC", "AF", "AG", "AH", "AI", "AK", "AL", "AM",
        ))
        assert main["F4"].value == "合成产品 VN-C"
        assert main["J4"].value == 10
        assert main["K4"].value == 30
        assert main["AN4"].value == 2
        assert main["AA4"].value == "2026-09-20 10:00"
    finally:
        book.close()


@pytest.mark.skipif(
    not (os.environ.get("LXE_OFFICE_NODE") and os.environ.get("LXE_OFFICE_CLI")),
    reason="Host Office Kit paths are not configured",
)
def test_template_with_no_matching_sku_keeps_erp_rows_and_blank_mapping_results(
    tmp_path: Path, isolated_state: Path,
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.delenv("LXESKILL_SKILL_SCOPE", raising=False)
    base = _sources()
    sku = "VN-NEW"
    current = replace(
        base, skus=(sku,),
        sales={sku: {**base.sales["VN-A"], "SKU": sku}},
        inventory={sku: {**base.inventory["VN-A"], "SKU": sku}},
        products={sku: {**base.products["VN-A"], "SKU": sku}},
        in_transit={sku: Decimal("3")}, in_transit_mismatch=(sku,),
    )
    template = _template(tmp_path / "九表模板.xlsx")
    code = lxeskill.main([
        "vietnam", "replenishment", "calculate", *_report_arguments(tmp_path, current),
        "--template-file", str(template),
    ])
    result = json.loads(capsys.readouterr().out.splitlines()[-1])
    assert code == 0 and result["ok"], result
    assert result["data"]["validation"]["missing_mapping_count"] == 1
    output = Path(result["data"]["output_xlsx"])
    book = load_workbook(output, data_only=True)
    try:
        assert tuple(book.sheetnames) == REQUIRED_ORDER
        main = book["越南备货清单"]
        assert main["E2"].value == sku
        assert main["F2"].value == "合成产品 A"
        assert main["J2"].value == 5
        assert all(main[f"{column}2"].value in (None, "") for column in
                   ("B", "G", "AE", "AJ", "H", "AC", "AF", "AG", "AH", "AI", "AK", "AL", "AM"))
    finally:
        book.close()
    with ZipFile(template) as original, ZipFile(output) as complete:
        parts = _sheet_parts(original)
        assert all(original.read(parts[name]) == complete.read(parts[name]) for name in UNCHANGED)


@pytest.mark.skipif(
    not (os.environ.get("LXE_OFFICE_NODE") and os.environ.get("LXE_OFFICE_CLI")),
    reason="Host Office Kit paths are not configured",
)
def test_sales_delivery_keeps_recalculated_results_and_formulas(
    tmp_path, isolated_state, monkeypatch, capsys,
):
    from services.yacang.delivery import remove_inventory_sales_creation_date

    monkeypatch.delenv("LXESKILL_SKILL_SCOPE", raising=False)
    sources = _three_sources()
    sources = replace(sources, sales={
        sku: {**row, "创建日期": "2026-10-01"} for sku, row in sources.sales.items()
    })
    report_args = _report_arguments(tmp_path, sources)
    template = _template(tmp_path / "完整模板.xlsx", rows=[
        ("VN-A", 2, 10, 20, 15), ("VN-B", 2, 5, 25, 18),
    ])
    settings.upload_template(template)
    original_sales = Path(report_args[1])
    original_bytes = original_sales.read_bytes()
    delivered = tmp_path / "库存动销 交付.xlsx"
    delivered.write_bytes(original_bytes)
    remove_inventory_sales_creation_date(delivered)

    outputs = []
    for index, sales in enumerate((original_sales, delivered)):
        report_args[1] = str(sales)
        code = lxeskill.main([
            "vietnam", "replenishment", "calculate", *report_args,
            *(["--template-file", str(template)] if index == 0 else []),
        ])
        result = json.loads(capsys.readouterr().out.splitlines()[-1])
        assert code == 0 and result["ok"], result
        assert result["data"]["validation"]["status"] == "passed"
        assert result["data"]["sku_count"] == 3
        outputs.append(Path(result["data"]["output_xlsx"]))

    def cell_values(sheet):
        # openpyxl represents array formulas as objects, so compare their
        # expressions and ranges rather than identities from separate loads.
        return [tuple((value.text, value.ref) if isinstance(value, ArrayFormula) else value for value in row)
                for row in sheet.values]

    for data_only in (True, False):
        with closing(load_workbook(outputs[0], data_only=data_only)) as before, closing(
            load_workbook(outputs[1], data_only=data_only)
        ) as after:
            assert after.sheetnames == before.sheetnames
            assert tuple(after.sheetnames) == REQUIRED_ORDER
            assert cell_values(after["越南备货清单"]) == cell_values(before["越南备货清单"])
            assert list(after["库存商品信息"].values) == list(before["库存商品信息"].values)
            assert before["雅仓动销"]["P2"].value == "2026-10-01"
            assert after["雅仓动销"]["P2"].value is None
            if data_only:
                assert after["越南备货清单"]["H3"].value is not None
                assert after["越南备货清单"]["AA3"].value == "2026-09-20 10:00"
            else:
                assert after["越南备货清单"]["H3"].data_type == "f"
    assert original_sales.read_bytes() == original_bytes


@pytest.mark.skipif(
    not (os.environ.get("LXE_OFFICE_NODE") and os.environ.get("LXE_OFFICE_CLI")),
    reason="Host Office Kit paths are not configured",
)
def test_default_formula_parity_and_configurable_weights_and_exchange(
    tmp_path, isolated_state, monkeypatch,
):
    from services.vietnam_replenishment import workbook, recalculation
    from services.vietnam_replenishment.asset_contract import load_template_sku_parameters
    source = _sources()
    source = replace(source, sales={"VN-A": {**source.sales["VN-A"], "7天销量": 70, "15天销量": 90, "30天销量": 120}},
        products={"VN-A": {**source.products["VN-A"], "创建时间": "2025-01-01 10:00"}})
    template = _template(tmp_path / "template.xlsx", rows=[("VN-A", 2, 10, 20, 15)])
    parameters = load_template_sku_parameters(template, {"VN-A"})
    default = RecommendationConfig()
    def generate(name, config):
        output = tmp_path / f"{name}.xlsx"
        recalculation.generate_vietnam_workbook(output, sources=source, config=config, template=template, parameters=parameters)
        with output.open("rb") as stream:
            book = load_workbook(stream, data_only=True)
            try:
                return [cell.value for cell in book["越南备货清单"][2]], book["数据更改"]["G2"].value
            finally:
                book.close()
    original_load = workbook._load_skeleton
    def old_fixed_weights():
        book = original_load()
        book["越南备货清单"]["S2"] = "=K2*0.1/(30+AV2)+L2*0.3/(15+AW2)+M2*0.6/(7+AX2)"
        return book
    with monkeypatch.context() as legacy:
        legacy.setattr(workbook, "_load_skeleton", old_fixed_weights)
        legacy.setattr(recalculation, "_load_skeleton", old_fixed_weights)
        reference, _ = generate("reference", default)
    actual, cached_weight = generate("default", default)
    assert actual == reference
    assert cached_weight == 0.6
    recent, cached_weight = generate("recent", replace(default, sales_weight_7d=Decimal(1), sales_weight_15d=Decimal(0), sales_weight_30d=Decimal(0)))
    assert cached_weight == 1
    assert recent[18] == 10  # S: only the seven-day daily sales.
    assert recent[7] != actual[7]  # H: final replenishment quantity.
    fx, _ = generate("exchange", replace(default, exchange_rate=Decimal(4200)))
    assert fx[7] == actual[7]  # Price conversion never changes units to replenish.
    assert fx[31:39] != actual[31:39]  # Price/profit results reflect the new exchange rate.


@pytest.mark.skipif(
    not (os.environ.get("LXE_OFFICE_NODE") and os.environ.get("LXE_OFFICE_CLI")),
    reason="Host Office Kit paths are not configured",
)
@pytest.mark.parametrize("explicit", [False, True])
def test_optional_amounts_keep_quantities_and_only_blank_dependent_results(
    tmp_path, isolated_state, monkeypatch, capsys, explicit,
):
    from services.vietnam_replenishment.recalculation import recalculate_with_office

    monkeypatch.delenv("LXESKILL_SKILL_SCOPE", raising=False)
    combinations = list(product((False, True), repeat=3))
    skus = tuple(f"VN-{index:03}" for index in range(len(combinations)))
    base = _sources()
    sources = replace(
        base, skus=skus,
        sales={sku: {**base.sales["VN-A"], "SKU": sku, "7天销量": 70, "15天销量": 90, "30天销量": 120} for sku in skus},
        inventory={sku: {**base.inventory["VN-A"], "SKU": sku} for sku in skus},
        products={sku: {**base.products["VN-A"], "SKU": sku, "创建时间": "2025-01-01 10:00"} for sku in skus},
        in_transit={sku: Decimal(3) for sku in skus}, in_transit_mismatch=skus,
    )
    amounts = (10, 200000, 150000)
    template_rows = [
        (sku, 2, *(None if absent else value for absent, value in zip(missing, amounts, strict=True)))
        for sku, missing in zip(skus, combinations, strict=True)
    ]
    template = _template(tmp_path / "完整模板.xlsx", rows=template_rows)
    if not explicit:
        settings.upload_template(template)
    args = ["vietnam", "replenishment", "calculate", *_report_arguments(tmp_path, sources)]
    if explicit:
        args.extend(("--template-file", str(template)))
    assert lxeskill.main(args) == 0
    [result] = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line.strip()]
    assert result["ok"] is True and result["data"]["validation"]["status"] == "passed"
    assert result["data"]["validation"]["missing_mapping_count"] == 0
    output = Path(result["data"]["output_xlsx"])
    cached = load_workbook(output, data_only=True)
    formulas = load_workbook(output, data_only=False)
    try:
        main = cached["越南备货清单"]
        quantity, days = main["H2"].value, main["AC2"].value
        assert quantity > 0 and days > 0
        financial_columns = ("AF", "AG", "AH", "AI", "AK", "AL", "AM")
        reference = {column: main[f"{column}2"].value for column in financial_columns}
        for row, (cost_missing, cross_missing, discount_missing) in enumerate(combinations, 2):
            assert main[f"H{row}"].value == quantity
            assert main[f"AC{row}"].value == days
            blank = {
                "AF": cross_missing, "AI": cross_missing,
                "AG": cost_missing or cross_missing, "AH": cost_missing or cross_missing,
                "AK": discount_missing,
                "AL": cost_missing or discount_missing, "AM": cost_missing or discount_missing,
            }
            for column in financial_columns:
                assert formulas["越南备货清单"][f"{column}{row}"].data_type == "f"
                value = main[f"{column}{row}"].value
                if blank[column]:
                    assert value in (None, ""), (row, column, value)
                else:
                    assert value == reference[column], (row, column, value)
            for column, absent in zip(("G", "AE", "AJ"), combinations[row - 2], strict=True):
                if absent:
                    assert main[f"{column}{row}"].value is None

        if explicit:
            # Filling values in the delivered workbook must reactivate retained formulas.
            for row in range(2, len(skus) + 2):
                for column, amount in zip(("G", "AE", "AJ"), amounts, strict=True):
                    formulas["越南备货清单"][f"{column}{row}"] = amount
            filled = tmp_path / "补齐金额.xlsx"
            formulas.save(filled)
            updated = tmp_path / "补齐后 重算.xlsx"
            recalculate_with_office(filled, updated)
            recalculated = load_workbook(updated, data_only=True)
            try:
                for row in range(2, len(skus) + 2):
                    sheet = recalculated["越南备货清单"]
                    assert sheet[f"H{row}"].value == quantity
                    for column in financial_columns:
                        assert sheet[f"{column}{row}"].value == reference[column]
            finally:
                recalculated.close()
    finally:
        cached.close()
        formulas.close()
