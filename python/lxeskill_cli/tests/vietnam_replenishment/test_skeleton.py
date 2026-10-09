"""The published Vietnam workbook layout contains rules, never prior business data."""

from __future__ import annotations

from importlib.resources import files
from pathlib import Path
import runpy
import tomllib
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.formula import ArrayFormula
import pytest

from services.vietnam_replenishment.asset_contract import REQUIRED_SHEETS


RESOURCE = files("services.vietnam_replenishment").joinpath("resources/skeleton.xlsx")
MAIN_FORMULAS = (
    "H",
    *(chr(column) for column in range(ord("J"), ord("Z") + 1)),
    "AA", "AB", "AC",
    "AF", "AG", "AH", "AI",
    "AK", "AL", "AM",
    "AV", "AW", "AX", "AY",
)
MAIN_INPUTS = ("A", "B", "C", "D", "E", "F", "G", "I", "AD", "AE", "AJ", "AN", "AO", "AP", "AQ", "AR", "AS", "AT", "AU")
HEADERS = {
    "越南备货清单": (
        "品牌", "判断热销", "链接", "产品图片", "SKU", "名称", "成本", "陆运", None,
        "库存", "30天", "15天", "7天", "30天日销", "15天日销", "7天日销",
        "增量趋势", "调节系数", "日均", "调节后日均", "综合判定-是否清货",
        "去波段预处理", "款号", "款内最大日销", "款内SKU数", "款内动销SKU数",
        "上架时间", "可售天数", "备货天数", "传过的店铺", "跨境价",
        "闪购价", "毛利润", "毛利率", "前台价", "折扣价", "闪购价",
        "毛利润", "毛利率", "总在途", "9/18陆运", None, None, None, None, None, None,
        "30天", "15天", "7天", "汇率",
    ),
    "雅仓库存": ("条码", "SKU", "映射条码", "仓库", "规格", "库存数量", "占用数量", "在途数量", "冻结库存", "可用库存", "中文标题", "英文标题", "图片链接"),
    "雅仓动销": ("SKU", "商品名", "仓库", "3天销量", "7天销量", "15天销量", "30天销量", "60天销量", "90天销量", "库存", "占用", "在途", "冻结", "可用", "缺货数量", "创建日期"),
    "数据更改": ("30天", "15天", "7天", "汇率", "30天销量权重", "15天销量权重", "7天销量权重"),
    "库存商品信息": ("条码", "SKU", "规格", "中文标题", "英文标题", "长", "宽", "高", "重量", "图片链接", "创建时间"),
}


def test_skeleton_has_only_audited_five_sheet_layout() -> None:
    with RESOURCE.open("rb") as source:
        workbook = load_workbook(source, data_only=False)
    try:
        assert tuple(workbook.sheetnames) == REQUIRED_SHEETS
        assert not workbook.defined_names
        assert not workbook._external_links
        for name, headers in HEADERS.items():
            sheet = workbook[name]
            assert tuple(sheet.cell(1, column).value for column in range(1, len(headers) + 1)) == headers
            assert sheet.max_row <= 2
            assert not sheet._images
            assert not sheet._charts
            assert not sheet.tables
            for row in sheet:
                for cell in row:
                    assert cell.comment is None
                    assert cell.hyperlink is None

        main = workbook["越南备货清单"]
        for column in MAIN_FORMULAS:
            assert main[f"{column}2"].data_type == "f", column
        assert isinstance(main["AC2"].value, ArrayFormula)
        assert main["AC2"].value.ref == "AC2"
        assert "库存商品信息!B:K" in main["AA2"].value
        assert main["S2"].value == "=K2*数据更改!$E$2/(30+AV2)+L2*数据更改!$F$2/(15+AW2)+M2*数据更改!$G$2/(7+AX2)"
        for column in MAIN_INPUTS:
            assert main[f"{column}2"].value is None, column
        for column, source in zip(("AV", "AW", "AX", "AY"), ("A", "B", "C", "D")):
            assert main[f"{column}2"].value == f"=数据更改!{source}2"

        for name in ("雅仓库存", "雅仓动销", "库存商品信息"):
            assert all(cell.value is None for cell in workbook[name][2])
        assert [workbook["数据更改"].cell(2, column).value for column in range(1, 8)] == [0.8, 0.8, 0, 3900, 0.1, 0.3, 0.6]
    finally:
        workbook.close()


def test_skeleton_archive_contains_no_recoverable_business_objects() -> None:
    with RESOURCE.open("rb") as source, ZipFile(source) as archive:
        names = archive.namelist()
        assert not any(name.startswith(("xl/media/", "xl/drawings/", "xl/externalLinks/", "customXml/")) for name in names)
        assert not any("comment" in name.casefold() or "person" in name.casefold() for name in names)
        assert "docProps/custom.xml" not in names
        assert "xl/sharedStrings.xml" not in names
        assert not any("vbaProject.bin" in name for name in names)
        for name in names:
            if name.endswith(".rels"):
                assert b'TargetMode="External"' not in archive.read(name)
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                contents = archive.read(name)
                assert b"<hyperlink" not in contents
                assert b"<drawing" not in contents


def test_skeleton_is_explicitly_included_in_python_wheel() -> None:
    root = Path(__file__).resolve().parents[4]
    config = tomllib.loads((root / "pyproject.toml").read_text(encoding="utf-8"))
    force_include = config["tool"]["hatch"]["build"]["targets"]["wheel"]["force-include"]
    assert force_include["python/lxeskill_cli/services/vietnam_replenishment/resources/skeleton.xlsx"] == (
        "services/vietnam_replenishment/resources/skeleton.xlsx"
    )


def test_builder_drops_synthetic_history_and_objects(tmp_path: Path) -> None:
    with RESOURCE.open("rb") as packaged:
        source = load_workbook(packaged)
    main = source["越南备货清单"]
    # The builder audits the original supplied template, then applies the
    # reviewed configurable-weights formula to its data-free output.
    main["S2"] = "=K2*0.1/(30+AV2)+L2*0.3/(15+AW2)+M2*0.6/(7+AX2)"
    main["E2"] = "CANARY-PRIVATE-SKU"
    main["G2"] = 987654321
    main["D2"].comment = Comment("CANARY-PRIVATE-COMMENT", "Synthetic")
    main["C2"].hyperlink = "https://example.invalid/CANARY-PRIVATE-LINK"
    main["E3"] = "CANARY-HISTORY-SKU"
    source.create_sheet("旧批次")["A1"] = "CANARY-PRIVATE-SHEET"
    template = tmp_path / "synthetic-source.xlsx"
    source.save(template)
    source.close()

    root = Path(__file__).resolve().parents[4]
    build = runpy.run_path(str(root / "scripts/build-vietnam-skeleton.py"))["build_skeleton"]
    output = build(template, tmp_path / "rebuilt.xlsx")
    rebuilt = load_workbook(output)
    try:
        assert rebuilt["越南备货清单"]["S2"].value == "=K2*数据更改!$E$2/(30+AV2)+L2*数据更改!$F$2/(15+AW2)+M2*数据更改!$G$2/(7+AX2)"
        assert [rebuilt["数据更改"].cell(1, i).value for i in range(1, 8)] == list(HEADERS["数据更改"])
        assert [rebuilt["数据更改"].cell(2, i).value for i in range(1, 8)] == [0.8, 0.8, 0, 3900, 0.1, 0.3, 0.6]
    finally:
        rebuilt.close()
    with ZipFile(output) as archive:
        for member in archive.namelist():
            contents = archive.read(member)
            assert b"CANARY-PRIVATE" not in contents
            assert b"CANARY-HISTORY" not in contents
        assert not any(member.startswith("xl/comments") for member in archive.namelist())

    changed = load_workbook(template)
    changed["越南备货清单"]["E1"] = "Unreviewed header"
    changed.save(template)
    changed.close()
    with pytest.raises(ValueError, match="headers differ"):
        build(template, tmp_path / "rejected.xlsx")
    assert not (tmp_path / "rejected.xlsx").exists()
