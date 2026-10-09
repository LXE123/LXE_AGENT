"""A current export, explicit map, and template history form one operator map."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from services.vietnam_replenishment.asset_contract import AssetContractError, SkuParameters
from services.vietnam_replenishment.yacang_sources import VietnamSources


def _template(path: Path) -> Path:
    workbook = Workbook()
    main = workbook.active
    main.title = "越南备货清单"
    for name in ("雅仓库存", "雅仓动销", "数据更改", "库存商品信息"):
        workbook.create_sheet(name)
    for coordinate, header in {
        "B1": "判断热销", "E1": "SKU", "G1": "成本", "H1": "陆运",
        "AA1": "上架时间", "AE1": "跨境价", "AJ1": "折扣价", "AN1": "总在途",
        "AV1": "30天", "AW1": "15天", "AX1": "7天", "AY1": "汇率",
    }.items():
        main[coordinate] = header
    for name, headers in {
        "雅仓库存": {"B1": "SKU", "F1": "库存数量", "H1": "在途数量", "J1": "可用库存"},
        "雅仓动销": {"A1": "SKU", "E1": "7天销量", "F1": "15天销量", "G1": "30天销量"},
        "数据更改": {"A1": "30天", "B1": "15天", "C1": "7天", "D1": "汇率"},
        "库存商品信息": {"B1": "SKU", "K1": "创建时间"},
    }.items():
        for coordinate, header in headers.items():
            workbook[name][coordinate] = header
    for coordinate, value in zip(("A2", "B2", "C2", "D2"), (0.8, 0.8, 0, 3900)):
        workbook["数据更改"][coordinate] = value
    for coordinate, source in zip(("AV2", "AW2", "AX2", "AY2"), ("A2", "B2", "C2", "D2")):
        main[coordinate] = f"=数据更改!{source}"
    main["E2"], main["G2"], main["B2"] = "VN-A", 10, 1
    main["E3"], main["G3"] = "VN-B", 20
    main["E4"], main["G4"] = "GLOBAL-OLD", 99
    workbook.save(path)
    workbook.close()
    return path


def _explicit_map(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("SKU", "成本", "跨境价", "折扣价", "热销标记", "上架时间"))
    sheet.append(("VN-A", 0, None, None, 2, None))
    sheet.append(("GLOBAL-OLD", 500, None, None, None, None))
    workbook.save(path)
    workbook.close()
    return path


def _sources() -> VietnamSources:
    return VietnamSources(
        skus=("VN-A", "VN-B"),
        sales={"VN-A": {"SKU": "VN-A"}, "VN-B": {"SKU": "VN-B"}},
        inventory={"VN-A": {"SKU": "VN-A"}, "VN-B": {"SKU": "VN-B"}},
        products={"VN-A": {"SKU": "VN-A"}, "VN-B": {"SKU": "VN-B"}},
        missing_sales=(), missing_inventory=(), missing_products=(),
        missing_in_transit=(),
        in_transit={"VN-A": Decimal("2"), "VN-B": Decimal("0")},
        in_transit_mismatch=(), artifacts={},
    )


def test_preparation_keeps_upload_explicit_and_resolution_separate(tmp_path: Path, monkeypatch) -> None:
    from services.vietnam_replenishment import preparation

    sources = _sources()
    export_calls = []
    monkeypatch.setattr(
        preparation, "export_vietnam_sources", lambda: export_calls.append(True) or sources
    )
    output = tmp_path / "operator-upload.xlsx"

    result = preparation.prepare_operator_sku_map(
        _template(tmp_path / "template.xlsx"), output,
        _explicit_map(tmp_path / "current-map.xlsx"),
    )

    assert export_calls == [True]
    assert result.path == output
    assert result.sources is sources
    assert result.resolved["VN-A"].values == SkuParameters(cost=Decimal("0"), hot_flag=2)
    assert result.resolved["VN-A"].sources["cost"] == "explicit"
    assert result.resolved["VN-B"].values == SkuParameters(cost=Decimal("20"), hot_flag=2)
    assert result.resolved["VN-B"].sources == {"cost": "template", "hot_flag": "default"}

    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        upload = list(workbook["SKU参数映射"].values)
        reference = list(workbook["核对参考"].values)
    finally:
        workbook.close()
    assert tuple(row[0] for row in upload[1:]) == sources.skus
    assert upload[1][1:5] == (0, None, None, 2)
    assert upload[2][1:5] == (None, None, None, None)
    assert reference[2][1] == 20
    assert "GLOBAL-OLD" not in tuple(row[0] for row in reference)


def test_injected_current_sources_do_not_start_another_export(tmp_path: Path, monkeypatch) -> None:
    from services.vietnam_replenishment import preparation

    def unexpected_export():
        raise AssertionError("extra 雅仓 export")

    monkeypatch.setattr(preparation, "export_vietnam_sources", unexpected_export)
    result = preparation.prepare_operator_sku_map(
        _template(tmp_path / "template.xlsx"), tmp_path / "operator-upload.xlsx", sources=_sources(),
    )
    assert result.path.is_file()
    assert result.resolved["VN-A"].sources["cost"] == "template"


def test_invalid_local_input_stops_before_live_export(tmp_path: Path, monkeypatch) -> None:
    from services.vietnam_replenishment import preparation

    def unexpected_export():
        raise AssertionError("invalid local input triggered a live 雅仓 export")

    monkeypatch.setattr(preparation, "export_vietnam_sources", unexpected_export)
    template = _template(tmp_path / "template.xlsx")
    workbook = load_workbook(template)
    del workbook["雅仓库存"]
    workbook.save(template)
    workbook.close()

    with pytest.raises(AssetContractError, match="缺少工作表: 雅仓库存"):
        preparation.prepare_operator_sku_map(template, tmp_path / "output.xlsx")
