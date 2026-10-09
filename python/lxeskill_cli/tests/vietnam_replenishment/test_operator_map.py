"""Synthetic checks for the operations SKU upload workbook."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from services.vietnam_replenishment.asset_contract import SkuParameters, load_sku_parameters
from services.vietnam_replenishment.operator_map import write_operator_sku_map
from services.vietnam_replenishment.sku_parameters import HistoricalSkuParameters
from services.vietnam_replenishment.yacang_sources import VietnamSources


def _sources() -> VietnamSources:
    return VietnamSources(
        skus=("VN-A", "VN-B", "VN-C"),
        sales={"VN-A": {}, "VN-B": {}},
        inventory={"VN-A": {}, "VN-C": {}},
        products={
            "VN-A": {"创建时间": "2026-09-01 08:00"},
            "VN-B": {"创建时间": "2026-09-02 08:00"},
            "GLOBAL-OTHER": {"创建时间": "2026-09-03 08:00"},
        },
        missing_sales=("VN-C",),
        missing_inventory=("VN-B",),
        missing_products=("VN-C",),
        in_transit={"VN-A": Decimal("1"), "VN-B": None, "VN-C": None},
        missing_in_transit=("VN-B", "VN-C"),
        in_transit_mismatch=("VN-A",),
        artifacts={},
    )


def _history() -> dict[str, HistoricalSkuParameters]:
    return {
        "VN-A": HistoricalSkuParameters(
            values=SkuParameters(cost=Decimal("99"), hot_flag=2),
            unavailable={"cross_border_price": "模板为公式，不能作为历史值"},
        ),
        "VN-B": HistoricalSkuParameters(
            values=SkuParameters(
                cost=Decimal("3.5"),
                cross_border_price=Decimal("11"),
                discount_price=Decimal("9"),
                hot_flag=1,
                listed_at="2026-08-20",
            ),
            unavailable={},
        ),
        "VN-C": HistoricalSkuParameters(
            values=SkuParameters(),
            unavailable={"cost": "同一 SKU 的历史成本冲突"},
        ),
        "GLOBAL-OTHER": HistoricalSkuParameters(
            values=SkuParameters(cost=Decimal("1000")),
            unavailable={},
        ),
    }


def test_upload_sheet_round_trips_only_current_explicit_values(tmp_path: Path) -> None:
    explicit = {
        "VN-A": SkuParameters(
            cost=Decimal("0"),
            cross_border_price=Decimal("12.5"),
            hot_flag=1,
            listed_at="2026-09-19",
        ),
        "GLOBAL-OTHER": SkuParameters(cost=Decimal("500")),
    }
    output = tmp_path / "operator-map.xlsx"

    assert write_operator_sku_map(output, _sources(), explicit, _history()) == output

    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["SKU参数映射", "核对参考"]
        rows = list(workbook.worksheets[0].values)
        assert rows[0] == (
            "SKU", "成本", "跨境价", "折扣价", "热销标记", "上架时间"
        )
        assert rows[1] == ("VN-A", 0, 12.5, None, 1, "2026-09-19")
        assert rows[2] == ("VN-B", None, None, None, None, None)
        assert rows[3] == ("VN-C", None, None, None, None, None)
        assert len(rows) == 4
    finally:
        workbook.close()

    assert load_sku_parameters(output) == {
        "VN-A": explicit["VN-A"],
        "VN-B": SkuParameters(),
        "VN-C": SkuParameters(),
    }


def test_reference_sheet_shows_history_issues_and_missing_sources(tmp_path: Path) -> None:
    output = write_operator_sku_map(tmp_path / "operator-map.xlsx", _sources(), {}, _history())

    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        rows = list(workbook["核对参考"].values)
    finally:
        workbook.close()

    headers = rows[0]
    assert headers == (
        "SKU",
        "模板成本",
        "模板跨境价",
        "模板折扣价",
        "模板热销标记",
        "模板上架时间",
        "模板问题",
        "缺动销",
        "缺库存",
        "缺商品",
        "缺总在途",
        "在途不一致",
    )
    data = {row[0]: dict(zip(headers, row)) for row in rows[1:]}
    assert set(data) == {"VN-A", "VN-B", "VN-C"}
    assert data["VN-A"]["模板成本"] == 99
    assert data["VN-A"]["模板跨境价"] is None
    assert "模板跨境价" in data["VN-A"]["模板问题"]
    assert "模板为公式" in data["VN-A"]["模板问题"]
    assert data["VN-A"]["在途不一致"] == "是"
    assert data["VN-B"]["模板成本"] == 3.5
    assert data["VN-B"]["模板上架时间"] == "2026-08-20"
    assert data["VN-B"]["缺库存"] == "是"
    assert data["VN-B"]["缺总在途"] == "是"
    assert data["VN-C"]["缺动销"] == "是"
    assert data["VN-C"]["缺商品"] == "是"
    assert data["VN-C"]["缺总在途"] == "是"
    assert "历史成本冲突" in data["VN-C"]["模板问题"]


def test_literal_sku_starting_with_equals_is_not_written_as_formula(tmp_path: Path) -> None:
    sources = VietnamSources(
        skus=("=VN-SKU",),
        sales={},
        inventory={},
        products={},
        missing_sales=(),
        missing_inventory=(),
        missing_products=(),
        missing_in_transit=(),
        in_transit={},
        in_transit_mismatch=(),
        artifacts={},
    )

    output = write_operator_sku_map(tmp_path / "literal-sku.xlsx", sources, {}, {})

    assert set(load_sku_parameters(output)) == {"=VN-SKU"}
    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        assert workbook.worksheets[0]["A2"].data_type == "s"
        assert workbook["核对参考"]["G2"].value == "模板无此 SKU"
    finally:
        workbook.close()


@pytest.mark.parametrize("precise_price", [
    Decimal("0.1234567890123456789012345678"),
    Decimal("1E+1000"),
    Decimal("1E-1000"),
])
def test_high_precision_price_round_trips_without_excel_rounding(
    tmp_path: Path, precise_price: Decimal
) -> None:
    output = write_operator_sku_map(
        tmp_path / "precise-map.xlsx",
        _sources(),
        {"VN-A": SkuParameters(cross_border_price=precise_price)},
        {},
    )

    assert load_sku_parameters(output)["VN-A"].cross_border_price == precise_price
    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        assert workbook.worksheets[0]["C2"].value == str(precise_price)
        assert workbook.worksheets[0]["C2"].data_type == "s"
    finally:
        workbook.close()


def test_existing_operator_map_is_never_overwritten(tmp_path: Path) -> None:
    output = tmp_path / "current-map.xlsx"
    original = b"existing operator workbook"
    output.write_bytes(original)

    with pytest.raises(FileExistsError):
        write_operator_sku_map(output, _sources(), {}, {})

    assert output.read_bytes() == original


def test_partial_output_is_removed_if_workbook_save_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    output = tmp_path / "incomplete-map.xlsx"

    def fail_after_write(_workbook: Workbook, destination) -> None:
        destination.write(b"partial")
        raise OSError("simulated disk error")

    monkeypatch.setattr(Workbook, "save", fail_after_write)
    with pytest.raises(OSError, match="simulated disk error"):
        write_operator_sku_map(output, _sources(), {}, {})

    assert not output.exists()
