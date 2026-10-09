"""Current-SKU resolution of Vietnam business inputs from template history."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
import pytest

from services.vietnam_replenishment.asset_contract import AssetContractError, SkuParameters
from services.vietnam_replenishment.sku_parameters import (
    HistoricalSkuParameters,
    load_template_sku_parameters,
    resolve_sku_parameters,
)


FIELD_COLUMNS = {
    "hot_flag": "B",
    "cost": "G",
    "listed_at": "AA",
    "cross_border_price": "AE",
    "discount_price": "AJ",
}


def _template(path: Path, *rows: dict[str, object]) -> Path:
    workbook = Workbook()
    main = workbook.active
    main.title = "越南备货清单"
    for name in ("雅仓库存", "雅仓动销", "数据更改", "库存商品信息"):
        workbook.create_sheet(name)
    for coordinate, header in {
        "B1": "判断热销", "E1": "SKU", "G1": "成本", "H1": "陆运",
        "AA1": "上架时间", "AE1": "跨境价", "AJ1": "折扣价", "AN1": "总在途",
        "AV1": "30天", "AW1": "15天", "AX1": "7天", "AY1": "汇率",
    }.items():
        main[coordinate] = header
    for name, headers in {
        "雅仓库存": {"B1": "SKU", "F1": "库存数量", "H1": "在途数量", "J1": "可用库存"},
        "雅仓动销": {"A1": "SKU", "E1": "7天销量", "F1": "15天销量", "G1": "30天销量"},
        "数据更改": {"A1": "30天", "B1": "15天", "C1": "7天", "D1": "汇率"},
        "库存商品信息": {"B1": "SKU", "K1": "创建时间"},
    }.items():
        for coordinate, header in headers.items():
            workbook[name][coordinate] = header
    for coordinate, value in zip(("A2", "B2", "C2", "D2"), (0.8, 0.8, 0, 3900)):
        workbook["数据更改"][coordinate] = value
    for coordinate, source in zip(("AV2", "AW2", "AX2", "AY2"), ("A2", "B2", "C2", "D2")):
        main[coordinate] = f"=数据更改!{source}"
    for row_number, values in enumerate(rows, 2):
        main[f"E{row_number}"] = values["sku"]
        for field, column in FIELD_COLUMNS.items():
            if field in values:
                main[f"{column}{row_number}"] = values[field]
    workbook.save(path)
    workbook.close()
    return path


def test_template_history_reads_only_exact_current_skus_and_agreeing_duplicates(tmp_path: Path) -> None:
    path = _template(
        tmp_path / "history.xlsx",
        {"sku": " VN-A ", "hot_flag": 1, "cost": 0, "cross_border_price": 9.5,
         "discount_price": 7, "listed_at": "2026-09-19"},
        {"sku": "VN-A", "hot_flag": 1, "cost": None, "cross_border_price": 9.5,
         "discount_price": 7, "listed_at": "2026-09-19"},
        {"sku": "VN-B", "cost": 100},
    )

    history = load_template_sku_parameters(path, ["VN-A", "VN-NEW"])

    assert set(history) == {"VN-A", "VN-NEW"}
    assert history["VN-A"].values == SkuParameters(
        cost=Decimal("0"), cross_border_price=Decimal("9.5"),
        discount_price=Decimal("7"), hot_flag=1, listed_at="2026-09-19",
    )
    assert history["VN-A"].unavailable == {}
    assert history["VN-NEW"].values == SkuParameters()
    assert set(history["VN-NEW"].unavailable) == set(FIELD_COLUMNS)


def test_template_history_rejects_each_unusable_field_without_first_row_wins(tmp_path: Path) -> None:
    path = _template(
        tmp_path / "conflicts.xlsx",
        {"sku": "VN-A", "hot_flag": 1, "cost": 10, "cross_border_price": 15,
         "discount_price": 5, "listed_at": "2026-09-19"},
        {"sku": "VN-A", "hot_flag": 2, "cost": 11, "cross_border_price": "=10+5",
         "discount_price": -1, "listed_at": "invalid-date"},
    )

    historical = load_template_sku_parameters(path, ["VN-A"])["VN-A"]

    assert historical.values == SkuParameters()
    assert set(historical.unavailable) == set(FIELD_COLUMNS)
    assert "冲突" in historical.unavailable["cost"]
    assert "G2" in historical.unavailable["cost"] and "G3" in historical.unavailable["cost"]
    assert "公式" in historical.unavailable["cross_border_price"]
    assert "无效" in historical.unavailable["discount_price"]
    assert "无效" in historical.unavailable["listed_at"]


def test_resolution_uses_explicit_then_history_then_hot_default_and_keeps_zero() -> None:
    history = {
        "VN-A": HistoricalSkuParameters(
            values=SkuParameters(
                cost=Decimal("12"), cross_border_price=Decimal("20"),
                discount_price=Decimal("15"), hot_flag=1, listed_at="2026-09-19",
            ),
            unavailable={},
        ),
        "VN-NEW": HistoricalSkuParameters(values=SkuParameters(), unavailable={}),
    }
    explicit = {
        "VN-A": SkuParameters(cost=Decimal("0"), hot_flag=2),
        "VN-EXTRA": SkuParameters(cost=Decimal("99")),
    }

    result = resolve_sku_parameters(["VN-A", "VN-NEW"], explicit, history)

    assert set(result) == {"VN-A", "VN-NEW"}
    assert result["VN-A"].values == SkuParameters(
        cost=Decimal("0"), cross_border_price=Decimal("20"),
        discount_price=Decimal("15"), hot_flag=2, listed_at="2026-09-19",
    )
    assert result["VN-A"].sources == {
        "cost": "explicit", "cross_border_price": "template", "discount_price": "template",
        "hot_flag": "explicit", "listed_at": "template",
    }
    assert result["VN-A"].unavailable == {}
    assert result["VN-NEW"].values == SkuParameters(hot_flag=2)
    assert result["VN-NEW"].sources == {"hot_flag": "default"}
    assert set(result["VN-NEW"].unavailable) == {
        "cost", "cross_border_price", "discount_price", "listed_at",
    }


def test_explicit_value_resolves_conflicted_history_but_other_fields_remain_missing() -> None:
    history = {
        "VN-A": HistoricalSkuParameters(
            values=SkuParameters(), unavailable={
                "cost": "模板历史值冲突: G2, G3",
                "cross_border_price": "模板历史公式不可用: AE3",
            }),
    }

    result = resolve_sku_parameters(
        ["VN-A"], {"VN-A": SkuParameters(cost=Decimal("0"))}, history,
    )["VN-A"]

    assert result.values.cost == Decimal("0")
    assert result.values.hot_flag == 2
    assert result.sources["cost"] == "explicit"
    assert "cost" not in result.unavailable
    assert result.unavailable["cross_border_price"] == "模板历史公式不可用: AE3"


def test_invalid_template_is_reported_before_historical_read(tmp_path: Path) -> None:
    path = _template(tmp_path / "invalid-template.xlsx", {"sku": "VN-A"})
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    del workbook["雅仓库存"]
    workbook.save(path)
    workbook.close()

    with pytest.raises(AssetContractError, match="缺少工作表: 雅仓库存"):
        load_template_sku_parameters(path, ["VN-A"])
