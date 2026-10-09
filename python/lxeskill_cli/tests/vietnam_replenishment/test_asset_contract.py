"""Synthetic contracts for the Vietnam business template."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from services.vietnam_replenishment.asset_contract import (
    AssetContractError,
    SkuParameters,
    load_sku_parameters,
    validate_complete_sku_parameters,
    validate_template,
)


REQUIRED_SHEETS = (
    "越南备货清单",
    "雅仓库存",
    "雅仓动销",
    "数据更改",
    "库存商品信息",
)

AUXILIARY_HEADERS = {
    "雅仓库存": {"B1": "SKU", "F1": "库存数量", "H1": "在途数量", "J1": "可用库存"},
    "雅仓动销": {"A1": "SKU", "E1": "7天销量", "F1": "15天销量", "G1": "30天销量"},
    "数据更改": {"A1": "30天", "B1": "15天", "C1": "7天", "D1": "汇率"},
    "库存商品信息": {"B1": "SKU", "K1": "创建时间"},
}


def _template(path: Path, *, extra_sheet: bool = False) -> Path:
    workbook = Workbook()
    main = workbook.active
    main.title = REQUIRED_SHEETS[0]
    for name in REQUIRED_SHEETS[1:]:
        workbook.create_sheet(name)
    if extra_sheet:
        workbook.create_sheet("历史辅助表")
    for coordinate, header in {
        "B1": "判断热销",
        "E1": "SKU",
        "G1": "成本",
        "H1": "陆运",
        "AA1": "上架时间",
        "AE1": "跨境价",
        "AJ1": "折扣价",
        "AN1": "总在途",
        "AV1": "30天",
        "AW1": "15天",
        "AX1": "7天",
        "AY1": "汇率",
    }.items():
        main[coordinate] = header
    for sheet_name, headers in AUXILIARY_HEADERS.items():
        sheet = workbook[sheet_name]
        for coordinate, header in headers.items():
            sheet[coordinate] = header
    change = workbook["数据更改"]
    for coordinate, value in zip(("A2", "B2", "C2", "D2"), (0.8, 0.8, 0, 3900)):
        change[coordinate] = value
    for coordinate, source in zip(("AV2", "AW2", "AX2", "AY2"), ("A2", "B2", "C2", "D2")):
        main[coordinate] = f"=数据更改!{source}"
    main["E2"] = "VN-SKU-1"
    main["E4"] = "VN-SKU-2"
    workbook.save(path)
    workbook.close()
    return path


@pytest.mark.parametrize("extra_sheet", [False, True])
def test_full_template_accepts_required_sheets_and_optional_extra(
    tmp_path: Path, extra_sheet: bool
) -> None:
    path = _template(tmp_path / "synthetic-template.xlsx", extra_sheet=extra_sheet)

    contract = validate_template(path)

    assert set(REQUIRED_SHEETS).issubset(contract.sheet_names)
    assert len(contract.sheet_names) == 5 + int(extra_sheet)
    assert contract.main_rows == 2


def test_full_template_allows_repeated_historical_sku_rows(tmp_path: Path) -> None:
    path = _template(tmp_path / "repeated-history.xlsx")
    workbook = load_workbook(path)
    workbook[REQUIRED_SHEETS[0]]["E4"] = "VN-SKU-1"
    workbook.save(path)
    workbook.close()

    assert validate_template(path).main_rows == 2


def test_full_template_reports_missing_required_sheet(tmp_path: Path) -> None:
    path = _template(tmp_path / "missing-sheet.xlsx")
    workbook = load_workbook(path)
    del workbook["数据更改"]
    workbook.save(path)
    workbook.close()

    with pytest.raises(AssetContractError, match="缺少工作表: 数据更改"):
        validate_template(path)


def test_full_template_reports_moved_sku_header_coordinate(tmp_path: Path) -> None:
    path = _template(tmp_path / "moved-header.xlsx")
    workbook = load_workbook(path)
    main = workbook[REQUIRED_SHEETS[0]]
    main["D1"] = "SKU"
    main["E1"] = None
    workbook.save(path)
    workbook.close()

    with pytest.raises(AssetContractError, match=r"越南备货清单!E1:.*SKU"):
        validate_template(path)


@pytest.mark.parametrize("coordinate", ["AA1", "AE1", "AJ1"])
def test_full_template_requires_historical_parameter_headers(
    tmp_path: Path, coordinate: str
) -> None:
    path = _template(tmp_path / "missing-historical-header.xlsx")
    workbook = load_workbook(path)
    workbook[REQUIRED_SHEETS[0]][coordinate] = None
    workbook.save(path)
    workbook.close()

    with pytest.raises(AssetContractError) as error:
        validate_template(path)
    assert f"越南备货清单!{coordinate}" in str(error.value)


def test_full_template_reports_actual_malformed_xlsx_error(tmp_path: Path) -> None:
    path = tmp_path / "corrupted.xlsx"
    path.write_bytes(b"not an xlsx archive")

    with pytest.raises(AssetContractError, match="BadZipFile: File is not a zip file"):
        validate_template(path)


@pytest.mark.parametrize(
    ("sheet_name", "coordinate"),
    [
        (sheet_name, coordinate)
        for sheet_name, headers in AUXILIARY_HEADERS.items()
        for coordinate in headers
    ],
)
def test_full_template_reports_missing_auxiliary_header(
    tmp_path: Path, sheet_name: str, coordinate: str
) -> None:
    path = _template(tmp_path / "missing-auxiliary-header.xlsx")
    workbook = load_workbook(path)
    workbook[sheet_name][coordinate] = None
    workbook.save(path)
    workbook.close()

    with pytest.raises(AssetContractError) as error:
        validate_template(path)
    assert f"{sheet_name}!{coordinate}" in str(error.value)


@pytest.mark.parametrize(
    ("main_coordinate", "source_coordinate"),
    list(zip(("AV2", "AW2", "AX2", "AY2"), ("A2", "B2", "C2", "D2"))),
)
def test_full_template_requires_each_parameter_reference(
    tmp_path: Path, main_coordinate: str, source_coordinate: str
) -> None:
    path = _template(tmp_path / "wrong-reference.xlsx")
    workbook = load_workbook(path)
    workbook[REQUIRED_SHEETS[0]][main_coordinate] = "=1+1"
    workbook.save(path)
    workbook.close()

    with pytest.raises(AssetContractError) as error:
        validate_template(path)
    assert f"越南备货清单!{main_coordinate}" in str(error.value)
    assert source_coordinate in str(error.value)


@pytest.mark.parametrize("coordinate", ["A2", "B2", "C2", "D2"])
def test_parameter_input_area_rejects_formula(tmp_path: Path, coordinate: str) -> None:
    path = _template(tmp_path / "formula-in-input.xlsx")
    workbook = load_workbook(path)
    workbook["数据更改"][coordinate] = "=1+1"
    workbook.save(path)
    workbook.close()

    with pytest.raises(AssetContractError) as error:
        validate_template(path)
    assert f"数据更改!{coordinate}" in str(error.value)


MAP_HEADERS = ("折扣价", "热销标记", "SKU", "成本", "跨境价", "上架时间")


def _sku_map(path: Path, *rows: tuple[object, ...], headers=MAP_HEADERS) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SKU参数"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


def test_sku_parameters_match_exact_text_sku_and_preserve_explicit_values(tmp_path: Path) -> None:
    path = _sku_map(
        tmp_path / "parameters.xlsx",
        (None, 1, " VN-001 ", 0, 16.5, date(2026, 9, 19)),
        (9.9, 2, "VN-002", 10.25, 20.5, None),
        (None, None, None, None, None, None),
        ("", "", "", "", "", ""),
    )

    result = load_sku_parameters(path)

    assert set(result) == {"VN-001", "VN-002"}
    assert result["VN-001"] == SkuParameters(
        cost=Decimal("0"),
        cross_border_price=Decimal("16.5"),
        discount_price=None,
        hot_flag=1,
        listed_at="2026-09-19",
    )
    assert result["VN-002"] == SkuParameters(
        cost=Decimal("10.25"),
        cross_border_price=Decimal("20.5"),
        discount_price=Decimal("9.9"),
        hot_flag=2,
        listed_at=None,
    )


def test_sku_parameters_keep_internal_sku_characters_and_optional_date(tmp_path: Path) -> None:
    path = _sku_map(
        tmp_path / "parameters.xlsx",
        (None, None, " VN A  01 ", None, None),
        headers=MAP_HEADERS[:-1],
    )

    result = load_sku_parameters(path)

    assert set(result) == {"VN A  01"}
    assert result["VN A  01"].listed_at is None


@pytest.mark.parametrize(
    "listed_at", ["2026-09-19", "2026-09-19 10:00", "2026-09-19T10:00:00+07:00"]
)
def test_sku_parameters_accept_explicit_iso_listed_at_text(
    tmp_path: Path, listed_at: str
) -> None:
    path = _sku_map(
        tmp_path / "listed-at-text.xlsx",
        (None, 1, "VN-001", 3, 5, listed_at),
    )

    assert load_sku_parameters(path)["VN-001"].listed_at == listed_at


@pytest.mark.parametrize(
    "listed_at",
    ["tomorrow", "2026-02-30 10:00", "2026-01-01 nonsense", "2026-01-01 25:00"],
)
def test_sku_parameters_reject_invalid_listed_at_text(
    tmp_path: Path, listed_at: str
) -> None:
    path = _sku_map(tmp_path / "bad-listed-at.xlsx", (None, 1, "VN-001", 3, 5, listed_at))

    with pytest.raises(AssetContractError) as error:
        load_sku_parameters(path)
    assert "F2" in str(error.value)


def test_sku_parameters_reject_duplicate_sku_after_trim(tmp_path: Path) -> None:
    path = _sku_map(
        tmp_path / "duplicate.xlsx",
        (None, 1, "VN-001", 3, 5, None),
        (None, 2, " VN-001 ", 4, 6, None),
    )

    with pytest.raises(AssetContractError) as error:
        load_sku_parameters(path)
    assert "C3" in str(error.value) and "C2" in str(error.value)


@pytest.mark.parametrize(
    ("sku", "expected"),
    [(None, "C2"), (" ", "C2"), (123, "C2")],
)
def test_sku_parameters_reject_nonblank_row_without_text_sku(
    tmp_path: Path, sku: object, expected: str
) -> None:
    path = _sku_map(tmp_path / "bad-sku.xlsx", (None, 1, sku, 3, 5, None))

    with pytest.raises(AssetContractError) as error:
        load_sku_parameters(path)
    assert expected in str(error.value)


@pytest.mark.parametrize("cost", ["not-a-number", -1, "NaN", "Infinity"])
def test_sku_parameters_reject_invalid_numeric_input(tmp_path: Path, cost: object) -> None:
    path = _sku_map(tmp_path / "bad-cost.xlsx", (None, 1, "VN-001", cost, 5, None))

    with pytest.raises(AssetContractError) as error:
        load_sku_parameters(path)
    assert "D2" in str(error.value)


def test_sku_parameters_reject_formula_in_input_cell(tmp_path: Path) -> None:
    path = _sku_map(tmp_path / "formula.xlsx", (None, 1, "VN-001", "=1+1", 5, None))

    with pytest.raises(AssetContractError) as error:
        load_sku_parameters(path)
    assert "D2" in str(error.value)


@pytest.mark.parametrize("hot_flag", [0, 3, "hot"])
def test_sku_parameters_reject_hot_flag_outside_one_or_two(
    tmp_path: Path, hot_flag: object
) -> None:
    path = _sku_map(tmp_path / "bad-hot.xlsx", (None, hot_flag, "VN-001", 3, 5, None))

    with pytest.raises(AssetContractError) as error:
        load_sku_parameters(path)
    assert "B2" in str(error.value)


def test_sku_parameters_require_unique_named_headers(tmp_path: Path) -> None:
    path = _sku_map(
        tmp_path / "duplicate-header.xlsx",
        (None, 1, "VN-001", 3, 5, None),
        headers=("折扣价", "热销标记", "SKU", "成本", "跨境价", "成本"),
    )

    with pytest.raises(AssetContractError) as error:
        load_sku_parameters(path)
    assert "成本" in str(error.value)


def test_sku_parameters_require_every_named_header(tmp_path: Path) -> None:
    path = _sku_map(
        tmp_path / "missing-header.xlsx",
        (None, 1, "VN-001", 3),
        headers=("折扣价", "热销标记", "SKU", "成本"),
    )

    with pytest.raises(AssetContractError, match="缺少表头 '跨境价'"):
        load_sku_parameters(path)


def test_sku_parameters_keep_actual_workbook_read_error(tmp_path: Path) -> None:
    path = tmp_path / "corrupted-map.xlsx"
    path.write_bytes(b"not an xlsx archive")

    with pytest.raises(AssetContractError, match="BadZipFile: File is not a zip file"):
        load_sku_parameters(path)



def test_complete_sku_map_accepts_explicit_zero_prices(tmp_path: Path) -> None:
    path = _sku_map(tmp_path / "zero.xlsx", (0, None, "VN-A", 0, 0))
    values = validate_complete_sku_parameters(path)
    assert values["VN-A"].cost == Decimal("0")
    assert values["VN-A"].discount_price == Decimal("0")


@pytest.mark.parametrize(
    ("row", "message"),
    [
        ((None, None, "VN-A", 1, 2), "折扣价"),
        ((1, None, "VN-A", "1234567890123456", 2), "Excel 精度"),
        ((1, None, "VN-A", "1e-400", 2), "精确写入"),
    ],
)
def test_complete_sku_map_rejects_missing_or_inexact_prices(
    tmp_path: Path, row: tuple[object, ...], message: str,
) -> None:
    path = _sku_map(tmp_path / "bad.xlsx", row)
    with pytest.raises(AssetContractError, match=message):
        validate_complete_sku_parameters(path)


def test_complete_sku_map_rejects_empty_first_sheet(tmp_path: Path) -> None:
    path = _sku_map(tmp_path / "empty.xlsx")
    with pytest.raises(AssetContractError, match="没有 SKU"):
        validate_complete_sku_parameters(path)
