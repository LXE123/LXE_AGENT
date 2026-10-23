from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from hashlib import sha256
from importlib.resources import files
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
import pytest

from services.vietnam_replenishment import settings, workflow
from services.vietnam_replenishment.full_workbook import UNCHANGED
from services.vietnam_replenishment.workbook import RecommendationConfig
from shared.input_assets import current_asset, previous_asset
from shared import input_assets


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "app"))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path))
    from shared import workspace
    monkeypatch.setattr(workspace, "_workspace_root", tmp_path)
    monkeypatch.setattr(workspace, "_input_root", tmp_path / "app" / "inputs")


def nine_sheet_template(path: Path, *, marker: str = "first") -> Path:
    resource = files("services.vietnam_replenishment").joinpath("resources/skeleton.xlsx")
    with resource.open("rb") as stream:
        book = load_workbook(stream)
    try:
        book["越南备货清单"]["B2"] = 2
        book["越南备货清单"]["E2"] = "VN-A"
        book["越南备货清单"]["G2"] = 10
        for index, name in enumerate(UNCHANGED, 1):
            book.create_sheet(name, index)["A1"] = marker
        book.save(path)
    finally:
        book.close()
    return path


def test_nine_sheet_template_upload_export_and_invalid_replacement(tmp_path):
    source = nine_sheet_template(tmp_path / "完整九表.xlsx")
    state = settings.upload_template(source)
    assert state["template"]["file_name"] == source.name
    saved = current_asset("vietnam_full_workbook")
    assert saved is not None and saved.path.read_bytes() == source.read_bytes()
    exported = tmp_path / "导出当前模板.xlsx"
    assert settings.export_template(exported) == {"path": str(exported)}
    assert exported.read_bytes() == source.read_bytes()
    with pytest.raises(ValueError, match="不能覆盖"):
        settings.export_template(saved.path)
    assert saved.path.read_bytes() == source.read_bytes()

    invalid = tmp_path / "invalid.xlsx"
    invalid.write_bytes(b"not a workbook")
    with pytest.raises(ValueError):
        settings.upload_template(invalid)
    assert current_asset("vietnam_full_workbook").path.read_bytes() == source.read_bytes()

    replacement = nine_sheet_template(tmp_path / "新版九表.xlsx", marker="second")
    assert settings.upload_template(replacement)["template"]["file_name"] == replacement.name
    assert previous_asset("vietnam_full_workbook").path.read_bytes() == source.read_bytes()


def test_nine_sheet_template_export_requires_saved_version(tmp_path):
    with pytest.raises(ValueError, match="模板"):
        settings.export_template(tmp_path / "missing.xlsx")


def test_1904_upload_preserves_current_previous_and_export(tmp_path):
    previous = nine_sheet_template(tmp_path / "previous.xlsx", marker="previous")
    current = nine_sheet_template(tmp_path / "current.xlsx", marker="current")
    settings.upload_template(previous)
    settings.upload_template(current)
    versions = (current_asset("vietnam_full_workbook"), previous_asset("vietnam_full_workbook"))
    before = {version.path: sha256(version.path.read_bytes()).hexdigest() for version in versions}
    unsupported = nine_sheet_template(tmp_path / "1904.xlsx", marker="unsupported")
    book = load_workbook(unsupported)
    book.epoch = CALENDAR_MAC_1904
    book.save(unsupported)
    book.close()
    with pytest.raises(ValueError, match="1904.*1900"):
        settings.upload_template(unsupported)
    assert (current_asset("vietnam_full_workbook"), previous_asset("vietnam_full_workbook")) == versions
    assert all(sha256(path.read_bytes()).hexdigest() == digest for path, digest in before.items())
    exported = tmp_path / "exported.xlsx"
    settings.export_template(exported)
    assert exported.read_bytes() == current.read_bytes()


@pytest.mark.parametrize("failure", ["copy", "partial_copy", "short_copy", "fsync", "metadata", "before_replace", "previous_replace"])
@pytest.mark.parametrize("same_name", [False, True])
def test_failed_template_promotion_preserves_both_versions_and_can_retry(
    tmp_path, monkeypatch, failure, same_name,
):
    oldest = nine_sheet_template(tmp_path / "oldest.xlsx", marker="previous")
    old = nine_sheet_template(tmp_path / "current.xlsx", marker="current")
    settings.upload_template(oldest)
    settings.upload_template(old)
    incoming_dir = tmp_path / "incoming"
    incoming_dir.mkdir()
    incoming = nine_sheet_template(incoming_dir / (old.name if same_name else "new.xlsx"), marker="new")
    root = input_assets.slot_dir("vietnam_full_workbook")

    def versions():
        return tuple((version.file_name, sha256(version.path.read_bytes()).hexdigest())
                     for version in (current_asset("vietnam_full_workbook"), previous_asset("vietnam_full_workbook")))

    before = versions()
    replace_file = input_assets.os.replace
    copy_file = input_assets.shutil.copyfile
    failed = False

    def fail_copy(source, destination, *args, **kwargs):
        if not Path(destination).is_relative_to(root):
            return copy_file(source, destination, *args, **kwargs)
        if failure in ("partial_copy", "short_copy"):
            Path(destination).write_bytes(Path(source).read_bytes()[:100])
        if failure == "short_copy":
            return str(destination)
        raise OSError(5, "injected copy I/O error")

    def fail_replace(source, destination):
        nonlocal failed
        destination = Path(destination)
        target = "current" if failure == "before_replace" else "previous"
        if not failed and (destination == root / target or destination.parent == root / target):
            failed = True
            raise OSError(5, f"injected {target} replace error")
        return replace_file(source, destination)

    for _ in range(2):
        failed = False
        with monkeypatch.context() as fault:
            if failure in ("copy", "partial_copy", "short_copy"):
                fault.setattr(input_assets.shutil, "copyfile", fail_copy)
            elif failure == "metadata":
                fault.setattr(input_assets.shutil, "copystat", fail_copy)
            elif failure == "fsync":
                def fail_sync(_descriptor):
                    raise OSError(5, "injected fsync I/O error")
                fault.setattr(input_assets.os, "fsync", fail_sync)
            else:
                fault.setattr(input_assets.os, "replace", fail_replace)
            expected = input_assets.InputAssetError if failure == "short_copy" else OSError
            with pytest.raises(expected, match="changed while staging" if failure == "short_copy" else "injected"):
                settings.upload_template(incoming)
        assert versions() == before
        assert {path.name for path in root.iterdir()} == {"current", "previous"}
        state = settings.read_state()
        assert state["template_error"] is None
        exported = tmp_path / "after-failure.xlsx"
        settings.export_template(exported)
        assert exported.read_bytes() == old.read_bytes()

    assert settings.upload_template(incoming)["template_error"] is None
    assert current_asset("vietnam_full_workbook").path.read_bytes() == incoming.read_bytes()
    assert previous_asset("vietnam_full_workbook").path.read_bytes() == old.read_bytes()
    assert {path.name for path in root.iterdir()} == {"current", "previous"}


def test_defaults_persist_and_reload_in_another_process():
    state = settings.read_state()
    assert state["parameters"] == settings.config_json(RecommendationConfig())
    assert "sku_map" not in state and "sku_map_error" not in state
    raw = {**state["parameters"], "exchange_rate": "4100", "sales_weight_7d": "1", "sales_weight_15d": "0", "sales_weight_30d": "0"}
    settings.save_parameters(raw)
    child = subprocess.run([sys.executable, "-I", "-B", "-m", "lxeskill", "vietnam", "replenishment", "settings", "--action", "read"], capture_output=True, text=True)
    assert child.returncode == 0, child.stderr
    assert json.loads(child.stdout)["data"]["data"]["parameters"] == raw


@pytest.mark.parametrize("change", [
    {"sales_weight_7d": "0.5"}, {"sales_weight_7d": "-0.1"}, {"exchange_rate": "0"},
    {"day_adjustment_15d": "-1"}, {"day_adjustment_7d": "NaN"}, {"exchange_rate": "1e9999999999"},
    {"exchange_rate": "1e-9999"}, {"exchange_rate": "0.1234567890123456"}, {"sales_weight_7d": 0.6},
    {"extra": "1"},
])
def test_invalid_save_preserves_previous_file(change):
    old = settings.read_state()["parameters"]
    file = settings.data_directory() / "parameters.json"
    before = file.read_bytes()
    with pytest.raises((ValueError, ArithmeticError)):
        settings.save_parameters({**old, **change})
    assert file.read_bytes() == before


def test_corrupt_existing_configuration_is_reported_not_reset():
    settings.read_state()
    file = settings.data_directory() / "parameters.json"
    file.write_text("{broken")
    state = settings.read_state()
    assert "JSONDecodeError" in state["parameters_error"]
    assert state["parameters"] is None
    assert file.read_text() == "{broken"


def test_config_snapshot_is_stable_during_run(tmp_path):
    with settings.run_inputs() as (config, source):
        assert source == str(settings.data_directory() / "parameters.json")
        settings.save_parameters(settings.config_json(replace(config, exchange_rate=Decimal(4200))))
        assert config.exchange_rate == Decimal(3900)
    with settings.run_inputs() as (updated, _):
        assert updated.exchange_rate == Decimal(4200)


def test_invalid_configuration_stops_before_reports(tmp_path, monkeypatch):
    template = nine_sheet_template(tmp_path / "template.xlsx")
    settings.read_state()
    path = settings.data_directory() / "parameters.json"
    path.write_text('{"exchange_rate":"0"}')
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: pytest.fail("Source loading must not run"))
    with pytest.raises(ValueError):
        workflow.generate_current_vietnam_recommendation(
            sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx", template=str(template),
        )


@pytest.mark.parametrize("old_parameters", ["{broken", json.dumps({"vietnam_recommendation": {"exchange_rate": "4200"}})])
def test_initialization_uses_only_skill_data_directory(tmp_path, old_parameters):
    app = tmp_path / "app"
    (app / "config").mkdir(parents=True)
    (app / "config" / "settings.json").write_text(old_parameters)
    state = settings.read_state()
    assert state["parameters"] == settings.config_json(RecommendationConfig())
    assert state["parameters_error"] is None
    assert "sku_map" not in state and "sku_map_error" not in state
    assert (app / "config" / "settings.json").read_text() == old_parameters


def test_saved_template_corruption_is_reported_and_upload_can_repair(tmp_path):
    state = settings.upload_template(nine_sheet_template(tmp_path / "initial.xlsx"))
    target = Path(state["template"]["path"])
    target.write_bytes(b"broken")
    state = settings.read_state()
    assert "BadZipFile" in state["template_error"]
    assert state["template"] is None
    assert target.read_bytes() == b"broken"
    replacement = nine_sheet_template(tmp_path / "replacement.xlsx", marker="repaired")
    state = settings.upload_template(replacement)
    assert state["template_error"] is None
    assert Path(state["template"]["path"]).read_bytes() == replacement.read_bytes()


def test_settings_cli_and_saved_template_read_preserve_real_errors_without_replacing_data(tmp_path):
    state = settings.upload_template(nine_sheet_template(tmp_path / "good.xlsx"))
    target = Path(state["template"]["path"])
    parameters = settings.data_directory() / "parameters.json"
    before = {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in (target, parameters)}
    invalid = nine_sheet_template(tmp_path / "中文 不完整.xlsx")
    book = load_workbook(invalid)
    book["越南备货清单"]["B1"] = "错误标题"
    book.save(invalid)
    book.close()
    child = subprocess.run(
        [sys.executable, "-I", "-B", "-m", "lxeskill", "vietnam", "replenishment", "settings", "--action", "upload", "--path", str(invalid)],
        capture_output=True, text=True, encoding="utf-8",
    )
    assert child.returncode == 4, child.stderr
    result = json.loads(child.stdout)
    assert result["ok"] is False
    assert "第 2 列表头应为 判断热销" in result["error"]["message"]
    assert {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in before} == before
    assert not list(target.parent.parent.glob(".promote-*"))

    # A template edited outside the app uses exactly the same validation on read.
    target.write_bytes(invalid.read_bytes())
    state = settings.read_state()
    assert state["template"] is None
    assert state["template_error"] == result["error"]["message"]
    assert target.read_bytes() == invalid.read_bytes()
    assert (parameters.read_bytes(), parameters.stat().st_mtime_ns) == before[parameters]


def test_template_snapshot_lock_covers_selection_and_copy_only(tmp_path, monkeypatch):
    import time

    a = nine_sheet_template(tmp_path / "A.xlsx", marker="A")
    b = nine_sheet_template(tmp_path / "B.xlsx", marker="B")
    settings.upload_template(a)
    before = current_asset("vietnam_full_workbook")
    ready = tmp_path / "upload-at-lock"
    done = tmp_path / "upload-completed"
    destination = tmp_path / "run-snapshot.xlsx"
    child = None
    copy = settings.shutil.copyfile

    def copy_with_concurrent_upload(source, target):
        nonlocal child
        child = subprocess.Popen(
            [sys.executable, "-I", "-B", "-c", '''
import sys
from pathlib import Path
from contextlib import contextmanager
from services.vietnam_replenishment import settings
from shared.process_lock import _try_lock, _unlock
locked = settings._locked
@contextmanager
def signal_lock_attempt():
    with (settings.data_directory() / ".settings.lock").open("r+b") as stream:
        acquired = _try_lock(stream)
        Path(sys.argv[2]).write_text("unexpectedly-acquired" if acquired else "blocked")
        if acquired:
            _unlock(stream)
    with locked() as root:
        yield root
settings._locked = signal_lock_attempt
settings.upload_template(Path(sys.argv[1]))
Path(sys.argv[3]).write_text("completed")
''', str(b), str(ready), str(done)],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
        )
        deadline = time.monotonic() + 10
        while not ready.exists() and time.monotonic() < deadline:
            time.sleep(0.01)
        assert ready.exists(), "Uploader did not reach the existing settings lock"
        assert ready.read_text() == "blocked"
        assert child.poll() is None and not done.exists()
        assert current_asset("vietnam_full_workbook") == before
        return copy(source, target)

    monkeypatch.setattr(settings.shutil, "copyfile", copy_with_concurrent_upload)
    try:
        source, provenance = settings.snapshot_template(None, destination)
        assert source == before.path and provenance["file_name"] == "A.xlsx"
        assert destination.read_bytes() == a.read_bytes()
        # The snapshot has returned, so the uploader must now be able to finish.
        out, err = child.communicate(timeout=30)
        assert child.returncode == 0, (out, err)
        assert done.exists()
        assert destination.read_bytes() == a.read_bytes()
        assert current_asset("vietnam_full_workbook").path.read_bytes() == b.read_bytes()
        assert previous_asset("vietnam_full_workbook").path.read_bytes() == a.read_bytes()
    finally:
        if child is not None and child.poll() is None:
            child.kill()
            child.communicate()
