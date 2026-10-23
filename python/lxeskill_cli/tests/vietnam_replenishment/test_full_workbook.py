"""The exported nine-sheet workbook retains the user's four business tabs."""

from __future__ import annotations

from datetime import datetime
from hashlib import sha256
from importlib.resources import files
import os
from pathlib import Path
from xml.dom import minidom
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900
import pytest

from services.vietnam_replenishment.asset_contract import AssetContractError, load_template_sku_parameters
from services.vietnam_replenishment.full_workbook import (
    MAIN, CURRENT, REQUIRED_ORDER, UNCHANGED, FullWorkbookError, _sheet_parts,
    combine_full_workbook, validate_full_template,
)


def _inputs(directory: Path) -> tuple[Path, Path]:
    resource = files("services.vietnam_replenishment").joinpath("resources/skeleton.xlsx")
    with resource.open("rb") as handle:
        original = load_workbook(handle)
    for index, name in enumerate(UNCHANGED, 1):
        sheet = original.create_sheet(name, index)
        sheet["A1"] = f"original {name}"
        sheet["A1"].font = Font(color="FF336699")
        sheet["B2"] = "=越南备货清单!H2"
        sheet["C3"] = "SKU-PREVIOUS"
    template = directory / "template.xlsx"
    original.save(template)
    original.close()

    with resource.open("rb") as handle:
        fresh = load_workbook(handle)
    fresh["越南备货清单"]["E2"] = "SKU-CURRENT"
    fresh["越南备货清单"]["E2"].font = Font(bold=True, color="FF112233")
    fresh["雅仓库存"]["B2"] = "SKU-CURRENT"
    current = directory / "current.xlsx"
    fresh.save(current)
    fresh.close()
    return template, current


def test_four_unchanged_sheets_keep_exact_xml_and_current_five_are_inserted(tmp_path: Path) -> None:
    template, current = _inputs(tmp_path)
    before = sha256(template.read_bytes()).digest()
    output = tmp_path / "final.xlsx"

    combine_full_workbook(template, current, output)

    assert sha256(template.read_bytes()).digest() == before
    with ZipFile(template) as original, ZipFile(output) as result:
        names = _sheet_parts(original)
        assert result.testzip() is None
        for name in UNCHANGED:
            assert result.read(names[name]) == original.read(names[name])

    book = load_workbook(output, read_only=False, data_only=False)
    try:
        assert tuple(book.sheetnames) == REQUIRED_ORDER
        assert book["越南备货清单"]["E2"].value == "SKU-CURRENT"
        assert book["越南备货清单"]["E2"].font.bold
        assert book["雅仓库存"]["B2"].value == "SKU-CURRENT"
        for name in UNCHANGED:
            assert book[name]["A1"].value == f"original {name}"
            assert book[name]["A1"].font.color.rgb == "FF336699"
            assert book[name]["B2"].value == "=越南备货清单!H2"
            assert book[name]["C3"].value == "SKU-PREVIOUS"
    finally:
        book.close()


def test_bad_sheet_order_is_rejected_without_partial_output(tmp_path: Path) -> None:
    template, current = _inputs(tmp_path)
    book = load_workbook(template)
    book.move_sheet("改库存", offset=1)
    book.save(template)
    book.close()
    output = tmp_path / "final.xlsx"

    with pytest.raises(FullWorkbookError, match="指定顺序的九表"):
        combine_full_workbook(template, current, output)
    assert not output.exists()


@pytest.mark.parametrize("flag", ["1", "true"])
def test_1904_template_is_rejected_during_validation(tmp_path, flag):
    template, _ = _inputs(tmp_path)
    with ZipFile(template) as archive:
        root = ET.fromstring(archive.read("xl/workbook.xml"))
    root.find(f"{{{MAIN}}}workbookPr").set("date1904", flag)
    _rewrite_parts(template, {"xl/workbook.xml": ET.tostring(root)})
    book = load_workbook(template, read_only=True)
    assert book.epoch == CALENDAR_MAC_1904
    book.close()
    with pytest.raises(FullWorkbookError, match="1904.*1900"):
        validate_full_template(template)


@pytest.mark.parametrize("template_epoch,current_epoch", [
    (CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900),
    (CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904),
    (CALENDAR_MAC_1904, CALENDAR_MAC_1904),
])
def test_unsupported_date_systems_never_publish_shifted_dates(tmp_path, template_epoch, current_epoch):
    template, current = _inputs(tmp_path)
    for path, epoch in ((template, template_epoch), (current, current_epoch)):
        book = load_workbook(path)
        book.epoch = epoch
        if path == current:
            book["库存商品信息"]["K2"] = datetime(2026, 10, 1)
        book.save(path)
        book.close()
    before = {path: path.read_bytes() for path in (template, current)}
    output = tmp_path / "final.xlsx"
    with pytest.raises(FullWorkbookError, match="日期系统"):
        combine_full_workbook(template, current, output)
    assert not output.exists()
    assert all(path.read_bytes() == data for path, data in before.items())


@pytest.mark.parametrize("flag", [None, "0", "false"])
def test_1900_dates_keep_their_values_and_original_business_parts(tmp_path, flag):
    template, current = _inputs(tmp_path)
    for path in (template, current):
        book = load_workbook(path)
        if path == current:
            book["库存商品信息"]["K2"] = datetime(2026, 10, 1)
        else:
            for name in UNCHANGED:
                book[name]["A2"] = datetime(2020, 2, 3)
        book.save(path)
        book.close()
        with ZipFile(path) as archive:
            root = ET.fromstring(archive.read("xl/workbook.xml"))
        properties = root.find(f"{{{MAIN}}}workbookPr")
        if flag is None:
            properties.attrib.pop("date1904", None)
        else:
            properties.set("date1904", flag)
        _rewrite_parts(path, {"xl/workbook.xml": ET.tostring(root)})
    validate_full_template(template)
    output = tmp_path / "final.xlsx"
    combine_full_workbook(template, current, output)
    with ZipFile(template) as original, ZipFile(output) as exported:
        parts = _sheet_parts(original)
        assert all(original.read(parts[name]) == exported.read(parts[name]) for name in UNCHANGED)
    book = load_workbook(output, data_only=True)
    assert book.epoch == CALENDAR_WINDOWS_1900
    assert book["库存商品信息"]["K2"].value == datetime(2026, 10, 1)
    assert all(book[name]["A2"].value == datetime(2020, 2, 3) for name in UNCHANGED)
    book.close()


def _builtin_template(path: Path) -> Path:
    """A valid nine-sheet workbook with built-in formats and no optional numFmts."""
    book = Workbook()
    book.remove(book.active)
    for name in REQUIRED_ORDER:
        book.create_sheet(name)
    main = book[REQUIRED_ORDER[0]]
    for cell, value in {"B1": "判断热销", "E1": "SKU", "G1": "成本", "AE1": "跨境价", "AJ1": "折扣价",
                        "B2": 2, "E2": "VN-A", "G2": 12, "AE2": 22, "AJ2": 16}.items():
        main[cell] = value
    for name in UNCHANGED:
        book[name]["A1"] = f"original {name}"
        book[name]["A2"] = 1234.5
        book[name]["A2"].number_format = "0.00"
        book[name]["B2"] = "=越南备货清单!H2"
    book.save(path)
    book.close()
    staged = path.with_suffix(".tmp")
    with ZipFile(path) as source, ZipFile(staged, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/styles.xml":
                styles = ET.fromstring(data)
                formats = styles.find(f"{{{MAIN}}}numFmts")
                assert formats is not None and len(formats) == 0
                styles.remove(formats)
                data = ET.tostring(styles)
            target.writestr(item, data)
    staged.replace(path)
    return path


def _style_namespace(path: Path, prefix: str | None, section_default: str | None = None) -> None:
    """Use an explicit SpreadsheetML binding without changing style semantics."""
    staged = path.with_suffix(".tmp")
    with ZipFile(path) as source, ZipFile(staged, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/styles.xml":
                document = minidom.parseString(data)
                try:
                    for node in list(document.getElementsByTagNameNS(MAIN, "*")):
                        document.renameNode(node, MAIN, f"{prefix}:{node.localName}" if prefix else node.localName)
                    root = document.documentElement
                    for attribute in list(root.attributes.values()):
                        if attribute.value == MAIN and attribute.name.startswith("xmlns"):
                            root.removeAttribute(attribute.name)
                    root.setAttribute(f"xmlns:{prefix}" if prefix else "xmlns", MAIN)
                    if section_default is not None:
                        root.setAttribute("xmlns", MAIN)
                        for name in ("fonts", "cellXfs"):
                            document.getElementsByTagNameNS(MAIN, name)[0].setAttribute("xmlns", section_default)
                    data = document.toxml(encoding="utf-8")
                finally:
                    document.unlink()
                assert all(node.tag.startswith(f"{{{MAIN}}}") for node in ET.fromstring(data).iter())
            target.writestr(item, data)
    staged.replace(path)


@pytest.mark.parametrize("template_prefix,section_default", [
    (None, None), ("ns0", None), ("sheet", "urn:unused-style-default"),
])
@pytest.mark.parametrize("calculated_prefix", [None, "calc"])
@pytest.mark.parametrize("missing_formats", [False, True])
def test_style_merge_accepts_optional_numfmts_and_keeps_all_style_references_valid(
    tmp_path, missing_formats, template_prefix, section_default, calculated_prefix,
):
    template, current = _inputs(tmp_path)
    if missing_formats:
        _builtin_template(template)
    book = load_workbook(current)
    book[CURRENT[0]].conditional_formatting.add(
        "E2", CellIsRule(operator="equal", formula=['"SKU-CURRENT"'], font=Font(bold=True)),
    )
    book.save(current)
    book.close()
    _style_namespace(template, template_prefix, section_default)
    _style_namespace(current, calculated_prefix)
    validate_full_template(template)
    output = tmp_path / "merged.xlsx"
    combine_full_workbook(template, current, output)

    with ZipFile(template) as original, ZipFile(output) as merged:
        before = ET.fromstring(original.read("xl/styles.xml"))
        after = ET.fromstring(merged.read("xl/styles.xml"))
        assert all(node.tag.startswith(f"{{{MAIN}}}") for node in after.iter())
        assert [node.tag for node in after][:2] == [f"{{{MAIN}}}numFmts", f"{{{MAIN}}}fonts"]
        assert [node.tag for node in after].index(f"{{{MAIN}}}dxfs") < [node.tag for node in after].index(f"{{{MAIN}}}tableStyles")
        formats = after.find(f"{{{MAIN}}}numFmts")
        assert formats is not None and all(node.tag == f"{{{MAIN}}}numFmt" for node in formats)
        identifiers = [int(node.attrib["numFmtId"]) for node in formats]
        assert len(identifiers) == len(set(identifiers))
        assert int(formats.attrib["count"]) == len(identifiers)
        for section in ("cellStyleXfs", "cellXfs"):
            old = before.find(f"{{{MAIN}}}{section}")
            new = after.find(f"{{{MAIN}}}{section}")
            assert [ET.tostring(node) for node in new[:len(old)]] == [ET.tostring(node) for node in old]
            for node in new:
                assert node.tag == f"{{{MAIN}}}xf"
                number = int(node.attrib.get("numFmtId", 0))
                assert number < 164 or number in identifiers
                for attribute, name in (("fontId", "fonts"), ("fillId", "fills"), ("borderId", "borders"), ("xfId", "cellStyleXfs")):
                    if attribute in node.attrib:
                        assert int(node.attrib[attribute]) < len(after.find(f"{{{MAIN}}}{name}"))
        parts = _sheet_parts(original)
        assert all(original.read(parts[name]) == merged.read(parts[name]) for name in UNCHANGED)

    result, reference = load_workbook(output), load_workbook(current)
    try:
        assert tuple(result.sheetnames) == REQUIRED_ORDER
        for name in CURRENT:
            assert result[name].max_row == reference[name].max_row
            assert result[name].max_column == reference[name].max_column
            for row in reference[name]:
                for cell in row:
                    assert result[name][cell.coordinate].number_format == cell.number_format
        if missing_formats:
            assert result[UNCHANGED[0]]["A2"].number_format == "0.00"
    finally:
        result.close()
        reference.close()


@pytest.mark.parametrize("template_prefix,calculated_prefix,missing_formats", [
    (None, None, False), ("ns0", None, True),
    ("sheet", "calc", False), ("sheet", "calc", True),
])
@pytest.mark.skipif(
    not (os.environ.get("LXE_OFFICE_NODE") and os.environ.get("LXE_OFFICE_CLI")),
    reason="Host Office Kit paths are not configured",
)
def test_template_styles_can_be_reopened_by_office(tmp_path, template_prefix, calculated_prefix, missing_formats):
    from services.vietnam_replenishment.recalculation import recalculate_with_office

    template, current = _inputs(tmp_path)
    if missing_formats:
        _builtin_template(template)
    _style_namespace(template, template_prefix)
    _style_namespace(current, calculated_prefix)
    validate_full_template(template)
    merged = tmp_path / "merged.xlsx"
    combine_full_workbook(template, current, merged)
    reopened = tmp_path / "office.xlsx"
    recalculate_with_office(merged, reopened)
    book = load_workbook(reopened, data_only=True)
    try:
        assert tuple(book.sheetnames) == REQUIRED_ORDER
        assert book[REQUIRED_ORDER[0]]["E2"].value == "SKU-CURRENT"
        if missing_formats:
            for name in UNCHANGED:
                assert book[name]["A2"].value == 1234.5
                assert book[name]["A2"].number_format == "0.00"
    finally:
        book.close()


def test_template_parameters_use_saved_formula_result_without_a_sku_map(tmp_path: Path) -> None:
    template, _ = _inputs(tmp_path)
    book = load_workbook(template)
    main = book["越南备货清单"]
    for cell, value in {"B2": 2, "E2": "SKU-CURRENT", "G2": 10,
                        "AE2": "=G2*20", "AJ2": None}.items():
        main[cell] = value
    book.save(template)
    book.close()
    cached = tmp_path / "cached.xlsx"
    with ZipFile(template) as source, ZipFile(cached, "w") as target:
        part = _sheet_parts(source)["越南备货清单"]
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == part:
                xml = ET.fromstring(payload)
                cell = xml.find(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c[@r='AE2']")
                cell.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v").text = "200"
                payload = ET.tostring(xml)
            target.writestr(info, payload)
    actual = load_template_sku_parameters(cached, {"SKU-CURRENT", "SKU-NEW"})
    assert actual["SKU-CURRENT"].cost == 10
    assert actual["SKU-CURRENT"].cross_border_price == 200
    assert actual["SKU-CURRENT"].discount_price is None
    assert "SKU-NEW" not in actual


def test_template_conflicting_duplicate_current_sku_is_rejected(tmp_path: Path) -> None:
    template, _ = _inputs(tmp_path)
    book = load_workbook(template)
    main = book["越南备货清单"]
    for row, cost in ((2, 10), (3, 11)):
        for column, value in (("B", 2), ("E", "SKU-CURRENT"), ("G", cost),
                              ("AE", 20), ("AJ", 15)):
            main[f"{column}{row}"] = value
    book.save(template)
    book.close()
    with pytest.raises(AssetContractError, match="参数冲突"):
        load_template_sku_parameters(template, {"SKU-CURRENT"})


# These inputs are built from Workbook, independently of the packaged skeleton
# and the earlier namespace-renaming helper that missed the review counterexamples.
def _independent_inputs(directory):
    paths = []
    for label, names in (("business", REQUIRED_ORDER), ("calculated", CURRENT)):
        book = Workbook()
        book.remove(book.active)
        for name in names:
            book.create_sheet(name)["A1"] = name
        main = book[CURRENT[0]]
        for cell, value in {"B1": "判断热销", "E1": "SKU", "G1": "成本", "AE1": "跨境价", "AJ1": "折扣价",
                            "B2": 2, "E2": "001-A", "G2": 12, "AE2": 22, "AJ2": 16, "H2": "=G2*2"}.items():
            main[cell] = value
        main["G2"].number_format = "0.00000"
        if names == CURRENT:
            main.conditional_formatting.add("H2", CellIsRule(operator="greaterThan", formula=["10"], font=Font(bold=True)))
        path = directory / f"{label}.xlsx"
        book.save(path)
        book.close()
        paths.append(path)
    return tuple(paths)


def _rewrite_parts(path, updates):
    staged = path.with_suffix(".rewrite")
    with ZipFile(path) as source, ZipFile(staged, "w") as target:
        for item in source.infolist():
            target.writestr(item, updates.get(item.filename, source.read(item.filename)))
        for name, raw in updates.items():
            if name not in source.namelist():
                target.writestr(name, raw)
    staged.replace(path)


def _one_shared_string(path, prefix):
    from services.vietnam_replenishment.full_workbook import CONTENT, PACKAGE_REL, REL

    def tag(local):
        return f"{prefix}:{local}" if prefix else local

    declaration = f'xmlns:{prefix}' if prefix else 'xmlns'
    sst = f'<{tag("sst")} {declaration}="{MAIN}" count="1" uniqueCount="1"><{tag("si")}><{tag("r")}><{tag("rPr")}><{tag("b")}/></{tag("rPr")}><{tag("t")}>001-A</{tag("t")}></{tag("r")}></{tag("si")}></{tag("sst")}>'
    with ZipFile(path) as archive:
        part = _sheet_parts(archive)[CURRENT[0]]
        sheet = ET.fromstring(archive.read(part))
        cell = sheet.find(f".//{{{MAIN}}}c[@r='E2']")
        cell.clear()
        cell.attrib.update(r="E2", t="s")
        ET.SubElement(cell, f"{{{MAIN}}}v").text = "0"
        types = ET.fromstring(archive.read("[Content_Types].xml"))
        ET.SubElement(types, f"{{{CONTENT}}}Override", PartName="/xl/sharedStrings.xml",
                      ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml")
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        ET.SubElement(relations, f"{{{PACKAGE_REL}}}Relationship", Id="independentSst", Type=REL+"/sharedStrings", Target="sharedStrings.xml")
    _rewrite_parts(path, {part: ET.tostring(sheet), "xl/sharedStrings.xml": sst.encode(),
                          "[Content_Types].xml": ET.tostring(types), "xl/_rels/workbook.xml.rels": ET.tostring(relations)})


@pytest.mark.parametrize("target_prefix,source_prefix", [("ss", ""), ("", "result"), ("one", "two")])
def test_independent_shared_string_qnames_and_rich_text(tmp_path, target_prefix, source_prefix):
    template, current = _independent_inputs(tmp_path)
    _one_shared_string(template, target_prefix)
    _one_shared_string(current, source_prefix)
    validate_full_template(template)
    output = tmp_path / "merged.xlsx"
    combine_full_workbook(template, current, output)
    with ZipFile(output) as archive:
        root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
        assert all(node.tag.startswith(f"{{{MAIN}}}") for node in root.iter())
        assert len(root.findall(f"{{{MAIN}}}si")) == 2
    book = load_workbook(output)
    assert book[CURRENT[0]]["E2"].value == "001-A"
    book.close()


def test_independent_prefixed_package_declarations(tmp_path):
    from services.vietnam_replenishment.full_workbook import CONTENT, PACKAGE_REL, REL

    template, current = _independent_inputs(tmp_path)
    _one_shared_string(current, "fresh")
    updates = {}
    for part, uri, prefix in (("[Content_Types].xml", CONTENT, "ct"), ("xl/_rels/workbook.xml.rels", PACKAGE_REL, "package")):
        with ZipFile(template) as archive:
            document = minidom.parseString(archive.read(part))
        for node in list(document.getElementsByTagNameNS(uri, "*")):
            document.renameNode(node, uri, prefix+":"+node.localName)
        document.documentElement.removeAttribute("xmlns")
        document.documentElement.setAttribute("xmlns:"+prefix, uri)
        updates[part] = document.toxml(encoding="utf-8")
        document.unlink()
    _rewrite_parts(template, updates)
    validate_full_template(template)
    output = tmp_path / "merged.xlsx"
    combine_full_workbook(template, current, output)
    with ZipFile(output) as archive:
        types = ET.fromstring(archive.read("[Content_Types].xml"))
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        assert all(n.tag in (f"{{{CONTENT}}}Default", f"{{{CONTENT}}}Override") for n in types)
        assert all(n.tag == f"{{{PACKAGE_REL}}}Relationship" for n in relations)
        assert len([n for n in types if n.get("PartName") == "/xl/sharedStrings.xml"]) == 1
        assert len([n for n in relations if n.get("Type") == REL+"/sharedStrings" and n.get("Target") == "sharedStrings.xml"]) == 1
    book = load_workbook(output)
    assert book[CURRENT[0]]["E2"].value == "001-A"
    book.close()


@pytest.mark.parametrize("local_rebinding", [False, True])
@pytest.mark.parametrize("inherited_ignorable", [False, True])
def test_independent_style_attributes_preserve_namespace_scope(tmp_path, local_rebinding, inherited_ignorable):
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    extension = "urn:pr90-independent-ignorable"
    template, current = _independent_inputs(tmp_path)
    with ZipFile(current) as archive:
        document = minidom.parseString(archive.read("xl/styles.xml"))
    root = document.documentElement
    root.setAttribute("xmlns:mc", mc)
    root.setAttribute("xmlns:q", extension)
    font = document.getElementsByTagNameNS(MAIN, "font")[0]
    if local_rebinding:
        font.setAttribute("xmlns:q", extension)
        child = document.createElementNS(MAIN, "q:b")
        child.setAttribute("xmlns:q", MAIN)
        font.appendChild(child)
    (root if inherited_ignorable else font).setAttributeNS(mc, "mc:Ignorable", "q")
    font.setAttributeNS(extension, "q:marker", "keep-uri")
    _rewrite_parts(current, {"xl/styles.xml": document.toxml(encoding="utf-8")})
    document.unlink()
    book = load_workbook(current)
    book.close()
    output = tmp_path / "merged.xlsx"
    combine_full_workbook(template, current, output)
    with ZipFile(output) as archive:
        raw = archive.read("xl/styles.xml")
        styles = ET.fromstring(raw)
        imported = styles.find(f"{{{MAIN}}}fonts")[1]
        assert imported.attrib[f"{{{extension}}}marker"] == "keep-uri"
        assert imported.attrib[f"{{{mc}}}Ignorable"] == "q"
        assert all(n.tag.startswith(f"{{{MAIN}}}") for n in styles.iter())
        document = minidom.parseString(raw)
        imported = document.getElementsByTagNameNS(MAIN, "fonts")[0].childNodes[1]
        assert imported.getAttribute("xmlns:q") == extension
        if local_rebinding:
            bold = imported.getElementsByTagNameNS(MAIN, "b")[-1]
            assert bold.getAttribute("xmlns:q") == MAIN
        document.unlink()
    book = load_workbook(output)
    book.close()


@pytest.mark.parametrize("table_styles", [False, True])
@pytest.mark.parametrize("colors", [False, True])
def test_independent_dxfs_schema_order(tmp_path, table_styles, colors):
    template, current = _independent_inputs(tmp_path)
    with ZipFile(template) as archive:
        root = ET.fromstring(archive.read("xl/styles.xml"))
    for name, present in (("tableStyles", table_styles), ("colors", colors), ("dxfs", False)):
        node = root.find(f"{{{MAIN}}}{name}")
        if not present and node is not None:
            root.remove(node)
        elif present and node is None:
            ET.SubElement(root, f"{{{MAIN}}}{name}")
    _rewrite_parts(template, {"xl/styles.xml": ET.tostring(root)})
    output = tmp_path / "merged.xlsx"
    combine_full_workbook(template, current, output)
    with ZipFile(output) as archive:
        root = ET.fromstring(archive.read("xl/styles.xml"))
        order = ("numFmts", "fonts", "fills", "borders", "cellStyleXfs", "cellXfs", "cellStyles", "dxfs", "tableStyles", "colors", "extLst")
        positions = [order.index(n.tag.removeprefix(f"{{{MAIN}}}")) for n in root]
        assert positions == sorted(positions)
        assert root.find(f"{{{MAIN}}}dxfs") is not None


@pytest.mark.parametrize("corruption", ["sst_namespace", "package_namespace", "style_order", "style_reference"])
def test_final_validation_rejects_confirmed_xml_corruption(tmp_path, monkeypatch, corruption):
    from services.vietnam_replenishment import full_workbook as merger

    template, current = _independent_inputs(tmp_path)
    _one_shared_string(current, "")
    if corruption == "sst_namespace":
        append = merger._append_strings
        def broken_strings(*args):
            data, offset, new = append(*args)
            return data.replace(b"<r>", b'<r xmlns="">'), offset, new
        monkeypatch.setattr(merger, "_append_strings", broken_strings)
    elif corruption == "package_namespace":
        declarations = merger._string_part_declarations
        def broken_declarations(*args):
            parts = declarations(*args)
            parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(b"<Override ", b'<Override xmlns="" ', 1)
            return parts
        monkeypatch.setattr(merger, "_string_part_declarations", broken_declarations)
    else:
        append = merger._append_styles
        def broken_styles(*args):
            data, offsets = append(*args)
            root = ET.fromstring(data)
            if corruption == "style_order":
                dxfs = root.find(f"{{{MAIN}}}dxfs")
                root.remove(dxfs)
                root.append(dxfs)
            else:
                root.find(f"{{{MAIN}}}cellXfs")[0].set("fontId", "999999")
            return ET.tostring(root), offsets
        monkeypatch.setattr(merger, "_append_styles", broken_styles)
    output = tmp_path / "bad.xlsx"
    with pytest.raises(FullWorkbookError):
        combine_full_workbook(template, current, output)
    assert not output.exists()


def test_serialization_rejects_attribute_namespace_rebinding():
    from services.vietnam_replenishment.full_workbook import _serialize_xml

    document = minidom.parseString(f'<styleSheet xmlns="{MAIN}" xmlns:q="urn:extension"><fonts><font q:marker="keep"/></fonts></styleSheet>')
    font = document.getElementsByTagNameNS(MAIN, "font")[0]
    # The DOM retains the original attribute URI even after its binding is damaged.
    font.setAttribute("xmlns:q", MAIN)
    try:
        assert font.getAttributeNode("q:marker").namespaceURI == "urn:extension"
        with pytest.raises(FullWorkbookError, match="属性命名空间"):
            _serialize_xml(document)
    finally:
        document.unlink()


def test_inherited_ignorable_resolves_uri_before_local_prefix_rebinding(tmp_path):
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    extension = "urn:pr90-independent-ignorable"
    template, current = _independent_inputs(tmp_path)
    with ZipFile(current) as archive:
        document = minidom.parseString(archive.read("xl/styles.xml"))
    root = document.documentElement
    root.setAttribute("xmlns:mc", mc)
    root.setAttribute("xmlns:q", extension)
    root.setAttributeNS(mc, "mc:Ignorable", "q")
    font = document.getElementsByTagNameNS(MAIN, "font")[0]
    # The inherited Ignorable token refers to EXT at its declaration site,
    # although q now denotes SpreadsheetML inside the imported font.
    font.setAttribute("xmlns:q", MAIN)
    extension_node = document.createElementNS(extension, "aud:marker")
    extension_node.setAttribute("xmlns:aud", extension)
    font.appendChild(extension_node)
    _rewrite_parts(current, {"xl/styles.xml": document.toxml(encoding="utf-8")})
    document.unlink()
    output = tmp_path / "merged.xlsx"
    combine_full_workbook(template, current, output)
    with ZipFile(output) as archive:
        document = minidom.parseString(archive.read("xl/styles.xml"))
    try:
        imported = document.getElementsByTagNameNS(MAIN, "fonts")[0].childNodes[1]
        prefixes = imported.getAttributeNS(mc, "Ignorable").split()
        assert prefixes
        assert {imported.getAttribute("xmlns:"+prefix) for prefix in prefixes} == {extension}
        assert imported.getAttribute("xmlns:q") == MAIN
        assert imported.getElementsByTagNameNS(extension, "marker")[0].namespaceURI == extension
    finally:
        document.unlink()
