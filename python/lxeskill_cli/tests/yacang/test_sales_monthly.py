from __future__ import annotations

import json
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pytest
from openpyxl import Workbook, load_workbook

from services.agent_cli.yacang.export_sales_90d import run as run_sales_90d
from services.agent_cli.yacang.export_sales_monthly import run as run_sales_monthly
from services.yacang.export_intent import normalize_export_intent
from services.yacang.export_workflow import plan_export_workflow
from services.yacang.exports.sales_90d import export_sales_90d, project_sales_90d_sources
from services.yacang.exports.sales_monthly import (
    export_sales_monthly,
    project_sales_monthly_sources,
)
from services.yacang.exports.sales_source import (
    InventorySalesSourceBatch,
    acquire_inventory_sales_sources,
)
from services.yacang.projection import (
    SALES_90D_HEADERS,
    SALES_MONTHLY_HEADERS,
    publish_inventory_sales_workbook,
)
from services.yacang.validation import INVENTORY_SALES_HEADERS
from services.yacang.warehouses import WAREHOUSES


def _write_xlsx(path: Path, warehouse: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(INVENTORY_SALES_HEADERS)
    sheet.append(("SKU-1", "商品", warehouse, 1, 7, 15, 30, 60, 90, 10, 0, 0, 0, 10, 0, "2026-09-01"))
    workbook.save(path)
    workbook.close()


def _headers(path: str) -> tuple[str, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return tuple(str(value or "") for value in next(sheet.iter_rows(values_only=True)))
    finally:
        workbook.close()


def _first_data_row(path: str) -> tuple[Any, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        next(rows)
        return tuple(next(rows))
    finally:
        workbook.close()


def test_single_warehouse_inventory_sales_preserves_source_bytes(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "published.xlsx"
    _write_xlsx(source, "MY8801")
    expected = source.read_bytes()

    row_count = publish_inventory_sales_workbook(
        [("MY8801", source)],
        destination,
    )

    assert row_count == 1
    assert destination.read_bytes() == expected
    assert _headers(str(destination)) == INVENTORY_SALES_HEADERS


def test_multi_warehouse_inventory_sales_merges_complete_rows_in_fixed_order(tmp_path: Path) -> None:
    thailand = tmp_path / "thailand.xlsx"
    malaysia = tmp_path / "malaysia.xlsx"
    destination = tmp_path / "combined.xlsx"
    _write_xlsx(thailand, "TH8802")
    _write_xlsx(malaysia, "MY8801")

    row_count = publish_inventory_sales_workbook(
        [("TH8802", thailand), ("MY8801", malaysia)],
        destination,
    )

    workbook = load_workbook(destination, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["库存动销"]
        rows = list(workbook["库存动销"].iter_rows(values_only=True))
    finally:
        workbook.close()
    assert row_count == 2
    assert tuple(rows[0]) == INVENTORY_SALES_HEADERS
    assert [row[2] for row in rows[1:]] == ["MY8801", "TH8802"]
    assert all(len(row) == 16 for row in rows)


class FakeClient:
    def __init__(self) -> None:
        self.login_calls: list[tuple[str, str]] = []
        self.submissions: list[tuple[int, str, str]] = []
        self.download_calls: list[str] = []
        self.current: tuple[int, str, str] | None = None

    def login(self, mobile: str, password: str) -> None:
        self.login_calls.append((mobile, password))

    def list_downloads(self, *, limit: int = 50) -> list[dict[str, Any]]:
        if self.current is None:
            return []
        warehouse_id, start_date, end_date = self.current
        yacang_timezone = ZoneInfo("Asia/Shanghai")
        start_timestamp = int(
            datetime.combine(date.fromisoformat(start_date), time.min, tzinfo=yacang_timezone).timestamp()
        )
        end_exclusive_timestamp = int(
            datetime.combine(
                date.fromisoformat(end_date) + timedelta(days=1),
                time.min,
                tzinfo=yacang_timezone,
            ).timestamp()
        )
        return [{
            "id": f"{warehouse_id}-{start_date}-{end_date}",
            "name": "库存动销导出",
            "type": "10",
            "path": f"https://oss-accelerate.seaya.cn/example/{warehouse_id}.xlsx",
            "param_where": json.dumps([
                ["warehouse_id", "in", [str(warehouse_id)]],
                ["create_time", ">=", start_timestamp],
                ["create_time", "<", end_exclusive_timestamp],
            ]),
        }]

    def create_inventory_sales_export(self, *, warehouse_id: int, start_date: str, end_date: str) -> str:
        self.current = (warehouse_id, start_date, end_date)
        self.submissions.append(self.current)
        return "fake-request-id"

    def download_xlsx(self, url: str, destination: Path) -> None:
        self.download_calls.append(url)
        warehouse = next(item.code for item in WAREHOUSES if f"/{item.warehouse_id}.xlsx" in url)
        _write_xlsx(destination, warehouse)


def test_planned_source_fetch_is_acquired_once_and_projected_twice(tmp_path: Path) -> None:
    normalized = normalize_export_intent(
        "导出 MY8801 的两种销量",
        today=lambda: date(2026, 9, 14),
    )
    plan = plan_export_workflow(normalized, execution_date="2026-09-14")
    client = FakeClient()

    batch = acquire_inventory_sales_sources(
        plan["source_fetches"],
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=0,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
    )
    monthly = project_sales_monthly_sources(batch, output_dir=tmp_path)
    daily = project_sales_90d_sources(batch, output_dir=tmp_path)

    assert len(plan["source_fetches"]) == 1
    assert client.login_calls == [("account", "password")]
    assert client.submissions == [(26, "2026-09-14", "2026-09-14")]
    assert len(client.download_calls) == 1
    source_fetch_id = plan["source_fetches"][0]["source_fetch_id"]
    assert batch.results[0]["source_fetch_id"] == source_fetch_id
    assert monthly["exports"][0]["source_fetch_id"] == source_fetch_id
    assert daily["exports"][0]["source_fetch_id"] == source_fetch_id
    assert _headers(monthly["xlsx_paths"][0]) == SALES_MONTHLY_HEADERS
    assert _headers(daily["xlsx_paths"][0]) == SALES_90D_HEADERS


def test_failed_shared_source_fans_out_to_both_projections_without_artifacts(tmp_path: Path) -> None:
    source_fetch_id = "inventory-sales-source:MY8801:2026-09-07:2026-09-14"
    batch = InventorySalesSourceBatch(
        created_start_date="2026-09-07",
        created_end_date="2026-09-14",
        results=(
            {
                "warehouse": "MY8801",
                "status": "failed",
                "source_fetch_id": source_fetch_id,
                "error_code": "EXPORT_POLL_TIMEOUT",
                "stage": "轮询导出队列",
                "error_type": "YacangError",
                "error": "导出任务轮询超时",
            },
        ),
    )

    monthly = project_sales_monthly_sources(batch, output_dir=tmp_path)
    daily = project_sales_90d_sources(batch, output_dir=tmp_path)

    assert monthly["xlsx_paths"] == []
    assert daily["xlsx_paths"] == []
    assert monthly["exports"][0]["source_fetch_id"] == source_fetch_id
    assert daily["exports"][0]["source_fetch_id"] == source_fetch_id
    assert monthly["exports"][0]["error_code"] == "EXPORT_POLL_TIMEOUT"
    assert daily["exports"][0]["error_code"] == "EXPORT_POLL_TIMEOUT"


def test_exports_one_7_15_30_workbook_per_warehouse_from_four_source_requests(tmp_path: Path) -> None:
    client = FakeClient()
    result = export_sales_monthly(
        as_of_date="2026-09-13",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=0,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
    )

    assert client.login_calls == [("account", "password")]
    assert result["sales_window_days"] == [7, 15, 30]
    assert result["export_count"] == 4
    assert client.submissions == [
        (warehouse.warehouse_id, "2026-09-13", "2026-09-13")
        for warehouse in WAREHOUSES
    ]
    assert Path(result["xlsx_paths"][0]).name == "雅仓系统-库存动销_马来西亚仓_2026-09-13.xlsx"
    assert Path(result["xlsx_paths"][-1]).name == "雅仓系统-库存动销_越南仓_2026-09-13.xlsx"
    assert all(_headers(path) == SALES_MONTHLY_HEADERS for path in result["xlsx_paths"])
    assert _first_data_row(result["xlsx_paths"][0]) == (
        "SKU-1", "商品", "MY8801", 1, 7, 15, 30, 60, 90, 10, 0, 0, 0, 10, 0, "2026-09-01",
    )


def test_exports_90_day_column_as_four_separate_workbooks_with_dynamic_date(tmp_path: Path) -> None:
    client = FakeClient()
    result = export_sales_90d(
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=0,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
        today=lambda: date(2026, 9, 13),
    )

    assert client.login_calls == [("account", "password")]
    assert client.submissions == [
        (warehouse.warehouse_id, "2026-09-13", "2026-09-13")
        for warehouse in WAREHOUSES
    ]
    assert result["as_of_date"] == "2026-09-13"
    assert result["sales_window_days"] == 90
    assert result["export_count"] == 4
    assert Path(result["xlsx_paths"][0]).name == "雅仓系统-库存动销_马来西亚仓_2026-09-13.xlsx"
    assert Path(result["xlsx_paths"][-1]).name == "雅仓系统-库存动销_越南仓_2026-09-13.xlsx"
    assert all(_headers(path) == SALES_90D_HEADERS for path in result["xlsx_paths"])
    assert _first_data_row(result["xlsx_paths"][0]) == (
        "SKU-1", "商品", "MY8801", 1, 7, 15, 30, 60, 90, 10, 0, 0, 0, 10, 0, "2026-09-01",
    )


def test_second_sales_projection_reuses_the_same_four_source_workbooks(tmp_path: Path) -> None:
    client = FakeClient()
    export_sales_monthly(
        as_of_date="2026-09-13",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=300,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
    )
    result = export_sales_90d(
        as_of_date="2026-09-13",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=300,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
    )

    assert len(client.login_calls) == 1
    assert len(client.submissions) == 4
    assert result["sales_window_days"] == 90
    assert {item["source"] for item in result["exports"]} == {"cache"}


def test_explicit_source_range_and_warehouse_subset_are_preserved(tmp_path: Path) -> None:
    client = FakeClient()
    result = export_sales_monthly(
        start_date="2026-08-01",
        end_date="2026-09-13",
        warehouses=["TH8802", "MY8801"],
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=0,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
    )

    assert client.submissions == [
        (26, "2026-08-01", "2026-09-13"),
        (47, "2026-08-01", "2026-09-13"),
    ]
    assert result["warehouse_count"] == 2
    assert [item["warehouse"] for item in result["exports"]] == ["MY8801", "TH8802"]


@pytest.mark.parametrize(
    ("runner", "business_type"),
    [(run_sales_monthly, "sales-monthly"), (run_sales_90d, "sales-90d")],
)
def test_cli_failure_preserves_high_level_context_without_credentials(
    monkeypatch: pytest.MonkeyPatch,
    runner: Any,
    business_type: str,
) -> None:
    monkeypatch.delenv("LXE_YACANG_MOBILE", raising=False)
    monkeypatch.delenv("LXE_YACANG_PASSWORD", raising=False)
    result = runner({"as_of_date": "2026-09-13"})
    assert result["success"] is False
    assert result["business_type"] == business_type
    assert "缺少雅仓账号或密码" in result["exception"]
    assert "password" not in result["exception"]
