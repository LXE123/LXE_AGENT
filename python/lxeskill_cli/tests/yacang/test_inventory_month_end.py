from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from services.agent_cli.yacang.export_inventory_month_end import run as run_legacy_inventory
from services.yacang.errors import YacangError
from services.yacang.exports.inventory_month_end import (
    InventoryCurrentSnapshotRequest,
    export_inventory_current_snapshot,
    export_inventory_month_end,
    task_matches_inventory_current_snapshot,
    task_matches_inventory_month_end,
)
from services.yacang.validation import INVENTORY_LIST_HEADERS, validate_inventory_list_workbook
from services.yacang.warehouses import WAREHOUSES


def _write_xlsx(path: Path, warehouse: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(INVENTORY_LIST_HEADERS)
    sheet.append((
        "0000000000000", "SKU-1", "标签", "MAP-1", warehouse, "规格", "1*1*1", 10,
        20, 1, 2, 0, 17, "中文标题", "English title", "https://example.invalid/image.jpg",
        "上架", "在售",
    ))
    workbook.save(path)
    workbook.close()


class FakeClient:
    def __init__(self) -> None:
        self.login_calls: list[tuple[str, str]] = []
        self.submissions: list[int] = []
        self.current: int | None = None

    def login(self, mobile: str, password: str) -> None:
        self.login_calls.append((mobile, password))

    def list_downloads(self, *, limit: int = 50) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = [{"id": "old", "name": "库存导出", "type": "1"}]
        if self.current is not None:
            rows.insert(0, {
                "id": f"new-{self.current}",
                "name": "库存导出",
                "type": "1",
                "path": f"https://oss-accelerate.seaya.cn/test/{self.current}.xlsx",
                "param_where": json.dumps({
                    "warehouse_id": str(self.current),
                    "goods_sku_condition": "2",
                    "user_id": "fixture-user",
                }),
            })
        return rows

    def create_inventory_list_export(self, *, warehouse_id: int) -> str:
        self.current = warehouse_id
        self.submissions.append(warehouse_id)
        return f"request-{warehouse_id}"

    def download_xlsx(self, url: str, destination: Path) -> None:
        warehouse = next(
            item.code for item in WAREHOUSES if f"/{item.warehouse_id}.xlsx" in url
        )
        _write_xlsx(destination, warehouse)


class FailingFakeClient(FakeClient):
    def __init__(self) -> None:
        super().__init__()
        self.download_errors: dict[int, Exception] = {}

    def download_xlsx(self, url: str, destination: Path) -> None:
        warehouse_id = int(url.rsplit("/", 1)[-1].removesuffix(".xlsx"))
        error = self.download_errors.get(warehouse_id)
        if error is not None:
            raise error
        super().download_xlsx(url, destination)


def _headers(path: str) -> tuple[str, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return tuple(str(value or "") for value in next(sheet.iter_rows(values_only=True)))
    finally:
        workbook.close()


def test_queue_match_uses_inventory_type_name_warehouse_and_condition() -> None:
    request = InventoryCurrentSnapshotRequest("MY8801", 26, "2026-09-13")
    task = {
        "name": "库存导出",
        "type": "1",
        "warehouse_id": "0",
        "param_where": '{"warehouse_id":"26","goods_sku_condition":"2"}',
    }
    assert task_matches_inventory_current_snapshot(task, request)
    assert not task_matches_inventory_current_snapshot(
        {**task, "param_where": '{"warehouse_id":"46","goods_sku_condition":"2"}'},
        request,
    )
    assert not task_matches_inventory_current_snapshot({**task, "type": "10"}, request)
    assert task_matches_inventory_month_end(task, request)


def test_exports_four_current_inventory_workbooks_serially(tmp_path: Path) -> None:
    client = FakeClient()
    result = export_inventory_current_snapshot(
        as_of_date="2026-09-13",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=0,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
        today=lambda: date(2026, 9, 13),
    )

    assert client.login_calls == [("account", "password")]
    assert client.submissions == [26, 46, 47, 80]
    assert result["business_type"] == "inventory-current-snapshot"
    assert result["snapshot_semantics"] == "current-at-execution"
    assert result["export_count"] == 4
    assert Path(result["xlsx_paths"][0]).name == "雅仓系统-库存列表_马来西亚仓_2026-09-13.xlsx"
    assert Path(result["xlsx_paths"][-1]).name == "雅仓系统-库存列表_越南仓_2026-09-13.xlsx"
    assert all(_headers(path) == INVENTORY_LIST_HEADERS for path in result["xlsx_paths"])


def test_can_export_one_allowlisted_warehouse(tmp_path: Path) -> None:
    client = FakeClient()
    result = export_inventory_current_snapshot(
        as_of_date="2026-09-13",
        warehouse="ph8805",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=0,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
        today=lambda: date(2026, 9, 13),
    )

    assert client.submissions == [46]
    assert result["export_count"] == 1
    assert result["exports"][0]["warehouse"] == "PH8805"
    assert Path(result["xlsx_paths"][0]).name == "雅仓系统-库存列表_菲律宾仓_2026-09-13.xlsx"


def test_inventory_partial_download_failure_continues_later_warehouses(tmp_path: Path) -> None:
    client = FailingFakeClient()
    client.download_errors[47] = YacangError(
        "下载 XLSX", "fixture failure", code="XLSX_DOWNLOAD_NETWORK_ERROR"
    )
    result = export_inventory_current_snapshot(
        as_of_date="2026-09-13",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=0,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
        today=lambda: date(2026, 9, 13),
    )

    assert result["overall_status"] == "partial_success"
    assert [item["status"] for item in result["exports"]] == [
        "success", "success", "failed", "success",
    ]
    assert client.submissions == [26, 46, 47, 80]
    assert len(result["xlsx_paths"]) == 3


def test_rejects_warehouse_outside_the_allowlist(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="warehouse 必须是"):
        export_inventory_current_snapshot(
            warehouse="UNKNOWN",
            output_dir=tmp_path,
            today=lambda: date(2026, 9, 13),
        )


def test_rejects_historical_date_because_remote_export_has_no_date_filter(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="不支持历史快照日期"):
        export_inventory_current_snapshot(
            as_of_date="2026-08-31",
            output_dir=tmp_path,
            today=lambda: date(2026, 9, 13),
        )


def test_validator_rejects_cross_warehouse_rows(tmp_path: Path) -> None:
    path = tmp_path / "wrong.xlsx"
    _write_xlsx(path, "PH8805")
    with pytest.raises(YacangError, match="其他仓库"):
        validate_inventory_list_workbook(path, warehouse_code="MY8801")


def test_legacy_export_entry_returns_canonical_current_snapshot_type(tmp_path: Path) -> None:
    client = FakeClient()
    result = export_inventory_month_end(
        as_of_date="2026-09-13",
        warehouse="MY8801",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=0,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
        today=lambda: date(2026, 9, 13),
    )

    assert result["business_type"] == "inventory-current-snapshot"
    assert result["exports"][0]["business_type"] == "inventory-current-snapshot"


def test_legacy_cli_failure_uses_canonical_business_type(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("LXE_YACANG_MOBILE", raising=False)
    monkeypatch.delenv("LXE_YACANG_PASSWORD", raising=False)
    result = run_legacy_inventory({"as_of_date": date.today().isoformat()})
    assert result["success"] is False
    assert result["business_type"] == "inventory-current-snapshot"
    assert "缺少雅仓账号或密码" in result["exception"]
    assert "password" not in result["exception"]
