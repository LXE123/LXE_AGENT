from __future__ import annotations

from datetime import date
import inspect
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import services.yacang.export_executor as executor_module
from services.agent_cli.yacang.export_workflow import run as run_export_workflow_cli
from services.agent_cli.yacang.preview_workflow import run as run_preview_workflow_cli
from services.yacang.auth import login_from_environment
from services.yacang.errors import YacangError
from services.yacang.export_executor import execute_export_plan
from services.yacang.export_intent import normalize_export_intent
from services.yacang.export_workflow import (
    ALL_DATA_TYPES as WORKFLOW_DATA_TYPES,
    EXPORT_PLAN_SCHEMA_VERSION,
    normalize_export_request,
    plan_export_workflow,
    run_export_workflow,
    sales_source_fetch_id,
)
from services.yacang.exports.sales_source import InventorySalesSourceBatch
from services.yacang.submission import YacangSubmissionStore
from services.yacang.validation import INVENTORY_SALES_HEADERS


FIXED_TODAY = lambda: date(2026, 9, 14)


def normalized(text: str) -> dict:
    return normalize_export_intent(text, today=FIXED_TODAY)


def test_preview_command_returns_plan_without_executing(monkeypatch) -> None:
    def unexpected_executor(_plan):
        raise AssertionError("preview must not execute")

    monkeypatch.setattr(executor_module, "execute_export_plan", unexpected_executor)

    result = run_preview_workflow_cli({"request_text": "导出 MY8801 的月度销量"})

    assert result["success"] is True
    assert result["plan"]["requires_clarification"] is False
    assert result["plan"]["logical_tasks"][0]["task_id"] == "inventory-sales:MY8801"


def test_preview_command_keeps_clarification_without_remote_tasks() -> None:
    result = run_preview_workflow_cli({"request_text": "最近卖得怎么样"})

    assert result["success"] is False
    assert result["plan"]["requires_clarification"] is True
    assert result["plan"]["logical_tasks"] == []
    assert result["plan"]["source_fetches"] == []


def test_workflow_compatibility_wrapper_returns_canonical_inventory_type() -> None:
    request = normalize_export_request("导出当前库存", today=FIXED_TODAY)

    assert WORKFLOW_DATA_TYPES == (
        "inventory-sales",
        "inventory-current-snapshot",
        "inbound-listing-time",
    )
    assert request["data_types"] == ["inventory-current-snapshot"]


def test_inventory_sales_plan_keeps_one_task_and_source_per_warehouse() -> None:
    plan = plan_export_workflow(
        normalized("导出马来和泰国的库存动销"),
        execution_date="2026-09-14",
    )

    assert [(task["task_id"], task["data_type"]) for task in plan["logical_tasks"]] == [
        ("inventory-sales:MY8801", "inventory-sales"),
        ("inventory-sales:TH8802", "inventory-sales"),
    ]
    assert [fetch["warehouse"] for fetch in plan["source_fetches"]] == ["MY8801", "TH8802"]
    assert all(task["source_fetch_id"] == fetch["source_fetch_id"] for task, fetch in zip(
        plan["logical_tasks"], plan["source_fetches"], strict=True
    ))
    assert all(
        task["data_type"] not in {"sales-monthly", "sales-90d"}
        for task in plan["logical_tasks"]
    )


def test_plan_orders_tasks_scopes_parameters_and_shared_sales_sources() -> None:
    plan = plan_export_workflow(
        normalized("导出 MY8801、TH8802 的两种销量、当前库存和上架时间"),
        execution_date="2026-09-14",
    )

    assert [(task["data_type"], task.get("warehouse")) for task in plan["logical_tasks"]] == [
        ("inventory-sales", "MY8801"),
        ("inventory-sales", "TH8802"),
        ("inventory-current-snapshot", "MY8801"),
        ("inventory-current-snapshot", "TH8802"),
        ("inbound-listing-time", None),
    ]
    assert set(plan) == {
        "schema_version",
        "execution_date",
        "requires_clarification",
        "questions",
        "diagnostics",
        "logical_tasks",
        "source_fetches",
    }
    assert plan["schema_version"] == EXPORT_PLAN_SCHEMA_VERSION == "1"
    assert plan["execution_date"] == "2026-09-14"
    assert len(plan["source_fetches"]) == 2

    inventory_sales_my = plan["logical_tasks"][0]
    assert inventory_sales_my["effective_parameters"] == {
        "warehouse": "MY8801",
        "created_date_filter": {
            "mode": "default",
            "created_start_date": "2026-09-14",
            "created_end_date": "2026-09-14",
            "source": "system_default",
            "input_fragments": [],
            "normalization_rules": ["default_execution_day"],
        },
    }
    assert plan["logical_tasks"][2]["effective_parameters"] == {"warehouse": "MY8801"}
    assert plan["logical_tasks"][2]["snapshot"] == "current"
    assert plan["logical_tasks"][4]["effective_parameters"] == {}
    assert "warehouse" not in plan["logical_tasks"][4]
    assert "created_date_filter" not in plan["logical_tasks"][4]
    assert plan["logical_tasks"][4]["warehouse_scope"] == "all"
    assert plan["source_fetches"][0]["source_fetch_id"] == sales_source_fetch_id(
        "MY8801", "2026-09-14", "2026-09-14"
    )


def test_ambiguous_intent_stops_before_planning_any_remote_source() -> None:
    plan = plan_export_workflow(normalized("最近卖得怎么样"), execution_date="2026-09-14")

    assert plan["requires_clarification"] is True
    assert plan["questions"]
    assert plan["logical_tasks"] == []
    assert plan["source_fetches"] == []


def test_default_plan_splits_inventory_sales_and_inventory_across_all_four_warehouses() -> None:
    plan = plan_export_workflow(normalized("导出雅仓数据"), execution_date="2026-09-14")

    assert [
        task["warehouse"]
        for task in plan["logical_tasks"]
        if task["data_type"] == "inventory-sales"
    ] == ["MY8801", "PH8805", "TH8802", "VN8806"]
    assert [
        task["warehouse"]
        for task in plan["logical_tasks"]
        if task["data_type"] == "inventory-current-snapshot"
    ] == ["MY8801", "PH8805", "TH8802", "VN8806"]
    assert len(plan["source_fetches"]) == 4


def test_unsupported_sales_window_keeps_diagnostic_without_creating_sales_task() -> None:
    plan = plan_export_workflow(normalized("导出56天销量"), execution_date="2026-09-14")

    assert plan["requires_clarification"] is False
    assert plan["logical_tasks"] == []
    assert plan["source_fetches"] == []
    assert plan["diagnostics"] == [
        {
            "kind": "unsupported",
            "code": "UNSUPPORTED_SALES_WINDOW",
            "requested_value": {"sales_window_days": 56},
            "message": "当前完整库存动销报表只包含固定的销量窗口，不支持自定义销量天数",
        }
    ]


@pytest.mark.parametrize("text", ["90天逐日销量", "90天每天销量", "日销量明细", "逐日销量"])
def test_daily_detail_does_not_plan_cumulative_sales_source(text: str) -> None:
    plan = plan_export_workflow(normalized(text), execution_date="2026-09-14")

    assert plan["logical_tasks"] == []
    assert plan["source_fetches"] == []
    assert [item["code"] for item in plan["diagnostics"]] == ["UNSUPPORTED_DATA_TYPE"]


def test_unsupported_daily_detail_does_not_block_independent_current_inventory() -> None:
    plan = plan_export_workflow(
        normalized("导出 MY8801 当前库存和90天每天销量"),
        execution_date="2026-09-14",
    )

    assert [(task["data_type"], task["warehouse"]) for task in plan["logical_tasks"]] == [
        ("inventory-current-snapshot", "MY8801"),
    ]
    assert plan["source_fetches"] == []
    assert [item["code"] for item in plan["diagnostics"]] == ["UNSUPPORTED_DATA_TYPE"]


def test_historical_inventory_keeps_supported_sales_without_planning_inventory() -> None:
    plan = plan_export_workflow(
        normalized("导出 MY8801 的上个月月底库存和月度销量"),
        execution_date="2026-09-14",
    )

    assert [(task["data_type"], task["status"]) for task in plan["logical_tasks"]] == [
        ("inventory-sales", "not_run"),
    ]
    assert plan["diagnostics"] == [
        {
            "kind": "unsupported",
            "code": "UNSUPPORTED_HISTORICAL_INVENTORY",
            "message": "当前雅仓能力不支持指定历史日期或历史月末库存快照。",
        }
    ]
    assert len(plan["source_fetches"]) == 1
    assert plan["source_fetches"][0]["warehouse"] == "MY8801"


def test_historical_inventory_only_generates_no_tasks_or_source_fetch() -> None:
    plan = plan_export_workflow(
        normalized("导出8月31日库存"),
        execution_date="2026-09-14",
    )

    assert plan["logical_tasks"] == []
    assert plan["source_fetches"] == []
    assert plan["diagnostics"][0]["code"] == "UNSUPPORTED_HISTORICAL_INVENTORY"


def test_bare_month_end_inventory_stops_planning_for_clarification() -> None:
    plan = plan_export_workflow(
        normalized("导出月末库存"),
        execution_date="2026-09-14",
    )

    assert plan["requires_clarification"] is True
    assert plan["logical_tasks"] == []
    assert plan["source_fetches"] == []
    assert plan["questions"][0]["code"] == "AMBIGUOUS_INVENTORY_SNAPSHOT"


def test_inbound_listing_time_is_one_global_task_without_date_or_warehouse_parameters() -> None:
    plan = plan_export_workflow(normalized("导出什么时候上架"), execution_date="2026-09-14")

    assert plan["logical_tasks"] == [
        {
            "task_id": "inbound-listing-time:global",
            "data_type": "inbound-listing-time",
            "scope": "global",
            "warehouse_scope": "all",
            "status": "not_run",
            "effective_parameters": {},
        }
    ]
    assert plan["source_fetches"] == []


def test_inbound_with_chinese_warehouse_alias_remains_one_global_task() -> None:
    plan = plan_export_workflow(
        normalized("导出马来仓入库时间"),
        execution_date="2026-09-14",
    )

    assert plan["requires_clarification"] is False
    assert plan["diagnostics"] == []
    assert plan["logical_tasks"] == [
        {
            "task_id": "inbound-listing-time:global",
            "data_type": "inbound-listing-time",
            "scope": "global",
            "warehouse_scope": "all",
            "status": "not_run",
            "effective_parameters": {},
        }
    ]
    assert plan["source_fetches"] == []


def test_short_warehouse_aliases_stay_canonical_in_tasks_and_source_fetch_ids() -> None:
    plan = plan_export_workflow(
        normalized("导出 MY 和 PH 最近一个月销量"),
        execution_date="2026-09-14",
    )

    assert [task["warehouse"] for task in plan["logical_tasks"]] == ["MY8801", "PH8805"]
    assert [fetch["warehouse"] for fetch in plan["source_fetches"]] == ["MY8801", "PH8805"]
    assert [fetch["source_fetch_id"] for fetch in plan["source_fetches"]] == [
        "inventory-sales-source:MY8801:2026-09-14:2026-09-14",
        "inventory-sales-source:PH8805:2026-09-14:2026-09-14",
    ]


def test_warehouse_product_phrase_plans_only_the_global_inbound_task() -> None:
    plan = plan_export_workflow(
        normalized("一次性导出仓库产品"),
        execution_date="2026-09-14",
    )

    assert [(task["data_type"], task.get("warehouse")) for task in plan["logical_tasks"]] == [
        ("inbound-listing-time", None),
    ]
    assert plan["logical_tasks"][0]["effective_parameters"] == {}
    assert plan["source_fetches"] == []


def test_chinese_warehouse_aliases_reach_planner_as_canonical_codes_only() -> None:
    plan = plan_export_workflow(
        normalized("马来和泰国当前库存"),
        execution_date="2026-09-14",
    )

    assert [task["warehouse"] for task in plan["logical_tasks"]] == ["MY8801", "TH8802"]
    assert [task["effective_parameters"] for task in plan["logical_tasks"]] == [
        {"warehouse": "MY8801"},
        {"warehouse": "TH8802"},
    ]
    assert plan["source_fetches"] == []
    assert all(
        value not in str(plan)
        for value in ("马来", "马来仓", "马来西亚", "泰国", "泰仓")
    )


def test_current_inventory_is_per_warehouse_without_created_dates_or_sales_source() -> None:
    plan = plan_export_workflow(normalized("四个仓库还剩多少货"), execution_date="2026-09-14")

    assert [task["task_id"] for task in plan["logical_tasks"]] == [
        "inventory-current-snapshot:MY8801",
        "inventory-current-snapshot:PH8805",
        "inventory-current-snapshot:TH8802",
        "inventory-current-snapshot:VN8806",
    ]
    assert all(
        task["effective_parameters"] == {"warehouse": task["warehouse"]}
        for task in plan["logical_tasks"]
    )
    assert plan["source_fetches"] == []


def test_clarification_short_circuits_before_executor() -> None:
    def unexpected_executor(_plan: dict) -> dict:
        raise AssertionError("executor must not run while clarification is required")

    result = run_export_workflow(
        "最近卖得怎么样",
        today=FIXED_TODAY,
        executor=unexpected_executor,
    )

    assert result["overall_status"] == "needs_clarification"
    assert result["tasks"] == []
    assert result["artifacts"] == []
    assert result["questions"]


def test_executor_receives_only_the_frozen_export_plan() -> None:
    captured: dict = {}

    def recording_executor(plan: dict) -> dict:
        captured["plan"] = plan
        return {
            "schema_version": "1",
            "overall_status": "success",
            "tasks": [],
            "artifacts": [],
            "questions": [],
            "diagnostics": [],
        }

    result = run_export_workflow(
        "导出 MY8801 当前库存",
        today=FIXED_TODAY,
        executor=recording_executor,
    )

    assert result["overall_status"] == "success"
    assert set(result) == {
        "schema_version", "overall_status", "tasks", "artifacts", "questions", "diagnostics",
    }
    assert "success" not in result
    assert set(captured["plan"]) == {
        "schema_version",
        "execution_date",
        "requires_clarification",
        "questions",
        "diagnostics",
        "logical_tasks",
        "source_fetches",
    }
    assert "request_text" not in json.dumps(captured["plan"], ensure_ascii=False)


def _successful_export(data_type: str, warehouse: str | None, path: Path) -> dict:
    export = {
        "business_type": data_type,
        "status": "success",
        "output_filename": path.name,
        "xlsx_path": str(path),
        "row_count": 1,
        "source": "fixture",
    }
    if warehouse is not None:
        export["warehouse"] = warehouse
    return {
        "overall_status": "success",
        "exports": [export],
        "xlsx_paths": [str(path)],
    }


def _source_batch(
    plan: dict,
    *,
    source_dir: Path | None = None,
    failed_warehouse: str | None = None,
) -> InventorySalesSourceBatch:
    results = []
    for fetch in plan["source_fetches"]:
        failed = fetch["warehouse"] == failed_warehouse
        source_path = None
        if not failed:
            source_path = (source_dir or Path.cwd()) / f"raw-{fetch['warehouse']}.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(INVENTORY_SALES_HEADERS)
            sheet.append((
                "SKU-1", "商品", fetch["warehouse"], 3, 7, 15, 30, 60, 90,
                10, 0, 0, 0, 10, 0, "2026-09-01",
            ))
            workbook.save(source_path)
            workbook.close()
        results.append({
            "source_fetch_id": fetch["source_fetch_id"],
            "warehouse": fetch["warehouse"],
            "status": "failed" if failed else "success",
            "xlsx_path": str(source_path) if source_path is not None else None,
            "source": "fixture",
            **(
                {
                    "error_code": "EXPORT_POLL_TIMEOUT",
                    "stage": "轮询导出队列",
                    "error_type": "YacangError",
                    "error": "fixture poll timeout",
                }
                if failed
                else {}
            ),
        })
    created = plan["source_fetches"][0]["effective_parameters"]["created_date_filter"]
    return InventorySalesSourceBatch(
        created_start_date=created["created_start_date"],
        created_end_date=created["created_end_date"],
        results=tuple(results),
    )


def test_executor_acquires_each_source_once_and_returns_one_complete_artifact(
    monkeypatch,
    tmp_path: Path,
) -> None:
    assert list(inspect.signature(execute_export_plan).parameters) == ["plan"]
    plan = plan_export_workflow(
        normalized("导出 MY8801 两种销量"),
        execution_date="2026-09-14",
    )
    calls: list[object] = []

    def acquire(source_fetches, **_kwargs):
        calls.append(("source", [item["source_fetch_id"] for item in source_fetches]))
        return _source_batch(plan, source_dir=tmp_path)

    def publish(sources, destination):
        calls.append(("publish", [warehouse for warehouse, _path in sources]))
        destination.write_bytes(Path(sources[0][1]).read_bytes())
        return 1

    monkeypatch.setattr(executor_module, "acquire_inventory_sales_sources", acquire)
    monkeypatch.setattr(executor_module, "publish_inventory_sales_workbook", publish)

    result = execute_export_plan(plan)

    assert calls == [
        ("source", [plan["source_fetches"][0]["source_fetch_id"]]),
        ("publish", ["MY8801"]),
    ]
    assert result["overall_status"] == "success"
    assert [task["status"] for task in result["tasks"]] == ["success"]
    assert [artifact["data_type"] for artifact in result["artifacts"]] == ["inventory-sales"]
    assert len(result["artifacts"]) == 1
    assert all("xlsx_path" not in artifact for artifact in result["artifacts"])
    assert all(".source-cache" not in artifact["path"] for artifact in result["artifacts"])


def test_executor_merges_multi_warehouse_sources_into_one_complete_artifact(
    monkeypatch,
    tmp_path: Path,
) -> None:
    plan = plan_export_workflow(
        normalized("导出马来和泰国的库存动销"),
        execution_date="2026-09-14",
    )
    monkeypatch.setattr(executor_module, "dataset_dir", lambda _dataset: tmp_path)
    monkeypatch.setattr(
        executor_module,
        "acquire_inventory_sales_sources",
        lambda _fetches, **_kwargs: _source_batch(plan, source_dir=tmp_path),
    )

    result = execute_export_plan(plan)

    assert result["overall_status"] == "success"
    assert [task["status"] for task in result["tasks"]] == ["success", "success"]
    assert len(result["artifacts"]) == 1
    artifact = result["artifacts"][0]
    assert artifact["data_type"] == "inventory-sales"
    assert artifact["warehouses"] == ["MY8801", "TH8802"]
    assert artifact["filename"] == "雅仓系统-库存动销_马来西亚仓-泰国仓_2026-09-14.xlsx"
    assert all(task["artifact_ids"] == [artifact["artifact_id"]] for task in result["tasks"])
    workbook = load_workbook(artifact["path"], read_only=True, data_only=True)
    try:
        rows = list(workbook["库存动销"].iter_rows(values_only=True))
    finally:
        workbook.close()
    assert tuple(rows[0]) == INVENTORY_SALES_HEADERS
    assert [row[2] for row in rows[1:]] == ["MY8801", "TH8802"]


@pytest.mark.parametrize(
    ("request_text", "expected_context_uses"),
    [
        ("导出四个仓库最近一个月销量", 1),
        ("导出四个仓库近90天销量", 1),
        ("导出四个仓库两种销量", 1),
        ("导出四个仓库当前库存", 4),
        ("导出入库和上架时间", 1),
        ("导出雅仓数据", 6),
        ("导出 MY8801 全部数据", 3),
    ],
)
def test_executor_reuses_one_authenticated_client_and_persistent_submission_store(
    monkeypatch,
    tmp_path: Path,
    request_text: str,
    expected_context_uses: int,
) -> None:
    plan = plan_export_workflow(normalized(request_text), execution_date="2026-09-14")
    created_clients: list[object] = []
    context_clients: list[object] = []
    context_stores: list[object] = []

    class ClientProbe:
        def __init__(self) -> None:
            self.authenticated = False
            self.login_calls = 0

        @property
        def is_authenticated(self) -> bool:
            return self.authenticated

        def login(self, _mobile: str, _password: str) -> None:
            self.login_calls += 1
            self.authenticated = True

    def create_client() -> ClientProbe:
        client = ClientProbe()
        created_clients.append(client)
        return client

    def authenticate(client: ClientProbe, submission_store: object) -> None:
        context_clients.append(client)
        context_stores.append(submission_store)
        assert isinstance(submission_store, YacangSubmissionStore)
        login_from_environment(client, mobile="account", password="password")

    def acquire(_source_fetches, *, client, submission_store, **_kwargs):
        authenticate(client, submission_store)
        return _source_batch(plan, source_dir=tmp_path)

    def inventory(*, warehouse, client, submission_store, **_kwargs):
        authenticate(client, submission_store)
        return _successful_export(
            "inventory-current-snapshot",
            warehouse,
            tmp_path / f"inventory-{warehouse}.xlsx",
        )

    def inbound(*, client, submission_store, **_kwargs):
        authenticate(client, submission_store)
        return _successful_export("inbound-listing-time", None, tmp_path / "inbound.xlsx")

    monkeypatch.setattr(executor_module, "YacangClient", create_client, raising=False)
    monkeypatch.setattr(executor_module, "acquire_inventory_sales_sources", acquire)
    monkeypatch.setattr(
        executor_module,
        "publish_inventory_sales_workbook",
        lambda sources, destination: destination.write_bytes(Path(sources[0][1]).read_bytes()) or 1,
    )
    monkeypatch.setattr(executor_module, "export_inventory_current_snapshot", inventory)
    monkeypatch.setattr(executor_module, "export_inbound_listing_time", inbound)

    result = execute_export_plan(plan)

    assert result["overall_status"] == "success"
    assert len(created_clients) == 1
    assert len(context_clients) == expected_context_uses
    assert len({id(client) for client in context_clients}) == 1
    assert len({id(store) for store in context_stores}) == 1
    assert created_clients[0].login_calls == 1


def test_executor_preserves_fixed_type_and_warehouse_execution_order(monkeypatch, tmp_path: Path) -> None:
    plan = plan_export_workflow(normalized("导出雅仓数据"), execution_date="2026-09-14")
    calls: list[str] = []

    monkeypatch.setattr(
        executor_module,
        "acquire_inventory_sales_sources",
        lambda _fetches, **_kwargs: (
            calls.append("source") or _source_batch(plan, source_dir=tmp_path)
        ),
    )
    def publish(sources, destination):
        calls.append("publish-inventory-sales")
        destination.write_bytes(Path(sources[0][1]).read_bytes())
        return len(sources)

    monkeypatch.setattr(executor_module, "publish_inventory_sales_workbook", publish)
    monkeypatch.setattr(
        executor_module,
        "export_inventory_current_snapshot",
        lambda *, warehouse, **_kwargs: (
            calls.append(f"inventory-current-snapshot:{warehouse}")
            or _successful_export(
                "inventory-current-snapshot", warehouse, tmp_path / f"inventory-{warehouse}.xlsx"
            )
        ),
    )
    monkeypatch.setattr(
        executor_module,
        "export_inbound_listing_time",
        lambda **_kwargs: (
            calls.append("inbound-listing-time:global")
            or _successful_export("inbound-listing-time", None, tmp_path / "inbound.xlsx")
        ),
    )

    result = execute_export_plan(plan)

    assert calls == [
        "source",
        "inventory-current-snapshot:MY8801", "inventory-current-snapshot:PH8805",
        "inventory-current-snapshot:TH8802", "inventory-current-snapshot:VN8806",
        "inbound-listing-time:global",
        "publish-inventory-sales",
    ]
    assert result["overall_status"] == "success"
    assert [task["data_type"] for task in result["tasks"]] == [
        "inventory-sales", "inventory-sales", "inventory-sales", "inventory-sales",
        "inventory-current-snapshot", "inventory-current-snapshot",
        "inventory-current-snapshot", "inventory-current-snapshot",
        "inbound-listing-time",
    ]
    assert len(result["artifacts"]) == 6


def test_local_source_failure_continues_and_preserves_success_artifacts(monkeypatch, tmp_path: Path) -> None:
    plan = plan_export_workflow(
        normalized("导出 MY8801、TH8802 的两种销量"), execution_date="2026-09-14"
    )
    calls: list[str] = []
    monkeypatch.setattr(
        executor_module,
        "acquire_inventory_sales_sources",
        lambda _fetches, **_kwargs: _source_batch(
            plan,
            source_dir=tmp_path,
            failed_warehouse="MY8801",
        ),
    )

    def publish(sources, destination):
        calls.extend(warehouse for warehouse, _path in sources)
        destination.write_bytes(Path(sources[0][1]).read_bytes())
        return len(sources)

    monkeypatch.setattr(executor_module, "publish_inventory_sales_workbook", publish)

    result = execute_export_plan(plan)

    assert calls == ["TH8802"]
    assert [task["status"] for task in result["tasks"]] == ["failed", "success"]
    assert result["overall_status"] == "partial_success"
    assert len(result["artifacts"]) == 1


def test_global_source_status_stops_later_logical_tasks(monkeypatch, tmp_path: Path) -> None:
    plan = plan_export_workflow(
        normalized("导出 MY8801、TH8802 的库存动销和当前库存"), execution_date="2026-09-14"
    )
    batch = _source_batch(plan, source_dir=tmp_path)
    source_results = [dict(item) for item in batch.results]
    source_results[1].update({
        "status": "failed",
        "xlsx_path": None,
        "error_code": "EXPORT_STATUS_UNKNOWN",
        "stage": "读取导出队列",
        "error_type": "YacangError",
        "error": "fixture unknown queue status",
    })
    batch = InventorySalesSourceBatch(
        created_start_date=batch.created_start_date,
        created_end_date=batch.created_end_date,
        results=tuple(source_results),
    )
    calls: list[str] = []
    monkeypatch.setattr(
        executor_module,
        "acquire_inventory_sales_sources",
        lambda _fetches, **_kwargs: batch,
    )

    def publish(sources, destination):
        calls.extend(warehouse for warehouse, _path in sources)
        destination.write_bytes(Path(sources[0][1]).read_bytes())
        return len(sources)

    monkeypatch.setattr(executor_module, "publish_inventory_sales_workbook", publish)
    monkeypatch.setattr(
        executor_module,
        "export_inventory_current_snapshot",
        lambda **_kwargs: calls.append("unexpected-inventory"),
    )

    result = execute_export_plan(plan)

    assert calls == ["MY8801"]
    assert [task["status"] for task in result["tasks"]] == [
        "success", "failed", "skipped", "skipped",
    ]
    assert result["overall_status"] == "partial_success"
    assert len(result["artifacts"]) == 1


@pytest.mark.parametrize(
    "error_code",
    [
        "YACANG_CREDENTIALS_MISSING",
        "YACANG_AUTH_EXPIRED",
        "YACANG_FORBIDDEN",
        "YACANG_RATE_LIMITED",
        "EXPORT_SUBMIT_UNKNOWN",
        "EXPORT_STATUS_UNKNOWN",
        "YACANG_REMOTE_ERROR",
    ],
)
def test_global_error_stops_later_tasks_and_keeps_prior_artifact(
    monkeypatch,
    tmp_path: Path,
    error_code: str,
) -> None:
    plan = plan_export_workflow(
        normalized("导出 MY8801、TH8802 当前库存和上架时间"), execution_date="2026-09-14"
    )
    calls: list[str] = []

    def inventory(*, warehouse, **_kwargs):
        calls.append(warehouse)
        if warehouse == "TH8802":
            raise YacangError(
                "提交库存导出",
                "HTTP 403",
                code=error_code,
                http_status=403,
                scope="global",
            )
        return _successful_export(
            "inventory-current-snapshot", warehouse, tmp_path / f"inventory-{warehouse}.xlsx"
        )

    monkeypatch.setattr(executor_module, "export_inventory_current_snapshot", inventory)
    monkeypatch.setattr(
        executor_module,
        "export_inbound_listing_time",
        lambda **_kwargs: calls.append("inbound") or _successful_export(
            "inbound-listing-time", None, tmp_path / "inbound.xlsx"
        ),
    )

    result = execute_export_plan(plan)

    assert calls == ["MY8801", "TH8802"]
    assert [task["status"] for task in result["tasks"]] == ["success", "failed", "skipped"]
    assert result["overall_status"] == "partial_success"
    assert len(result["artifacts"]) == 1
    assert result["diagnostics"][-1]["code"] == "EXPORT_SKIPPED"


@pytest.mark.parametrize(
    "error_code",
    ["EXPORT_POLL_TIMEOUT", "XLSX_DOWNLOAD_NETWORK_ERROR", "XLSX_MIME_INVALID"],
)
def test_local_inventory_error_continues_later_tasks(monkeypatch, tmp_path: Path, error_code: str) -> None:
    plan = plan_export_workflow(
        normalized("导出 MY8801、TH8802 当前库存和上架时间"), execution_date="2026-09-14"
    )
    calls: list[str] = []

    def inventory(*, warehouse, **_kwargs):
        calls.append(warehouse)
        if warehouse == "MY8801":
            raise YacangError("库存导出", "fixture local failure", code=error_code)
        return _successful_export(
            "inventory-current-snapshot", warehouse, tmp_path / f"inventory-{warehouse}.xlsx"
        )

    monkeypatch.setattr(executor_module, "export_inventory_current_snapshot", inventory)
    monkeypatch.setattr(
        executor_module,
        "export_inbound_listing_time",
        lambda **_kwargs: (
            calls.append("inbound")
            or _successful_export("inbound-listing-time", None, tmp_path / "inbound.xlsx")
        ),
    )

    result = execute_export_plan(plan)

    assert calls == ["MY8801", "TH8802", "inbound"]
    assert [task["status"] for task in result["tasks"]] == ["failed", "success", "success"]
    assert result["overall_status"] == "partial_success"
    assert len(result["artifacts"]) == 2


def test_unsupported_only_fails_without_calling_any_exporter(monkeypatch) -> None:
    plan = plan_export_workflow(normalized("导出56天销量"), execution_date="2026-09-14")
    monkeypatch.setattr(
        executor_module,
        "acquire_inventory_sales_sources",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("must not fetch")),
    )

    result = execute_export_plan(plan)

    assert result["overall_status"] == "failed"
    assert result["tasks"] == []
    assert result["artifacts"] == []
    assert result["diagnostics"][0]["code"] == "UNSUPPORTED_SALES_WINDOW"


def test_mixed_supported_and_unsupported_is_partial_success(monkeypatch, tmp_path: Path) -> None:
    plan = plan_export_workflow(
        normalized("导出 MY8801 当前库存和56天销量"), execution_date="2026-09-14"
    )
    monkeypatch.setattr(
        executor_module,
        "export_inventory_current_snapshot",
        lambda *, warehouse, **_kwargs: _successful_export(
            "inventory-current-snapshot", warehouse, tmp_path / "inventory.xlsx"
        ),
    )

    result = execute_export_plan(plan)

    assert result["overall_status"] == "partial_success"
    assert [task["status"] for task in result["tasks"]] == ["success"]
    assert result["artifacts"][0]["path"] == str(tmp_path / "inventory.xlsx")
    assert "xlsx_path" not in result["artifacts"][0]
    assert result["diagnostics"][0]["code"] == "UNSUPPORTED_SALES_WINDOW"
    assert result["diagnostics"][0]["requested_value"] == {"sales_window_days": 56}
    assert all(
        "requested_value" not in task.get("effective_parameters", {})
        for task in result["tasks"]
    )


def test_unified_cli_returns_canonical_clarification_envelope() -> None:
    result = run_export_workflow_cli({"request_text": "最近卖得怎么样"})

    assert result["success"] is False
    assert set(result).issuperset({
        "success",
        "schema_version",
        "overall_status",
        "tasks",
        "artifacts",
        "questions",
        "diagnostics",
    })
    assert result["overall_status"] == "needs_clarification"
    assert result["tasks"] == []
    assert result["artifacts"] == []
