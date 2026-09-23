from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from services.agent_cli.yacang.export_inbound_listing_time import run
from services.yacang.errors import YacangError
from services.yacang.exports import inbound_listing_time as inbound_listing_time_module
from services.yacang.exports.inbound_listing_time import (
    InboundListingTimeRequest,
    export_inbound_listing_time,
    task_matches_inbound_listing_time,
)
from services.yacang.risk import YacangRiskController
from services.yacang.validation import (
    INBOUND_LISTING_TIME_HEADERS,
    validate_inbound_listing_time_workbook,
)


def _write_xlsx(path: Path, *, created_at: str = "2026-09-13 19:13") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(INBOUND_LISTING_TIME_HEADERS)
    sheet.append((
        "0000000000000", "SKU-1", "规格", "中文标题", "English title",
        10, 20, 30, 40, "https://example.invalid/image.jpg", created_at,
    ))
    workbook.save(path)
    workbook.close()


class FakeClient:
    def __init__(self) -> None:
        self.login_calls: list[tuple[str, str]] = []
        self.submission_count = 0

    def login(self, mobile: str, password: str) -> None:
        self.login_calls.append((mobile, password))

    def list_downloads(self, *, limit: int = 50) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = [{"id": "old", "name": "资料导出", "type": "6"}]
        if self.submission_count:
            rows.insert(0, {
                "id": "new-fixture-task",
                "name": "资料导出",
                "type": "6",
                "path": "https://oss-accelerate.seaya.cn/test/products.xlsx",
                "param_where": json.dumps([
                    ["user_id", "=", "fixture-user"],
                    ["status", "=", "1"],
                ]),
            })
        return rows

    def create_inbound_listing_time_export(self) -> str:
        self.submission_count += 1
        return "fixture-request"

    def download_xlsx(self, url: str, destination: Path) -> None:
        assert url.endswith("/products.xlsx")
        _write_xlsx(destination)


def _headers(path: str) -> tuple[str, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return tuple(str(value or "") for value in next(sheet.iter_rows(values_only=True)))
    finally:
        workbook.close()


def test_queue_match_uses_name_type_and_active_status() -> None:
    request = InboundListingTimeRequest("2026-09-13")
    task = {
        "name": "资料导出",
        "type": "6",
        "param_where": '[["user_id","=","fixture-user"],["status","=","1"]]',
    }
    assert task_matches_inbound_listing_time(task, request)
    assert not task_matches_inbound_listing_time({**task, "type": "1"}, request)
    assert not task_matches_inbound_listing_time(
        {**task, "param_where": '[["status","=","0"]]'},
        request,
    )


def test_exports_one_combined_warehouse_product_workbook(tmp_path: Path) -> None:
    client = FakeClient()
    result = export_inbound_listing_time(
        as_of_date="2026-09-13",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=0,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
        today=lambda: date(2026, 9, 13),
    )

    assert client.login_calls == [("account", "password")]
    assert client.submission_count == 1
    assert result["business_type"] == "inbound-listing-time"
    assert result["creation_time_field"] == "创建时间"
    assert result["export_count"] == 1
    assert Path(result["xlsx_paths"][0]).name == "雅仓系统-产品-仓库产品_2026-09-13.xlsx"
    assert _headers(result["xlsx_paths"][0]) == INBOUND_LISTING_TIME_HEADERS


def test_fresh_valid_file_is_reused_without_login(tmp_path: Path) -> None:
    path = tmp_path / "雅仓系统-产品-仓库产品_2026-09-13.xlsx"
    _write_xlsx(path)
    client = FakeClient()
    result = export_inbound_listing_time(
        as_of_date="2026-09-13",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=300,
        client=client,  # type: ignore[arg-type]
        today=lambda: date(2026, 9, 13),
    )
    assert client.login_calls == []
    assert client.submission_count == 0
    assert result["exports"][0]["source"] == "cache"


def test_rejects_historical_file_date(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="必须是执行当天"):
        export_inbound_listing_time(
            as_of_date="2026-08-31",
            output_dir=tmp_path,
            today=lambda: date(2026, 9, 13),
        )


def test_validator_rejects_invalid_creation_time(tmp_path: Path) -> None:
    path = tmp_path / "wrong-time.xlsx"
    _write_xlsx(path, created_at="2026/09/13")
    with pytest.raises(YacangError, match="创建时间必须为"):
        validate_inbound_listing_time_workbook(path)


def test_global_risk_controller_rejects_duplicate_submission() -> None:
    risk = YacangRiskController(sleep=lambda _seconds: None)
    risk.before_global_submission(
        business_key="inbound-listing-time",
        request_key="2026-09-13",
    )
    with pytest.raises(YacangError, match="已提交相同任务"):
        risk.before_global_submission(
            business_key="inbound-listing-time",
            request_key="2026-09-13",
        )


def test_cli_failure_is_factual_without_credentials(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("LXE_YACANG_MOBILE", raising=False)
    monkeypatch.delenv("LXE_YACANG_PASSWORD", raising=False)
    monkeypatch.setattr(inbound_listing_time_module, "dataset_dir", lambda _dataset: tmp_path)
    result = run({"as_of_date": date.today().isoformat()})
    assert result["success"] is False
    assert result["business_type"] == "inbound-listing-time"
    assert "缺少雅仓账号或密码" in result["exception"]
    assert "password" not in result["exception"]
