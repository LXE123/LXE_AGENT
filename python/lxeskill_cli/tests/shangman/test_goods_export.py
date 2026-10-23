from __future__ import annotations

import io
import json
from pathlib import Path
import re
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from zipfile import ZipFile
import asyncio

import pytest
from openpyxl import Workbook

from services.shangman import goods_export as exports
from services.shangman.auth import AuthError, Credentials, LoginToken
from services.shangman.state import AuthStore, atomic_json, read_json


def workbook_bytes(*, empty=False, bilingual=False, bad_dimensions=False, bad_headers=False):
    wb = Workbook()
    sheet = wb.active
    headers = sorted(exports.REQUIRED_HEADERS)
    if bad_headers:
        headers.remove("SKU")
    translations = {"SKU": "商品编码", "商品名": "商品名称"}
    sheet.append([f"English\n({translations.get(h, h)})"
                  if bilingual else h for h in headers])
    if not empty:
        row = ["source data"] * len(headers)
        values = {"SKU": "SKU-1", "创建时间": "2026-01-01 00:00:00", "有效库存": 17, "7天销量": 3}
        for header, value in values.items():
            if header in headers:
                row[headers.index(header)] = value
        sheet.append(row)
        if "有效库存" in headers:
            sheet.cell(2, headers.index("有效库存") + 1).number_format = "0.00"
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    if not bad_dimensions:
        return output.getvalue()
    rewritten = io.BytesIO()
    with ZipFile(output) as source, ZipFile(rewritten, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', content)
            target.writestr(item, content)
    return rewritten.getvalue()


def identity_record(record_id, sku, created, *, status=2):
    return {"id": record_id, "goodsNo": sku, "goodsSku": sku,
            "createTime": created, "status": status}


def identity_workbook(rows, *, formula=False, second_product_sheet=False):
    wb = Workbook()
    sheet = wb.active
    headers = sorted(exports.REQUIRED_HEADERS)
    sheet.append(headers)
    for sku, created, stock, sale7 in rows:
        row = ["source data"] * len(headers)
        row[headers.index("SKU")] = sku
        row[headers.index("创建时间")] = created
        row[headers.index("有效库存")] = stock
        row[headers.index("7天销量")] = sale7
        row[headers.index("商品名")] = ""
        if formula:
            row[headers.index("预警库存")] = "=1+1"
        sheet.append(row)
    if second_product_sheet:
        duplicate_sheet = wb.create_sheet("第二商品表")
        duplicate_sheet.append(headers)
        duplicate_sheet.append(["SKU-OTHER" if header == "SKU" else
                                "2026-02-01 00:00:00" if header == "创建时间" else None
                                for header in headers])
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


@pytest.fixture
def context(tmp_path, monkeypatch):
    from shared import workspace
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path))
    monkeypatch.delenv("LXE_SHANGMAN_CONFIG_REVISION", raising=False)
    for key, value in zip(("TENANT_ID", "USERNAME", "PROCESSED_PASSWORD"), ("tenant-123", "operator-456", "password-789")):
        monkeypatch.setenv("LXE_SHANGMAN_" + key, value)
    monkeypatch.setattr(workspace, "_artifact_root", tmp_path / "artifacts")
    credentials = Credentials.from_environment()
    store = AuthStore(credentials)
    store.save(LoginToken("private-token", 300))
    return credentials, store, tmp_path


@pytest.fixture
def server(context, monkeypatch):
    state = {"post_status": 200, "get_status": 200, "list_status": 200,
             "xlsx": workbook_bytes(), "calls": [],
             "list_rows": [{"id": "goods-1", "goodsNo": "SKU-1", "goodsSku": "SKU-1",
                            "createTime": "2026-01-01 00:00:00", "status": 2}]}

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            state["calls"].append(("POST", self.path, dict(self.headers)))
            if state.get("before_response"):
                state["before_response"]()
            self.send_response(state["post_status"])
            self.send_header("Location", "/must-not-follow")
            self.end_headers()
            body = state.get("payload", {"code": 200, "success": True, "data": state["url"]})
            self.wfile.write((body if isinstance(body, str) else json.dumps(body)).encode())

        def do_GET(self):
            state["calls"].append(("GET", self.path, dict(self.headers)))
            if self.path.startswith(exports.LIST_PATH):
                self.send_response(state["list_status"])
                self.send_header("Location", "/must-not-follow")
                self.end_headers()
                page = int(self.path.split("current=", 1)[1].split("&", 1)[0])
                rows = state["list_rows"]
                pages = max(1, (len(rows) + exports.LIST_PAGE_SIZE - 1) // exports.LIST_PAGE_SIZE)
                start = (page - 1) * exports.LIST_PAGE_SIZE
                payload = state.get("list_payloads", {}).get(page, state.get("list_payload", {
                    "code": 200, "success": True,
                    "data": {"current": page, "pages": pages, "total": len(rows),
                             "records": rows[start:start + exports.LIST_PAGE_SIZE]},
                }))
                self.wfile.write(json.dumps(payload).encode())
            else:
                self.send_response(state["get_status"])
                self.send_header("Location", "/must-not-follow")
                self.send_header("Set-Cookie", "ERP_SESSION=do-not-reuse")
                self.end_headers()
                self.wfile.write(state["xlsx"])

        def log_message(self, *args):
            pass

    http = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=http.serve_forever, daemon=True)
    thread.start()
    state["url"] = f"http://localhost:{http.server_port}/goods.xlsx?signature=signed-private-value"
    monkeypatch.setattr(exports, "BASE_URL", f"http://127.0.0.1:{http.server_port}")
    # Only the approved-URL boundary is bypassed for this loopback integration.
    # Its production behavior is independently tested below.
    monkeypatch.setattr(exports.GoodsExporter, "_validate_download_url", lambda self, url: url)
    try:
        yield state
    finally:
        http.shutdown()
        http.server_close()
        thread.join()


@pytest.mark.parametrize("options", [{}, {"bilingual": True}, {"bad_dimensions": True}, {"empty": True}])
def test_real_http_chain_preserves_official_rows_and_does_not_leak_auth(context, server, options):
    server["xlsx"] = workbook_bytes(**options)
    if options.get("empty"):
        server["list_rows"] = []
    result = exports.run({})
    assert result["success"], result
    path = Path(result["artifact_path"])
    from openpyxl import load_workbook
    output = load_workbook(path, data_only=False)
    sheet = output[output.sheetnames[0]]
    headers = [exports._header(cell.value) for cell in sheet[1]]
    assert sheet.max_column == len(sorted(exports.REQUIRED_HEADERS))
    if not options.get("empty"):
        assert sheet.cell(2, headers.index("有效库存") + 1).value == 17
        assert sheet.cell(2, headers.index("7天销量") + 1).value == 3
        assert sheet.cell(2, headers.index("有效库存") + 1).number_format == "0.00"
    output.close()
    assert path.name.startswith("上马-商品-")
    assert path.parent.parent == context[2] / "artifacts/shangman/indonesia"
    assert result["row_count"] == (0 if options.get("empty") else 1)
    assert bool(result["notice"]) == bool(options.get("empty"))
    assert len(server["calls"]) == 3
    method, endpoint, headers = server["calls"][0]
    assert (method, endpoint) == ("POST", exports.EXPORT_PATH)
    assert headers["Blade-Auth"] == "bearer private-token"
    assert headers["Authorization"] == context[0].basic_auth
    assert headers["Tenant-Id"] == context[0].tenant_id
    download_headers = server["calls"][1][2]
    assert not {"Authorization", "Blade-Auth", "Tenant-Id", "Cookie"} & download_headers.keys()
    method, endpoint, list_headers = server["calls"][2]
    assert method == "GET" and endpoint.startswith(exports.LIST_PATH + "?status=2&current=1&size=500")
    assert list_headers["Blade-Auth"] == "bearer private-token"
    assert not list(path.parent.glob(".*.xlsx"))
    assert "private-token" not in json.dumps(result)
    assert "signed-private-value" not in json.dumps(result)


def test_enabled_list_paginates_all_records_and_preserves_duplicate_skus(context, server):
    rows = [identity_record(f"id-{i}", f"SKU-{i // 2}", f"2026-01-01 00:{i // 60:02d}:{i % 60:02d}")
            for i in range(501)]
    server["list_rows"] = rows
    client = exports.GoodsExporter(context[0], context[1])

    async def query():
        async with exports.aiohttp.ClientSession() as session:
            return await client._enabled_goods(session, {
                "Authorization": context[0].basic_auth,
                "Blade-Auth": "bearer private-token",
                "Tenant-Id": context[0].tenant_id,
            }, "private-token")

    result = asyncio.run(query())
    assert len(result) == 501
    assert len({row["id"] for row in result}) == 501
    assert sum(row["goodsSku"] == "SKU-0" for row in result) == 2
    page_calls = [call for call in server["calls"] if call[1].startswith(exports.LIST_PATH)]
    assert [call[1].split("current=", 1)[1].split("&", 1)[0] for call in page_calls] == ["1", "2"]


def test_empty_enabled_list_accepts_zero_pages(context, server):
    server["list_rows"] = []
    server["list_payloads"] = {1: {"code": 200, "success": True,
                                  "data": {"current": 1, "pages": 0, "total": 0, "records": []}}}
    client = exports.GoodsExporter(context[0], context[1])

    async def query():
        async with exports.aiohttp.ClientSession() as session:
            return await client._enabled_goods(session, {"Blade-Auth": "bearer private-token"}, "private-token")

    assert asyncio.run(query()) == []


@pytest.mark.parametrize("mutation,code", [
    ("missing_record", "pagination_incomplete"),
    ("changed_total", "pagination_incomplete"),
    ("wrong_status", "invalid_response"),
    ("duplicate_id", "pagination_incomplete"),
])
def test_enabled_list_rejects_incomplete_or_invalid_pages(context, server, mutation, code):
    rows = [identity_record(f"id-{i}", f"SKU-{i}", "2026-01-01 00:00:00") for i in range(501)]
    if mutation == "wrong_status":
        rows[0]["status"] = 1
    elif mutation == "duplicate_id":
        rows[-1]["id"] = rows[0]["id"]
    server["list_rows"] = rows
    if mutation in {"missing_record", "changed_total"}:
        payload = {"code": 200, "success": True,
                   "data": {"current": 2, "pages": 2, "total": 501,
                            "records": rows[500:]}}
        if mutation == "missing_record":
            payload["data"]["records"] = []
        else:
            payload["data"]["total"] = 502
        server["list_payloads"] = {2: payload}
    client = exports.GoodsExporter(context[0], context[1])

    async def query():
        async with exports.aiohttp.ClientSession() as session:
            return await client._enabled_goods(session, {"Blade-Auth": "bearer private-token"}, "private-token")

    with pytest.raises(AuthError) as error:
        asyncio.run(query())
    assert error.value.code == code


def test_filter_requires_one_to_one_code_and_create_time_mapping(tmp_path):
    source, target = tmp_path / "source.xlsx", tmp_path / "target.xlsx"
    source.write_bytes(identity_workbook([
        ("SAME-SKU", "2026-01-01 00:00:00", 11, 4),
        ("SAME-SKU", "2026-01-02 00:00:00", 22, 8),
        ("DISABLED", "2026-01-03 00:00:00", 33, 12),
    ]))
    enabled = [identity_record("id-1", "SAME-SKU", "2026-01-01 00:00:00"),
               identity_record("id-2", "SAME-SKU", "2026-01-02 00:00:00")]
    result = exports.filter_enabled_rows(source, target, enabled)
    assert result == {"row_count": 2, "filtered_out_count": 1, "matched_enabled_count": 2}
    from openpyxl import load_workbook
    output = load_workbook(target, data_only=False)
    sheet = output[output.sheetnames[0]]
    headers = [cell.value for cell in sheet[1]]
    assert [sheet.cell(row, headers.index("SKU") + 1).value for row in (2, 3)] == ["SAME-SKU", "SAME-SKU"]
    assert [sheet.cell(row, headers.index("有效库存") + 1).value for row in (2, 3)] == [11, 22]
    assert [sheet.cell(row, headers.index("7天销量") + 1).value for row in (2, 3)] == [4, 8]
    assert sheet.cell(2, headers.index("商品名") + 1).value is None
    output.close()

    source.write_bytes(identity_workbook([
        ("DUP", "2026-01-01 00:00:00", 1, 1),
        ("DUP", "2026-01-01 00:00:00", 2, 2),
    ]))
    with pytest.raises(AuthError, match="唯一匹配"):
        exports.filter_enabled_rows(source, target, [identity_record("id", "DUP", "2026-01-01 00:00:00")])

    with pytest.raises(AuthError, match="唯一关联"):
        exports.filter_enabled_rows(source, target, [
            identity_record("id-1", "DUP", "2026-01-01 00:00:00"),
            identity_record("id-2", "DUP", "2026-01-01 00:00:00"),
        ])

    with pytest.raises(AuthError, match="未能.*唯一匹配"):
        exports.filter_enabled_rows(source, target, [identity_record("missing", "NOT-IN-EXPORT", "2026-01-01 00:00:00")])

    with pytest.raises(AuthError, match="goodsNo 与 goodsSku 不一致"):
        exports.filter_enabled_rows(source, target, [
            {**identity_record("id", "DUP", "2026-01-01 00:00:00"), "goodsNo": "OTHER"}])


@pytest.mark.parametrize("column_count", [len(exports.REQUIRED_HEADERS), 36])
@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("empty", [False, True])
def test_filtered_dimension_uses_retained_cells(tmp_path, column_count, read_only, empty):
    from openpyxl import load_workbook

    headers = sorted(exports.REQUIRED_HEADERS)
    headers += [f"业务字段{i}" for i in range(column_count - len(headers))]
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    created = "2026-01-02 00:00:00"
    for sku in ("DISABLED", "ENABLED"):
        values = {"SKU": sku, "创建时间": created, "有效库存": 17,
                  "7天销量": 3, "15天销量": 5, "30天销量": 8}
        sheet.append([values.get(header, f"{sku}:{header}") for header in headers])
    sheet.cell(2, column_count + 2, "disabled-only")
    sheet.cell(3, headers.index("有效库存") + 1).number_format = "0.00"
    original = io.BytesIO()
    workbook.save(original)
    workbook.close()

    source, target = tmp_path / "source.xlsx", tmp_path / "filtered.xlsx"
    # The official export declares A1 even though its cells span many columns.
    with ZipFile(original) as archive, ZipFile(source, "w") as rewritten:
        for item in archive.infolist():
            content = archive.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', content)
            rewritten.writestr(item, content)

    records = [] if empty else [identity_record("enabled-id", "ENABLED", created)]
    result = exports.filter_enabled_rows(source, target, records)
    assert result["row_count"] == (0 if empty else 1)
    output = load_workbook(target, read_only=read_only, data_only=False)
    try:
        sheet = output.worksheets[0]
        assert sheet.max_column == column_count
        assert sheet.max_row == (1 if empty else 2)
        rows = list(sheet.iter_rows(values_only=True))
        assert rows[0] == tuple(headers)
        assert len(rows) == (1 if empty else 2)
        if not empty:
            expected = {"SKU": "ENABLED", "创建时间": created, "有效库存": 17,
                        "7天销量": 3, "15天销量": 5, "30天销量": 8}
            assert rows[1] == tuple(expected.get(header, f"ENABLED:{header}") for header in headers)
            assert sheet.cell(2, headers.index("有效库存") + 1).number_format == "0.00"
    finally:
        output.close()


@pytest.mark.parametrize(("enabled_code", "workbook_code", "enabled_time", "workbook_time"), [
    ("SKU A", "SKUA", "2026-01-01 00:00:00", "2026-01-01 00:00:00"),
    ("SKU-A", "ＳＫＵ-A", "2026-01-01 00:00:00", "2026-01-01 00:00:00"),
    ("SKU-A", "SKU-A", "2026-01-01 00:00:00.000001", "2026-01-01 00:00:00.000002"),
    ("SKU-A", "SKU-A", "2026-01-01 00:00:00Z", "2026-01-01 00:00:00"),
])
def test_filter_does_not_normalize_distinct_identity_values(
        tmp_path, enabled_code, workbook_code, enabled_time, workbook_time):
    source, target = tmp_path / "source.xlsx", tmp_path / "target.xlsx"
    source.write_bytes(identity_workbook([(workbook_code, workbook_time, 17, 3)]))

    with pytest.raises(AuthError, match="未能.*唯一匹配"):
        exports.filter_enabled_rows(
            source, target, [identity_record("id-1", enabled_code, enabled_time)])

    assert not target.exists()


def test_filter_preserves_distinct_unicode_and_whitespace_skus_and_rejects_key_collision(tmp_path):
    source, target = tmp_path / "source.xlsx", tmp_path / "target.xlsx"
    source.write_bytes(identity_workbook([
        ("SKU A", "2026-01-01 00:00:00", 11, 4),
        ("SKUA", "2026-01-01 00:00:00", 22, 8),
        ("SKU-A", "2026-01-01 00:00:00", 33, 12),
        ("ＳＫＵ-A", "2026-01-01 00:00:00", 44, 16),
    ]))
    enabled = [
        identity_record("id-1", "SKU A", "2026-01-01 00:00:00"),
        identity_record("id-2", "SKUA", "2026-01-01 00:00:00"),
        identity_record("id-3", "SKU-A", "2026-01-01 00:00:00"),
        identity_record("id-4", "ＳＫＵ-A", "2026-01-01 00:00:00"),
    ]
    result = exports.filter_enabled_rows(source, target, enabled)
    assert result["row_count"] == 4
    assert result["matched_enabled_count"] == 4

    # Two enabled records with a single shared key cannot be distinguished by
    # their SKU and creation time, even when their API IDs differ.
    with pytest.raises(AuthError, match="无法唯一关联"):
        exports.filter_enabled_rows(source, target, [
            identity_record("id-5", "SKU A", "2026-01-01 00:00:00"),
            identity_record("id-6", "SKU A", "2026-01-01 00:00:00"),
        ])


def test_filter_rejects_invalid_creation_time_even_when_values_match(tmp_path):
    source, target = tmp_path / "source.xlsx", tmp_path / "target.xlsx"
    source.write_bytes(identity_workbook([("SKU-1", "not-a-time", 17, 3)]))

    with pytest.raises(AuthError, match="无法解析商品创建时间"):
        exports.filter_enabled_rows(
            source, target, [identity_record("id-1", "SKU-1", "not-a-time")])


@pytest.mark.parametrize("options,error", [
    ({"formula": True}, "包含公式"),
    ({"second_product_sheet": True}, "多个商品数据工作表"),
])
def test_filter_fails_closed_for_unsupported_workbook_structures(tmp_path, options, error):
    source, target = tmp_path / "source.xlsx", tmp_path / "target.xlsx"
    source.write_bytes(identity_workbook([
        ("SKU-1", "2026-01-01 00:00:00", 17, 3),
    ], **options))
    with pytest.raises(AuthError, match=error):
        exports.filter_enabled_rows(source, target, [identity_record("id-1", "SKU-1", "2026-01-01 00:00:00")])
    assert not target.exists()


@pytest.mark.parametrize("status,expected", [(401, "login_required"), (403, "enabled_list_http_failed"),
                                               (429, "enabled_list_http_failed")])
def test_enabled_list_http_failures_publish_no_workbook(context, server, status, expected):
    server["list_status"] = status
    result = exports.run({})
    assert not result["success"]
    assert result["error"]["code"] == expected
    assert result.get("files", []) == []
    assert not list((context[2] / "artifacts").rglob("*.xlsx"))
    assert len([call for call in server["calls"] if call[1].startswith(exports.LIST_PATH)]) == 1
    assert context[1].status()["authenticated"] is (status != 401)


def test_enabled_list_business_failure_publishes_nothing(context, server):
    server["list_payload"] = {"code": 500, "success": False, "msg": "actual list business failure"}
    result = exports.run({})
    assert not result["success"]
    assert result["error"]["code"] == "enabled_list_business_failed"
    assert "actual list business failure" in result["error"]["message"]
    assert not list((context[2] / "artifacts").rglob("*.xlsx"))
    assert len([call for call in server["calls"] if call[1].startswith(exports.LIST_PATH)]) == 1


def test_sequential_invocations_use_new_loops_and_unique_output_directories(context, server):
    results = [exports.run({}) for _ in range(3)]
    assert all(r["success"] for r in results), results
    assert len({r["artifact_path"] for r in results}) == 3
    assert len(server["calls"]) == 9
    server["get_status"] = 500
    assert exports.run({})["error"]["code"] == "download_failed"
    server["get_status"] = 200
    assert exports.run({})["success"]


@pytest.mark.parametrize("state", ["missing", "expired", "changed_credentials"])
def test_unusable_login_makes_no_http_request(context, server, state, monkeypatch):
    _, store, _ = context
    if state == "missing":
        store.clear()
    elif state == "expired":
        payload = read_json(store.state_path)
        payload["expires_at"] = time.time() - 1
        atomic_json(store.state_path, payload)
    else:
        monkeypatch.setenv("LXE_SHANGMAN_PROCESSED_PASSWORD", "changed-password")
    result = exports.run({})
    assert not result["success"]
    assert result["error"]["code"] == "login_required"
    assert not server["calls"]


@pytest.mark.parametrize("http_status,payload,invalidates", [
    (401, {"msg": "actual token rejection"}, True),
    (200, {"code": 401, "msg": "actual token rejection"}, True),
    (403, {"msg": "permission denied"}, False),
    (429, {"msg": "rate limited"}, False),
    (500, {"msg": "server failed"}, False),
    (302, {"msg": "redirect"}, False),
    (200, {"code": 500, "success": False, "msg": "business failed"}, False),
    (200, "not-json", False),
    (200, [1, 2], False),
])
def test_export_failures_keep_diagnostics_without_retries(context, server, http_status, payload, invalidates):
    server.update(post_status=http_status, payload=payload)
    result = exports.run({})
    assert not result["success"]
    assert len(server["calls"]) == 1
    expected = payload.get("msg") if isinstance(payload, dict) else payload if isinstance(payload, str) else "[1, 2]"
    assert expected in result["error"]["message"]
    assert context[1].status()["authenticated"] is not invalidates
    if invalidates:
        assert result["error"]["code"] == "login_required"


def test_rejected_old_token_does_not_erase_concurrent_login(context, server):
    server.update(post_status=401, payload={"msg": "expired"},
                  before_response=lambda: context[1].save(LoginToken("new-token", 300)))
    assert exports.run({})["error"]["code"] == "login_required"
    assert context[1].read_token() == "new-token"


@pytest.mark.parametrize("status", [401, 403, 429, 500, 302])
def test_oss_download_failures_do_not_invalidate_erp_login(context, server, status):
    server.update(get_status=status, xlsx=b"original OSS failure")
    result = exports.run({})
    assert not result["success"]
    assert "original OSS failure" in result["error"]["message"]
    assert len(server["calls"]) == 2
    assert context[1].status()["authenticated"]
    assert not list(context[2].rglob("*.xlsx"))


@pytest.mark.parametrize("content", [b"", b"<html>login page</html>", workbook_bytes(bad_headers=True)])
def test_invalid_workbooks_are_not_published(context, server, content):
    server["xlsx"] = content
    assert exports.run({})["error"]["code"] == "invalid_workbook"
    assert not list(context[2].rglob("*.xlsx"))


@pytest.mark.parametrize("url", [
    "http://oss.erp.shangmanet.com/file.xlsx", "https://evil.test/file.xlsx",
    "https://oss.erp.shangmanet.com.evil.test/file.xlsx", "https://user:pass@erp.shangmanet.com/file",
    "https://oss.erp.shangmanet.com:8443/file", "https://oss.erp.shangmanet.com/fi\nle",
    None,
])
def test_download_host_restrictions(context, url):
    client = exports.GoodsExporter(context[0], context[1])
    with pytest.raises(AuthError, match="HTTPS"):
        client._validate_download_url(url)


@pytest.mark.parametrize("host", ["erp.shangmanet.com", "oss.erp.shangmanet.com"])
def test_approved_download_hosts(context, host):
    url = f"https://{host}/goods.xlsx?sign=secret"
    assert exports.GoodsExporter(context[0], context[1])._validate_download_url(url) == url


def test_error_diagnostics_redact_token_credentials_and_signed_query(context, server):
    server.update(post_status=500, payload={"msg": "private-token password-789 " + server["url"]})
    message = exports.run({})["error"]["message"]
    assert "HTTP 500" in message
    for secret in ("private-token", "password-789", "signed-private-value"):
        assert secret not in message


def test_cli_delivers_exactly_one_original_workbook(context, server, capsys, monkeypatch):
    from lxeskill import cli
    from lxeskill.business import load_catalog
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", "shangman-products-export")
    entry = load_catalog()["shangman_goods_export"]
    assert entry["input_schema"] == {"type": "object", "properties": {}, "additionalProperties": False}
    assert cli.main(["shangman", "products", "export"]) == 0
    terminal = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert terminal["ok"]
    assert terminal["files"] == [terminal["data"]["artifact_path"]]
    assert Path(terminal["files"][0]).is_file()
    assert len(list(Path(terminal["files"][0]).parent.glob("*.xlsx"))) == 1
    server["list_status"] = 403
    assert cli.main(["shangman", "products", "export"]) != 0
    failed_list = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert not failed_list["ok"] and failed_list["files"] == []
    server["list_status"] = 200
    server["post_status"] = 403
    assert cli.main(["shangman", "products", "export"]) != 0
    failed = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert not failed["ok"] and failed["files"] == []


def test_network_timeout_preserves_exception_and_does_not_retry(context, monkeypatch):
    calls = []
    def timeout(*args, **kwargs):
        calls.append(1)
        raise TimeoutError("actual connection deadline")
    monkeypatch.setattr(exports.aiohttp, "ClientSession", timeout)
    result = exports.run({})
    assert "TimeoutError: actual connection deadline" in result["error"]["message"]
    assert calls == [1]
    assert context[1].status()["authenticated"]
