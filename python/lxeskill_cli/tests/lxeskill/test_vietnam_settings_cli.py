from __future__ import annotations

from importlib.resources import files
import json
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from lxeskill import cli
from lxeskill.business import load_catalog
from services.agent_cli.vietnam_replenishment import settings as command
from services.vietnam_replenishment import settings
from services.vietnam_replenishment.full_workbook import UNCHANGED


COMMAND = ["vietnam", "replenishment", "settings"]


def test_settings_command_is_internal_and_hidden_from_listing(capsys):
    entry = load_catalog()["vietnam_replenishment_settings"]
    assert entry["command_path"] == COMMAND
    assert entry["visibility"] == "internal"
    assert entry["exposed"] is False
    assert entry["session_mode"] == "none"
    assert entry["owner_skills"] == []
    assert cli.main(["list"]) == 0
    result = json.loads(capsys.readouterr().out)
    assert " ".join(COMMAND) not in {item["command"] for item in result["data"]["commands"]}


@pytest.mark.parametrize("arguments", [
    {}, {"action": "invalid"}, {"action": []},
    {"action": "read", "path": "/unexpected.xlsx"},
    {"action": "save"}, {"action": "upload"}, {"action": "export"},
    {"action": "export", "kind": "current", "path": "/unexpected-kind.xlsx"},
    *[{"action": "upload", "path": path} for path in (None, 1, "", " ", "relative.xlsx")],
])
def test_invalid_action_arguments_never_reach_storage(arguments, monkeypatch):
    def unexpected(*_args):
        pytest.fail("invalid arguments must not reach settings storage")

    for name in ("read_state", "save_parameters", "upload_template", "export_template"):
        monkeypatch.setattr(settings, name, unexpected)
    result = command.run(arguments)
    assert result["success"] is False
    assert result["error"]["code"] == "ValueError"
    assert result["error"]["message"]


def test_cli_round_trip_uses_shared_storage_and_preserves_failed_upload(tmp_path, monkeypatch):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "app"))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path))
    monkeypatch.delenv("LXE_AGENT_SESSION_ID", raising=False)

    def invoke(action, *flags, expected=0):
        child = subprocess.run(
            [sys.executable, "-I", "-B", "-m", "lxeskill", *COMMAND, "--action", action, *flags],
            capture_output=True, text=True, encoding="utf-8",
        )
        assert child.returncode == expected, child.stderr + child.stdout
        result = json.loads(child.stdout)
        assert result["type"] == "result"
        assert result["protocol_version"] == "1"
        assert result["ok"] is (expected == 0)
        return result

    parameters = invoke("read")["data"]["data"]["parameters"]
    parameters["exchange_rate"] = "4200"
    invoke("save", "--parameters", json.dumps(parameters))
    assert invoke("read")["data"]["data"]["parameters"] == parameters

    source = tmp_path / "中文 九表模板.xlsx"
    resource = files("services.vietnam_replenishment").joinpath("resources/skeleton.xlsx")
    with resource.open("rb") as stream:
        book = load_workbook(stream)
    for index, name in enumerate(UNCHANGED, 1):
        book.create_sheet(name, index)["A1"] = f"original {name}"
    book.save(source)
    book.close()
    state = invoke("upload", "--path", str(source))["data"]["data"]
    uploaded = state["template"]
    assert state["template"]["file_name"] == source.name
    assert "sku_map" not in state
    assert Path(state["template"]["path"]).read_bytes() == source.read_bytes()
    with settings.run_inputs() as (config, _):
        assert str(config.exchange_rate) == "4200"
    invalid = tmp_path / "坏表.xlsx"
    invalid.write_bytes(b"not an xlsx")
    failed = invoke("upload", "--path", str(invalid), expected=cli.EXIT_BUSINESS)
    assert "BadZipFile" in failed["error"]["message"]
    assert failed["error"]["message"] == failed["data"]["error"]["message"]
    assert invoke("read")["data"]["data"]["template"] == uploaded
    output = tmp_path / "导出 当前模板.xlsx"
    assert invoke("export", "--path", str(output))["data"]["data"] == {"path": str(output)}
    assert output.read_bytes() == source.read_bytes()


def test_settings_cli_rejects_retired_export_kind(tmp_path, capsys):
    assert cli.main([*COMMAND, "--action", "export", "--kind", "template", "--path", str(tmp_path / "out.xlsx")]) == cli.EXIT_USAGE
    result = json.loads(capsys.readouterr().out)
    assert "unknown option: --kind" in result["error"]["message"]
    assert not (tmp_path / "out.xlsx").exists()


@pytest.mark.parametrize("action", ["read", "invalid"])
def test_settings_cli_emits_utf8_even_with_legacy_stdio(tmp_path, monkeypatch, action):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "中文 数据"))
    script = """
import runpy, sys
sys.stdout.reconfigure(encoding="gbk")
sys.stderr.reconfigure(encoding="gbk")
sys.argv = ["lxeskill", "vietnam", "replenishment", "settings", "--action", sys.argv[1]]
runpy.run_module("lxeskill", run_name="__main__")
"""
    child = subprocess.run([sys.executable, "-I", "-B", "-c", script, action], capture_output=True)
    result = json.loads(child.stdout.decode("utf-8"))
    if action == "read":
        assert child.returncode == 0, child.stderr
        assert "中文 数据" in result["data"]["data"]["directory"]
    else:
        assert child.returncode == cli.EXIT_BUSINESS, child.stderr
        assert result["ok"] is False
        assert "无效的桌面设置操作" in result["error"]["message"]
