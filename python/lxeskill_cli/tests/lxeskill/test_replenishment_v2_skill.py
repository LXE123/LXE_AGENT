from __future__ import annotations

import json
import re
from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import Workbook

from lxeskill import cli as lxeskill
from lxeskill.business import load_catalog
from services.agent_cli.mabang import analyze_store_msku_sales as sales_adapter
from services.agent_cli.mabang import calculate_store_msku_replenishment as ordinary_adapter
from services.agent_cli.mabang import generate_replenishment_v2 as adapter
from services.mabang.amazon.fba.replenishment_v2_workbook import ReplenishmentV2WorkbookResult, SOURCE_REQUIRED_COLUMNS
from services.mabang.amazon.fba.sku_catalog import LocalSkuDefinition, SkuCatalogSnapshot
from services.mabang.amazon.fba.source_verification import annotate_source, load_verified_source, stamp_report


COMMAND = ["replenish", "v2", "generate"]
WORKFLOW_SKILL = "replenishment-workflow-map"
INPUTS = {
    "store_name": "Amazon-Test",
    "msku_source_xlsx_path": "source.xlsx",
    "shenzhen_inventory_report_xlsx_path": "inventory.xlsx",
    "unlinked_shipments_snapshot_path": "snapshot.xlsx",
}


def _args(inputs: dict[str, str]) -> list[str]:
    return [*COMMAND, *(item for key, value in inputs.items() for item in (f"--{key.replace('_', '-')}", value))]


def _terminal(capsys) -> dict:
    records = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line.strip()]
    assert len(records) == 1
    assert records[0]["type"] == "result"
    return records[0]


@pytest.fixture(autouse=True)
def isolated_scope(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", WORKFLOW_SKILL)
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path / "workspace"))


def test_replenishment_v2_is_exposed_with_one_owner_and_the_existing_inputs(capsys) -> None:
    entry = load_catalog()["mabang_generate_replenishment_v2"]
    assert entry["visibility"] == "business"
    assert entry["exposed"] is True
    assert entry["owner_skills"] == [WORKFLOW_SKILL]
    assert entry["input_modes"] == ["flags"]
    schema = entry["input_schema"]
    assert schema["required"] == list(INPUTS) and schema["additionalProperties"] is False
    assert {key: schema["properties"][key] for key in INPUTS} == {key: {"type": "string", "minLength": 1} for key in INPUTS}
    assert schema["properties"]["peak_boost_enabled"]["type"] == "boolean"
    assert schema["properties"]["peak_boost_days"]["minimum"] == 0
    assert schema["properties"]["peak_boost_days"]["maximum"] == 120
    assert entry["deliver_artifacts_on_failure"] is True
    assert lxeskill.main(["list"]) == 0
    commands = _terminal(capsys)["data"]["commands"]
    assert any(item["command"] == " ".join(COMMAND) for item in commands)
    assert all(item["command"] != "replenish calculate" for item in commands)
    assert all(item["command"] != "replenish v2 settings" for item in commands)
    assert lxeskill.main(["describe", *COMMAND]) == 0
    described = _terminal(capsys)["data"]
    assert described["owner_skills"] == [WORKFLOW_SKILL]
    assert described["input_schema"]["required"] == list(INPUTS)
    assert "--peak-boost-enabled=false" in described["input_schema"]["properties"]["peak_boost_enabled"]["description"]


@pytest.mark.parametrize("operation", ["generate", "settings"])
@pytest.mark.parametrize("prefix", [[], ["describe"]])
def test_retired_v2_command_names_are_not_aliases(operation, prefix, monkeypatch, capsys):
    execute = Mock(side_effect=AssertionError("retired commands must not execute business modules"))
    monkeypatch.setattr(lxeskill, "execute_module_json", execute)
    assert lxeskill.main([*prefix, "replenish", "final-restock", operation]) == lxeskill.EXIT_USAGE
    assert _terminal(capsys)["error"]["code"] == "unknown_command"
    execute.assert_not_called()


@pytest.mark.parametrize("workflow,scope,command", [
    ("V1", "replenishment-calculate", ["replenish", "calculate"]),
    ("V2", WORKFLOW_SKILL, COMMAND),
])
def test_main_workflow_passes_the_current_snapshot_to_the_right_calculator(
    workflow, scope, command, tmp_path, monkeypatch, capsys,
):
    policy = (Path(__file__).resolve().parents[4] / "skills" / "replenishment-workflow-map" / "SKILL.md").read_text(encoding="utf-8")
    routes = [line for line in policy.splitlines() if line.startswith(f"| {workflow} |") and "lxeskill" in line]
    assert len(routes) == 1, "the main workflow must distinguish the two calculation contracts"
    documented_command = re.search(r"`lxeskill ([^`]+?) --", routes[0])
    documented_flag = re.search(r"(--unlinked-shipments-[a-z-]+)", routes[0])
    assert documented_command and documented_flag
    assert documented_command.group(1).split() == command
    if workflow == "V1":
        assert scope in routes[0]
    else:
        assert load_catalog()["mabang_generate_replenishment_v2"]["owner_skills"] == [scope]

    snapshot = tmp_path / "本轮货件 快照.xlsx"
    output = tmp_path / "workspace/.lxeagent/artifacts/replenish/result.xlsx"
    output.parent.mkdir(parents=True)
    book = Workbook()
    book.save(snapshot)
    book.save(output)
    book.close()
    output_field = "output_xlsx_path" if scope == WORKFLOW_SKILL else "report_xlsx_path"
    result = Mock()
    result.to_payload.return_value = {"success": True, output_field: str(output)}
    final_calculate = Mock(return_value=result)
    ordinary_calculate = Mock(return_value=result)
    sales = Mock(side_effect=AssertionError("calculation handoff must not start sales analysis"))
    monkeypatch.setattr(adapter, "generate_replenishment_v2", final_calculate)
    monkeypatch.setattr(ordinary_adapter, "calculate_store_msku_replenishment", ordinary_calculate)
    monkeypatch.setattr(sales_adapter, "analyze_store_msku_sales", sales)
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", scope)
    arguments = {"store_name": INPUTS["store_name"]}
    if scope == WORKFLOW_SKILL:
        arguments.update({key: INPUTS[key] for key in ("msku_source_xlsx_path", "shenzhen_inventory_report_xlsx_path")})
    argv = [*documented_command.group(1).split(),
            *(item for key, value in arguments.items() for item in (f"--{key.replace('_', '-')}", value)),
            documented_flag.group(1), str(snapshot)]
    assert lxeskill.main(argv) == 0
    terminal = _terminal(capsys)
    assert terminal["ok"] is True and terminal["files"] == [str(output)]
    selected, other = (final_calculate, ordinary_calculate) if scope == WORKFLOW_SKILL else (ordinary_calculate, final_calculate)
    assert selected.call_count == 1
    assert selected.call_args.kwargs["unlinked_shipments_snapshot_path"] == str(snapshot)
    other.assert_not_called()
    sales.assert_not_called()
    assert "lxeskill describe <命令路径>" in policy


@pytest.mark.parametrize("scope,command,wrong_flag", [
    (WORKFLOW_SKILL, COMMAND, "--unlinked-shipments-snapshot"),
    ("replenishment-calculate", ["replenish", "calculate"], "--unlinked-shipments-snapshot-path"),
])
def test_shipment_snapshot_flags_are_not_interchangeable(scope, command, wrong_flag, monkeypatch, capsys):
    final_calculate = Mock(side_effect=AssertionError("invalid flag must not execute calculation"))
    ordinary_calculate = Mock(side_effect=AssertionError("invalid flag must not execute calculation"))
    monkeypatch.setattr(adapter, "generate_replenishment_v2", final_calculate)
    monkeypatch.setattr(ordinary_adapter, "calculate_store_msku_replenishment", ordinary_calculate)
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", scope)
    assert lxeskill.main([*command, "--store-name", "Amazon-Test", wrong_flag, "snapshot.xlsx"]) == lxeskill.EXIT_USAGE
    terminal = _terminal(capsys)
    assert terminal["ok"] is False and terminal["files"] == []
    assert terminal["error"]["message"] == f"unknown option: {wrong_flag}"
    final_calculate.assert_not_called()
    ordinary_calculate.assert_not_called()


@pytest.mark.parametrize("name,owner,fields", [
    ("mabang_resolve_fba_store", "replenishment-store-resolve", {"store_name"}),
    ("mabang_download_store_msku_excel", "mabang-store-msku-export", {"store_name", "store_id", "id_type"}),
    ("mabang_export_store_msku_actual_inventory", "mabang-store-shenzhen-inventory-export", {"store_name"}),
    ("mabang_download_store_unlinked_shipments", "mabang-store-unlinked-shipments-export", {"store_name"}),
])
def test_store_only_acquisition_reuses_exposed_existing_commands(name, owner, fields) -> None:
    entry = load_catalog()[name]
    assert entry["visibility"] == "business"
    assert entry["exposed"] is True
    assert owner in entry["owner_skills"]
    assert fields <= set(entry["input_schema"]["properties"])


def test_shenzhen_source_is_optional_for_legacy_and_explicit_in_main_workflow() -> None:
    entry = load_catalog()["mabang_export_store_msku_actual_inventory"]
    schema = entry["input_schema"]
    assert schema["required"] == ["store_name"]
    assert schema["properties"]["msku_source_xlsx_path"]["type"] == "string"
    assert schema["properties"]["msku_source_xlsx_path"]["minLength"] == 1
    skill = (Path(__file__).resolve().parents[4] / "skills" / "replenishment-workflow-map" / "SKILL.md").read_text(encoding="utf-8")
    assert "--msku-source-xlsx-path" in skill
    assert "source_msku_xlsx_path" in skill


def test_replenishment_v2_uses_shenzhen_stock_without_native_msku_inventory_opt_in() -> None:
    entry = load_catalog()["mabang_download_store_msku_excel"]
    assert "include_local_inventory" not in entry["input_schema"]["properties"]
    assert "本地库存" not in SOURCE_REQUIRED_COLUMNS
    skill = (Path(__file__).resolve().parents[4] / "skills" / WORKFLOW_SKILL / "SKILL.md").read_text(encoding="utf-8")
    workflow = (Path(__file__).resolve().parents[4] / 'skills' / 'replenishment-workflow-map' / 'SKILL.md').read_text(encoding='utf-8')
    assert 'lxeskill mabang store msku export --store-id' in workflow
    assert '--include-local-inventory' not in skill
    assert '真实库存（深圳仓库）数量' in skill


@pytest.mark.parametrize("scope,command", [
    ("replenishment-calculate", COMMAND),
    (WORKFLOW_SKILL, ["replenish", "calculate"]),
    ("", COMMAND),
    ("replenishment-final-restock", COMMAND),
])
def test_new_and_old_commands_respect_their_own_skill_scope(scope, command, monkeypatch, capsys) -> None:
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", scope)
    assert lxeskill.main(command) == lxeskill.EXIT_ENVIRONMENT
    terminal = _terminal(capsys)
    assert terminal["error"]["code"] == "skill_not_in_scope"
    assert terminal["files"] == []


@pytest.mark.parametrize("missing,diagnostic", [
    ("store_name", "store_name 不能为空"),
    ("msku_source_xlsx_path", "MSKU 源文件不存在"),
    ("shenzhen_inventory_report_xlsx_path", "深圳库存报告不存在"),
    ("unlinked_shipments_snapshot_path", "未关联货件快照不存在"),
])
def test_incomplete_inputs_return_actual_service_diagnostic(missing, diagnostic, tmp_path: Path, capsys) -> None:
    source = tmp_path / "existing-input.xlsx"
    source.touch()
    inputs = {key: str(source) for key in INPUTS if key != "store_name"}
    inputs["store_name"] = "Amazon-Test"
    assert lxeskill.main(_args({key: value for key, value in inputs.items() if key != missing})) == lxeskill.EXIT_BUSINESS
    terminal = _terminal(capsys)
    assert terminal["ok"] is False
    assert terminal["error"]["code"] == "business_cli_failed"
    assert diagnostic in terminal["error"]["message"]
    assert terminal["data"]["exception"] == terminal["error"]["message"]
    assert terminal["files"] == []


def test_removed_source_path_flag_is_rejected_in_the_new_skill_scope(capsys) -> None:
    assert lxeskill.main([
        *_args(INPUTS), "--shenzhen-source-msku-xlsx-path", "source.xlsx",
    ]) == lxeskill.EXIT_USAGE
    terminal = _terminal(capsys)
    assert terminal["error"]["code"] == "invalid_arguments"
    assert terminal["error"]["message"] == "unknown option: --shenzhen-source-msku-xlsx-path"
    assert terminal["files"] == []


def test_source_without_native_local_inventory_passes_cli_source_validation(tmp_path: Path, monkeypatch, capsys) -> None:
    from services.mabang.amazon.fba import replenishment_v2_workbook as service

    source = tmp_path / "202610151000-Amazon-Test_店铺MSKU数据.xlsx"
    workbook = Workbook()
    workbook.active.append(SOURCE_REQUIRED_COLUMNS)
    workbook.active.append([
        "MSKU-1" if column == "MSKU" else "SKU-1" if column == "本地SKU" else 0
        for column in SOURCE_REQUIRED_COLUMNS
    ])
    workbook.save(source)
    workbook.close()
    annotate_source(
        source,
        SkuCatalogSnapshot("Amazon-Test", "100", "shopId", (LocalSkuDefinition("SKU-1", 1),)),
        requested_store_name="Amazon-Test",
    )
    inventory = tmp_path / "inventory.xlsx"
    workbook = Workbook()
    workbook.save(inventory)
    workbook.close()
    stamp_report(inventory, load_verified_source(source, store_name="Amazon-Test").metadata)
    snapshot = tmp_path / "snapshot.xlsx"
    snapshot.touch()
    # A sentinel at the next stage proves the CLI did not reject the missing uq122 column.
    read_inventory = Mock(side_effect=service.ReplenishmentV2WorkbookError("已进入深圳报告数量读取"))
    monkeypatch.setattr(service, "_load_shenzhen_rows", read_inventory)
    assert lxeskill.main(_args({
        "store_name": "Amazon-Test", "msku_source_xlsx_path": str(source),
        "shenzhen_inventory_report_xlsx_path": str(inventory),
        "unlinked_shipments_snapshot_path": str(snapshot),
    })) == lxeskill.EXIT_BUSINESS
    terminal = _terminal(capsys)
    assert terminal["error"]["message"] == "已进入深圳报告数量读取"
    read_inventory.assert_called_once()


def test_unknown_store_returns_actual_local_source_mismatch(tmp_path: Path, capsys) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.append(["MSKU", "本地SKU"])
    workbook.active.append(["MSKU-1", "SKU-1"])
    workbook.save(source)
    workbook.close()
    annotate_source(
        source,
        SkuCatalogSnapshot("Amazon-Test", "100", "shopId", (LocalSkuDefinition("SKU-1", 1),)),
        requested_store_name="Amazon-Test",
    )
    # Store verification must fail before reading the inventory or shipment inputs.
    inputs = {key: str(source) for key in INPUTS if key != "store_name"}
    assert lxeskill.main(_args({"store_name": "Unknown-Store", **inputs})) == lxeskill.EXIT_BUSINESS
    terminal = _terminal(capsys)
    assert terminal["ok"] is False
    assert terminal["error"]["code"] == "business_cli_failed"
    assert terminal["error"]["message"] == "源数据核验店铺不一致: 请求=Unknown-Store, 快照=Amazon-Test"
    assert terminal["data"]["exception"] == terminal["error"]["message"]
    assert terminal["files"] == []


@pytest.mark.parametrize("formula_errors", [0, 1])
def test_terminal_preserves_success_or_formula_failure_and_file(
    formula_errors, tmp_path: Path, monkeypatch, capsys,
) -> None:
    output = tmp_path / "workspace" / ".lxeagent" / "artifacts" / "replenish" / "final_restock_recommendation" / "test.xlsx"
    output.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.save(output)
    workbook.close()
    output = output.resolve()
    result = ReplenishmentV2WorkbookResult(
        store_name="Amazon-Test", source_data_time="202610131000", output_xlsx_path=str(output),
        row_count=1, formula_error_count=formula_errors,
        formula_error_examples=("AC2=#VALUE!",) if formula_errors else (),
    )
    generate = Mock(return_value=result)
    monkeypatch.setattr(adapter, "generate_replenishment_v2", generate)
    assert lxeskill.main(_args(INPUTS)) == (lxeskill.EXIT_BUSINESS if formula_errors else 0)
    terminal = _terminal(capsys)
    assert terminal["ok"] is (formula_errors == 0)
    assert terminal["data"] == result.to_payload()
    assert terminal["files"] == [str(output)]
    if formula_errors:
        assert terminal["error"]["message"] == result.to_payload()["exception"]
        assert "AC2=#VALUE!" in terminal["error"]["message"]
    generate.assert_called_once_with("Amazon-Test", **{key: value for key, value in INPUTS.items() if key != "store_name"}, peak_boost_enabled=None, peak_boost_days=None)


@pytest.mark.parametrize("days_flags,expected_days", [([], None), (["--peak-boost-days", "20"], 20)])
def test_explicit_current_request_boost_flags_reach_existing_generator(
    tmp_path, monkeypatch, capsys, days_flags, expected_days,
):
    from services.agent_cli.mabang import generate_replenishment_v2 as adapter
    output = tmp_path / "workspace/.lxeagent/artifacts/replenish/final_restock_recommendation/result.xlsx"
    output.parent.mkdir(parents=True)
    book = Workbook()
    book.save(output)
    book.close()
    result = ReplenishmentV2WorkbookResult("Amazon-Test", "202610201000", str(output), 1, 0, ())
    generate = Mock(return_value=result)
    monkeypatch.setattr(adapter, "generate_replenishment_v2", generate)
    assert lxeskill.main([*_args(INPUTS), "--peak-boost-enabled", *days_flags]) == 0
    assert _terminal(capsys)["ok"] is True
    assert generate.call_args.kwargs["peak_boost_enabled"] is True
    assert generate.call_args.kwargs["peak_boost_days"] == expected_days


@pytest.mark.parametrize("flags", [
    ["--peak-boost-enabled=invalid"],
    ["--peak-boost-enabled=0"],
    ["--peak-boost-enabled="],
    ["--peak-boost-enabled", "false"],
])
def test_boost_flag_values_fail_before_business_execution(flags, monkeypatch, capsys):
    generate = Mock(side_effect=AssertionError("invalid flag values must not generate a workbook"))
    monkeypatch.setattr(adapter, "generate_replenishment_v2", generate)
    assert lxeskill.main([*_args(INPUTS), *flags]) == lxeskill.EXIT_USAGE
    terminal = _terminal(capsys)
    assert terminal["ok"] is False
    assert terminal["error"]["code"] == "invalid_arguments"
    assert terminal["files"] == []
    generate.assert_not_called()


@pytest.mark.parametrize("days", ["-1", "121", "1.5", "invalid"])
def test_invalid_request_days_stop_before_reading_business_inputs(days, monkeypatch, capsys):
    from services.mabang.amazon.fba import replenishment_v2_workbook as service
    read_input = Mock(side_effect=AssertionError("invalid days must not read input files"))
    monkeypatch.setattr(service, "_path", read_input)
    assert lxeskill.main([*_args(INPUTS), "--peak-boost-days", days]) != 0
    assert _terminal(capsys)["ok"] is False
    read_input.assert_not_called()


def test_unknown_inventory_keeps_review_file_but_returns_cli_failure(tmp_path, monkeypatch, capsys) -> None:
    output = tmp_path / "workspace/.lxeagent/artifacts/replenish/final_restock_recommendation/result-库存待核对.xlsx"
    output.parent.mkdir(parents=True)
    output.touch()
    result = ReplenishmentV2WorkbookResult(
        "Amazon-Test", "202610201000", str(output), 1, 0, (), 1,
        ("MSKU=A；类别=无本地SKU；原因=源表本地SKU为空",),
    )
    generate = Mock(return_value=result)
    monkeypatch.setattr(adapter, "generate_replenishment_v2", generate)
    assert lxeskill.main(_args(INPUTS)) == lxeskill.EXIT_BUSINESS
    terminal = _terminal(capsys)
    assert terminal["ok"] is False
    assert terminal["files"] == [str(output)]
    assert terminal["data"]["success"] is False
    assert terminal["data"]["unknown_inventory_count"] == 1
    assert terminal["data"]["formula_error_count"] == 0
    assert "须运营确认后才能正式交付" in terminal["error"]["message"]
    generate.assert_called_once()


def test_main_skill_is_the_single_store_name_workflow_for_v1_v2_and_both() -> None:
    skills = Path(__file__).resolve().parents[4] / "skills"
    workflow = (skills / WORKFLOW_SKILL / "SKILL.md").read_text(encoding="utf-8")
    assert not (skills / "replenishment-final-restock" / "SKILL.md").exists()
    assert "replenishment-final-restock" not in workflow
    assert load_catalog()["mabang_generate_replenishment_v2"]["owner_skills"] == [WORKFLOW_SKILL]
    for command in (
        "lxeskill replenish store resolve", "lxeskill mabang store msku export",
        "lxeskill mabang store shenzhen-inventory export", "lxeskill mabang store unlinked-shipments export",
        "lxeskill replenish calculate", "lxeskill replenish v2 generate", "send_files",
    ):
        assert command in workflow
    # One static policy check owns stable command/data terms. Runtime tests cover
    # discovery and permissions. Real Agent version decisions and stopping still
    # require acceptance through agentcli/Desktop; these assertions do not prove them.
    for term in (
        "lxeskill list", "lxeskill describe", "V1", "V2", "replenishment-sales-analyze",
        "--msku-source-xlsx-path", "--shenzhen-inventory-report-xlsx-path",
        "--unlinked-shipments-snapshot-path", "--peak-boost-enabled", "--peak-boost-days",
        "真实库存（深圳仓库）数量", "35列", "14个公式字段", "Office Kit",
        "库存待核对", "success=false", "ok=false",
    ):
        assert term in workflow
    for retired in ("catalog.json", "--include-local-inventory", "lxeskill replenish v2 settings"):
        assert retired not in workflow
    assert load_catalog()["mabang_analyze_store_msku_sales"]["owner_skills"] == ["replenishment-sales-analyze"]
    assert "require_complete_inventory" not in load_catalog()["mabang_export_store_msku_actual_inventory"]["input_schema"]["properties"]
