from __future__ import annotations

import json
import re
from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import Workbook

from lxeskill import cli as lxeskill
from lxeskill.business import load_catalog
from services.agent_cli.mabang import analyze_store_msku_sales as sales_adapter
from services.agent_cli.mabang import calculate_store_msku_replenishment as ordinary_adapter
from services.agent_cli.mabang import generate_final_restock_recommendation as adapter
from services.mabang.amazon.fba.final_restock_workbook import FinalRestockWorkbookResult, SOURCE_REQUIRED_COLUMNS
from services.mabang.amazon.fba.sku_catalog import LocalSkuDefinition, SkuCatalogSnapshot
from services.mabang.amazon.fba.source_verification import annotate_source


COMMAND = ["replenish", "final-restock", "generate"]
SKILL = "replenishment-final-restock"
INPUTS = {
    "store_name": "Amazon-Test",
    "msku_source_xlsx_path": "source.xlsx",
    "shenzhen_inventory_report_xlsx_path": "inventory.xlsx",
    "unlinked_shipments_snapshot_path": "snapshot.xlsx",
}


def _args(inputs: dict[str, str]) -> list[str]:
    return [*COMMAND, *(item for key, value in inputs.items() for item in (f"--{key.replace('_', '-')}", value))]


def _terminal(capsys) -> dict:
    records = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line.strip()]
    assert len(records) == 1
    assert records[0]["type"] == "result"
    return records[0]


@pytest.fixture(autouse=True)
def isolated_scope(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", SKILL)
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path / "workspace"))


def test_final_restock_is_exposed_with_one_owner_and_the_existing_inputs(capsys) -> None:
    entry = load_catalog()["mabang_generate_final_restock_recommendation"]
    assert entry["visibility"] == "business"
    assert entry["exposed"] is True
    assert entry["owner_skills"] == [SKILL]
    assert entry["input_modes"] == ["flags"]
    schema = entry["input_schema"]
    assert schema["required"] == list(INPUTS) and schema["additionalProperties"] is False
    assert {key: schema["properties"][key] for key in INPUTS} == {key: {"type": "string", "minLength": 1} for key in INPUTS}
    assert schema["properties"]["peak_boost_enabled"]["type"] == "boolean"
    assert schema["properties"]["peak_boost_days"]["minimum"] == 0
    assert schema["properties"]["peak_boost_days"]["maximum"] == 120
    assert entry["deliver_artifacts_on_failure"] is True
    assert lxeskill.main(["list"]) == 0
    commands = _terminal(capsys)["data"]["commands"]
    assert any(item["command"] == " ".join(COMMAND) for item in commands)
    assert all(item["command"] != "replenish calculate" for item in commands)
    assert all(item["command"] != "replenish final-restock settings" for item in commands)
    assert lxeskill.main(["describe", *COMMAND]) == 0
    described = _terminal(capsys)["data"]
    assert described["owner_skills"] == [SKILL]
    assert described["input_schema"]["required"] == list(INPUTS)
    assert "传裸开关 --peak-boost-enabled，不附带值；否则省略" in described["input_schema"]["properties"]["peak_boost_enabled"]["description"]


def test_skills_use_public_queries_and_distinguish_unverified_calculation_rules() -> None:
    skills = Path(__file__).resolve().parents[4] / "skills"
    final_policy = (skills / SKILL / "SKILL.md").read_text(encoding="utf-8")
    inventory_policy = (skills / "mabang-store-shenzhen-inventory-export" / "SKILL.md").read_text(encoding="utf-8")
    assert "lxeskill list" in final_policy and "lxeskill describe" in final_policy
    assert "catalog.json" not in final_policy
    assert "`internal` 命令默认不出现在 `list` 中" in final_policy
    assert "桌面设置命令不属于本 Skill 的模型业务入口" in final_policy
    assert "lxeskill replenish final-restock settings" not in final_policy
    assert "普通备货按原核验规则排除未通过商品" in inventory_policy
    assert "保留全部 MSKU 并写入完整 14 个公式字段" in inventory_policy
    assert "商品类型核验未通过本身不阻断最终备货" in inventory_policy
    assert "文件级来源异常仍须失败停止" in inventory_policy
    msku_policy = (skills / "mabang-store-msku-export" / "SKILL.md").read_text(encoding="utf-8")
    assert "普通备货流程中" in msku_policy
    assert "普通备货完整任务下载成功后继续销量与库存" in msku_policy
    assert "普通备货没有可计算记录时" in msku_policy
    assert "保留全部 MSKU，并为所有行生成完整 14 个公式字段" in msku_policy
    assert "按 `replenishment-final-restock` 继续" in msku_policy
    assert "不额外运行普通备货销量分析" in msku_policy


@pytest.mark.parametrize("workflow,scope,command", [
    ("普通备货", "replenishment-calculate", ["replenish", "calculate"]),
    ("Amazon 最终备货", SKILL, COMMAND),
])
def test_shipment_skill_routes_pass_the_current_snapshot_to_the_right_calculator(
    workflow, scope, command, tmp_path, monkeypatch, capsys,
):
    policy = (Path(__file__).resolve().parents[4] / "skills" / "mabang-store-unlinked-shipments-export" / "SKILL.md").read_text(encoding="utf-8")
    routes = [line for line in policy.splitlines() if line.startswith(f"| {workflow} |")]
    assert len(routes) == 1, "shipment Skill must distinguish the two calculation contracts"
    documented_command = re.search(r"`lxeskill ([^`]+)`", routes[0])
    documented_flag = re.search(r"`(--unlinked-shipments-[^`]+)`", routes[0])
    assert documented_command and documented_flag
    assert documented_command.group(1).split() == command
    assert scope in routes[0]

    snapshot = tmp_path / "本轮货件 快照.xlsx"
    output = tmp_path / "workspace/.lxeagent/artifacts/replenish/result.xlsx"
    output.parent.mkdir(parents=True)
    book = Workbook()
    book.save(snapshot)
    book.save(output)
    book.close()
    output_field = "output_xlsx_path" if scope == SKILL else "report_xlsx_path"
    result = Mock()
    result.to_payload.return_value = {"success": True, output_field: str(output)}
    final_calculate = Mock(return_value=result)
    ordinary_calculate = Mock(return_value=result)
    sales = Mock(side_effect=AssertionError("calculation handoff must not start sales analysis"))
    monkeypatch.setattr(adapter, "generate_final_restock_recommendation", final_calculate)
    monkeypatch.setattr(ordinary_adapter, "calculate_store_msku_replenishment", ordinary_calculate)
    monkeypatch.setattr(sales_adapter, "analyze_store_msku_sales", sales)
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", scope)
    arguments = {"store_name": INPUTS["store_name"]}
    if scope == SKILL:
        arguments.update({key: INPUTS[key] for key in ("msku_source_xlsx_path", "shenzhen_inventory_report_xlsx_path")})
    argv = [*documented_command.group(1).split(),
            *(item for key, value in arguments.items() for item in (f"--{key.replace('_', '-')}", value)),
            documented_flag.group(1), str(snapshot)]
    assert lxeskill.main(argv) == 0
    terminal = _terminal(capsys)
    assert terminal["ok"] is True and terminal["files"] == [str(output)]
    selected, other = (final_calculate, ordinary_calculate) if scope == SKILL else (ordinary_calculate, final_calculate)
    assert selected.call_count == 1
    assert selected.call_args.kwargs["unlinked_shipments_snapshot_path"] == str(snapshot)
    other.assert_not_called()
    sales.assert_not_called()
    assert "lxeskill describe " + " ".join(command) in policy


@pytest.mark.parametrize("scope,command,wrong_flag", [
    (SKILL, COMMAND, "--unlinked-shipments-snapshot"),
    ("replenishment-calculate", ["replenish", "calculate"], "--unlinked-shipments-snapshot-path"),
])
def test_shipment_snapshot_flags_are_not_interchangeable(scope, command, wrong_flag, monkeypatch, capsys):
    final_calculate = Mock(side_effect=AssertionError("invalid flag must not execute calculation"))
    ordinary_calculate = Mock(side_effect=AssertionError("invalid flag must not execute calculation"))
    monkeypatch.setattr(adapter, "generate_final_restock_recommendation", final_calculate)
    monkeypatch.setattr(ordinary_adapter, "calculate_store_msku_replenishment", ordinary_calculate)
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", scope)
    assert lxeskill.main([*command, "--store-name", "Amazon-Test", wrong_flag, "snapshot.xlsx"]) == lxeskill.EXIT_USAGE
    terminal = _terminal(capsys)
    assert terminal["ok"] is False and terminal["files"] == []
    assert terminal["error"]["message"] == f"unknown option: {wrong_flag}"
    final_calculate.assert_not_called()
    ordinary_calculate.assert_not_called()


@pytest.mark.parametrize("name,owner,fields", [
    ("mabang_resolve_fba_store", "replenishment-store-resolve", {"store_name"}),
    ("mabang_download_store_msku_excel", "mabang-store-msku-export", {"store_name", "store_id", "id_type"}),
    ("mabang_export_store_msku_actual_inventory", "mabang-store-shenzhen-inventory-export", {"store_name"}),
    ("mabang_download_store_unlinked_shipments", "mabang-store-unlinked-shipments-export", {"store_name"}),
])
def test_store_only_acquisition_reuses_exposed_existing_commands(name, owner, fields) -> None:
    entry = load_catalog()[name]
    assert entry["visibility"] == "business"
    assert entry["exposed"] is True
    assert owner in entry["owner_skills"]
    assert fields <= set(entry["input_schema"]["properties"])


def test_shenzhen_source_is_optional_for_legacy_and_explicit_in_final_skill() -> None:
    entry = load_catalog()["mabang_export_store_msku_actual_inventory"]
    schema = entry["input_schema"]
    assert schema["required"] == ["store_name"]
    assert schema["properties"]["msku_source_xlsx_path"]["type"] == "string"
    assert schema["properties"]["msku_source_xlsx_path"]["minLength"] == 1
    skill = (Path(__file__).resolve().parents[4] / "skills" / SKILL / "SKILL.md").read_text(encoding="utf-8")
    assert '--msku-source-xlsx-path "<步骤2返回的xlsx_path>"' in skill
    assert "不回退历史文件或省略参数重试" in skill
    assert "库存导出自身会查找该店铺本地最新源表" not in skill


def test_final_restock_opts_in_to_native_msku_local_inventory() -> None:
    entry = load_catalog()["mabang_download_store_msku_excel"]
    assert entry["input_schema"]["properties"]["include_local_inventory"]["type"] == "boolean"
    skill = (Path(__file__).resolve().parents[4] / "skills" / SKILL / "SKILL.md").read_text(encoding="utf-8")
    assert 'lxeskill mabang store msku export --store-id' in skill
    assert '--include-local-inventory' in skill
    assert '步骤 2 源表中每个 MSKU 自身的“本地库存”' in skill


@pytest.mark.parametrize("scope,command", [
    ("replenishment-calculate", COMMAND),
    (SKILL, ["replenish", "calculate"]),
    ("", COMMAND),
])
def test_new_and_old_commands_respect_their_own_skill_scope(scope, command, monkeypatch, capsys) -> None:
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", scope)
    assert lxeskill.main(command) == lxeskill.EXIT_ENVIRONMENT
    terminal = _terminal(capsys)
    assert terminal["error"]["code"] == "skill_not_in_scope"
    assert terminal["files"] == []


@pytest.mark.parametrize("missing,diagnostic", [
    ("store_name", "store_name 不能为空"),
    ("msku_source_xlsx_path", "MSKU 源文件不存在"),
    ("shenzhen_inventory_report_xlsx_path", "深圳库存报告不存在"),
    ("unlinked_shipments_snapshot_path", "未关联货件快照不存在"),
])
def test_incomplete_inputs_return_actual_service_diagnostic(missing, diagnostic, tmp_path: Path, capsys) -> None:
    source = tmp_path / "existing-input.xlsx"
    source.touch()
    inputs = {key: str(source) for key in INPUTS if key != "store_name"}
    inputs["store_name"] = "Amazon-Test"
    assert lxeskill.main(_args({key: value for key, value in inputs.items() if key != missing})) == lxeskill.EXIT_BUSINESS
    terminal = _terminal(capsys)
    assert terminal["ok"] is False
    assert terminal["error"]["code"] == "business_cli_failed"
    assert diagnostic in terminal["error"]["message"]
    assert terminal["data"]["exception"] == terminal["error"]["message"]
    assert terminal["files"] == []


def test_removed_source_path_flag_is_rejected_in_the_new_skill_scope(capsys) -> None:
    assert lxeskill.main([
        *_args(INPUTS), "--shenzhen-source-msku-xlsx-path", "source.xlsx",
    ]) == lxeskill.EXIT_USAGE
    terminal = _terminal(capsys)
    assert terminal["error"]["code"] == "invalid_arguments"
    assert terminal["error"]["message"] == "unknown option: --shenzhen-source-msku-xlsx-path"
    assert terminal["files"] == []


def test_source_without_native_local_inventory_is_rejected_by_cli(tmp_path: Path, capsys) -> None:
    source = tmp_path / "202610151000-Amazon-Test_店铺MSKU数据.xlsx"
    workbook = Workbook()
    columns = [column for column in SOURCE_REQUIRED_COLUMNS if column != "本地库存"]
    workbook.active.append(columns)
    workbook.active.append([
        "MSKU-1" if column == "MSKU" else "SKU-1" if column == "本地SKU" else 0
        for column in columns
    ])
    workbook.save(source)
    workbook.close()
    annotate_source(
        source,
        SkuCatalogSnapshot("Amazon-Test", "100", "shopId", (LocalSkuDefinition("SKU-1", 1),)),
        requested_store_name="Amazon-Test",
    )
    inputs = {key: str(source) for key in INPUTS if key != "store_name"}
    assert lxeskill.main(_args({"store_name": "Amazon-Test", **inputs})) == lxeskill.EXIT_BUSINESS
    terminal = _terminal(capsys)
    assert terminal["ok"] is False
    assert terminal["error"]["code"] == "business_cli_failed"
    assert "本地库存" in terminal["error"]["message"]
    assert "缺少列" in terminal["error"]["message"]
    assert terminal["files"] == []


def test_skill_describes_a_rules_without_suppressing_unverified_advice() -> None:
    skill = (Path(__file__).resolve().parents[4] / "skills" / SKILL / "SKILL.md").read_text()
    assert "`uq122`" in skill
    assert "14 个公式字段" in skill
    assert "--require-complete-inventory" not in skill
    assert "缺失、空白、非法或非有限" in skill
    assert "合法占位行不因此阻断" in skill
    assert "require_complete_inventory" not in load_catalog()["mabang_export_store_msku_actual_inventory"]["input_schema"]["properties"]
    assert "每轮默认关闭" in skill
    assert "不自动弹旺季询问窗口" in skill
    assert "--peak-boost-enabled" in skill
    assert "未明确要求时必须省略该参数" in skill
    assert "裸开关" in skill
    assert "省略或传 false" not in skill
    assert "U–AC" in skill
    assert "13 列公式" in skill
    assert "绑定核验未通过" in skill
    assert "商品类型核验失败不等于没有库存" in skill


def test_unknown_store_returns_actual_local_source_mismatch(tmp_path: Path, capsys) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.append(["MSKU", "本地SKU"])
    workbook.active.append(["MSKU-1", "SKU-1"])
    workbook.save(source)
    workbook.close()
    annotate_source(
        source,
        SkuCatalogSnapshot("Amazon-Test", "100", "shopId", (LocalSkuDefinition("SKU-1", 1),)),
        requested_store_name="Amazon-Test",
    )
    # Store verification must fail before reading the inventory or shipment inputs.
    inputs = {key: str(source) for key in INPUTS if key != "store_name"}
    assert lxeskill.main(_args({"store_name": "Unknown-Store", **inputs})) == lxeskill.EXIT_BUSINESS
    terminal = _terminal(capsys)
    assert terminal["ok"] is False
    assert terminal["error"]["code"] == "business_cli_failed"
    assert terminal["error"]["message"] == "源数据核验店铺不一致: 请求=Unknown-Store, 快照=Amazon-Test"
    assert terminal["data"]["exception"] == terminal["error"]["message"]
    assert terminal["files"] == []


@pytest.mark.parametrize("formula_errors", [0, 1])
def test_terminal_preserves_success_or_formula_failure_and_file(
    formula_errors, tmp_path: Path, monkeypatch, capsys,
) -> None:
    output = tmp_path / "workspace" / ".lxeagent" / "artifacts" / "replenish" / "final_restock_recommendation" / "test.xlsx"
    output.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.save(output)
    workbook.close()
    output = output.resolve()
    result = FinalRestockWorkbookResult(
        store_name="Amazon-Test", source_data_time="202610131000", output_xlsx_path=str(output),
        row_count=1, formula_error_count=formula_errors,
        formula_error_examples=("AC2=#VALUE!",) if formula_errors else (),
    )
    generate = Mock(return_value=result)
    monkeypatch.setattr(adapter, "generate_final_restock_recommendation", generate)
    assert lxeskill.main(_args(INPUTS)) == (lxeskill.EXIT_BUSINESS if formula_errors else 0)
    terminal = _terminal(capsys)
    assert terminal["ok"] is (formula_errors == 0)
    assert terminal["data"] == result.to_payload()
    assert terminal["files"] == [str(output)]
    if formula_errors:
        assert terminal["error"]["message"] == result.to_payload()["exception"]
        assert "AC2=#VALUE!" in terminal["error"]["message"]
    generate.assert_called_once_with("Amazon-Test", **{key: value for key, value in INPUTS.items() if key != "store_name"}, peak_boost_enabled=False, peak_boost_days=None)


@pytest.mark.parametrize("days_flags,expected_days", [([], None), (["--peak-boost-days", "20"], 20)])
def test_explicit_current_request_boost_flags_reach_existing_generator(
    tmp_path, monkeypatch, capsys, days_flags, expected_days,
):
    from services.agent_cli.mabang import generate_final_restock_recommendation as adapter
    output = tmp_path / "workspace/.lxeagent/artifacts/replenish/final_restock_recommendation/result.xlsx"
    output.parent.mkdir(parents=True)
    book = Workbook()
    book.save(output)
    book.close()
    result = FinalRestockWorkbookResult("Amazon-Test", "202610201000", str(output), 1, 0, ())
    generate = Mock(return_value=result)
    monkeypatch.setattr(adapter, "generate_final_restock_recommendation", generate)
    assert lxeskill.main([*_args(INPUTS), "--peak-boost-enabled", *days_flags]) == 0
    assert _terminal(capsys)["ok"] is True
    assert generate.call_args.kwargs["peak_boost_enabled"] is True
    assert generate.call_args.kwargs["peak_boost_days"] == expected_days


@pytest.mark.parametrize("flags", [
    ["--peak-boost-enabled=false"],
    ["--peak-boost-enabled=true"],
    ["--peak-boost-enabled="],
    ["--peak-boost-enabled", "false"],
])
def test_boost_flag_values_fail_before_business_execution(flags, monkeypatch, capsys):
    generate = Mock(side_effect=AssertionError("invalid flag values must not generate a workbook"))
    monkeypatch.setattr(adapter, "generate_final_restock_recommendation", generate)
    assert lxeskill.main([*_args(INPUTS), *flags]) == lxeskill.EXIT_USAGE
    terminal = _terminal(capsys)
    assert terminal["ok"] is False
    assert terminal["error"]["code"] == "invalid_arguments"
    assert terminal["files"] == []
    generate.assert_not_called()


@pytest.mark.parametrize("days", ["-1", "121", "1.5", "invalid"])
def test_invalid_request_days_stop_before_reading_business_inputs(days, monkeypatch, capsys):
    from services.mabang.amazon.fba import final_restock_workbook as service
    read_input = Mock(side_effect=AssertionError("invalid days must not read input files"))
    monkeypatch.setattr(service, "_path", read_input)
    assert lxeskill.main([*_args(INPUTS), "--peak-boost-days", days]) != 0
    assert _terminal(capsys)["ok"] is False
    read_input.assert_not_called()
