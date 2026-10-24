from __future__ import annotations

import io
import json
from decimal import Decimal
from importlib.resources import files
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from lxeskill import cli as lxeskill
from lxeskill.business import load_catalog
from services.agent_cli.vietnam_replenishment import generate
from services.vietnam_replenishment import settings, workflow
from services.vietnam_replenishment.workflow import VietnamRecommendationRun
from services.vietnam_replenishment.full_workbook import UNCHANGED
from services.vietnam_replenishment.workbook import RecommendationConfig
from shared.input_assets import current_asset, previous_asset, load_input_assets
from shared.workspace import activate_project_workspace


COMMAND = ["vietnam", "replenishment", "calculate"]
INPUTS = {"sales_file": "sales.xlsx", "inventory_file": "inventory.xlsx", "products_file": "products.xlsx"}
ARGS = [item for name, value in INPUTS.items() for item in (f"--{name.replace('_', '-')}", value)]


def _args_with_template(tmp_path: Path) -> list[str]:
    template = tmp_path / "business-template.xlsx"
    if not template.exists():
        book = Workbook()
        book.active.title = "越南备货清单"
        for name in ("改库存", "改折扣", "报官方活动", "竞价报名", "雅仓库存", "雅仓动销", "数据更改", "库存商品信息"):
            book.create_sheet(name)
        book.save(template)
        book.close()
    return [*ARGS, "--template-file", str(template)]


def _valid_template(path: Path, marker: str) -> Path:
    resource = files("services.vietnam_replenishment").joinpath("resources/skeleton.xlsx")
    with resource.open("rb") as stream:
        book = load_workbook(stream)
    try:
        book["越南备货清单"]["B2"] = 2
        book["越南备货清单"]["E2"] = "VN-A"
        book["越南备货清单"]["G2"] = 10
        for index, name in enumerate(UNCHANGED, 1):
            book.create_sheet(name, index)["A1"] = marker
        book.save(path)
    finally:
        book.close()
    return path


def _record(capsys) -> dict:
    lines = [line for line in capsys.readouterr().out.splitlines() if line.strip()]
    assert len(lines) == 1
    return json.loads(lines[0])


def _no_workflow(monkeypatch: pytest.MonkeyPatch) -> None:
    def unexpected(**_arguments) -> None:
        raise AssertionError("workflow must not run")

    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", unexpected)


def test_catalog_requires_three_reports_and_allows_one_run_template(monkeypatch, capsys) -> None:
    entry = load_catalog()["vietnam_replenishment_generate"]
    assert entry["module"] == "services.agent_cli.vietnam_replenishment.generate"
    assert entry["command_path"] == COMMAND
    assert entry["session_mode"] == "none"
    assert entry["owner_skills"] == ["vietnam-replenishment"]
    assert entry["input_schema"] == {
        "type": "object", "properties": {
            **{name: {"type": "string", "minLength": 1} for name in INPUTS},
            "template_file": {"type": "string", "minLength": 1, "x-lxe-file-input": {
                "accepted_extensions": [".xlsx"],
                "instruction": "请上传完整的东南亚越南备货九表 XLSX 模板，或使用已保存的当前模板。",
            }},
        },
        "required": list(INPUTS), "additionalProperties": False,
    }
    assert entry["artifact_paths"] == [{"field": "output_xlsx", "role": "deliverable"}]
    assert entry.get("deliver_artifacts_on_failure") is not True
    assert load_input_assets()["vietnam_full_workbook"].management == "command"
    monkeypatch.delenv("LXESKILL_SKILL_SCOPE", raising=False)
    assert lxeskill.main([*COMMAND, "--help"]) == 0
    public = _record(capsys)["data"]
    assert public["input_schema"] == entry["input_schema"]
    assert public["input_modes"] == ["flags"]
    assert public["owner_skills"] == ["vietnam-replenishment"]
    assert public["attribution_skill"] == "vietnam-replenishment"


@pytest.mark.parametrize("arguments", [
    {}, {"map_path": "/tmp/other.xlsx"}, {**INPUTS, "sku_map_file": "old.xlsx"}, {**INPUTS, "cost_rate": 0.7},
    *[{key: value for key, value in INPUTS.items() if key != missing} for missing in INPUTS],
    *[{**INPUTS, key: value} for key in (*INPUTS, "template_file") for value in ("", " ", 2, None)],
])
def test_adapter_rejects_missing_invalid_and_unknown_arguments(arguments: dict, monkeypatch: pytest.MonkeyPatch) -> None:
    _no_workflow(monkeypatch)
    result = generate.run(arguments)
    assert result["success"] is False
    assert result["error"]["code"] == "invalid_arguments"
    assert "output_xlsx" not in result


@pytest.mark.parametrize("mode", ["stdin", "file"])
def test_json_input_rejects_unknown_fields(
    mode: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys
) -> None:
    _no_workflow(monkeypatch)
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    input_json = json.dumps({**INPUTS, 'map_path': '/tmp/other.xlsx'})
    if mode == "stdin":
        monkeypatch.setattr("sys.stdin", io.StringIO(input_json))
        args = [*COMMAND, "--stdin-json"]
    else:
        source = tmp_path / "input.json"
        source.write_text(input_json, encoding="utf-8")
        args = [*COMMAND, "--input-json", str(source)]

    assert lxeskill.main(args) == lxeskill.EXIT_USAGE
    record = _record(capsys)
    assert record["ok"] is False
    assert record["error"]["code"] == "invalid_arguments"
    assert record["files"] == []


def test_success_delivers_only_final_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys
) -> None:
    workspace = tmp_path / "workspace"
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(workspace))
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    output = workspace / ".lxeagent" / "artifacts" / "vietnam" / "recommendations" / "run" / "越南备货清单.xlsx"
    output.parent.mkdir(parents=True)
    output.write_bytes(b"synthetic workbook")
    output = output.resolve()  # Windows may restore the directory's on-disk casing.
    monkeypatch.setattr(
        generate,
        "generate_current_vietnam_recommendation",
        lambda **_arguments: VietnamRecommendationRun(
            output_xlsx=output, sku_count=2,
            config=RecommendationConfig(
                day_adjustment_30d=Decimal("0.7"), day_adjustment_15d=Decimal("0.6"),
                day_adjustment_7d=Decimal("0.1"), exchange_rate=Decimal("4000"),
            ),
            config_source="parameters.json", template_source="business-template.xlsx",
            source_files=({"report": "inventory-sales", "path": "sales.xlsx", "sha256": "0" * 64},),
            validation={"status": "passed", "sku_count": 2},
            template_asset_source={"slot": "vietnam_full_workbook", "from": "upload", "file_name": "business-template.xlsx"},
        ),
    )

    assert lxeskill.main([*COMMAND, *_args_with_template(tmp_path)]) == 0
    record = _record(capsys)
    assert record["ok"] is True
    assert record["files"] == [str(output)]
    template_source = record["data"].pop("asset_sources")["template_file"]
    assert template_source["slot"] == "vietnam_full_workbook"
    assert template_source["from"] == "upload"
    assert template_source["file_name"] == "business-template.xlsx"
    assert record["data"] == {
        "success": True,
        "status": "completed",
        "warehouse": "VN8806",
        "sku_count": 2,
        "output_xlsx": str(output),
        "config": {
            "day_adjustment_30d": "0.7", "day_adjustment_15d": "0.6",
            "day_adjustment_7d": "0.1", "exchange_rate": "4000",
            "sales_weight_7d": "0.6", "sales_weight_15d": "0.3", "sales_weight_30d": "0.1",
        },
        "config_source": "parameters.json",
        "template_source": "business-template.xlsx",
        "source_files": [{"report": "inventory-sales", "path": "sales.xlsx", "sha256": "0" * 64}],
        "validation": {"status": "passed", "sku_count": 2},
    }
    saved = current_asset("vietnam_full_workbook")
    assert saved is None
    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", workflow.generate_current_vietnam_recommendation)
    assert lxeskill.main([*COMMAND, *ARGS]) == lxeskill.EXIT_USAGE
    missing = _record(capsys)
    assert missing["error"]["code"] == "input_required"
    assert missing["recovery"]["field"] == "template_file"


def test_chat_template_is_temporary_across_runs_and_sessions(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys,
) -> None:
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path / "workspace"))
    activate_project_workspace()
    template_a = _valid_template(tmp_path / "current-a.xlsx", "A")
    template_b = _valid_template(tmp_path / "one-run-b.xlsx", "B")
    settings.upload_template(template_a)
    saved_a = current_asset("vietnam_full_workbook")
    assert saved_a is not None
    saved_mtime = saved_a.path.stat().st_mtime_ns
    assert previous_asset("vietnam_full_workbook") is None
    output = tmp_path / "workspace" / ".lxeagent" / "artifacts" / "vietnam" / "result.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(b"synthetic workbook")
    selected: list[str | None] = []

    def succeed(**arguments):
        selected.append(arguments.get("template"))
        version = current_asset("vietnam_full_workbook")
        explicit = arguments.get("template")
        provenance = {"slot": "vietnam_full_workbook", "from": "upload" if explicit else "stored",
                      "file_name": Path(explicit).name if explicit else version.file_name}
        if not explicit:
            provenance["updated_at"] = version.updated_at
        return VietnamRecommendationRun(
            output_xlsx=output, sku_count=1, config=RecommendationConfig(),
            config_source="parameters.json",
            template_source=arguments.get("template") or str(current_asset("vietnam_full_workbook").path),
            source_files=(), validation={"status": "passed", "sku_count": 1},
            template_asset_source=provenance,
        )

    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", succeed)
    monkeypatch.setenv("LXE_AGENT_SESSION_ID", "chat-one")
    assert lxeskill.main([*COMMAND, *ARGS, "--template-file", str(template_b)]) == 0
    _record(capsys)
    assert selected[-1] == str(template_b)
    assert current_asset("vietnam_full_workbook") == saved_a
    assert saved_a.path.read_bytes() == template_a.read_bytes()
    assert saved_a.path.stat().st_mtime_ns == saved_mtime
    assert previous_asset("vietnam_full_workbook") is None

    monkeypatch.setenv("LXE_AGENT_SESSION_ID", "chat-two")
    assert lxeskill.main([*COMMAND, *ARGS]) == 0
    _record(capsys)
    assert selected[-1] is None  # Default selection belongs to the protected workflow snapshot.
    assert current_asset("vietnam_full_workbook").path.read_bytes() == template_a.read_bytes()
    assert previous_asset("vietnam_full_workbook") is None

    def fail(**_arguments):
        raise RuntimeError("synthetic calculation failure")

    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", fail)
    assert lxeskill.main([*COMMAND, *ARGS, "--template-file", str(template_b)]) == lxeskill.EXIT_BUSINESS
    _record(capsys)
    assert current_asset("vietnam_full_workbook").path.read_bytes() == template_a.read_bytes()
    assert previous_asset("vietnam_full_workbook") is None

    settings.upload_template(template_b)
    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", succeed)
    assert lxeskill.main([*COMMAND, *ARGS]) == 0
    _record(capsys)
    assert selected[-1] is None
    assert current_asset("vietnam_full_workbook").path.read_bytes() == template_b.read_bytes()


def test_business_failure_preserves_real_redacted_diagnostic_and_delivers_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys
) -> None:
    class ObservedFailure(RuntimeError):
        code = "template_missing"

    def fail(**_arguments) -> None:
        raise ObservedFailure("请先上传越南九表模板: password=topsecret")

    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("LXE_YACANG_PASSWORD", "topsecret")
    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", fail)

    assert lxeskill.main([*COMMAND, *_args_with_template(tmp_path)]) == lxeskill.EXIT_BUSINESS
    record = _record(capsys)
    assert record["ok"] is False
    assert record["files"] == []
    assert record["data"]["error"]["code"] == "template_missing"
    assert "ObservedFailure" in record["error"]["message"]
    assert "请先上传" in record["error"]["message"]
    assert "topsecret" not in record["error"]["message"]
    assert "output_xlsx" not in record["data"]


def test_missing_generated_file_fails_without_deliverable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys
) -> None:
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    output = tmp_path / "state" / "artifacts" / "vietnam" / "recommendations" / "missing.xlsx"
    monkeypatch.setattr(
        generate,
        "generate_current_vietnam_recommendation",
        lambda **_arguments: SimpleNamespace(output_xlsx=output, sku_count=2),
    )

    assert lxeskill.main([*COMMAND, *_args_with_template(tmp_path)]) == lxeskill.EXIT_BUSINESS
    record = _record(capsys)
    assert record["ok"] is False
    assert record["files"] == []
    assert "FileNotFoundError" in record["error"]["message"]


def test_reports_and_template_reach_workflow(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    calls = []

    def stop_after_arguments(**arguments):
        calls.append(arguments)
        raise RuntimeError("stopped after argument forwarding")

    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", stop_after_arguments)
    assert lxeskill.main([*COMMAND, *_args_with_template(tmp_path)]) == lxeskill.EXIT_BUSINESS
    assert calls == [{**{name.removesuffix("_file"): value for name, value in INPUTS.items()},
                      "template": str(tmp_path / "business-template.xlsx")}]
    assert _record(capsys)["files"] == []


def test_old_no_argument_invocation_reports_missing_paths(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    _no_workflow(monkeypatch)
    template = _args_with_template(tmp_path)[-1]
    assert lxeskill.main([*COMMAND, "--template-file", template]) == lxeskill.EXIT_BUSINESS
    result = _record(capsys)
    assert result["files"] == []
    for name in INPUTS:
        assert name in result["error"]["message"]


def test_first_calculation_requires_template_upload(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_: pytest.fail("Reports must not load without a template"))
    assert lxeskill.main([*COMMAND, *ARGS]) == lxeskill.EXIT_USAGE
    result = _record(capsys)
    assert result["error"]["code"] == "input_required"
    assert result["recovery"]["field"] == "template_file"
    assert result["files"] == []


@pytest.mark.parametrize("upload_window", ["after_cli_check", "during_reports"])
def test_current_template_snapshot_survives_real_concurrent_upload(
    tmp_path, monkeypatch, capsys, upload_window,
):
    import os
    import subprocess
    import sys
    from services.vietnam_replenishment import workflow

    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path / "workspace"))
    activate_project_workspace()
    a = _valid_template(tmp_path / "A.xlsx", "A")
    b = _valid_template(tmp_path / "B.xlsx", "B")
    book = load_workbook(b)
    book["越南备货清单"]["G2"] = 77
    book.save(b)
    book.close()
    settings.upload_template(a)
    old = current_asset("vietnam_full_workbook")
    monkeypatch.setattr(lxeskill, "current_asset", lambda _: pytest.fail("Shared CLI resolved a business-owned template"))

    def upload_from_another_process():
        child = subprocess.run(
            [sys.executable, "-I", "-B", "-m", "lxeskill", "vietnam", "replenishment", "settings",
             "--action", "upload", "--path", str(b)],
            env=os.environ.copy(), capture_output=True, text=True, timeout=30,
        )
        assert child.returncode == 0, (child.stdout, child.stderr)

    if upload_window == "after_cli_check":
        require = lxeskill._require_uploaded_file_inputs
        def check_then_replace(*args, **kwargs):
            require(*args, **kwargs)
            upload_from_another_process()
        monkeypatch.setattr(lxeskill, "_require_uploaded_file_inputs", check_then_replace)

    def load_reports(**_paths):
        if upload_window == "during_reports":
            # A separate uploader must be able to finish while this run is reading
            # reports, proving the settings lock is no longer held here.
            upload_from_another_process()
        return SimpleNamespace(skus=("VN-A",), artifacts={}, file_hashes={}, in_transit_mismatch=())

    expected = b if upload_window == "after_cli_check" else a
    seen_snapshots = []
    def generate_from_snapshot(output, *, template, parameters, **_kwargs):
        seen_snapshots.append(template)
        assert template.read_bytes() == expected.read_bytes()
        assert parameters["VN-A"].cost == (77 if expected == b else 10)
        assert current_asset("vietnam_full_workbook").file_name == "B.xlsx"
        assert not old.path.exists()
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(b"synthetic result")
        return output

    monkeypatch.setattr(workflow, "load_vietnam_files", load_reports)
    monkeypatch.setattr(workflow, "generate_vietnam_workbook", generate_from_snapshot)
    assert lxeskill.main([*COMMAND, *ARGS]) == 0
    record = _record(capsys)
    assert record["ok"] is True
    assert Path(record["data"]["template_source"]).name == expected.name
    assert record["data"]["asset_sources"]["template_file"]["file_name"] == expected.name
    assert record["data"]["asset_sources"]["template_file"]["from"] == "stored"
    assert len(seen_snapshots) == 1 and not seen_snapshots[0].exists()
    assert current_asset("vietnam_full_workbook").path.read_bytes() == b.read_bytes()
    assert previous_asset("vietnam_full_workbook").path.read_bytes() == a.read_bytes()


def test_shared_cli_has_no_vietnam_business_dependency():
    import ast

    tree = ast.parse(Path(lxeskill.__file__).read_text(encoding="utf-8"))
    assert not [node.module for node in ast.walk(tree) if isinstance(node, ast.ImportFrom)
                and (node.module or "").startswith("services.vietnam_replenishment")]
    assert not [node.value for node in ast.walk(tree) if isinstance(node, ast.Constant)
                and node.value == "vietnam_full_workbook"]
