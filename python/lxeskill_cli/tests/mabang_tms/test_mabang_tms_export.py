from __future__ import annotations

import io
import json
import re
import ssl
from pathlib import Path
from unittest.mock import MagicMock, Mock

import pytest
import requests
from openpyxl import Workbook, load_workbook

from services.mabang_tms import workflow
from services.mabang_tms.client import FIELDS, TmsClient, TmsError, diagnostic
from services.mabang_tms.workbooks import coverage, read_workbook, write_merged

HEADERS = ['库存SKU', '30天销量', '7天销量', '可用库存', '15天销量', '仓库',
           '库存量', '在途量', '成本价', '采购价', '入库时间', '平台额外销量']
WAREHOUSE_CASES = [('PH1', 12107), ('VN1', 12831)]


def workbook(rows, headers=HEADERS):
    wb = Workbook()
    wb.active.append(headers)
    for row in rows:
        wb.active.append(row)
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def record(i, **kw):
    return {'id': i, 'stockSku': f'SKU-{i}', 'warehouseName': '测试仓', **kw}


def values(r):
    return [r['stockSku'], '97', '21', 0, '71', r['warehouseName'], -1, 0, 0, 0, '2026-08-25 11:35:29', '123']


class FakeClient:
    def __init__(self, pages):
        self.pages = pages
        self.login_count = 0
        self.calls = []
        self.contents = []
        self.warehouse_requests = []
        self.token = 'test-token'

    def login(self):
        self.login_count += 1

    def check_budget(self):
        pass

    def list_page(self, page, *, warehouse_id=None):
        self.warehouse_requests.append(('list', warehouse_id))
        self.calls.append(('list', page))
        self.current = self.pages[page - 1]
        return self.current

    def export(self, ids, *, warehouse_id=None):
        self.warehouse_requests.append(('export', warehouse_id))
        self.calls.append(('export', ids))
        return 'https://tms-cos.mabangerp.com/test?signature=secret'

    def download(self, url):
        self.calls.append(('download', None))
        content = workbook([values(r) for r in self.current['datas']])
        self.contents.append(content)
        return content

    def safe(self, value):
        return diagnostic(value, ('test-token', 'password'))

    def close(self):
        self.token = ''


def execute(monkeypatch, tmp_path, pages, client=None, arguments=None):
    client = client or FakeClient(pages)
    monkeypatch.setenv('LXE_MABANG_TMS_ACCOUNT', 'account')
    monkeypatch.setenv('LXE_MABANG_TMS_PASSWORD', 'password')
    monkeypatch.delenv('LXESKILL_SKILL_SCOPE', raising=False)
    monkeypatch.setattr(workflow, 'TmsClient', lambda *args: client)
    monkeypatch.setattr(workflow, 'dataset_dir', lambda *args: tmp_path.joinpath(*args[1:]))
    return workflow.run(arguments or {}), client


def page(records, total, number=1):
    return {'code': '200', 'datas': records, 'totalNum': total, 'page': number}


def test_short_pages_continue_until_total_and_original_bytes_preserved(monkeypatch, tmp_path):
    result, client = execute(monkeypatch, tmp_path, [page([record(1)], 2), page([record(2)], 2, 2)])
    assert result['success'] and result['read_unique_ids'] == result['reported_total'] == 2
    assert result['completed_batches'] == 2 and client.login_count == 1
    assert [c[0] for c in client.calls] == ['list', 'export', 'download'] * 2
    assert client.token == ''
    for batch, content in zip(result['batches'], client.contents):
        from pathlib import Path
        assert Path(batch['raw_path']).read_bytes() == content
    assert len(result['artifacts']) == 1
    wb = load_workbook(result['artifacts'][0]['path'])
    assert wb.active.max_row == 3
    assert wb.active['B2'].value == '97' and wb.active['G2'].value == -1
    wb.close()


def test_8386_records_nine_batches(monkeypatch, tmp_path):
    pages = [page([record(j) for j in range(i, min(i + 1000, 8386))], 8386, i // 1000 + 1)
             for i in range(0, 8386, 1000)]
    result, client = execute(monkeypatch, tmp_path, pages)
    assert result['success'] and result['row_count'] == 8386
    assert result['completed_batches'] == 9 and client.login_count == 1


@pytest.mark.parametrize('warehouse,warehouse_id,total,effective_page_size', [
    ('PH1', 12107, 4325, 1000),
    ('VN1', 12831, 4208, 1000),
    ('all', None, 8533, 1000),
    ('all', None, 8533, 50),
])
def test_real_client_request_budget_at_known_warehouse_scale(monkeypatch, tmp_path, warehouse, warehouse_id, total, effective_page_size):
    # Synthetic responses at the known business scale. Exercise real client
    # payloads; all's effective production page size is not captured evidence.
    # A simulated server cap must yield partial failure, never false completion.
    records = [record(i, stockSku=f'SKU-{i // 3}',
                      warehouseId=warehouse_id or (12107 if i < 4325 else 12831),
                      warehouseName=warehouse if warehouse_id else ('PH1' if i < 4325 else 'VN1'))
               for i in range(total)]
    current = []
    session = Mock()

    def post(url, *, json, **kwargs):
        nonlocal current
        if url.endswith('/login'):
            return response(body={'code': '200', 'data': {'apiToken': 'offline-token'}})
        assert json['pageSize'] == 1000
        assert json.get('warehouseId') == warehouse_id
        assert ('warehouseId' in json) == (warehouse_id is not None)
        if url.endswith('/findMyStockwarehouseList'):
            start = (json['page'] - 1) * effective_page_size
            current = records[start:start + effective_page_size]
            return response(body={'code': '200', 'data': None, 'datas': current, 'totalNum': total,
                                  'pop': {'pageSize': effective_page_size, 'totalPage': 1, 'totalNum': total}})
        assert url.endswith('/exportStockwarehouse')
        assert json['idsList'] == [r['id'] for r in current]
        return response(body={'code': '200', 'pop': 'https://tms-cos.mabangerp.com/offline'})

    session.post.side_effect = post
    client = TmsClient('account', 'password', session=session, sleeper=lambda _: None)
    client.download = Mock(side_effect=lambda _: workbook([values(r) for r in current]))
    result, _ = execute(monkeypatch, tmp_path, [], client, arguments={'warehouse': warehouse})
    expected_pages = min(100, (total + effective_page_size - 1) // effective_page_size)
    expected_records = min(total, expected_pages * effective_page_size)
    listings = [call.kwargs['json'] for call in session.post.call_args_list
                if call.args[0].endswith('/findMyStockwarehouseList')]
    exports = [call.kwargs['json'] for call in session.post.call_args_list
               if call.args[0].endswith('/exportStockwarehouse')]
    assert [p['page'] for p in listings] == list(range(1, expected_pages + 1))
    assert [i for p in exports for i in p['idsList']] == list(range(expected_records))
    assert result['read_unique_ids'] == result['exported_unique_ids'] == result['row_count'] == expected_records
    assert result['reported_total'] == total and result['remaining_records'] == total - expected_records
    assert result['completed_batches'] == client.download.call_count == len(exports) == expected_pages
    if expected_records == total:
        assert result['success'] and result['status'] == 'completed'
    else:
        assert not result['success'] and result['status'] == 'partial_success'
        assert result['error']['code'] == 'page_limit' and result['failed_page'] == 100
    headers, rows, _ = read_workbook(Path(result['artifacts'][0]['path']).read_bytes())
    assert headers == tuple(HEADERS)
    assert [list(row) for row in rows] == [values(r) for r in records[:expected_records]]
    assert len(rows) > len({row[0] for row in rows})  # Independent same-SKU rows survive.
    session.close.assert_called_once()


@pytest.mark.parametrize('second,code', [
    (page([], 2, 2), 'early_empty_page'),
    (page([record(2)], 3, 2), 'total_changed'),
    (page([record(2)], 1, 2), 'total_changed'),
    (page([record(1)], 2, 2), 'duplicate_id'),
    (page([record(2), record('2')], 2, 2), 'duplicate_id'),
    (page([record(2)], 2, 1), 'page_mismatch'),
    (page([record(2), record(3)], 2, 2), 'total_exceeded'),
])
@pytest.mark.parametrize('warehouse,warehouse_id', WAREHOUSE_CASES + [('all', None)])
def test_partial_not_success_on_incomplete_pagination(monkeypatch, tmp_path, second, code, warehouse, warehouse_id):
    first = {**page([record(1, warehouseId=warehouse_id, warehouseName=warehouse)], 2),
             'pop': {'totalPage': None}}
    second = {**second, 'datas': [{**r, 'warehouseId': warehouse_id, 'warehouseName': warehouse}
                                for r in second['datas']], 'pop': {'totalPage': 'unused'}}
    result, client = execute(monkeypatch, tmp_path, [first, second], arguments={'warehouse': warehouse})
    assert not result['success'] and result['error']['code'] == code
    assert result['status'] == 'partial_success' and len(result['artifacts']) == 1
    assert result['exported_unique_ids'] == 1 and result['remaining_records'] == 1
    assert len([c for c in client.calls if c[0] == 'export']) == 1


def test_zero_no_fake_attachment(monkeypatch, tmp_path):
    result, _ = execute(monkeypatch, tmp_path, [page([], 0)])
    assert result['success'] and result['artifacts'] == [] and result['row_count'] == 0


def test_sample_shape_fifty_rows_fortyone_skus_preserved(monkeypatch, tmp_path):
    records = [record(i, stockSku=f'SAMPLE-{i if i < 41 else i-41}') for i in range(50)]
    result, _ = execute(monkeypatch, tmp_path, [page(records, 50)])
    assert result['success'] and result['row_count'] == 50 and result['distinct_skus'] == 41
    assert result['headers'] == HEADERS


def test_eleven_column_sample_and_unknown_extra_column():
    for headers in (HEADERS[:11], HEADERS):
        content = workbook([values(record(1))[:len(headers)]], headers)
        actual, rows, _ = read_workbook(content)
        assert actual == tuple(headers) and rows[0][1] == '97'


@pytest.mark.parametrize('records,rows', [
    ([record(1)], [values(record(2))]),
    ([record(1)], [values(record(1, warehouseName='另一个仓'))]),
    ([{'id': 1}], [values(record(1))]),
])
def test_coverage_fails_closed(records, rows):
    with pytest.raises(TmsError):
        coverage(records, HEADERS, rows)


def test_one_id_may_expand_to_multiple_rows():
    coverage([record(1)], HEADERS, [values(record(1)), values(record(1))])


def test_save_failure_reports_both_original_and_publication_error(monkeypatch, tmp_path):
    def fail(*args):
        raise OSError('disk full')
    monkeypatch.setattr(workflow, 'write_merged', fail)
    result, _ = execute(monkeypatch, tmp_path, [page([record(1)], 2), page([], 2, 2)])
    assert not result['success'] and result['artifacts'] == []
    assert result['error']['code'] == 'early_empty_page'
    assert 'disk full' in result['error']['publication_error']


def test_atomic_publication_and_literal_strings(tmp_path):
    row = values(record(1, stockSku='=not-a-formula'))
    write_merged(tmp_path / 'out.xlsx', HEADERS, [row])
    wb = load_workbook(tmp_path / 'out.xlsx')
    assert wb.active['A2'].data_type == 's'
    wb.close()


def response(status=200, body=None, text='response'):
    r = Mock(status_code=status, text=text)
    r.json.return_value = body if body is not None else {'code': '200'}
    return r


@pytest.mark.parametrize('status', [401, 403, 429, 302])
def test_rejections_never_retry(status):
    session = Mock()
    session.post.return_value = response(status, text='actual platform error')
    client = TmsClient('account', 'password', session=session, sleeper=lambda _: None)
    with pytest.raises(TmsError, match='actual platform error'):
        client.list_page(1)
    assert session.post.call_count == 1


def test_only_list_retries_and_keeps_fixed_fields():
    session = Mock()
    session.post.side_effect = [response(503), response(body={'code': '200', 'datas': [], 'totalNum': 0})]
    client = TmsClient('account', 'password', session=session, sleeper=lambda _: None)
    client.list_page(1)
    assert session.post.call_count == 2
    session.post.reset_mock()
    session.post.side_effect = requests.Timeout('actual timeout')
    with pytest.raises(TmsError, match='actual timeout'):
        client.export([1])
    assert session.post.call_count == 1
    assert session.post.call_args.kwargs['json']['stockValueArr'] == FIELDS
    assert len(FIELDS) == 12 and FIELDS[-1] == 'sale5'


@pytest.mark.parametrize('body', [{'code': '200'}, {'code': '200', 'data': {'apiToken': ''}}, {'code': '401', 'msg': 'denied'}])
def test_bad_login_never_retries(body):
    session = Mock()
    session.post.return_value = response(body=body)
    with pytest.raises(TmsError):
        TmsClient('account', 'password', session=session).login()
    assert session.post.call_count == 1


def test_non_json_and_redaction_before_truncation():
    session = Mock()
    session.post.return_value = response(text='pwd=password https://host/x?sig=private ' + 'x'*5000)
    session.post.return_value.json.side_effect = ValueError('not json')
    with pytest.raises(TmsError) as caught:
        TmsClient('account', 'password', session=session).login()
    message = str(caught.value)
    assert 'password' not in message and 'sig=private' not in message and '[truncated]' in message


def test_account_lock_and_credential_absence(monkeypatch, tmp_path):
    monkeypatch.delenv('LXE_MABANG_TMS_ACCOUNT', raising=False)
    monkeypatch.delenv('LXE_MABANG_TMS_PASSWORD', raising=False)
    assert workflow.run({})['error']['code'] == 'credentials_missing'
    assert workflow.run({'extra': True})['error']['code'] == 'invalid_arguments'


def test_unsafe_download_url_never_requested():
    client = TmsClient('a', 'p')
    with pytest.raises(TmsError):
        client.download('https://untrusted.test/file')


def test_corrupt_file_and_missing_headers():
    with pytest.raises(TmsError):
        read_workbook(b'not excel')
    with pytest.raises(TmsError):
        read_workbook(workbook([['value']], ['wrong']))


@pytest.mark.parametrize('stage', ['login', 'export', 'download'])
def test_task_stops_after_uncertain_failure(monkeypatch, tmp_path, stage):
    client = FakeClient([page([record(1)], 2), page([record(2)], 2, 2)])
    def fail(*args, **kwargs):
        raise TmsError('uncertain', 'actual timeout password test-token')
    setattr(client, stage, fail)
    result, _ = execute(monkeypatch, tmp_path, [], client)
    assert not result['success'] and not result['artifacts']
    assert 'password' not in str(result) and 'test-token' not in str(result)
    assert len([c for c in client.calls if c[0] == 'list']) <= 1


def test_header_change_preserves_only_validated_batch(monkeypatch, tmp_path):
    client = FakeClient([page([record(1)], 2), page([record(2)], 2, 2)])
    original_download = client.download
    def download(url):
        if client.current['page'] == 2:
            return workbook([values(record(2))[:11]], HEADERS[:11])
        return original_download(url)
    client.download = download
    result, _ = execute(monkeypatch, tmp_path, [], client)
    assert result['status'] == 'partial_success' and result['error']['code'] == 'header_mismatch'
    assert result['completed_batches'] == 1


def test_account_directories_isolated_and_busy_lock_prevents_login(monkeypatch, tmp_path):
    import hashlib
    from shared.process_lock import interprocess_lock
    account_root = tmp_path / hashlib.sha256(b'account').hexdigest()
    account_root.mkdir()
    with interprocess_lock(account_root / '.export.lock', timeout_seconds=0):
        result, client = execute(monkeypatch, tmp_path, [page([], 0)])
    assert not result['success'] and client.login_count == 0
    result, _ = execute(monkeypatch, tmp_path, [page([record(1)], 1)])
    assert Path(result['artifacts'][0]['path']).is_relative_to(account_root)


@pytest.mark.parametrize('failure', [
    requests.Timeout('download timeout'),
    requests.ConnectionError('download connection reset'),
    requests.exceptions.SSLError(ssl.SSLEOFError(8, 'UNEXPECTED_EOF_WHILE_READING')),
    requests.exceptions.SSLError(ssl.SSLCertVerificationError(1, 'certificate verify failed')),
    requests.exceptions.ProxyError('download proxy failed'),
])
def test_download_no_retry_and_no_auth_forwarding(monkeypatch, failure):
    get = Mock(side_effect=failure)
    monkeypatch.setattr(requests, 'get', get)
    session = Mock()
    client = TmsClient('account', 'password', session=session, sleeper=lambda _: None)
    client.token = 'private-token'
    with pytest.raises(TmsError) as caught:
        client.download('https://tms-cos.mabangerp.com/file?sig=private')
    assert caught.value.code == 'download_network_error'
    assert caught.value.__cause__ is failure
    assert type(failure).__name__ in str(caught.value)
    assert get.call_count == 1
    assert get.call_args.args == ('https://tms-cos.mabangerp.com/file?sig=private',)
    assert get.call_args.kwargs == {'timeout': (5, 30), 'stream': True, 'allow_redirects': False}
    session.post.assert_not_called()


@pytest.mark.parametrize('failure', [
    requests.exceptions.SSLError(ssl.SSLEOFError(8, 'UNEXPECTED_EOF_WHILE_READING')),
    requests.exceptions.ChunkedEncodingError('incomplete response body'),
])
def test_download_stream_interruption_discards_partial_bytes_and_closes_response(monkeypatch, failure):
    streamed = MagicMock(status_code=200)
    streamed.__enter__.return_value = streamed
    def chunks(_size):
        yield b'only the first part of a workbook'
        raise failure
    streamed.iter_content.side_effect = chunks
    get = Mock(return_value=streamed)
    monkeypatch.setattr(requests, 'get', get)
    client = TmsClient('account', 'password', sleeper=lambda _: None)
    with pytest.raises(TmsError) as caught:
        client.download('https://tms-cos.mabangerp.com/file?signature=synthetic')
    assert caught.value.code == 'download_network_error' and caught.value.__cause__ is failure
    assert get.call_count == 1
    streamed.__exit__.assert_called_once()
    client.close()


def test_normal_download_uses_separate_session_and_complete_stream(monkeypatch):
    session = requests.Session()
    session.headers['token'] = 'private-token'
    session.cookies.set('platform_session', 'private-cookie')
    client = TmsClient('account', 'password', session=session, sleeper=lambda _: None)
    sent = []
    content = workbook([values(record(1))])
    def send(download_session, request, **kwargs):
        sent.append((download_session, request, kwargs))
        result = requests.Response()
        result.status_code = 200
        result.raw = io.BytesIO(content)
        return result
    monkeypatch.setattr(requests.Session, 'send', send)
    monkeypatch.setattr(requests.sessions, 'get_netrc_auth', lambda *_: None)
    url = 'https://tms-cos.mabangerp.com/file?signature=synthetic&expires=123'
    assert client.download(url) == content
    assert len(sent) == 1
    download_session, request, kwargs = sent[0]
    assert download_session is not session
    assert request.method == 'GET' and request.url == url
    assert not {'token', 'Cookie', 'Authorization'}.intersection(request.headers)
    assert kwargs['timeout'] == (5, 30) and kwargs['stream'] and not kwargs['allow_redirects']
    assert kwargs['verify'] is True
    client.close()


@pytest.mark.parametrize('status', [401, 403, 429, 302, 503])
def test_download_http_rejection_never_retries(monkeypatch, status):
    streamed = MagicMock(status_code=status, text='actual download rejection')
    streamed.__enter__.return_value = streamed
    get = Mock(return_value=streamed)
    monkeypatch.setattr(requests, 'get', get)
    client = TmsClient('account', 'password', sleeper=lambda _: None)
    with pytest.raises(TmsError) as caught:
        client.download('https://tms-cos.mabangerp.com/file?signature=synthetic')
    assert caught.value.code == f'download_http_{status}'
    assert 'actual download rejection' in str(caught.value)
    assert get.call_count == 1
    streamed.iter_content.assert_not_called()
    streamed.__exit__.assert_called_once()
    client.close()


@pytest.mark.parametrize('outcome, status, stage, exception_type, elapsed_ms', [
    ('completed', 200, 'completed', None, 750),
    ('tls', None, 'unknown', 'SSLError', 250),
    ('stream', 200, 'response_body', 'ChunkedEncodingError', 750),
    ('http', 429, 'response_headers', 'TmsError', 250),
])
def test_download_logs_safe_timing_status_and_actual_error(monkeypatch, caplog, outcome, status, stage, exception_type, elapsed_ms):
    caplog.set_level('INFO', logger='services.mabang_tms.client')
    now = [0.0]
    def sleep(seconds):
        now[0] += seconds
    client = TmsClient('private-account', 'private-password', clock=lambda: now[0], sleeper=sleep)
    client.token = 'private-token'
    client.last_request = 0
    relative = '/exportstock/synthetic.xls?sig=private-signature&expires=private-expiry'
    url = 'https://tms-cos.mabangerp.com' + relative
    reason = f'with url: {relative} (UNEXPECTED_EOF_WHILE_READING); {url}; GET {relative}; private-account private-password apiToken=private-token'
    streamed = MagicMock(status_code=429 if outcome == 'http' else 200, text=reason)
    streamed.__enter__.return_value = streamed
    def chunks(_size):
        now[0] += 0.5
        yield b'first'
        if outcome == 'stream':
            raise requests.exceptions.ChunkedEncodingError(reason)
        yield b'second'
    streamed.iter_content.side_effect = chunks
    def get(*args, **kwargs):
        now[0] += 0.25
        if outcome == 'tls':
            raise requests.exceptions.SSLError(reason)
        return streamed
    get_mock = Mock(side_effect=get)
    monkeypatch.setattr(requests, 'get', get_mock)
    if outcome == 'completed':
        assert client.download(url) == b'firstsecond'
    else:
        with pytest.raises(TmsError) as caught:
            client.download(url)
        assert caught.value.code == ('download_http_429' if outcome == 'http' else 'download_network_error')
    records = [r for r in caplog.records if r.name == 'services.mabang_tms.client']
    assert len(records) == 1 and records[0].exc_info is None
    text = records[0].getMessage()
    metadata = json.loads(text[text.index('{'):])
    assert metadata == {'service': 'mabang_tms_file', 'elapsed_ms': elapsed_ms, 'http_status': status,
                        'stage': stage, 'exception_type': exception_type, 'reason': metadata['reason']}
    assert metadata['reason'] is None if outcome == 'completed' else 'UNEXPECTED_EOF_WHILE_READING' in metadata['reason']
    for secret in ['private-account', 'private-password', 'private-token', 'private-signature', 'private-expiry', relative, url]:
        assert secret not in text and secret not in repr(records[0].args)
    assert get_mock.call_count == 1
    client.close()


@pytest.mark.parametrize('fails', [False, True])
def test_download_logging_failure_cannot_change_original_result(monkeypatch, fails):
    from services.mabang_tms import client as tms_client
    monkeypatch.setattr(tms_client.logger, 'info', Mock(side_effect=OSError('log write failed')))
    streamed = MagicMock(status_code=200)
    streamed.__enter__.return_value = streamed
    streamed.iter_content.return_value = iter([b'complete'])
    failure = requests.exceptions.SSLError(ssl.SSLEOFError(8, 'UNEXPECTED_EOF_WHILE_READING'))
    get = Mock(side_effect=failure) if fails else Mock(return_value=streamed)
    monkeypatch.setattr(requests, 'get', get)
    client = TmsClient('account', 'password', sleeper=lambda _: None)
    if fails:
        with pytest.raises(TmsError) as caught:
            client.download('https://tms-cos.mabangerp.com/synthetic.xls')
        assert caught.value.code == 'download_network_error' and caught.value.__cause__ is failure
    else:
        assert client.download('https://tms-cos.mabangerp.com/synthetic.xls') == b'complete'
    assert get.call_count == 1
    client.close()


@pytest.mark.parametrize('failure_mode, code', [
    ('tls_before_response', 'download_network_error'),
    ('truncated_stream', 'download_network_error'),
    ('truncated_workbook', 'invalid_workbook'),
    ('rate_limit', 'download_http_429'),
])
def test_second_download_failure_retains_only_verified_batch_without_resubmission(monkeypatch, tmp_path, caplog, failure_mode, code):
    caplog.set_level('INFO', logger='services.mabang_tms.client')
    first = [record(i, stockSku='shared-sku', warehouseName='越南VN1仓', warehouseId=12831) for i in (1, 2)]
    second = [record(3, warehouseName='越南VN1仓', warehouseId=12831)]
    urls = [f'https://tms-cos.mabangerp.com/synthetic-{i}.xlsx?signature=synthetic' for i in (1, 2)]
    session = Mock()
    session.post.side_effect = [
        response(body={'code': '200', 'data': {'apiToken': 'test-token'}}),
        response(body=page(first, 3)), response(body={'code': '200', 'pop': urls[0]}),
        response(body=page(second, 3, 2)), response(body={'code': '200', 'pop': urls[1]}),
    ]
    content = workbook([values(r) for r in first])
    first_response = MagicMock(status_code=200)
    first_response.__enter__.return_value = first_response
    first_response.iter_content.return_value = iter([content])
    if failure_mode == 'tls_before_response':
        failed_download = requests.exceptions.SSLError(ssl.SSLEOFError(8, 'UNEXPECTED_EOF_WHILE_READING'))
    else:
        failed_download = MagicMock(status_code=429 if failure_mode == 'rate_limit' else 200, text='actual download rejection')
        failed_download.__enter__.return_value = failed_download
        if failure_mode == 'truncated_stream':
            def interrupted(_size):
                yield content[:100]
                raise requests.exceptions.ChunkedEncodingError('incomplete response body')
            failed_download.iter_content.side_effect = interrupted
        else:
            failed_download.iter_content.return_value = iter([content[:100]])
    get = Mock(side_effect=[first_response, failed_download])
    monkeypatch.setattr(requests, 'get', get)
    client = TmsClient('account', 'password', session=session, sleeper=lambda _: None)
    result, _ = execute(monkeypatch, tmp_path, [], client, arguments={'warehouse': 'VN1'})
    assert not result['success'] and result['status'] == 'partial_success'
    assert result['error']['code'] == code and result['failed_page'] == 2
    assert result['read_unique_ids'] == 3 and result['exported_unique_ids'] == result['row_count'] == 2
    assert result['remaining_records'] == 1 and result['completed_batches'] == 1
    assert result['warehouses'] == ['越南VN1仓'] and result['distinct_skus'] == 1
    assert session.post.call_count == 5 and get.call_count == 2
    calls = session.post.call_args_list
    assert [c.args[0].rsplit('/', 1)[-1] for c in calls] == [
        'login', 'findMyStockwarehouseList', 'exportStockwarehouse', 'findMyStockwarehouseList', 'exportStockwarehouse']
    assert all(c.kwargs['json']['warehouseId'] == 12831 for c in calls[1:])
    assert [calls[i].kwargs['json']['idsList'] for i in (2, 4)] == [[1, 2], [3]]
    assert [c.args[0] for c in get.call_args_list] == urls
    assert client.token == ''
    session.close.assert_called_once()
    artifact = Path(result['artifacts'][0]['path'])
    assert '越南VN1仓-部分导出' in artifact.name
    assert list(artifact.parent.glob('原始-*')) == [Path(result['batches'][0]['raw_path'])]
    assert Path(result['batches'][0]['raw_path']).read_bytes() == content
    wb = load_workbook(artifact)
    assert list(wb.active.values) == [tuple(HEADERS), *(tuple(values(r)) for r in first)]
    wb.close()
    logs = [r.getMessage() for r in caplog.records if r.name == 'services.mabang_tms.client']
    metadata = [json.loads(line[line.index('{'):]) for line in logs]
    assert len(metadata) == 2
    assert metadata[0]['http_status'] == 200 and metadata[0]['exception_type'] is None
    assert metadata[1]['http_status'] == (None if failure_mode == 'tls_before_response' else 429 if failure_mode == 'rate_limit' else 200)
    assert metadata[1]['exception_type'] == {'tls_before_response': 'SSLError', 'truncated_stream': 'ChunkedEncodingError',
                                          'truncated_workbook': None, 'rate_limit': 'TmsError'}[failure_mode]


def test_execution_budgets(monkeypatch, tmp_path):
    result, client = execute(monkeypatch, tmp_path, [page([], 100001)])
    assert result['error']['code'] == 'record_limit'
    assert len(client.calls) == 1
    client = TmsClient('a', 'p', clock=lambda: 0)
    client.clock = lambda: 901
    with pytest.raises(TmsError, match='15 分钟'):
        client.check_budget()


def test_cli_retains_partial_artifact_without_auth_material(monkeypatch, tmp_path, capsys):
    import json
    from lxeskill import cli
    from lxeskill import business
    from services.commands.mabang_tms import export_run
    output = tmp_path / 'partial.xlsx'
    output.write_bytes(workbook([values(record(1))]))
    monkeypatch.setattr(business, 'artifact_root', lambda: tmp_path)
    payload = {'success': False, 'status': 'partial_success', 'error': {'code': 'early_empty_page', 'message': 'actual error'},
               'artifacts': [{'path': str(output)}], 'auth_refresh_required': False}
    monkeypatch.setattr(export_run, 'run_with_events', lambda *args, **kwargs: payload)
    assert cli.main(["mabang-tms", "products", "export"]) != 0
    result = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert result['ok'] is False and result['files']
    assert result['data']['status'] == 'partial_success'


# Synthetic rows based on the user-confirmed 2026-10-16 request contract.
# These are not captured production responses or a real Vietnam Excel sample.
@pytest.mark.parametrize('warehouse,warehouse_id', WAREHOUSE_CASES)
def test_selected_warehouse_requests_share_id_and_original_fields(warehouse, warehouse_id):
    session = Mock()
    session.post.side_effect = [response(), response(body={'code': '200', 'pop': 'file'})]
    client = TmsClient('account', 'password', session=session, sleeper=lambda _: None)
    client.list_page(2, warehouse_id=warehouse_id)
    client.export([41, '42'], warehouse_id=warehouse_id)
    listing, export = [call.kwargs['json'] for call in session.post.call_args_list]
    assert listing == {**workflow.FILTERS, 'page': 2, 'pageSize': 1000, 'warehouseId': warehouse_id}
    assert export == {'startNum': None, 'pageSize': 1000, 'classId': None, 'gridproperty': -1,
                      'idsList': [41, '42'], 'isCombo': 1, 'ischecked': False, 'orderBys': '2',
                      'status': 1, 'stockValueArr': FIELDS, 'warehouseId': warehouse_id}


def test_legacy_client_requests_do_not_add_warehouse_parameter():
    session = Mock()
    session.post.side_effect = [response(), response(body={'code': '200', 'pop': 'file'})]
    client = TmsClient('account', 'password', session=session, sleeper=lambda _: None)
    client.list_page(1)
    client.export([1])
    assert all('warehouseId' not in call.kwargs['json'] for call in session.post.call_args_list)


@pytest.mark.parametrize('warehouse,warehouse_id', WAREHOUSE_CASES)
def test_scoped_pagination_preserves_duplicate_skus_and_uses_top_level_ids(monkeypatch, tmp_path, warehouse, warehouse_id):
    other_id = 12831 if warehouse_id == 12107 else 12107
    records = [record(i, stockSku='same-sku', warehouseId=str(warehouse_id), warehouseName=warehouse,
                      stock={'id': 900 + i, 'warehouseId': other_id}) for i in (1, 2, 3)]
    result, client = execute(monkeypatch, tmp_path,
                             [page(records[:1], 3), page(records[1:], 3, 2)], arguments={'warehouse': warehouse})
    assert result['success'] and result['reported_total'] == result['exported_unique_ids'] == 3
    assert result['completed_batches'] == 2 and result['row_count'] == 3 and result['distinct_skus'] == 1
    assert result['warehouse_scope'] == 'selected_warehouse'
    assert result['warehouse'] == warehouse and result['warehouse_id'] == warehouse_id
    assert result['filters']['warehouseId'] == warehouse_id
    assert client.warehouse_requests == [('list', warehouse_id), ('export', warehouse_id)] * 2
    assert [value for stage, value in client.calls if stage == 'export'] == [[1], [2, 3]]
    assert client.login_count == 1
    headers, rows, _ = read_workbook(Path(result['artifacts'][0]['path']).read_bytes())
    assert {row[headers.index('仓库')] for row in rows} == {warehouse}
    assert [row[0] for row in rows] == ['same-sku'] * 3
    assert [list(row) for row in rows] == [values(r) for r in records]


@pytest.mark.parametrize('actual', [12107, None, True, {'id': 12831}])
def test_selected_list_identity_cannot_fall_back_to_nested_stock(monkeypatch, tmp_path, actual):
    row = record(1, warehouseId=actual, warehouseName='VN1', stock={'warehouseId': 12831})
    result, client = execute(monkeypatch, tmp_path, [page([row], 1)], arguments={'warehouse': 'VN1'})
    assert result['error']['code'] == 'warehouse_scope_mismatch'
    assert not result['artifacts'] and not result['batches']
    assert [stage for stage, _ in client.calls] == ['list']


@pytest.mark.parametrize('stage', ['list', 'excel'])
@pytest.mark.parametrize('failed_page', [1, 2])
def test_mixed_warehouse_batch_is_never_published(monkeypatch, tmp_path, stage, failed_page):
    good = record(1, warehouseId=12831, warehouseName='VN1')
    bad = record(2, warehouseId=12831 if stage == 'excel' else 12107, warehouseName='VN1')
    pages = [page([good], 2), page([bad], 2, 2)] if failed_page == 2 else [page([good, bad], 2)]
    client = FakeClient(pages)
    if stage == 'excel':
        original_download = client.download
        def download(url):
            if client.current['page'] == failed_page:
                return workbook([values(r) for r in client.current['datas']] + [values(record(2, warehouseName='PH1'))])
            return original_download(url)
        client.download = download
    result, _ = execute(monkeypatch, tmp_path, pages, client, arguments={'warehouse': 'VN1'})
    assert not result['success']
    assert result['error']['code'] == ('warehouse_scope_mismatch' if stage == 'list' else 'coverage_mismatch')
    assert result['completed_batches'] == failed_page - 1
    assert result['status'] == ('partial_success' if failed_page == 2 else 'failed')
    assert len(list(tmp_path.rglob('原始-*.xlsx'))) == failed_page - 1
    if failed_page == 2:
        assert result['remaining_records'] == 1
        headers, rows, _ = read_workbook(Path(result['artifacts'][0]['path']).read_bytes())
        assert len(rows) == 1 and rows[0][headers.index('仓库')] == 'VN1'
    else:
        assert not result['artifacts']


@pytest.mark.parametrize('arguments', [{}, {'warehouse': 'all'}])
def test_all_scope_preserves_legacy_multiwarehouse_export(monkeypatch, tmp_path, arguments):
    records = [record(1, stockSku='shared', warehouseId=12107, warehouseName='PH1'),
               record(2, stockSku='shared', warehouseId=12831, warehouseName='VN1')]
    result, client = execute(monkeypatch, tmp_path, [page(records, 2)], arguments=arguments)
    assert result['success'] and result['warehouses'] == ['PH1', 'VN1']
    assert result['warehouse_scope'] == 'current_account_all' and 'warehouseId' not in result['filters']
    assert client.warehouse_requests == [('list', None), ('export', None)]
    assert result['row_count'] == 2 and result['distinct_skus'] == 1


@pytest.mark.parametrize('arguments', [{'warehouse': ''}, {'warehouse': 'VN8806'}, {'warehouse': 12831},
                                      {'warehouse': ['PH1', 'VN1']}, {'warehouse': None}, {'country': 'VN'}])
def test_invalid_scope_is_rejected_before_authentication(monkeypatch, arguments):
    client = Mock()
    monkeypatch.setattr(workflow, 'TmsClient', client)
    result = workflow.run(arguments)
    assert result['error']['code'] == 'invalid_arguments'
    client.assert_not_called()


def test_agent_scoped_invocation_requires_explicit_scope_before_authentication(monkeypatch):
    monkeypatch.setenv('LXESKILL_SKILL_SCOPE', 'mabang-tms-products-export')
    client = Mock()
    monkeypatch.setattr(workflow, 'TmsClient', client)
    result = workflow.run({})
    assert result['error']['code'] == 'warehouse_required'
    assert not result['artifacts']
    client.assert_not_called()


@pytest.mark.parametrize('warehouse,warehouse_id', WAREHOUSE_CASES)
def test_selected_zero_records_does_not_fall_back_to_all(monkeypatch, tmp_path, warehouse, warehouse_id):
    result, client = execute(monkeypatch, tmp_path, [page([], 0)], arguments={'warehouse': warehouse})
    assert result['success'] and not result['artifacts'] and result['row_count'] == 0
    assert client.warehouse_requests == [('list', warehouse_id)]


def test_scoped_early_empty_page_retains_verified_partial(monkeypatch, tmp_path):
    result, _ = execute(monkeypatch, tmp_path,
                        [page([record(1, warehouseId=12831, warehouseName='VN1')], 2), page([], 2, 2)],
                        arguments={'warehouse': 'VN1'})
    assert result['status'] == 'partial_success' and result['error']['code'] == 'early_empty_page'
    assert result['reported_total'] == 2 and result['remaining_records'] == 1
    assert len(result['artifacts']) == 1


def test_selected_record_can_still_expand_to_multiple_excel_rows(monkeypatch, tmp_path):
    r = record(1, warehouseId=12831, warehouseName='越南 VN1 仓')
    client = FakeClient([page([r], 1)])
    client.download = lambda _: workbook([values(r), values(r)])
    result, _ = execute(monkeypatch, tmp_path, [], client, arguments={'warehouse': 'VN1'})
    assert result['success'] and result['exported_unique_ids'] == 1 and result['row_count'] == 2
    assert result['warehouses'] == ['越南 VN1 仓']


@pytest.mark.parametrize('warehouse,label', [('PH1', '菲律宾PH1仓'), ('VN1', '越南VN1仓'),
                                           ('all', '全部仓库'), (None, '全部仓库')])
@pytest.mark.parametrize('partial', [False, True])
def test_final_filename_identifies_requested_scope_and_retains_partial_marker(monkeypatch, tmp_path, warehouse, label, partial):
    # An all-warehouse request must retain that name even if only PH1 has rows.
    wid = dict(WAREHOUSE_CASES).get(warehouse, 12107)
    records = [record(1, warehouseId=wid, warehouseName=warehouse if warehouse in ('PH1', 'VN1') else 'PH1')]
    pages = [page(records, 2 if partial else 1)] + ([page([], 2, 2)] if partial else [])
    result, _ = execute(monkeypatch, tmp_path, pages, arguments={'warehouse': warehouse} if warehouse else {})
    assert result['status'] == ('partial_success' if partial else 'completed')
    path = Path(result['artifacts'][0]['path'])
    marker = '-部分导出' if partial else ''
    assert re.fullmatch(rf'马帮TMS-商品-{label}{marker}-\d{{8}}-\d{{6}}\.xlsx', path.name)
    assert '合并' not in path.name
    headers, rows, _ = read_workbook(path.read_bytes())
    assert list(headers) == HEADERS and [list(row) for row in rows] == [values(r) for r in records]


@pytest.mark.parametrize('warehouse,agent', [('PH1', True), ('VN1', True), ('all', True), (None, False)])
def test_cli_scope_dispatches_real_workflow_and_deliverable(monkeypatch, tmp_path, capsys, warehouse, agent):
    import json
    from lxeskill import cli, business
    chosen = WAREHOUSE_CASES if warehouse in (None, 'all') else [(warehouse, dict(WAREHOUSE_CASES)[warehouse])]
    records = [record(i, warehouseId=wid, warehouseName=name) for i, (name, wid) in enumerate(chosen, 1)]
    client = FakeClient([page(records, len(records))])
    monkeypatch.setenv('LXE_MABANG_TMS_ACCOUNT', 'account')
    monkeypatch.setenv('LXE_MABANG_TMS_PASSWORD', 'password')
    if agent:
        monkeypatch.setenv('LXESKILL_SKILL_SCOPE', 'mabang-tms-products-export')
    else:
        monkeypatch.delenv('LXESKILL_SKILL_SCOPE', raising=False)
    monkeypatch.setattr(workflow, 'TmsClient', lambda *args: client)
    monkeypatch.setattr(workflow, 'dataset_dir', lambda *args: tmp_path.joinpath(*args[1:]))
    monkeypatch.setattr(business, 'artifact_root', lambda: tmp_path)
    args = ['mabang-tms', 'products', 'export'] + (['--warehouse', warehouse] if warehouse else [])
    assert cli.main(args) == 0
    result = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert result['ok'] and len(result['files']) == 1 and Path(result['files'][0]).is_file()
    assert result['files'] == [result['data']['artifacts'][0]['path']]
    label = {'PH1': '菲律宾PH1仓', 'VN1': '越南VN1仓'}.get(warehouse, '全部仓库')
    assert re.fullmatch(rf'马帮TMS-商品-{label}-\d{{8}}-\d{{6}}\.xlsx', Path(result['files'][0]).name)
    assert result['data']['warehouses'] == [name for name, _ in chosen]
    assert result['data']['warehouse'] == (warehouse or 'all')
    selected_id = dict(WAREHOUSE_CASES).get(warehouse)
    assert client.warehouse_requests == [('list', selected_id), ('export', selected_id)]


@pytest.mark.parametrize('warehouse', [None, 'VN8806', '12831', ''])
def test_agent_cli_missing_or_invalid_scope_never_authenticates(monkeypatch, capsys, warehouse):
    import json
    from lxeskill import cli
    monkeypatch.setenv('LXESKILL_SKILL_SCOPE', 'mabang-tms-products-export')
    client = Mock()
    monkeypatch.setattr(workflow, 'TmsClient', client)
    args = ['mabang-tms', 'products', 'export'] + (['--warehouse', warehouse] if warehouse is not None else [])
    assert cli.main(args) != 0
    result = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert not result['ok'] and not result['files']
    assert result['data']['error']['code'] == ('warehouse_required' if warehouse is None else 'invalid_arguments')
    client.assert_not_called()


def test_agent_cli_scoped_partial_delivers_only_verified_warehouse(monkeypatch, tmp_path, capsys):
    import json
    from lxeskill import cli, business
    client = FakeClient([page([record(1, warehouseId=12831, warehouseName='VN1')], 2),
                         page([record(2, warehouseId=12107, warehouseName='PH1')], 2, 2)])
    monkeypatch.setenv('LXE_MABANG_TMS_ACCOUNT', 'account')
    monkeypatch.setenv('LXE_MABANG_TMS_PASSWORD', 'password')
    monkeypatch.setenv('LXESKILL_SKILL_SCOPE', 'mabang-tms-products-export')
    monkeypatch.setattr(workflow, 'TmsClient', lambda *args: client)
    monkeypatch.setattr(workflow, 'dataset_dir', lambda *args: tmp_path.joinpath(*args[1:]))
    monkeypatch.setattr(business, 'artifact_root', lambda: tmp_path)
    assert cli.main(['mabang-tms', 'products', 'export', '--warehouse', 'VN1']) != 0
    result = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert not result['ok'] and result['data']['status'] == 'partial_success' and len(result['files']) == 1
    headers, rows, _ = read_workbook(Path(result['files'][0]).read_bytes())
    assert {r[headers.index('仓库')] for r in rows} == {'VN1'}
    assert result['data']['error']['code'] == 'warehouse_scope_mismatch'


@pytest.mark.parametrize('failed_page,retained_rows', [(1, 1), (1, 2), (2, 1)])
def test_review_missing_same_sku_inventory_records_cannot_pass_set_coverage(monkeypatch, tmp_path, failed_page, retained_rows):
    records = [record(i, stockSku='shared', warehouseId=12831, warehouseName='VN1') for i in (1, 2, 3)]
    pages = ([page(records[:1], 3), page(records[1:], 3, 2)] if failed_page == 2
             else [page(records, 3)])
    client = FakeClient(pages)
    original_download = client.download
    def download(url):
        if client.current['page'] == failed_page:
            return workbook([values(r) for r in client.current['datas'][:retained_rows]])
        return original_download(url)
    client.download = download
    result, _ = execute(monkeypatch, tmp_path, pages, client, arguments={'warehouse': 'VN1'})
    assert not result['success']
    assert result['error']['code'] == 'coverage_count_mismatch'
    assert result['completed_batches'] == failed_page - 1
    assert result['status'] == ('partial_success' if failed_page == 2 else 'failed')
    assert len(list(tmp_path.rglob('原始-*.xlsx'))) == failed_page - 1


@pytest.mark.parametrize('total_page', [-1, True, 'invalid', None, 1.5])
@pytest.mark.parametrize('warehouse,warehouse_id', WAREHOUSE_CASES + [('all', None)])
def test_total_page_format_does_not_block_valid_records(monkeypatch, tmp_path, total_page, warehouse, warehouse_id):
    body = {**page([record(1, warehouseId=warehouse_id, warehouseName=warehouse)], 1),
            'totalPage': total_page, 'pop': {'totalPage': total_page}}
    result, client = execute(monkeypatch, tmp_path, [body], arguments={'warehouse': warehouse})
    assert result['success'] and result['status'] == 'completed'
    assert result['reported_total'] == result['read_unique_ids'] == result['exported_unique_ids'] == 1
    assert result['remaining_records'] == 0 and result['row_count'] == 1 and len(result['artifacts']) == 1
    assert [stage for stage, _ in client.calls] == ['list', 'export', 'download']
    assert client.warehouse_requests == [('list', warehouse_id), ('export', warehouse_id)]


@pytest.mark.parametrize('warehouse,warehouse_id', WAREHOUSE_CASES + [('all', None)])
def test_review_real_vn1_total_page_echo_does_not_gate_completion(monkeypatch, tmp_path, warehouse, warehouse_id):
    # Captured VN1 metadata, synthetic inventory/Excel rows. PH1/all are offline
    # scope permutations, not captured production responses from those scopes.
    evidence = json.loads(Path(__file__).with_name('fixtures').joinpath('vn1-pagination-metadata.json').read_text())
    assert evidence['response_fields'] == {'records': 'datas', 'total': 'totalNum', 'pagination': 'pop'}
    assert evidence['request'] == {'warehouseId': 12831, 'pageSize': 1000}
    pages, offset = [], 0
    for metadata in evidence['pages']:
        rows = [record(j, stockSku=f'SKU-{j // 2}',
                       warehouseId=warehouse_id or (12107 if j % 2 else 12831),
                       warehouseName=warehouse if warehouse_id else ('PH1' if j % 2 else 'VN1'))
                for j in range(offset, offset + metadata['row_count'])]
        pages.append({'code': '200', 'data': None, 'datas': rows,
                      'totalNum': evidence['totalNum'], 'pop': metadata['pop']})
        offset += metadata['row_count']
    result, client = execute(monkeypatch, tmp_path, pages, arguments={'warehouse': warehouse})
    assert result['success'] and result['reported_total'] == result['read_unique_ids'] == 4208
    assert result['exported_unique_ids'] == 4208 and result['remaining_records'] == 0
    assert result['completed_batches'] == 5 and result['row_count'] == 4208 and result['distinct_skus'] == 2104
    assert [value for stage, value in client.calls if stage == 'list'] == [1, 2, 3, 4, 5]
    assert [i for stage, ids in client.calls if stage == 'export' for i in ids] == list(range(4208))
    assert client.warehouse_requests == [('list', warehouse_id), ('export', warehouse_id)] * 5
    assert result['warehouses'] == ([warehouse] if warehouse_id else ['PH1', 'VN1'])
    assert 'reported_total_pages' not in result


def test_review_cross_page_total_page_change_no_longer_stops(monkeypatch, tmp_path):
    records = [record(i, warehouseId=12831, warehouseName='VN1') for i in (1, 2, 3)]
    pages = [{**page(records[:1], 3), 'totalPage': '2'},
             {**page(records[1:], 3, 2), 'totalPage': 3}]
    result, client = execute(monkeypatch, tmp_path, pages, arguments={'warehouse': 'VN1'})
    assert result['success'] and result['completed_batches'] == 2
    assert [value for stage, value in client.calls if stage == 'export'] == [[1], [2, 3]]


def test_review_single_page_total_page_echo_mismatch_no_longer_gates(monkeypatch, tmp_path):
    body = {**page([record(1, warehouseId=12831, warehouseName='VN1')], 1), 'totalPage': 2}
    result, _ = execute(monkeypatch, tmp_path, [body], arguments={'warehouse': 'VN1'})
    assert result['success'] and result['read_unique_ids'] == result['reported_total'] == 1


@pytest.mark.parametrize('warehouse,warehouse_id', WAREHOUSE_CASES)
def test_review_scoped_short_pages_continue_with_page_echo(monkeypatch, tmp_path, warehouse, warehouse_id):
    pages = [{**page([record(1, warehouseId=warehouse_id, warehouseName=warehouse)], 2), 'totalPage': 1},
             {**page([record(2, warehouseId=warehouse_id, warehouseName=warehouse)], 2, 2), 'totalPage': '2'}]
    result, _ = execute(monkeypatch, tmp_path, pages, arguments={'warehouse': warehouse})
    assert result['success'] and result['completed_batches'] == 2 and result['read_unique_ids'] == 2
    assert 'reported_total_pages' not in result


def test_review_coverage_accepts_real_vn1_warehouse_field(monkeypatch, tmp_path):
    # Real list rows carry warehouseName=None and warehouse='越南VN1仓'.
    row = record(1, warehouseId=12831, warehouseName=None, warehouse='越南VN1仓')
    client = FakeClient([page([row], 1)])
    client.download = lambda _: workbook([values({**row, 'warehouseName': '越南VN1仓'})])
    result, _ = execute(monkeypatch, tmp_path, [], client, arguments={'warehouse': 'VN1'})
    assert result['success'] and result['warehouses'] == ['越南VN1仓']


@pytest.mark.parametrize('total_page', [0, 1, 2, 101])
def test_review_zero_total_page_forms_remain_compatible(monkeypatch, tmp_path, total_page):
    result, _ = execute(monkeypatch, tmp_path, [{**page([], 0), 'totalPage': total_page}],
                        arguments={'warehouse': 'VN1'})
    assert result['success'] and not result['artifacts']


@pytest.mark.parametrize('warehouse,warehouse_id', WAREHOUSE_CASES)
def test_review_scoped_cross_page_id_repetition_stops(monkeypatch, tmp_path, warehouse, warehouse_id):
    pages = [{**page([record(1, warehouseId=warehouse_id, warehouseName=warehouse)], 2), 'totalPage': 2},
             {**page([record('1', warehouseId=warehouse_id, warehouseName=warehouse)], 2, 2), 'totalPage': 2}]
    result, client = execute(monkeypatch, tmp_path, pages, arguments={'warehouse': warehouse})
    assert result['status'] == 'partial_success' and result['error']['code'] == 'duplicate_id'
    assert [value for stage, value in client.calls if stage == 'export'] == [[1]]


@pytest.mark.parametrize('warehouse,warehouse_id', WAREHOUSE_CASES + [('all', None)])
@pytest.mark.parametrize('metadata', [
    {},  # Missing metadata must not terminate a short page.
    {'pop': {'totalPage': 1}},  # Stale first-page value must not terminate a short page.
    {'pop': {'totalPage': 101}},  # Metadata is not the actual request budget.
    {'pop': {'totalPage': 0}},  # No unverified relation to totalNum.
    {'totalPage': 5, 'pop': {'totalPage': 1}},  # No unverified cross-wrapper equality.
])
def test_total_page_metadata_does_not_override_record_completeness(monkeypatch, tmp_path, warehouse, warehouse_id, metadata):
    pages = [{**page([record(i, warehouseId=warehouse_id, warehouseName=warehouse)], 2, i), **metadata}
             for i in (1, 2)]
    result, client = execute(monkeypatch, tmp_path, pages, arguments={'warehouse': warehouse})
    assert result['success'] and result['read_unique_ids'] == result['exported_unique_ids'] == 2
    assert result['remaining_records'] == 0 and result['row_count'] == 2
    assert [value for stage, value in client.calls if stage == 'list'] == [1, 2]
    assert [value for stage, value in client.calls if stage == 'export'] == [[1], [2]]


@pytest.mark.parametrize('warehouse,warehouse_id', WAREHOUSE_CASES + [('all', None)])
def test_actual_request_page_budget_still_stops_incomplete_export(monkeypatch, tmp_path, warehouse, warehouse_id):
    pages = [{**page([record(i, warehouseId=warehouse_id, warehouseName=warehouse)], 101, i),
              'pop': {'totalPage': 1}}
             for i in range(1, 102)]
    result, client = execute(monkeypatch, tmp_path, pages, arguments={'warehouse': warehouse})
    assert not result['success'] and result['error']['code'] == 'page_limit'
    assert result['status'] == 'partial_success' and result['failed_page'] == 100
    assert result['read_unique_ids'] == result['exported_unique_ids'] == result['completed_batches'] == 100
    assert result['remaining_records'] == 1 and result['row_count'] == 100
    assert [value for stage, value in client.calls if stage == 'list'] == list(range(1, 101))
    assert len(result['artifacts']) == 1
