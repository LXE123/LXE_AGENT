from __future__ import annotations

import io
import os
from itertools import chain
from pathlib import Path
from zipfile import ZipFile, is_zipfile

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from .client import TmsError

REQUIRED_HEADERS = {"库存SKU", "30天销量", "7天销量", "可用库存", "15天销量", "仓库",
                    "库存量", "在途量", "成本价", "采购价", "入库时间"}
XLS_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def read_workbook(content):
    if content.startswith(XLS_SIGNATURE):
        wb = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            if wb.nsheets != 1:
                raise TmsError("unexpected_sheets", f"导出文件工作表数为 {wb.nsheets}，预期 1")
            sheet = wb.sheet_by_index(0)
            rows = []
            for i in range(sheet.nrows):
                row = []
                for cell in sheet.row(i):
                    if cell.ctype == xlrd.XL_CELL_ERROR:
                        raise TmsError("cell_error", f"源表第 {i + 1} 行包含 Excel 错误")
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, wb.datemode)
                               if cell.ctype == xlrd.XL_CELL_DATE else cell.value)
                rows.append(row)
        finally:
            wb.release_resources()
        extension = "xls"
    elif is_zipfile(io.BytesIO(content)):
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 250_000_000:
                raise TmsError("workbook_size_limit", "XLSX 解压后超过 250 MB")
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            if len(wb.sheetnames) != 1:
                raise TmsError("unexpected_sheets", f"导出文件工作表为 {wb.sheetnames}，预期 1 张")
            rows = []
            for row in wb.active.iter_rows():
                if any(cell.data_type in {"f", "e"} for cell in row):
                    raise TmsError("unexpected_formula", "平台原表含公式或错误单元格，无法作为原始值合并")
                rows.append([c.value for c in row])
        finally:
            wb.close()
        extension = "xlsx"
    else:
        raise TmsError("invalid_workbook", "下载文件头不是 XLS/XLSX")
    if not rows:
        raise TmsError("missing_headers", "导出文件为空")
    headers = rows[0]
    if any(not isinstance(h, str) or not h for h in headers) or len(set(headers)) != len(headers):
        raise TmsError("invalid_headers", f"导出表头为空或重复: {headers}")
    if not REQUIRED_HEADERS.issubset(headers):
        raise TmsError("missing_headers", f"实际表头 {headers}; 缺少 {sorted(REQUIRED_HEADERS - set(headers))}")
    values = [tuple(row) for row in rows[1:] if any(v not in (None, "") for v in row)]
    if any(len(row) != len(headers) for row in values):
        raise TmsError("invalid_row", "导出行宽与表头不一致")
    return tuple(headers), values, extension


def coverage(records, headers, rows):
    """Compare exported SKU/warehouse pairs, without assuming one row per ID."""
    expected = set()
    for record in records:
        sku = record.get("stockSku")
        warehouse = record.get("warehouseName") or record.get("warehouse")
        if not isinstance(sku, str) or not sku or not isinstance(warehouse, str) or not warehouse:
            raise TmsError("coverage_schema_unknown", f"无法核对列表 SKU/仓库字段，实际字段: {sorted(record)}")
        expected.add((sku, warehouse))
    sku_index, warehouse_index = headers.index("库存SKU"), headers.index("仓库")
    actual = {(row[sku_index], row[warehouse_index]) for row in rows}
    if actual != expected:
        raise TmsError("coverage_mismatch", f"批次商品/仓库覆盖不一致: 缺少 {len(expected - actual)}，多出 {len(actual - expected)}")


def atomic_bytes(path: Path, content: bytes):
    temporary = path.with_name("." + path.name + ".tmp")
    try:
        with temporary.open("xb") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_merged(path: Path, headers, rows):
    wb = Workbook(write_only=True)
    sheet = wb.create_sheet("商品")
    for row in chain((headers,), rows):
        cells = [WriteOnlyCell(sheet, value=value) for value in row]
        # Preserve text beginning with '=' as text, never turn platform data into formulas.
        for cell, value in zip(cells, row):
            if isinstance(value, str):
                cell.data_type = "s"
        sheet.append(cells)
    try:
        stream = io.BytesIO()
        wb.save(stream)
        atomic_bytes(path, stream.getvalue())
    finally:
        wb.close()
