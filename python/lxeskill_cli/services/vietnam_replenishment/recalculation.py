"""Recalculate and verify a current Vietnam workbook before publishing it."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_UP
import os
from pathlib import Path
import re
import subprocess
import sys
from tempfile import TemporaryDirectory
from typing import Mapping

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.formula import ArrayFormula

from services.yacang.errors import safe_remote_detail

from .asset_contract import (
    AUXILIARY_HEADERS,
    MAIN_HEADERS,
    REQUIRED_SHEETS,
    SkuParameters,
    load_sku_parameters,
)
from .workbook import (
    RecommendationConfig,
    _load_skeleton,
    _validated_rows,
    write_vietnam_workbook,
)
from .yacang_sources import VietnamSources


class WorkbookGenerationError(ValueError):
    """The current inputs or calculated workbook cannot be delivered."""


def _office_diagnostic(value: object) -> str:
    secrets = tuple(
        os.environ.get(name, "")
        for name in ("LXE_YACANG_MOBILE", "LXE_YACANG_PASSWORD")
    )
    return safe_remote_detail(value, secrets=secrets, limit=2000)


def recalculate_with_office(
    input_path: str | Path, output_path: str | Path, *, timeout_seconds: int = 300
) -> None:
    """Use the host's managed LibreOffice Kit through the existing launcher."""
    source = Path(input_path)
    output = Path(output_path)
    if source.resolve() == output.resolve():
        raise WorkbookGenerationError("Office 重算的输入与输出路径不能相同")
    command = [
        sys.executable,
        "-m",
        "shared.office",
        "recalculate",
        "--input",
        str(source),
        "--output",
        str(output),
    ]
    try:
        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout_seconds,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        detail = _office_diagnostic(exc.stderr or exc.stdout or exc)
        raise WorkbookGenerationError(f"Office 重算超时（{timeout_seconds} 秒）: {detail}") from exc
    except OSError as exc:
        raise WorkbookGenerationError(
            f"无法启动 Office 重算: {type(exc).__name__}: {_office_diagnostic(exc)}"
        ) from exc
    if result.returncode != 0:
        detail = _office_diagnostic(result.stderr or result.stdout or "无诊断输出")
        raise WorkbookGenerationError(f"Office 重算失败（退出码 {result.returncode}）: {detail}")
    if not output.is_file() or output.stat().st_size == 0:
        raise WorkbookGenerationError(
            f"Office 报告成功但没有生成有效文件: {output}；输出: {_office_diagnostic(result.stdout)}"
        )


def _decimal(value: object, *, coordinate: str) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise WorkbookGenerationError(f"{coordinate}: 结果不是有限数值，实际为 {value!r}")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise WorkbookGenerationError(
            f"{coordinate}: 结果不是数字: {type(exc).__name__}: {_office_diagnostic(exc)}"
        ) from exc
    if not number.is_finite():
        raise WorkbookGenerationError(f"{coordinate}: 结果不是有限数值，实际为 {value!r}")
    return number


def _sku_rows(sheet, column: int, *, name: str) -> dict[str, int]:
    rows: dict[str, int] = {}
    for row_number in range(2, sheet.max_row + 1):
        value = sheet.cell(row_number, column).value
        if value is None or value == "":
            if any(cell.value is not None for cell in sheet[row_number]):
                raise WorkbookGenerationError(f"{name} 第 {row_number} 行有数据但没有 SKU")
            continue
        if not isinstance(value, str) or not value.strip() or value != value.strip():
            raise WorkbookGenerationError(f"{name} 第 {row_number} 行 SKU 不是规范文本")
        if value in rows:
            raise WorkbookGenerationError(
                f"{name} SKU {value} 在第 {rows[value]}、{row_number} 行重复"
            )
        rows[value] = row_number
    return rows


def _check_sku_set(rows: Mapping[str, int], expected: set[str], *, name: str) -> None:
    if set(rows) != expected:
        missing = sorted(expected - set(rows))[:10]
        extra = sorted(set(rows) - expected)[:10]
        raise WorkbookGenerationError(f"{name} 本轮 SKU 不一致：缺少 {missing}；多出 {extra}")


def _check_headers(workbook, skeleton) -> None:
    for name in REQUIRED_SHEETS:
        reference = skeleton[name]
        actual = workbook[name]
        for column in range(1, max(reference.max_column, actual.max_column) + 1):
            expected_value = reference.cell(1, column).value
            actual_value = actual.cell(1, column).value
            if actual_value != expected_value:
                raise WorkbookGenerationError(
                    f"{name}!{actual.cell(1, column).coordinate}: "
                    f"表头应为 {expected_value!r}，实际 {actual_value!r}"
                )
    for coordinate, expected in MAIN_HEADERS.items():
        actual = workbook[REQUIRED_SHEETS[0]][coordinate].value
        if actual != expected:
            raise WorkbookGenerationError(
                f"越南备货清单!{coordinate}: 表头应为 {expected!r}，实际 {actual!r}"
            )
    for name, headers in AUXILIARY_HEADERS.items():
        for coordinate, expected in headers.items():
            actual = workbook[name][coordinate].value
            if actual != expected:
                raise WorkbookGenerationError(
                    f"{name}!{coordinate}: 表头应为 {expected!r}，实际 {actual!r}"
                )


def _same_literal(expected: object, actual: object) -> bool:
    if expected is None:
        return actual is None
    if isinstance(expected, bool):
        return actual is expected
    if isinstance(expected, (Decimal, int, float)):
        try:
            return not isinstance(actual, bool) and Decimal(str(expected)) == Decimal(str(actual))
        except (InvalidOperation, TypeError, ValueError):
            return False
    if isinstance(expected, (date, datetime)):
        return actual == expected
    return actual == expected and not isinstance(actual, (date, datetime))


def _check_projected_source(sheet, source_rows, row_numbers: Mapping[str, int], replacements) -> None:
    for sku, row_number in row_numbers.items():
        source = source_rows[sku]
        override = replacements[sku]
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row_number, column)
            header = sheet.cell(1, column).value
            if header is None:
                expected_value = None
            else:
                expected_value = override.get(header, source.get(header))
            if cell.data_type == "f" or not _same_literal(expected_value, cell.value):
                raise WorkbookGenerationError(
                    f"{sheet.title} SKU {sku} 的 {header or cell.coordinate} "
                    f"与本轮来源不一致: {cell.coordinate}={cell.value!r}"
                )


def _normal_formula(formula: str) -> str:
    # LibreOffice writes bare Excel TRUE/FALSE constants as TRUE()/FALSE().
    return re.sub(r"\b(TRUE|FALSE)\b(?!\s*\()", r"\1()", formula)


def _check_main_formulas(main, skeleton, row_number: int, sku: str) -> None:
    reference = skeleton["越南备货清单"]
    for original in reference[2]:
        if original.data_type != "f":
            continue
        actual = main.cell(row_number, original.column)
        if isinstance(original.value, ArrayFormula):
            if not isinstance(actual.value, ArrayFormula) or actual.value.ref != actual.coordinate:
                raise WorkbookGenerationError(f"SKU {sku} 备货公式丢失: {actual.coordinate}")
            expected_formula = Translator(original.value.text, origin=original.coordinate).translate_formula(
                actual.coordinate
            )
            actual_formula = actual.value.text
        else:
            if actual.data_type != "f" or not isinstance(actual.value, str):
                raise WorkbookGenerationError(f"SKU {sku} 备货公式丢失: {actual.coordinate}")
            expected_formula = (
                original.value if original.column in (48, 49, 50, 51)
                else Translator(original.value, origin=original.coordinate).translate_formula(actual.coordinate)
            )
            actual_formula = actual.value
        if not isinstance(actual_formula, str) or _normal_formula(actual_formula) != _normal_formula(expected_formula):
            raise WorkbookGenerationError(
                f"SKU {sku} 公式与内置骨架不一致: {actual.coordinate}={actual_formula!r}"
            )


def _expected_replenishment(status: str, daily: Decimal, days: Decimal,
                            available: Decimal, transit: Decimal) -> Decimal:
    if status == "无动销无库存":
        return Decimal(0)
    quantity = (daily * days - available - transit) / Decimal(10)
    rounded = quantity.to_integral_value(rounding=ROUND_UP) * 10
    return min(Decimal(0), rounded) if status.startswith("清货") else rounded


_TREND_MULTIPLIERS = {
    "无动销": None,
    "持续上升": Decimal("1.2"),
    "近期回升": Decimal("1.1"),
    "平稳": Decimal("1"),
    "近期回落": Decimal("0.9"),
    "持续下滑": Decimal("0.8"),
}


def _numeric_formula_result(cell, sku: str) -> Decimal:
    if cell.data_type != "n":
        raise WorkbookGenerationError(
            f"SKU {sku} 的 {cell.coordinate} 重算缓存应为有限数值，"
            f"实际类型 {cell.data_type!r}、值 {cell.value!r}"
        )
    try:
        return _decimal(cell.value, coordinate=cell.coordinate)
    except WorkbookGenerationError as exc:
        raise WorkbookGenerationError(f"SKU {sku} 的 {exc}") from exc


def _expected_ascii_v(sku: str) -> str:
    # The skeleton keeps the SKU for these markers, otherwise strips a short
    # suffix after its last dash. ASCII avoids Excel/Office Unicode LEN drift.
    if "&" in sku or "+" in sku or "x2" in sku.lower():
        return sku
    last_dash = sku.rfind("-")
    if last_dash >= 0 and len(sku) - last_dash - 1 <= 3:
        return sku[:last_dash]
    return sku


def validate_recalculated_workbook(
    path: str | Path,
    *,
    sources: VietnamSources,
    parameters: Mapping[str, SkuParameters],
    config: RecommendationConfig,
) -> None:
    """Check formulas, current data, and saved results of the Kit output."""
    formulas = None
    values = None
    skeleton = None
    try:
        formulas = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
        skeleton = _load_skeleton()
    except Exception as exc:
        for workbook in (formulas, values, skeleton):
            if workbook is not None:
                workbook.close()
        raise WorkbookGenerationError(
            f"读取重算文件或内置骨架失败: {type(exc).__name__}: {_office_diagnostic(exc)}"
        ) from exc
    try:
        if tuple(formulas.sheetnames) != REQUIRED_SHEETS or tuple(values.sheetnames) != REQUIRED_SHEETS:
            raise WorkbookGenerationError(
                f"重算结果应恰好包含五张目标表，实际为 {formulas.sheetnames!r}"
            )
        _check_headers(formulas, skeleton)
        current_rows = {row.sku: row for row in _validated_rows(sources, parameters)}
        expected = set(sources.skus)
        main = formulas["越南备货清单"]
        main_values = values["越南备货清单"]
        main_rows = _sku_rows(main, 5, name="越南备货清单")
        _check_sku_set(main_rows, expected, name="越南备货清单")

        auxiliary_rows = {}
        for name, column in (("雅仓库存", 2), ("雅仓动销", 1), ("库存商品信息", 2)):
            rows = _sku_rows(formulas[name], column, name=name)
            _check_sku_set(rows, expected, name=name)
            auxiliary_rows[name] = rows

        _check_projected_source(
            formulas["雅仓库存"], sources.inventory, auxiliary_rows["雅仓库存"],
            {sku: {"在途数量": row.in_transit, "可用库存": row.available}
             for sku, row in current_rows.items()},
        )
        _check_projected_source(
            formulas["雅仓动销"], sources.sales, auxiliary_rows["雅仓动销"],
            {sku: {"7天销量": row.sales_7d, "15天销量": row.sales_15d,
                   "30天销量": row.sales_30d} for sku, row in current_rows.items()},
        )
        _check_projected_source(
            formulas["库存商品信息"], sources.products, auxiliary_rows["库存商品信息"],
            {sku: {"创建时间": row.listed_at} for sku, row in current_rows.items()},
        )

        for sku in sources.skus:
            row_number = main_rows[sku]
            current = current_rows[sku]
            product_time = current.listed_at
            _check_main_formulas(main, skeleton, row_number, sku)

            mapped = parameters[sku]
            direct_values = {
                2: current.hot_flag, 5: sku, 6: current.product.get("中文标题"),
                7: current.cost, 31: current.cross_border_price,
                36: current.discount_price, 40: current.in_transit,
            }
            reference = skeleton["越南备货清单"]
            for column in range(1, max(main.max_column, reference.max_column) + 1):
                if reference.cell(2, column).data_type == "f":
                    continue
                cell = main.cell(row_number, column)
                expected_value = direct_values.get(column)
                if cell.data_type == "f" or not _same_literal(expected_value, cell.value):
                    raise WorkbookGenerationError(
                        f"SKU {sku} 的 {cell.coordinate} 与本轮来源/映射不一致: {cell.value!r}"
                    )

            if main_values.cell(row_number, 27).value != product_time:
                raise WorkbookGenerationError(
                    f"SKU {sku} 上架时间重算值与仓库产品创建时间不一致"
                )
            for column, expected_value, label in (
                (10, current.available, "可用库存"),
                (11, current.sales_30d, "30天销量"),
                (12, current.sales_15d, "15天销量"),
                (13, current.sales_7d, "7天销量"),
                (48, config.weight_30d, "30天参数"),
                (49, config.weight_15d, "15天参数"),
                (50, config.weight_7d, "7天参数"),
                (51, config.exchange_rate, "汇率"),
            ):
                actual = _decimal(main_values.cell(row_number, column).value, coordinate=f"{label} 第 {row_number} 行")
                if actual != expected_value:
                    raise WorkbookGenerationError(f"SKU {sku} 的 {label} 重算值与本轮输入不一致")
            calculated = {}
            for column in (8, 19, 20, 29):
                cell = main_values.cell(row_number, column)
                calculated[column] = _numeric_formula_result(cell, sku)
            status = main_values.cell(row_number, 21)
            if not isinstance(status.value, str) or not status.value.strip():
                raise WorkbookGenerationError(f"SKU {sku} 综合判定没有有效重算结果")
            expected_quantity = _expected_replenishment(
                status.value, calculated[20], calculated[29],
                current.available, current.in_transit,
            )
            if calculated[8] != expected_quantity:
                raise WorkbookGenerationError(
                    f"SKU {sku} 最终备货量与已重算输入不一致: "
                    f"H{row_number}={calculated[8]}，应为 {expected_quantity}"
                )

            for column in range(8, 40):
                cell = main_values.cell(row_number, column)
                if cell.data_type != "e":
                    continue
                if column == 28 and cell.value == "#DIV/0!":
                    daily = _decimal(main_values.cell(row_number, 19).value, coordinate=f"S{row_number}")
                    if daily == 0:
                        continue
                if column in (34, 39) and cell.value == "#DIV/0!":
                    label = "跨境价" if column == 34 else "折扣价"
                    price = current.cross_border_price if column == 34 else current.discount_price
                    if price == 0:
                        raise WorkbookGenerationError(
                            f"SKU {sku} 的{label}为显式 0，{cell.coordinate} 利润率无法计算；"
                            "请在 SKU 映射表提供可计算的价格"
                        )
                raise WorkbookGenerationError(
                    f"SKU {sku} 重算产生公式错误 {cell.coordinate}: {cell.value}"
                )

            # These outputs are required even when every SKU has a complete map.
            for column in ("N", "O", "P", "X", "Y", "Z", "AF", "AG", "AH", "AI", "AK", "AL", "AM"):
                _numeric_formula_result(main_values[f"{column}{row_number}"], sku)
            ab = main_values[f"AB{row_number}"]
            if ab.data_type != "e":
                _numeric_formula_result(ab, sku)

            trend = main_values[f"Q{row_number}"]
            if not isinstance(trend.value, str) or trend.value not in _TREND_MULTIPLIERS:
                raise WorkbookGenerationError(
                    f"SKU {sku} 的 {trend.coordinate} 趋势判定不是有效重算文本: {trend.value!r}"
                )
            multiplier = main_values[f"R{row_number}"]
            expected_multiplier = _TREND_MULTIPLIERS[trend.value]
            if expected_multiplier is None:
                if multiplier.value is not None and multiplier.value != "":
                    raise WorkbookGenerationError(
                        f"SKU {sku} 的 {multiplier.coordinate} 无动销时应为空，实际为 {multiplier.value!r}"
                    )
            elif _numeric_formula_result(multiplier, sku) != expected_multiplier:
                raise WorkbookGenerationError(
                    f"SKU {sku} 的 {multiplier.coordinate} 趋势系数与 {trend.value} 不一致"
                )

            if sku.isascii():
                category = main_values[f"V{row_number}"]
                expected_category = _expected_ascii_v(sku)
                if expected_category:
                    valid = isinstance(category.value, str) and category.value == expected_category
                else:
                    valid = category.value is None or category.value == ""
                if not valid:
                    raise WorkbookGenerationError(
                        f"SKU {sku} 的 {category.coordinate} 款号与本轮 SKU 不一致: "
                        f"{category.value!r}，应为 {expected_category!r}"
                    )

        for column, expected_value in enumerate(
            (config.weight_30d, config.weight_15d, config.weight_7d, config.exchange_rate), 1
        ):
            actual = _decimal(
                formulas["数据更改"].cell(2, column).value,
                coordinate=f"数据更改!{chr(64 + column)}2",
            )
            if actual != expected_value:
                raise WorkbookGenerationError(f"数据更改第 {column} 个参数与本次配置不一致")
    finally:
        formulas.close()
        values.close()
        skeleton.close()


def generate_vietnam_workbook(
    map_path: str | Path,
    output_path: str | Path,
    *,
    sources: VietnamSources,
    config: RecommendationConfig | None = None,
) -> Path:
    """Produce one validated output; never call Yacang or use template history."""
    output = Path(output_path)
    if os.path.lexists(output):
        raise FileExistsError(output)
    current_config = config if config is not None else RecommendationConfig()
    parameters = load_sku_parameters(map_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix=".vietnam-workbook-", dir=output.parent) as directory:
        stage = Path(directory)
        draft = stage / "draft.xlsx"
        recalculated = stage / "recalculated.xlsx"
        write_vietnam_workbook(draft, sources, parameters, current_config)
        recalculate_with_office(draft, recalculated)
        validate_recalculated_workbook(
            recalculated, sources=sources, parameters=parameters, config=current_config
        )
        os.link(recalculated, output)
    return output


__all__ = [
    "WorkbookGenerationError",
    "generate_vietnam_workbook",
    "recalculate_with_office",
    "validate_recalculated_workbook",
]
