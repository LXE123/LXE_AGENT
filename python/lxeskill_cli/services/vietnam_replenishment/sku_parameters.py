"""Resolve current Vietnam SKU inputs without promoting template history to operator input."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
from typing import Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .asset_contract import AssetContractError, SkuParameters, validate_template


_FIELD_COLUMNS = {
    "cost": "G",
    "cross_border_price": "AE",
    "discount_price": "AJ",
    "hot_flag": "B",
    "listed_at": "AA",
}
_FIELDS = tuple(_FIELD_COLUMNS)
_FIRST_COLUMN = 2
_LAST_COLUMN = column_index_from_string("AJ")


@dataclass(frozen=True)
class HistoricalSkuParameters:
    values: SkuParameters
    unavailable: Mapping[str, str]


@dataclass(frozen=True)
class ResolvedSkuParameters:
    values: SkuParameters
    sources: Mapping[str, str]
    unavailable: Mapping[str, str]


def _current_skus(skus: Iterable[str]) -> tuple[str, ...]:
    result: list[str] = []
    seen: set[str] = set()
    for raw in skus:
        if not isinstance(raw, str) or not raw.strip():
            raise AssetContractError("本轮 SKU 必须是非空文本")
        sku = raw.strip()
        if sku not in seen:
            result.append(sku)
            seen.add(sku)
    return tuple(result)


def _blank(value: object) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def _parse_history_cell(
    cell: object, field: str, coordinate: str
) -> tuple[object | None, str | None]:
    if cell.data_type == "f":
        return None, f"{coordinate} 是公式"
    value = cell.value
    if _blank(value):
        return None, None

    if field == "listed_at":
        if isinstance(value, datetime):
            if cell.number_format == "yyyy-mm-dd":
                return value.date().isoformat(), None
            return str(value).strip(), None
        if isinstance(value, date):
            return value.isoformat(), None
        if isinstance(value, str):
            result = value.strip()
            if re.match(r"^\d{4}-\d{2}-\d{2}(?:$|[ T])", result):
                try:
                    if len(result) == 10:
                        date.fromisoformat(result)
                    else:
                        datetime.fromisoformat(result)
                    return result, None
                except ValueError:
                    pass
        return None, f"{coordinate} 是无效的上架时间"

    if isinstance(value, bool):
        return None, f"{coordinate} 是无效的{field}布尔值"
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return None, f"{coordinate} 是无效的{field}数值"
    if not number.is_finite() or number < 0:
        return None, f"{coordinate} 是无效的{field}数值"
    if field == "hot_flag":
        if number not in (1, 2):
            return None, f"{coordinate} 是无效的热销标记（只能是 1 或 2）"
        return int(number), None
    return number, None


def _locations(cells: list[str]) -> str:
    displayed = ", ".join(cells[:5])
    return f"{displayed}（共 {len(cells)} 处）" if len(cells) > 5 else displayed


def load_template_sku_parameters(
    path: str | Path, skus: Iterable[str]
) -> dict[str, HistoricalSkuParameters]:
    """Read only the requested SKUs; ambiguous history is unavailable per field."""
    current = _current_skus(skus)
    validate_template(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise AssetContractError(f"读取模板历史参数失败: {type(exc).__name__}: {exc}") from exc

    try:
        main = workbook["越南备货清单"]
        requested = set(current)
        seen_rows: set[str] = set()
        found: dict[str, dict[str, list[tuple[object, str]]]] = {
            sku: {field: [] for field in _FIELDS} for sku in current
        }
        issues: dict[str, dict[str, list[str]]] = {
            sku: {field: [] for field in _FIELDS} for sku in current
        }
        for row_number, row in enumerate(
            main.iter_rows(min_row=2, min_col=_FIRST_COLUMN, max_col=_LAST_COLUMN), 2
        ):
            sku_cell = row[column_index_from_string("E") - _FIRST_COLUMN]
            raw_sku = sku_cell.value
            if sku_cell.data_type == "f" or not isinstance(raw_sku, str):
                continue
            sku = raw_sku.strip()
            if sku not in requested:
                continue
            seen_rows.add(sku)
            for field, column in _FIELD_COLUMNS.items():
                cell = row[column_index_from_string(column) - _FIRST_COLUMN]
                coordinate = f"越南备货清单!{column}{row_number}"
                value, issue = _parse_history_cell(cell, field, coordinate)
                if issue is not None:
                    issues[sku][field].append(issue)
                elif value is not None:
                    found[sku][field].append((value, coordinate))

        result: dict[str, HistoricalSkuParameters] = {}
        for sku in current:
            values: dict[str, object] = {}
            unavailable: dict[str, str] = {}
            for field in _FIELDS:
                candidates = found[sku][field]
                field_issues = issues[sku][field]
                if sku not in seen_rows:
                    unavailable[field] = "模板历史中没有该 SKU"
                elif field_issues:
                    unavailable[field] = "模板历史值不可用: " + _locations(field_issues)
                elif candidates and all(value == candidates[0][0] for value, _ in candidates):
                    values[field] = candidates[0][0]
                elif candidates:
                    unavailable[field] = "模板历史值冲突: " + _locations(
                        [coordinate for _, coordinate in candidates]
                    )
                else:
                    unavailable[field] = "模板历史值缺失"
            result[sku] = HistoricalSkuParameters(
                values=SkuParameters(**values), unavailable=unavailable,
            )
        return result
    except AssetContractError:
        raise
    except Exception as exc:
        raise AssetContractError(f"读取模板历史参数失败: {type(exc).__name__}: {exc}") from exc
    finally:
        workbook.close()


def resolve_sku_parameters(
    skus: Iterable[str],
    explicit: Mapping[str, SkuParameters],
    history: Mapping[str, HistoricalSkuParameters],
) -> dict[str, ResolvedSkuParameters]:
    """Choose each SKU field by its source, preserving explicit zero values."""
    result: dict[str, ResolvedSkuParameters] = {}
    for sku in _current_skus(skus):
        operator_values = explicit.get(sku, SkuParameters())
        historical = history.get(sku)
        selected: dict[str, object] = {}
        sources: dict[str, str] = {}
        unavailable: dict[str, str] = {}
        for field in _FIELDS:
            explicit_value = getattr(operator_values, field)
            historical_value = getattr(historical.values, field) if historical else None
            if explicit_value is not None:
                selected[field] = explicit_value
                sources[field] = "explicit"
            elif historical_value is not None:
                selected[field] = historical_value
                sources[field] = "template"
            elif field == "hot_flag":
                selected[field] = 2
                sources[field] = "default"
            else:
                unavailable[field] = (
                    historical.unavailable.get(field, "模板历史值缺失")
                    if historical else "模板历史中没有该 SKU"
                )
        result[sku] = ResolvedSkuParameters(
            values=SkuParameters(**selected), sources=sources, unavailable=unavailable,
        )
    return result


__all__ = [
    "HistoricalSkuParameters",
    "ResolvedSkuParameters",
    "load_template_sku_parameters",
    "resolve_sku_parameters",
]
