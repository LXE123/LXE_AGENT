"""Place validated current results into the user's nine-sheet business workbook.

The four operational worksheets are copied as OOXML parts, so their cells,
cached values, formatting, and shared-formula definitions are untouched.
"""

from __future__ import annotations

from contextlib import contextmanager
from copy import copy
import os
from pathlib import Path
import posixpath
from xml.dom import XMLNS_NAMESPACE, minidom
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook


MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT = "http://schemas.openxmlformats.org/package/2006/content-types"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
REQUIRED_ORDER = (
    "越南备货清单", "改库存", "改折扣", "报官方活动", "竞价报名",
    "雅仓库存", "雅仓动销", "数据更改", "库存商品信息",
)
UNCHANGED = REQUIRED_ORDER[1:5]
CURRENT = (REQUIRED_ORDER[0], *REQUIRED_ORDER[5:])


class FullWorkbookError(ValueError):
    """The supplied business workbook cannot safely preserve all nine tabs."""


def _date_system(archive: ZipFile) -> str:
    properties = ET.fromstring(archive.read("xl/workbook.xml")).find(f"{{{MAIN}}}workbookPr")
    flag = properties.get("date1904", "0").strip() if properties is not None else "0"
    if flag not in ("0", "1", "false", "true"):
        raise FullWorkbookError(f"工作簿日期系统标志 date1904 无效: {flag!r}")
    return "1904" if flag in ("1", "true") else "1900"


def validate_full_template(path: str | Path) -> None:
    """Check that a stored template can supply inputs and preserve four tabs."""
    source = Path(path)
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise FullWorkbookError(f"越南备货模板必须是现有的 .xlsx 文件: {source}")
    if source.stat().st_size > 40 * 1024 * 1024:
        raise FullWorkbookError("越南备货模板超过 40 MiB")
    try:
        with ZipFile(source) as archive:
            items = archive.infolist()
            if len(items) > 2000 or sum(item.file_size for item in items) > 250 * 1024 * 1024:
                raise FullWorkbookError("越南备货模板 ZIP 内容超过限制")
            if any(item.flag_bits & 1 for item in items):
                raise FullWorkbookError("越南备货模板包含加密内容")
            if archive.testzip() is not None:
                raise FullWorkbookError("越南备货模板 ZIP 内容损坏")
            if _date_system(archive) != "1900":
                raise FullWorkbookError("不支持 1904 日期系统，请使用 1900 日期系统的越南备货模板")
            parts = _sheet_parts(archive)
            if tuple(parts) != REQUIRED_ORDER:
                raise FullWorkbookError("越南备货模板须包含指定顺序的九张工作表")
            sheets = ET.fromstring(archive.read("xl/workbook.xml")).find(f"{{{MAIN}}}sheets")
            if any(sheet.attrib.get("state", "visible") != "visible" for sheet in sheets):
                raise FullWorkbookError("越南备货模板的九张工作表都必须可见")
            if any(f"xl/worksheets/_rels/{Path(parts[name]).name}.rels" in archive.namelist() for name in UNCHANGED):
                raise FullWorkbookError("越南备货模板的四张业务表含独立附件关系，无法原样保留")
        workbook = load_workbook(source, read_only=True, data_only=False)
        try:
            first = next(workbook[REQUIRED_ORDER[0]].iter_rows(min_row=1, max_row=1, max_col=36))
            for column, label in ((2, "判断热销"), (5, "SKU"), (7, "成本"),
                                  (31, "跨境价"), (36, "折扣价")):
                if first[column - 1].value != label:
                    raise FullWorkbookError(f"越南备货清单第 {column} 列表头应为 {label}")
        finally:
            workbook.close()
    except FullWorkbookError:
        raise
    except (BadZipFile, KeyError, ET.ParseError, OSError, ValueError, StopIteration) as exc:
        raise FullWorkbookError(f"读取越南备货模板失败: {type(exc).__name__}: {exc}") from exc


@contextmanager
def _parsed(data: bytes):
    document = minidom.parseString(data)
    try:
        yield document
    finally:
        document.unlink()


def _sheet_parts(archive: ZipFile) -> dict[str, str]:
    relationships = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    }
    sheets = ET.fromstring(archive.read("xl/workbook.xml")).find(f"{{{MAIN}}}sheets")
    if sheets is None:
        raise FullWorkbookError("业务模板缺少工作表列表")
    parts = {}
    for sheet in sheets:
        target = relationships[sheet.attrib[f"{{{REL}}}id"]]
        part = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
        if not part.startswith("xl/worksheets/") or part not in archive.namelist():
            raise FullWorkbookError(f"工作表 {sheet.attrib['name']} 的文件无效")
        parts[sheet.attrib["name"]] = part
    return parts


def _elements(node):
    return [child for child in node.childNodes if child.nodeType == child.ELEMENT_NODE]


def _section(document, name: str):
    found = document.getElementsByTagNameNS(MAIN, name)
    return found[0] if found else None


def _import_node(document, node, parent):
    """Keep the source subtree's namespace scope, including attribute/MC prefixes.

    importNode copies local declarations but omits inherited ones. Copy those
    only where the destination scope differs; never overwrite local bindings,
    which may legitimately be rebound again further down the subtree.
    """
    def scope(element):
        bindings = {"xmlns": ""}
        ancestors = []
        while element.nodeType == element.ELEMENT_NODE:
            ancestors.append(element)
            element = element.parentNode
        for ancestor in reversed(ancestors):
            bindings.update((attr.name, attr.value) for attr in ancestor.attributes.values()
                            if attr.namespaceURI == XMLNS_NAMESPACE)
        return bindings

    copied = document.importNode(node, deep=True)
    destination = scope(parent)
    for binding, uri in scope(node.parentNode).items():
        if not copied.hasAttribute(binding) and destination.get(binding, "") != uri:
            copied.setAttributeNS(XMLNS_NAMESPACE, binding, uri)
    # MC Ignorable is inherited by URI, not by a prefix that a descendant may
    # have rebound. Carry that context using bindings valid at the imported root.
    bindings = scope(node)

    def bound_prefix(uri):
        existing = next((name.split(":", 1)[1] for name, value in bindings.items()
                         if name.startswith("xmlns:") and value == uri), None)
        if existing is not None:
            return existing
        index = 1
        while f"xmlns:lxe{index}" in bindings:
            index += 1
        prefix = f"lxe{index}"
        bindings[f"xmlns:{prefix}"] = uri
        copied.setAttributeNS(XMLNS_NAMESPACE, f"xmlns:{prefix}", uri)
        return prefix

    ignored = []
    ancestor = node
    while ancestor.nodeType == ancestor.ELEMENT_NODE:
        if ancestor.hasAttributeNS(MC, "Ignorable"):
            context = scope(ancestor)
            ignored.extend(bound_prefix(context[f"xmlns:{prefix}"])
                           for prefix in ancestor.getAttributeNS(MC, "Ignorable").split())
        ancestor = ancestor.parentNode
    if ignored:
        copied.setAttributeNS(MC, f"{bound_prefix(MC)}:Ignorable", " ".join(dict.fromkeys(ignored)))
    return copied


def _serialize_xml(document):
    """Check real QNames after serialization, not just minidom's in-memory URIs."""
    data = document.toxml(encoding="utf-8")

    def qnames(tree):
        return [(node.namespaceURI, node.localName,
                 sorted((attr.namespaceURI or "", attr.localName, attr.value)
                        for attr in node.attributes.values() if attr.namespaceURI != XMLNS_NAMESPACE))
                for node in tree.getElementsByTagName("*")]

    with _parsed(data) as reparsed:
        if qnames(document) != qnames(reparsed):
            raise FullWorkbookError("九表 XML 序列化改变了元素或属性命名空间")
    return data


def _validate_merged_xml(archive):
    """Guard the XML contracts touched by this merger before file publication."""
    styles = ET.fromstring(archive.read("xl/styles.xml"))
    order = ("numFmts", "fonts", "fills", "borders", "cellStyleXfs", "cellXfs",
             "cellStyles", "dxfs", "tableStyles", "colors", "extLst")
    if styles.tag != f"{{{MAIN}}}styleSheet":
        raise FullWorkbookError("九表 styles.xml 根节点命名空间无效")
    positions = [order.index(node.tag.removeprefix(f"{{{MAIN}}}"))
                 for node in styles if node.tag.removeprefix(f"{{{MAIN}}}") in order]
    if positions != sorted(set(positions)):
        raise FullWorkbookError("九表 styles.xml 节点顺序无效")
    # Extensions may have their own namespaces; standard elements cannot lose theirs.
    for part in ("xl/styles.xml", "xl/sharedStrings.xml"):
        if part not in archive.namelist():
            continue
        root = ET.fromstring(archive.read(part))
        pending = [root]
        while pending:
            node = pending.pop()
            if not node.tag.startswith("{"):
                raise FullWorkbookError(f"九表 {part} 包含空命名空间节点: {node.tag}")
            # Extension payloads have their own contracts, outside this targeted guard.
            if node.tag.startswith(f"{{{MAIN}}}") and node.tag != f"{{{MAIN}}}extLst":
                pending.extend(node)
    for part, uri, names in (("[Content_Types].xml", CONTENT, ("Default", "Override")),
                             ("xl/_rels/workbook.xml.rels", PACKAGE_REL, ("Relationship",))):
        if any(node.tag not in {f"{{{uri}}}{name}" for name in names}
               for node in ET.fromstring(archive.read(part))):
            raise FullWorkbookError(f"九表 {part} 包节点命名空间无效")
    formats = styles.find(f"{{{MAIN}}}numFmts")
    ids = [int(node.attrib["numFmtId"]) for node in formats] if formats is not None else []
    if len(ids) != len(set(ids)):
        raise FullWorkbookError("九表自定义数字格式 ID 重复")
    for name in ("cellStyleXfs", "cellXfs"):
        for xf in styles.find(f"{{{MAIN}}}{name}"):
            number = int(xf.attrib.get("numFmtId", 0))
            if number >= 164 and number not in ids:
                raise FullWorkbookError(f"九表 {name} 引用了不存在的 numFmtId={number}")
            for attr, section in (("fontId", "fonts"), ("fillId", "fills"),
                                  ("borderId", "borders"), ("xfId", "cellStyleXfs")):
                if attr in xf.attrib and not 0 <= int(xf.attrib[attr]) < len(styles.find(f"{{{MAIN}}}{section}")):
                    raise FullWorkbookError(f"九表 {name} 的 {attr} 引用无效: {xf.attrib[attr]}")


def _append_styles(template: ZipFile, calculated: ZipFile) -> tuple[bytes, dict[str, int]]:
    """Append new style indices; source style indices used by unchanged tabs stay stable."""
    with _parsed(template.read("xl/styles.xml")) as original, _parsed(calculated.read("xl/styles.xml")) as fresh:
        prefix = original.documentElement.prefix

        offsets = {}
        for name in ("fonts", "fills", "borders", "cellStyleXfs", "cellXfs", "dxfs"):
            old = _section(original, name)
            new = _section(fresh, name)
            if name == "dxfs" and old is None and new is not None:
                old = original.createElementNS(MAIN, f"{prefix}:dxfs" if prefix else "dxfs")
                old.setAttribute("count", "0")
                successor = next((_section(original, name) for name in ("tableStyles", "colors", "extLst")
                                  if _section(original, name) is not None), None)
                original.documentElement.insertBefore(old, successor)
            if old is None or new is None:
                if name == "dxfs" and new is None:
                    offsets[name] = 0
                    continue
                raise FullWorkbookError(f"样式表缺少 {name}")
            offsets[name] = len(_elements(old))

        original_formats = _section(original, "numFmts")
        fresh_formats = _section(fresh, "numFmts")
        used_ids = {int(node.getAttribute("numFmtId")) for node in _elements(original_formats)} if original_formats else set()
        next_id = max(used_ids | {163}) + 1
        number_formats = {}
        if fresh_formats is not None:
            if original_formats is None:
                original_formats = original.createElementNS(MAIN, f"{prefix}:numFmts" if prefix else "numFmts")
                original_formats.setAttribute("count", "0")
                original.documentElement.insertBefore(original_formats, original.documentElement.firstChild)
            for node in _elements(fresh_formats):
                old_id = int(node.getAttribute("numFmtId"))
                if old_id <= 163:
                    number_formats[old_id] = old_id
                    continue
                number_formats[old_id] = next_id
                copied = _import_node(original, node, original_formats)
                copied.setAttribute("numFmtId", str(next_id))
                original_formats.appendChild(copied)
                next_id += 1
            original_formats.setAttribute("count", str(len(_elements(original_formats))))

        for name in ("fonts", "fills", "borders", "cellStyleXfs", "cellXfs", "dxfs"):
            old = _section(original, name)
            new = _section(fresh, name)
            if new is None:
                continue
            for node in _elements(new):
                copied = _import_node(original, node, old)
                for attr, target in (("fontId", "fonts"), ("fillId", "fills"), ("borderId", "borders")):
                    if copied.hasAttribute(attr):
                        copied.setAttribute(attr, str(int(copied.getAttribute(attr)) + offsets[target]))
                if copied.hasAttribute("numFmtId"):
                    number = int(copied.getAttribute("numFmtId"))
                    copied.setAttribute("numFmtId", str(number_formats.get(number, number)))
                if copied.hasAttribute("xfId"):
                    copied.setAttribute("xfId", str(int(copied.getAttribute("xfId")) + offsets["cellStyleXfs"]))
                old.appendChild(copied)
            old.setAttribute("count", str(len(_elements(old))))
        return _serialize_xml(original), offsets


def _append_strings(template: ZipFile, calculated: ZipFile) -> tuple[bytes | None, int, bool]:
    old_bytes = template.read("xl/sharedStrings.xml") if "xl/sharedStrings.xml" in template.namelist() else None
    new_bytes = calculated.read("xl/sharedStrings.xml") if "xl/sharedStrings.xml" in calculated.namelist() else None
    if not new_bytes:
        return old_bytes, 0, False
    if not old_bytes:
        return new_bytes, 0, True
    with _parsed(old_bytes) as original, _parsed(new_bytes) as fresh:
        root = original.documentElement
        offset = len(_elements(root))
        for item in _elements(fresh.documentElement):
            root.appendChild(_import_node(original, item, root))
        root.setAttribute("uniqueCount", str(len(_elements(root))))
        root.setAttribute("count", str(len(_elements(root))))
        return _serialize_xml(original), offset, False


def _string_part_declarations(template: ZipFile) -> dict[str, bytes]:
    """Add the standard package references when an inline-string template needs an SST."""
    with _parsed(template.read("[Content_Types].xml")) as types:
        prefix = types.documentElement.prefix
        item = types.createElementNS(CONTENT, f"{prefix}:Override" if prefix else "Override")
        item.setAttribute("PartName", "/xl/sharedStrings.xml")
        item.setAttribute("ContentType", "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml")
        types.documentElement.appendChild(item)
        content_types = _serialize_xml(types)
    with _parsed(template.read("xl/_rels/workbook.xml.rels")) as relations:
        existing = {node.getAttribute("Id") for node in relations.getElementsByTagNameNS(PACKAGE_REL, "Relationship")}
        identifier = "rIdLxeSharedStrings"
        if identifier in existing:
            raise FullWorkbookError("工作簿关系 ID 与共享字符串表冲突")
        prefix = relations.documentElement.prefix
        item = relations.createElementNS(PACKAGE_REL, f"{prefix}:Relationship" if prefix else "Relationship")
        item.setAttribute("Id", identifier)
        item.setAttribute("Type", REL + "/sharedStrings")
        item.setAttribute("Target", "sharedStrings.xml")
        relations.documentElement.appendChild(item)
        workbook_rels = _serialize_xml(relations)
    return {"[Content_Types].xml": content_types, "xl/_rels/workbook.xml.rels": workbook_rels}


def _current_sheet(data: bytes, style_offset: int, string_offset: int, dxf_offset: int) -> bytes:
    with _parsed(data) as sheet:
        for tag, attr in (("c", "s"), ("row", "s"), ("col", "style")):
            for item in sheet.getElementsByTagNameNS(MAIN, tag):
                if item.hasAttribute(attr):
                    item.setAttribute(attr, str(int(item.getAttribute(attr)) + style_offset))
                if tag == "c" and item.getAttribute("t") == "s":
                    values = item.getElementsByTagNameNS(MAIN, "v")
                    if len(values) != 1 or not values[0].firstChild:
                        raise FullWorkbookError("重算工作表的共享字符串索引无效")
                    values[0].firstChild.data = str(int(values[0].firstChild.data) + string_offset)
        for item in sheet.getElementsByTagNameNS(MAIN, "cfRule"):
            if item.hasAttribute("dxfId"):
                item.setAttribute("dxfId", str(int(item.getAttribute("dxfId")) + dxf_offset))
        return sheet.toxml(encoding="utf-8")


def _workbook_xml(data: bytes, row_count: int) -> bytes:
    with _parsed(data) as workbook:
        for name in workbook.getElementsByTagNameNS(MAIN, "definedName"):
            if name.getAttribute("name") == "_xlnm._FilterDatabase" and name.getAttribute("localSheetId") == "0":
                if name.firstChild:
                    name.firstChild.data = f"越南备货清单!$A$1:$AY${row_count}"
        calc = workbook.getElementsByTagNameNS(MAIN, "calcPr")
        if calc:
            calc[0].setAttribute("calcMode", "auto")
            calc[0].setAttribute("fullCalcOnLoad", "1")
        return workbook.toxml(encoding="utf-8")


def combine_full_workbook(template_path: str | Path, current_path: str | Path, output_path: str | Path) -> None:
    """Keep four original worksheet parts byte-for-byte and replace the five calculated parts."""
    destination = Path(output_path)
    if os.path.lexists(destination):
        raise FileExistsError(destination)
    try:
        with ZipFile(template_path) as template, ZipFile(current_path) as current:
            template_dates, current_dates = _date_system(template), _date_system(current)
            if template_dates != current_dates:
                raise FullWorkbookError(
                    f"九表模板与计算结果的日期系统不一致（模板 {template_dates}，计算结果 {current_dates}），无法安全合并"
                )
            if template_dates != "1900":
                raise FullWorkbookError("九表合并不支持 1904 日期系统，请使用 1900 模板和计算结果")
            source_parts = _sheet_parts(template)
            result_parts = _sheet_parts(current)
            if tuple(source_parts) != REQUIRED_ORDER or tuple(result_parts) != CURRENT:
                raise FullWorkbookError("业务模板须为指定顺序的九表，重算结果须为五表")
            template_sheets = ET.fromstring(template.read("xl/workbook.xml")).find(f"{{{MAIN}}}sheets")
            if any(sheet.attrib.get("state", "visible") != "visible" for sheet in template_sheets):
                raise FullWorkbookError("业务模板的九张工作表都必须可见")
            if any(f"xl/worksheets/_rels/{Path(source_parts[name]).name}.rels" in template.namelist() for name in UNCHANGED):
                raise FullWorkbookError("待保留的四张表包含独立附件关系，当前无法原样保留")
            styles, offsets = _append_styles(template, current)
            strings, string_offset, new_strings = _append_strings(template, current)
            replacements = {"xl/styles.xml": styles}
            if strings is not None:
                replacements["xl/sharedStrings.xml"] = strings
            if new_strings:
                replacements.update(_string_part_declarations(template))
            for name in CURRENT:
                replacements[source_parts[name]] = _current_sheet(
                    current.read(result_parts[name]), offsets["cellXfs"], string_offset, offsets["dxfs"]
                )
            first = ET.fromstring(current.read(result_parts[CURRENT[0]]))
            dimension = first.find(f"{{{MAIN}}}dimension")
            row_count = int((dimension.attrib["ref"] if dimension is not None else "A1").split(":")[-1].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
            replacements["xl/workbook.xml"] = _workbook_xml(template.read("xl/workbook.xml"), row_count)
            destination.parent.mkdir(parents=True, exist_ok=True)
            created = False
            try:
                with destination.open("xb") as handle:
                    created = True
                    with ZipFile(handle, "w") as output:
                        output.comment = template.comment
                        for item in template.infolist():
                            if item.is_dir():
                                continue
                            data = replacements.get(item.filename)
                            output.writestr(copy(item), data if data is not None else template.read(item.filename))
                        if new_strings:
                            output.writestr("xl/sharedStrings.xml", strings)
                with ZipFile(destination) as published:
                    if published.testzip() is not None or tuple(_sheet_parts(published)) != REQUIRED_ORDER:
                        raise FullWorkbookError("九表成品文件校验失败")
                    _validate_merged_xml(published)
                    if any(published.read(source_parts[name]) != template.read(source_parts[name]) for name in UNCHANGED):
                        raise FullWorkbookError("九表成品中的原始业务表发生改变")
                try:
                    workbook = load_workbook(destination, read_only=True, data_only=False)
                    try:
                        for sheet in workbook:
                            for _row in sheet.iter_rows():
                                pass
                    finally:
                        workbook.close()
                except Exception as exc:
                    raise FullWorkbookError(f"九表成品不可读: {type(exc).__name__}: {exc}") from exc
            except BaseException:
                if created:
                    destination.unlink(missing_ok=True)
                raise
    except (KeyError, ET.ParseError, ValueError) as exc:
        if isinstance(exc, FullWorkbookError):
            raise
        raise FullWorkbookError(f"读取或合并九表模板失败: {type(exc).__name__}: {exc}") from exc


__all__ = ["CURRENT", "REQUIRED_ORDER", "UNCHANGED", "FullWorkbookError", "combine_full_workbook", "validate_full_template"]
