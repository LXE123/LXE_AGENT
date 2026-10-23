"""Nine-sheet SKU inputs and the generated Vietnam workbook layout."""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Mapping

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .numeric_contract import WorkbookInputError, excel_number


REQUIRED_SHEETS = (
    "越南备货清单",
    "雅仓库存",
    "雅仓动销",
    "数据更改",
    "库存商品信息",
)

MAIN_HEADERS = {
    "B1": "判断热销",
    "E1": "SKU",
    "G1": "成本",
    "H1": "陆运",
    "AA1": "上架时间",
    "AE1": "跨境价",
    "AJ1": "折扣价",
    "AN1": "总在途",
    "AV1": "30天",
    "AW1": "15天",
    "AX1": "7天",
    "AY1": "汇率",
}

AUXILIARY_HEADERS = {
    "雅仓库存": {"B1": "SKU", "F1": "库存数量", "H1": "在途数量", "J1": "可用库存"},
    "雅仓动销": {"A1": "SKU", "E1": "7天销量", "F1": "15天销量", "G1": "30天销量"},
    "数据更改": {"A1": "30天", "B1": "15天", "C1": "7天", "D1": "汇率"},
    "库存商品信息": {"B1": "SKU", "K1": "创建时间"},
}


class AssetContractError(ValueError):
    """A supplied workbook does not meet the template input contract."""


@dataclass(frozen=True)
class SkuParameters:
    cost: Decimal | None = None
    cross_border_price: Decimal | None = None
    discount_price: Decimal | None = None
    hot_flag: int | None = None


def _blank(value: object) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def validate_sku_parameter_values(values: Mapping[str, SkuParameters], *, allow_empty: bool = False) -> None:
    """Require a hot flag and positive supplied prices; ERP coverage is optional."""
    if not values and not allow_empty:
        raise WorkbookInputError("当前九表模板没有匹配的 SKU 参数")
    for sku, row in values.items():
        if not isinstance(row, SkuParameters):
            raise WorkbookInputError(f"SKU {sku} 的模板 SKU 参数行无效")
        for field, label in (
            ("cost", "成本"),
            ("cross_border_price", "跨境价"),
            ("discount_price", "折扣价"),
        ):
            value = getattr(row, field)
            if value is not None:
                excel_number(value, sku, label, positive=True)
        if isinstance(row.hot_flag, bool) or row.hot_flag not in (1, 2):
            raise WorkbookInputError(f"SKU {sku} 的热销标记必填，且只能是 1 或 2")


def load_template_sku_parameters(path: str | Path, skus: set[str]) -> dict[str, SkuParameters]:
    """Read current SKU inputs from the nine-sheet business template's main tab.

    Formula prices use their saved Excel results. Conflicting duplicate rows for
    a current SKU are rejected instead of choosing one arbitrarily.
    """
    formulas = values = None
    try:
        formulas = load_workbook(path, read_only=True, data_only=False)
        values = load_workbook(path, read_only=True, data_only=True)
        if formulas.sheetnames[0] != "越南备货清单" or values.sheetnames != formulas.sheetnames:
            raise AssetContractError("九表模板缺少首张『越南备货清单』")
        source = formulas.worksheets[0]
        cached = values.worksheets[0]
        headers = next(source.iter_rows(min_row=1, max_row=1, max_col=36))
        expected = {2: "判断热销", 5: "SKU", 7: "成本", 31: "跨境价", 36: "折扣价"}
        for column, label in expected.items():
            if headers[column - 1].value != label:
                raise AssetContractError(f"越南备货清单!{get_column_letter(column)}1 应为 {label!r}")
        result: dict[str, SkuParameters] = {}
        first_rows: dict[str, int] = {}
        for row_number, (row, saved) in enumerate(zip(
            source.iter_rows(min_row=2, max_col=36),
            cached.iter_rows(min_row=2, max_col=36),
        ), 2):
            sku = row[4].value
            if not isinstance(sku, str) or sku.strip() not in skus:
                continue
            sku = sku.strip()
            extracted = []
            for column, label in ((2, "热销标记"), (7, "成本"), (31, "跨境价"), (36, "折扣价")):
                cell = row[column - 1]
                raw = saved[column - 1].value if cell.data_type == "f" else cell.value
                if cell.data_type == "f" and raw is None:
                    raise AssetContractError(f"越南备货清单!{cell.coordinate} 的公式没有保存结果，无法提取 {label}")
                if _blank(raw):
                    extracted.append(None)
                else:
                    try:
                        extracted.append(excel_number(raw, sku, label, positive=label != "热销标记"))
                    except (WorkbookInputError, ArithmeticError) as exc:
                        raise AssetContractError(f"越南备货清单!{cell.coordinate}: {exc}") from exc
            hot, cost, cross_border, discount = extracted
            if hot not in (1, 2):
                raise AssetContractError(f"越南备货清单!B{row_number} 的热销标记必须是 1 或 2")
            parameter = SkuParameters(cost=cost, cross_border_price=cross_border,
                                      discount_price=discount, hot_flag=int(hot))
            if sku in result and result[sku] != parameter:
                raise AssetContractError(
                    f"SKU {sku} 在越南备货清单第 {first_rows[sku]}、{row_number} 行的参数冲突"
                )
            result[sku] = parameter
            first_rows.setdefault(sku, row_number)
        return result
    except AssetContractError:
        raise
    except Exception as exc:
        raise AssetContractError(f"读取九表模板参数失败: {type(exc).__name__}: {exc}") from exc
    finally:
        if formulas is not None:
            formulas.close()
        if values is not None:
            values.close()


__all__ = [
    "AssetContractError",
    "SkuParameters",
    "load_template_sku_parameters",
    "validate_sku_parameter_values",
]
