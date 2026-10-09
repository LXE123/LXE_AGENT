"""Build an operations upload map from current Vietnam SKUs."""

from __future__ import annotations

from decimal import Decimal
from math import isfinite
from pathlib import Path
from typing import TYPE_CHECKING, Mapping

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from services.vietnam_replenishment.asset_contract import SkuParameters

if TYPE_CHECKING:
    from services.vietnam_replenishment.sku_parameters import HistoricalSkuParameters
    from services.vietnam_replenishment.yacang_sources import VietnamSources


_FIELDS = (
    ("cost", "成本"),
    ("cross_border_price", "跨境价"),
    ("discount_price", "折扣价"),
    ("hot_flag", "热销标记"),
    ("listed_at", "上架时间"),
)
_UPLOAD_HEADERS = ("SKU", *(label for _, label in _FIELDS))
_REFERENCE_HEADERS = (
    "SKU",
    *(f"模板{label}" for _, label in _FIELDS),
    "模板问题",
    "缺动销",
    "缺库存",
    "缺商品",
    "缺总在途",
    "在途不一致",
)


def _excel_value(value: object) -> object:
    if not isinstance(value, Decimal):
        return value
    if len(value.as_tuple().digits) <= 15:
        try:
            numeric = float(value)
            if isfinite(numeric) and Decimal(str(numeric)) == value:
                return value
        except (OverflowError, ValueError):
            pass
    return str(value)


def _append_literals(sheet, values: tuple[object, ...]) -> None:
    """Keep literal text and numbers Excel cannot round-trip intact."""
    sheet.append(tuple(_excel_value(value) for value in values))
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _style_sheet(sheet, widths: tuple[int, ...]) -> None:
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{sheet.max_row}"
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _history_issue(history: HistoricalSkuParameters | None) -> str | None:
    if history is None:
        return "模板无此 SKU"
    labels = {name: f"模板{label}" for name, label in _FIELDS}
    issues = [
        f"{labels.get(name, name)}：{reason}"
        for name, reason in sorted(history.unavailable.items())
    ]
    return "；".join(issues) or None


def write_operator_sku_map(
    path: str | Path,
    sources: VietnamSources,
    explicit: Mapping[str, SkuParameters],
    history: Mapping[str, HistoricalSkuParameters],
) -> Path:
    """Write current SKUs for editing, with template and source issues on a separate sheet.

    Only values from ``explicit`` enter the first sheet. Template history stays
    visible for human review and never becomes an explicit upload value.
    """
    output = Path(path)
    skus = tuple(sources.skus)
    if not skus:
        raise ValueError("本轮越南 SKU 集合为空")
    if any(not isinstance(sku, str) or not sku.strip() or sku != sku.strip() for sku in skus):
        raise ValueError("本轮越南 SKU 必须是已规范化的非空文本")
    if len(set(skus)) != len(skus):
        raise ValueError("本轮越南 SKU 存在重复")

    workbook = Workbook()
    try:
        upload = workbook.active
        upload.title = "SKU参数映射"
        reference = workbook.create_sheet("核对参考")
        _append_literals(upload, _UPLOAD_HEADERS)
        _append_literals(reference, _REFERENCE_HEADERS)

        missing_sales = set(sources.missing_sales)
        missing_inventory = set(sources.missing_inventory)
        missing_products = set(sources.missing_products)
        missing_in_transit = set(sources.missing_in_transit)
        in_transit_mismatch = set(sources.in_transit_mismatch)

        for sku in skus:
            values = explicit.get(sku, SkuParameters())
            _append_literals(upload, (sku, *(getattr(values, name) for name, _ in _FIELDS)))

            historical = history.get(sku)
            historical_values = historical.values if historical else SkuParameters()
            _append_literals(
                reference,
                (
                    sku,
                    *(getattr(historical_values, name) for name, _ in _FIELDS),
                    _history_issue(historical),
                    "是" if sku in missing_sales else None,
                    "是" if sku in missing_inventory else None,
                    "是" if sku in missing_products else None,
                    "是" if sku in missing_in_transit else None,
                    "是" if sku in in_transit_mismatch else None,
                ),
            )

        _style_sheet(upload, (24, 16, 16, 16, 16, 22))
        _style_sheet(reference, (24, 16, 16, 16, 16, 22, 58, 12, 12, 12, 14, 16))
        output.parent.mkdir(parents=True, exist_ok=True)
        created = False
        try:
            with output.open("xb") as destination:
                created = True
                workbook.save(destination)
        except BaseException:
            if created:
                output.unlink(missing_ok=True)
            raise
    finally:
        workbook.close()
    return output


__all__ = ["write_operator_sku_map"]
