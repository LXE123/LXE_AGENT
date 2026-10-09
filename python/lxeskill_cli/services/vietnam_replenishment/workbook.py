"""Fill the packaged Vietnam formula workbook from one current VN8806 snapshot."""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from importlib.resources import files
from math import isfinite
from pathlib import Path
import re
from typing import Mapping

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from .asset_contract import REQUIRED_SHEETS, SkuParameters
from .yacang_sources import VietnamSources


class WorkbookInputError(ValueError):
    """Current source data or explicit inputs cannot yield a complete workbook."""


@dataclass(frozen=True)
class RecommendationConfig:
    weight_30d: Decimal = Decimal("0.8")
    weight_15d: Decimal = Decimal("0.8")
    weight_7d: Decimal = Decimal("0")
    exchange_rate: Decimal = Decimal("3900")


@dataclass(frozen=True)
class _CurrentRow:
    sku: str
    sales: Mapping[str, object]
    inventory: Mapping[str, object]
    product: Mapping[str, object]
    parameters: SkuParameters
    listed_at: str
    in_transit: Decimal
    available: Decimal
    sales_7d: Decimal
    sales_15d: Decimal
    sales_30d: Decimal
    cost: Decimal
    cross_border_price: Decimal
    discount_price: Decimal
    hot_flag: int


_SOURCE_TIME = re.compile(r"^[0-9]{4}-[0-9]{1,2}-[0-9]{1,2} [0-9]{1,2}:[0-9]{1,2}$")
_FIXED_PARAMETER_COLUMNS = frozenset(("AV", "AW", "AX", "AY"))
_BLANK_MAIN_COLUMNS = frozenset(("A", "C", "D", "AD", "AO", "AP", "AQ", "AR", "AS", "AT", "AU"))


def canonical_product_time(value: object) -> str:
    """Return the verified Yacang wall time as text for DATEVALUE(LEFT(AA,10))."""
    if not isinstance(value, str) or not _SOURCE_TIME.fullmatch(value.strip()):
        raise WorkbookInputError("仓库产品 创建时间必须为 YYYY-MM-DD HH:MM 文本")
    try:
        return datetime.strptime(value.strip(), "%Y-%m-%d %H:%M").strftime("%Y-%m-%d %H:%M")
    except ValueError as exc:
        raise WorkbookInputError(f"仓库产品 创建时间无效: {exc}") from exc


def _number(value: object, sku: str, field: str, *, positive: bool = False) -> Decimal:
    if value is None or isinstance(value, str) and not value.strip():
        raise WorkbookInputError(f"SKU {sku} 的 {field} 缺失")
    if isinstance(value, bool):
        raise WorkbookInputError(f"SKU {sku} 的 {field} 不是有限非负数：布尔值")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise WorkbookInputError(f"SKU {sku} 的 {field} 不是数字: {type(exc).__name__}: {exc}") from exc
    if not result.is_finite() or result < 0 or positive and result == 0:
        constraint = "有限正数" if positive else "有限非负数"
        raise WorkbookInputError(f"SKU {sku} 的 {field} 必须是{constraint}")
    if len(result.normalize().as_tuple().digits) > 15:
        raise WorkbookInputError(f"SKU {sku} 的 {field} 超出 Excel 精度（15 位有效数字）")
    try:
        as_float = float(result)
    except (OverflowError, ValueError) as exc:
        raise WorkbookInputError(f"SKU {sku} 的 {field} 超出 Excel 数值范围") from exc
    if not isfinite(as_float) or result != 0 and as_float == 0 or Decimal(str(as_float)) != result:
        raise WorkbookInputError(f"SKU {sku} 的 {field} 无法精确写入 Excel 数值单元格")
    return result


def _map_time_agrees(sku: str, listed_at: object, canonical: str) -> None:
    if listed_at is None or isinstance(listed_at, str) and not listed_at.strip():
        return
    if not isinstance(listed_at, str):
        raise WorkbookInputError(f"SKU {sku} 的旧映射上架时间不是文本，无法与商品创建时间核对")
    raw = listed_at.strip()
    product = datetime.strptime(canonical, "%Y-%m-%d %H:%M")
    try:
        if len(raw) == 10:
            agrees = date.fromisoformat(raw) == product.date()
        else:
            mapped = datetime.fromisoformat(raw)
            if mapped.tzinfo is not None:
                raise WorkbookInputError(
                    f"SKU {sku} 的旧映射上架时间含时区，无法与不含时区的商品创建时间核对"
                )
            agrees = mapped == product
    except ValueError as exc:
        raise WorkbookInputError(f"SKU {sku} 的旧映射上架时间无效: {exc}") from exc
    if not agrees:
        raise WorkbookInputError(
            f"SKU {sku} 的旧映射上架时间 {raw!r} 与仓库产品创建时间 {canonical!r} 冲突"
        )


def _validated_rows(
    sources: VietnamSources, parameters: Mapping[str, SkuParameters]
) -> tuple[_CurrentRow, ...]:
    skus = tuple(sources.skus)
    if not skus:
        raise WorkbookInputError("本轮 VN8806 SKU 集合为空")
    if any(not isinstance(sku, str) or not sku.strip() or sku != sku.strip() for sku in skus):
        raise WorkbookInputError("本轮 SKU 必须是已规范化的非空文本")
    if len(set(skus)) != len(skus):
        raise WorkbookInputError("本轮 SKU 集合重复")
    union = set(sources.sales) | set(sources.inventory)
    if set(skus) != union:
        raise WorkbookInputError("本轮 SKU 集合与 VN8806 动销、当前库存的并集不一致")

    result: list[_CurrentRow] = []
    for sku in skus:
        sales = sources.sales.get(sku)
        inventory = sources.inventory.get(sku)
        product = sources.products.get(sku)
        if sales is None:
            raise WorkbookInputError(f"SKU {sku} 缺少 VN8806 库存动销")
        if inventory is None:
            raise WorkbookInputError(f"SKU {sku} 缺少 VN8806 当前库存")
        if product is None:
            raise WorkbookInputError(f"SKU {sku} 缺少仓库产品资料")
        for label, row in (("库存动销", sales), ("当前库存", inventory), ("仓库产品", product)):
            if row.get("SKU") != sku:
                raise WorkbookInputError(f"SKU {sku} 的{label}行 SKU 不匹配")
        for label, row in (("库存动销", sales), ("当前库存", inventory)):
            if row.get("仓库") != "VN8806":
                raise WorkbookInputError(f"SKU {sku} 的{label}仓库不是 VN8806")
        title = product.get("中文标题")
        if not isinstance(title, str) or not title.strip():
            raise WorkbookInputError(f"SKU {sku} 的仓库产品中文标题缺失")
        try:
            listed_at = canonical_product_time(product.get("创建时间"))
        except WorkbookInputError as exc:
            raise WorkbookInputError(f"SKU {sku} 的{exc}") from exc

        values = parameters.get(sku)
        if values is None:
            raise WorkbookInputError(f"SKU {sku} 缺少运营 SKU 映射行")
        _map_time_agrees(sku, values.listed_at, listed_at)

        actual_transit = _number(inventory.get("在途数量"), sku, "当前库存在途数量")
        mapped_transit = _number(sources.in_transit.get(sku), sku, "权威在途数量")
        if actual_transit != mapped_transit:
            raise WorkbookInputError(f"SKU {sku} 的权威在途数量与当前库存的在途数量不一致")
        available = _number(inventory.get("可用库存"), sku, "可用库存")
        sales_7d = _number(sales.get("7天销量"), sku, "7天销量")
        sales_15d = _number(sales.get("15天销量"), sku, "15天销量")
        sales_30d = _number(sales.get("30天销量"), sku, "30天销量")
        cost = _number(values.cost, sku, "成本(cost)")
        cross_border = _number(values.cross_border_price, sku, "跨境价(cross_border_price)")
        discount = _number(values.discount_price, sku, "折扣价(discount_price)")
        if values.hot_flag is None:
            hot_flag = 2
        elif isinstance(values.hot_flag, bool) or values.hot_flag not in (1, 2):
            raise WorkbookInputError(f"SKU {sku} 的热销标记只能是 1 或 2")
        else:
            hot_flag = values.hot_flag
        result.append(_CurrentRow(
            sku=sku, sales=sales, inventory=inventory, product=product,
            parameters=values, listed_at=listed_at, in_transit=actual_transit,
            available=available, sales_7d=sales_7d, sales_15d=sales_15d,
            sales_30d=sales_30d, cost=cost, cross_border_price=cross_border,
            discount_price=discount, hot_flag=hot_flag,
        ))
    return tuple(result)


def _validated_config(config: RecommendationConfig) -> tuple[Decimal, Decimal, Decimal, Decimal]:
    if not isinstance(config, RecommendationConfig):
        raise WorkbookInputError("推荐参数必须是 RecommendationConfig")
    return (
        _number(config.weight_30d, "运行参数", "30天"),
        _number(config.weight_15d, "运行参数", "15天"),
        _number(config.weight_7d, "运行参数", "7天"),
        _number(config.exchange_rate, "运行参数", "汇率", positive=True),
    )


def _load_skeleton():
    resource = files("services.vietnam_replenishment").joinpath("resources", "skeleton.xlsx")
    with resource.open("rb") as source:
        return load_workbook(source, data_only=False)


def _literal(cell, value: object) -> None:
    cell.value = value
    if isinstance(value, str):
        # A real SKU or title may begin with '='; it remains source text.
        cell.data_type = "s"


def _project(sheet, row_number: int, source: Mapping[str, object], replacements: Mapping[str, object]) -> None:
    seen: set[str] = set()
    for column in range(1, sheet.max_column + 1):
        header = sheet.cell(1, column).value
        if not isinstance(header, str) or not header:
            continue
        if header in seen:
            raise WorkbookInputError(f"内置骨架 {sheet.title} 有重复表头 {header!r}")
        seen.add(header)
        _literal(sheet.cell(row_number, column), replacements.get(header, source.get(header)))


def _fill_main(main, row_number: int, current: _CurrentRow, blueprint: tuple) -> None:
    for index, (value, data_type, style) in enumerate(blueprint, 1):
        cell = main.cell(row_number, index)
        cell._style = copy(style)
        column = get_column_letter(index)
        if isinstance(value, ArrayFormula):
            if not value.text:
                raise WorkbookInputError(f"内置骨架 {column}2 的数组公式为空")
            translated = Translator(value.text, origin=f"{column}2").translate_formula(cell.coordinate)
            cell.value = ArrayFormula(cell.coordinate, text=translated)
        elif data_type == "f":
            cell.value = (
                value if column in _FIXED_PARAMETER_COLUMNS
                else Translator(value, origin=f"{column}2").translate_formula(cell.coordinate)
            )
        else:
            cell.value = None

    for column, value in (
        ("B", current.hot_flag), ("E", current.sku),
        ("F", current.product.get("中文标题")), ("G", current.cost),
        ("AE", current.cross_border_price), ("AJ", current.discount_price),
        ("AN", current.in_transit),
    ):
        _literal(main[f"{column}{row_number}"], value)
    for column in _BLANK_MAIN_COLUMNS:
        main[f"{column}{row_number}"].value = None


def write_vietnam_workbook(
    path: str | Path,
    sources: VietnamSources,
    parameters: Mapping[str, SkuParameters],
    config: RecommendationConfig,
) -> None:
    """Write only current SKU literals and translated formulas; never overwrite output."""
    output = Path(path)
    if output.exists():
        raise FileExistsError(output)
    rows = _validated_rows(sources, parameters)
    input_config = _validated_config(config)
    book = _load_skeleton()
    try:
        if tuple(book.sheetnames) != REQUIRED_SHEETS:
            raise WorkbookInputError(f"内置骨架工作表不符: {book.sheetnames!r}")
        main = book["越南备货清单"]
        for sheet in book:
            allowed_rows = 2 if sheet.title in ("越南备货清单", "数据更改") else 1
            if sheet.max_row > allowed_rows:
                raise WorkbookInputError(f"内置骨架 {sheet.title} 含非预期的数据行")
        blueprint = tuple(
            (cell.value, cell.data_type, copy(cell._style))
            for cell in main[2]
        )
        if not blueprint or main["AA2"].data_type != "f" or not isinstance(main["AC2"].value, ArrayFormula):
            raise WorkbookInputError("内置骨架缺少主表上架时间或备货天数公式")
        if any(main[f"{column}2"].data_type != "f" for column in _FIXED_PARAMETER_COLUMNS):
            raise WorkbookInputError("内置骨架缺少四项参数公式")

        change = book["数据更改"]
        for index, value in enumerate(input_config, 1):
            _literal(change.cell(2, index), value)
        for row_number, current in enumerate(rows, 2):
            _project(
                book["雅仓库存"], row_number, current.inventory,
                {"在途数量": current.in_transit, "可用库存": current.available},
            )
            _project(
                book["雅仓动销"], row_number, current.sales,
                {"7天销量": current.sales_7d, "15天销量": current.sales_15d, "30天销量": current.sales_30d},
            )
            _project(
                book["库存商品信息"], row_number, current.product,
                {"创建时间": current.listed_at},
            )
            _fill_main(main, row_number, current, blueprint)

        output.parent.mkdir(parents=True, exist_ok=True)
        created = False
        try:
            with output.open("xb") as destination:
                created = True
                book.save(destination)
        except BaseException:
            if created:
                output.unlink(missing_ok=True)
            raise
    finally:
        book.close()


__all__ = [
    "RecommendationConfig", "WorkbookInputError", "canonical_product_time",
    "write_vietnam_workbook",
]
