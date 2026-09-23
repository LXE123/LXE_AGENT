from __future__ import annotations

from datetime import datetime
from contextlib import ExitStack
import uuid
import os
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from shared.datasets import dataset_dir
from shared.filesystem import display_path, filesystem_path
from .client import YacangClient
from .contracts import Credentials, REPORTS, WAREHOUSES, normalize, tasks_for, task_key, export_parameters
from .errors import YacangError, safe_remote_detail
from .queue import wait_for_file
from .state import ExportState
from .validation import validate_inventory_sales_workbook, validate_inventory_list_workbook, validate_warehouse_products_workbook


def validate(path, task, diagnostic):
    if task['report'] == 'inventory-sales':
        count = validate_inventory_sales_workbook(path, warehouse_code=task['warehouse'], diagnostic=diagnostic)
    elif task['report'] == 'inventory-current-snapshot':
        count = validate_inventory_list_workbook(path, warehouse_code=task['warehouse'], diagnostic=diagnostic)
    else:
        count = validate_warehouse_products_workbook(path, diagnostic=diagnostic)
    with filesystem_path(path).open("rb") as source:
        wb = load_workbook(source, read_only=True)
        try:
            return {'row_count': count, 'sheet_names': wb.sheetnames}
        finally:
            wb.close()


def run(arguments):
    artifacts, tasks, client = [], [], None
    locks = ExitStack()
    try:
        if set(arguments) != {'params'}:
            raise YacangError('参数校验', '必须提供 params，且不接受其他顶层参数', code='invalid_arguments')
        params = normalize(arguments['params'])
        credentials = Credentials.from_environment()
        state = ExportState(credentials.account_id)
        planned = tasks_for(params)
        tasks = [{**t, 'status': 'not_run', 'filters': export_parameters(t)} for t in planned]
        locks.enter_context(state.run_lock())
        client = YacangClient(credentials, state)
        client.login()
        folder = filesystem_path(dataset_dir("yacang_exports", credentials.account_id, uuid.uuid4().hex))
        folder.mkdir(parents=True, exist_ok=False)
        for task in tasks:
            spec = {k: task[k] for k in ('report', 'warehouse', 'created_date')}
            key = task_key(credentials.account_id, spec)
            temporary = None
            try:
                with state.task_lock(key):
                    baseline = state.pending(key)
                    if baseline is None:
                        baseline = {str(r['id']) for r in client.list_downloads()}
                        state.begin(key, baseline)
                        try:
                            client.submit(spec)
                        except YacangError as exc:
                            if exc.code in {'date_range_required', 'remote_error', 'login_required', 'forbidden', 'rate_limited'}:
                                state.clear(key)  # Explicit platform rejection, no accepted export.
                            raise
                    # An interrupted/uncertain submission resumes read-only queue checks.
                    url = wait_for_file(client, spec, baseline)
                    # A ready file proves the remote task completed, even if download later fails.
                    state.clear(key)
                    name = REPORTS[spec['report']][0]
                    warehouse_name = '_' + WAREHOUSES[spec['warehouse']][1] if spec['warehouse'] else ''
                    filename = f'雅仓-{name}{warehouse_name}-{datetime.now(ZoneInfo("Asia/Shanghai")):%Y%m%d-%H%M%S}.xlsx'
                    target = folder / filename
                    temporary = folder / ('.' + filename)
                    client.download(url, temporary)
                    metadata = validate(temporary, spec, client.diagnostic)
                    temporary.replace(target)
                    artifact = {**spec, **metadata, 'filters': export_parameters(spec), 'path': str(display_path(target.resolve())), 'filename': filename,
                                'source': 'yacang', 'notice': '没有数据行' if metadata['row_count'] == 0 else ''}
                    artifacts.append(artifact)
                    task.update(status='completed', artifact_path=artifact['path'], row_count=metadata['row_count'])
            except Exception as exc:
                error = {'code': getattr(exc, 'code', 'export_failed'), 'message': client.diagnostic(f'{type(exc).__name__}: {exc}')}
                if temporary is not None:
                    try:
                        temporary.unlink(missing_ok=True)
                    except OSError as cleanup:
                        error['cleanup_error'] = client.diagnostic(f'{type(cleanup).__name__}: {cleanup}')
                task.update(status='failed', error=error)
                if getattr(exc, 'scope', 'global') == 'global':
                    break
        failed = [t for t in tasks if t['status'] != 'completed']
        result = {'success': not failed, 'status': 'partial_success' if failed and artifacts else 'failed' if failed else 'completed',
                  'params': params, 'tasks': tasks, 'artifacts': artifacts, 'auth_refresh_required': False}
        if failed:
            result['error'] = next(t['error'] for t in tasks if t.get('error'))
        return result
    except Exception as exc:
        message = client.diagnostic(f'{type(exc).__name__}: {exc}') if client else safe_remote_detail(f'{type(exc).__name__}: {exc}', secrets=tuple(os.getenv(k, '') for k in ('LXE_YACANG_MOBILE', 'LXE_YACANG_PASSWORD')))
        return {'success': False, 'status': 'failed', 'tasks': tasks, 'artifacts': artifacts, 'auth_refresh_required': False,
                'error': {'code': getattr(exc, 'code', 'export_failed'), 'message': message}}
    finally:
        if client:
            client.close()
        locks.close()
