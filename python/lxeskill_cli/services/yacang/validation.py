from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from services.yacang.errors import YacangError, safe_remote_detail


INVENTORY_SALES_HEADERS = (
    "SKU", "商品名", "仓库", "3天销量", "7天销量", "15天销量", "30天销量", "60天销量",
    "90天销量", "库存", "占用", "在途", "冻结", "可用", "缺货数量", "创建日期",
)
INVENTORY_LIST_HEADERS = (
    "条码", "SKU", "商品标签", "映射条码", "仓库", "规格", "尺寸(cm)", "重量(g)",
    "库存数量", "占用数量", "在途数量", "冻结库存", "可用库存", "中文标题", "英文标题",
    "图片链接", "退货处理方式", "状态",
)
WAREHOUSE_PRODUCTS_HEADERS = (
    "条码", "SKU", "规格", "中文标题", "英文标题", "长", "宽", "高", "重量", "图片链接", "创建时间",
)


def validate_inventory_sales_workbook(path: Path, *, warehouse_code: str, diagnostic=safe_remote_detail) -> int:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - expose the real parser failure
        raise YacangError("校验 XLSX", f"{type(exc).__name__}: {diagnostic(exc)}") from exc
    try:
        if len(workbook.sheetnames) != 1:
            raise YacangError("校验 XLSX", f"工作表数量应为 1，实际为 {len(workbook.sheetnames)}")
        sheet = workbook[workbook.sheetnames[0]]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in next(rows, ()))
        if headers != INVENTORY_SALES_HEADERS:
            raise YacangError("校验 XLSX", f"表头不匹配: {diagnostic(headers)}")
        row_count = 0
        wrong_warehouses: set[str] = set()
        for row in rows:
            if not any(value is not None for value in row):
                continue
            row_count += 1
            actual = str(row[2] or "").strip() if len(row) > 2 else ""
            if actual != warehouse_code and len(wrong_warehouses) < 5:
                wrong_warehouses.add(actual or "[empty]")
        if wrong_warehouses:
            raise YacangError(
                "校验 XLSX",
                f"期望仓库 {warehouse_code}，文件包含其他仓库: {sorted(wrong_warehouses)}",
            )
        return row_count
    finally:
        workbook.close()


def validate_inventory_list_workbook(path: Path, *, warehouse_code: str, diagnostic=safe_remote_detail) -> int:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - expose the real parser failure
        raise YacangError("校验库存列表 XLSX", f"{type(exc).__name__}: {diagnostic(exc)}") from exc
    try:
        if len(workbook.sheetnames) != 1:
            raise YacangError(
                "校验库存列表 XLSX",
                f"工作表数量应为 1，实际为 {len(workbook.sheetnames)}",
            )
        sheet = workbook[workbook.sheetnames[0]]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in next(rows, ()))
        if headers != INVENTORY_LIST_HEADERS:
            raise YacangError("校验库存列表 XLSX", f"表头不匹配: {diagnostic(headers)}")
        row_count = 0
        wrong_warehouses: set[str] = set()
        for row in rows:
            if not any(value is not None for value in row):
                continue
            row_count += 1
            actual = str(row[4] or "").strip() if len(row) > 4 else ""
            if actual != warehouse_code and len(wrong_warehouses) < 5:
                wrong_warehouses.add(actual or "[empty]")
        if wrong_warehouses:
            raise YacangError(
                "校验库存列表 XLSX",
                f"期望仓库 {warehouse_code}，文件包含其他仓库: {sorted(wrong_warehouses)}",
            )
        return row_count
    finally:
        workbook.close()


def validate_warehouse_products_workbook(path: Path, *, diagnostic=safe_remote_detail) -> int:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - expose the real parser failure
        raise YacangError("校验仓库产品 XLSX", f"{type(exc).__name__}: {diagnostic(exc)}") from exc
    try:
        if len(workbook.sheetnames) != 1:
            raise YacangError(
                "校验仓库产品 XLSX",
                f"工作表数量应为 1，实际为 {len(workbook.sheetnames)}",
            )
        sheet = workbook[workbook.sheetnames[0]]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in next(rows, ()))
        if headers != WAREHOUSE_PRODUCTS_HEADERS:
            raise YacangError("校验仓库产品 XLSX", f"表头不匹配: {diagnostic(headers)}")
        row_count = 0
        invalid_times = 0
        for row in rows:
            if not any(value is not None for value in row):
                continue
            row_count += 1
            raw_time = row[10] if len(row) > 10 else None
            try:
                datetime.strptime(str(raw_time or "").strip(), "%Y-%m-%d %H:%M")
            except ValueError:
                invalid_times += 1
        if invalid_times:
            raise YacangError(
                "校验仓库产品 XLSX",
                f"创建时间必须为 YYYY-MM-DD HH:MM 文本，异常行数: {invalid_times}",
            )
        return row_count
    finally:
        workbook.close()


__all__ = [
    "WAREHOUSE_PRODUCTS_HEADERS",
    "INVENTORY_LIST_HEADERS",
    "INVENTORY_SALES_HEADERS",
    "validate_warehouse_products_workbook",
    "validate_inventory_list_workbook",
    "validate_inventory_sales_workbook",
]
