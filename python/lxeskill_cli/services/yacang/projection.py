from __future__ import annotations

import os
import shutil
from pathlib import Path
from typing import Sequence
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from services.yacang.errors import YacangError, safe_remote_detail
from services.yacang.validation import INVENTORY_SALES_HEADERS, validate_inventory_sales_workbook
from services.yacang.warehouses import select_warehouses


# Compatibility exports for legacy type-level callers. Both now mean the complete
# inventory-sales workbook and must never trim the source columns.
SALES_MONTHLY_HEADERS = INVENTORY_SALES_HEADERS
SALES_90D_HEADERS = INVENTORY_SALES_HEADERS


def validate_projected_workbook(
    path: Path,
    *,
    expected_headers: tuple[str, ...],
    warehouse_code: str,
) -> int:
    if expected_headers != INVENTORY_SALES_HEADERS:
        raise ValueError("库存动销输出必须保留完整原始表头")
    return validate_inventory_sales_workbook(path, warehouse_code=warehouse_code)


def project_inventory_sales_workbook(
    source: Path,
    destination: Path,
    *,
    headers: tuple[str, ...],
    warehouse_code: str,
) -> int:
    """Compatibility wrapper that now publishes the complete source workbook."""
    if headers != INVENTORY_SALES_HEADERS:
        raise ValueError("库存动销输出必须保留完整原始表头")
    return publish_inventory_sales_workbook([(warehouse_code, source)], destination)


def publish_inventory_sales_workbook(
    sources: Sequence[tuple[str, Path]],
    destination: Path,
) -> int:
    """Publish one complete inventory-sales workbook from validated source files.

    A single warehouse is copied byte-for-byte. Multiple warehouses are merged
    into one sheet with one unchanged 16-column header in canonical warehouse
    order.
    """
    if not sources:
        raise ValueError("库存动销输出至少需要一个成功物理源")

    source_by_warehouse: dict[str, Path] = {}
    for warehouse_code, source in sources:
        code = str(warehouse_code or "").strip().upper()
        if code in source_by_warehouse:
            raise ValueError(f"库存动销输出包含重复仓库: {code}")
        source_path = Path(source)
        validate_inventory_sales_workbook(source_path, warehouse_code=code)
        source_by_warehouse[code] = source_path

    ordered = select_warehouses(list(source_by_warehouse))
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(f".{destination.stem}.{uuid4().hex}.part.xlsx")
    try:
        if len(ordered) == 1:
            shutil.copyfile(source_by_warehouse[ordered[0].code], temporary)
            row_count = validate_inventory_sales_workbook(
                temporary,
                warehouse_code=ordered[0].code,
            )
        else:
            row_count = _merge_inventory_sales_workbooks(
                [(warehouse.code, source_by_warehouse[warehouse.code]) for warehouse in ordered],
                temporary,
            )
        os.replace(temporary, destination)
        return row_count
    finally:
        temporary.unlink(missing_ok=True)


def _merge_inventory_sales_workbooks(
    sources: Sequence[tuple[str, Path]],
    destination: Path,
) -> int:
    output_workbook = Workbook(write_only=True)
    row_count = 0
    try:
        output_sheet = output_workbook.create_sheet("库存动销")
        output_sheet.append(INVENTORY_SALES_HEADERS)
        for warehouse_code, source in sources:
            source_workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                source_sheet = source_workbook[source_workbook.sheetnames[0]]
                rows = source_sheet.iter_rows(values_only=True)
                headers = tuple(str(value or "").strip() for value in next(rows, ()))
                if headers != INVENTORY_SALES_HEADERS:
                    raise YacangError(
                        "合并库存动销 XLSX",
                        f"表头不匹配: {safe_remote_detail(headers)}",
                    )
                for row in rows:
                    if not any(value is not None for value in row):
                        continue
                    actual_warehouse = str(row[2] or "").strip() if len(row) > 2 else ""
                    if actual_warehouse != warehouse_code:
                        raise YacangError(
                            "合并库存动销 XLSX",
                            f"期望仓库 {warehouse_code}，实际为 {safe_remote_detail(actual_warehouse)}",
                        )
                    output_sheet.append(tuple(
                        row[index] if index < len(row) else None
                        for index in range(len(INVENTORY_SALES_HEADERS))
                    ))
                    row_count += 1
            finally:
                source_workbook.close()
        output_workbook.save(destination)
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    finally:
        output_workbook.close()
    _validate_merged_workbook(destination, [warehouse for warehouse, _source in sources], row_count)
    return row_count


def _validate_merged_workbook(
    path: Path,
    warehouse_codes: Sequence[str],
    expected_row_count: int,
) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - preserve the real parser failure
        raise YacangError("校验合并库存动销 XLSX", f"{type(exc).__name__}: {safe_remote_detail(exc)}") from exc
    try:
        if workbook.sheetnames != ["库存动销"]:
            raise YacangError("校验合并库存动销 XLSX", "工作簿必须只包含库存动销 Sheet")
        rows = workbook["库存动销"].iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in next(rows, ()))
        if headers != INVENTORY_SALES_HEADERS:
            raise YacangError("校验合并库存动销 XLSX", f"表头不匹配: {safe_remote_detail(headers)}")
        actual_warehouses: list[str] = []
        row_count = 0
        for row in rows:
            if not any(value is not None for value in row):
                continue
            row_count += 1
            warehouse = str(row[2] or "").strip() if len(row) > 2 else ""
            if not actual_warehouses or actual_warehouses[-1] != warehouse:
                actual_warehouses.append(warehouse)
        if row_count != expected_row_count:
            raise YacangError("校验合并库存动销 XLSX", "合并前后数据行数不一致")
        expected_order = [code for code in warehouse_codes if code in actual_warehouses]
        if actual_warehouses != expected_order:
            raise YacangError("校验合并库存动销 XLSX", "仓库顺序不符合固定顺序")
    finally:
        workbook.close()


__all__ = [
    "SALES_90D_HEADERS",
    "SALES_MONTHLY_HEADERS",
    "project_inventory_sales_workbook",
    "publish_inventory_sales_workbook",
    "validate_projected_workbook",
]
