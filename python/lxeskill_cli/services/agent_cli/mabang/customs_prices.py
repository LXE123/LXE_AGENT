"""Prices from the uploaded summary; quantities and product metadata belong to ERP."""
from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


def number(value, location: str, *, positive: bool = False) -> Decimal:
    try:
        result = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise ValueError(f"{location} 不是有效数字: {value}") from None
    if not result.is_finite() or result < 0 or (positive and result == 0):
        raise ValueError(f"{location} 必须是{'正数' if positive else '非负有限数'}: {value}")
    return result


def key(value) -> str:
    return str(value or "").strip().upper()


@dataclass(frozen=True)
class SkuRange:
    row_number: int
    purchase_price: Decimal | None
    skus: frozenset[str]


@dataclass(frozen=True)
class PriceRow:
    row_number: int
    representative: str
    source_kind: str
    purchase_price: Decimal | None
    sale_price: Decimal
    coverage: tuple[SkuRange, ...]

    @property
    def skus(self) -> frozenset[str]:
        return frozenset(sku for item in self.coverage for sku in item.skus)


@dataclass(frozen=True)
class ResolvedPrice:
    sale_price: Decimal
    candidates: tuple[PriceRow, ...]
    matched_rows: tuple[PriceRow, ...]
    method: str


class PriceMatchError(ValueError):
    def __init__(self, message: str, candidates: tuple[PriceRow, ...]):
        super().__init__(message)
        self.candidates = candidates


def candidate_details(rows: tuple[PriceRow, ...]) -> list[dict]:
    return [{
        "summary_row": row.row_number, "representative": row.representative,
        "source_kind": row.source_kind, "purchase_price": row.purchase_price,
        "sale_price": row.sale_price,
        "restock_candidates": [{"row_number": item.row_number,
                                "purchase_price": item.purchase_price,
                                "skus": sorted(item.skus)} for item in row.coverage],
    } for row in rows]


def _rows(sheet, required):
    headers = {str(cell.value or "").strip(): cell.column for cell in sheet[1]}
    missing = set(required) - headers.keys()
    if missing:
        raise ValueError(f"{sheet.title} 缺少列: {', '.join(sorted(missing))}")
    for index in range(2, sheet.max_row + 1):
        values = {name: sheet.cell(index, column).value for name, column in headers.items()}
        if str(values.get("日期") or "").strip() == "合计":
            continue
        if not any(value is not None and str(value).strip() for value in values.values()):
            continue
        sku = key(values.get("库存sku（第一行）"))
        if not sku:
            raise ValueError(f"{sheet.title} 第{index}行缺少库存sku（第一行）")
        kind = "carryover" if str(values.get("日期") or "").strip() == "走库存" else "current_purchase"
        cost_value = values.get("原价")
        cost = number(cost_value, f"{sheet.title} 第{index}行原价") if cost_value not in (None, "") else None
        yield index, sku, kind, cost, values, headers


def read_prices(path: Path) -> list[PriceRow]:
    workbook = load_workbook(path, data_only=True)
    try:
        for name in ("备货单", "汇总表"):
            if name not in workbook.sheetnames:
                raise ValueError(f"{path.name} 缺少 {name} sheet")
        required = ("日期", "库存sku（第一行）")
        coverage = []
        for index, sku, kind, cost, values, _ in _rows(workbook["备货单"], (*required, "库存sku")):
            skus = {sku}
            for raw in re.split(r"[，,;；\r\n]+", str(values.get("库存sku") or "")):
                raw = re.sub(r"\s*×\s*\d+(?:\.\d+)?\s*$", "", raw.strip())
                if raw:
                    skus.add(key(raw))
            coverage.append((sku, kind, SkuRange(index, cost, frozenset(skus))))
        result = []
        for index, sku, kind, cost, values, headers in _rows(workbook["汇总表"], (*required, "售价")):
            cell = workbook["汇总表"].cell(index, headers["售价"]).coordinate
            price = number(values["售价"], f"{path.name} 汇总表!{cell}", positive=True)
            candidates = [item for item in coverage if item[0:2] == (sku, kind)]
            if len(candidates) > 1 and cost is not None:
                candidates = [item for item in candidates if item[2].purchase_price == cost]
            if not candidates:
                raise ValueError(f"{path.name} 汇总表第{index}行未找到备货单 SKU 范围: {sku}, {kind}, 原价={cost}")
            # A representative SKU may describe several inventory lots. Keep all
            # possible ranges; a price is resolved only for an actual ERP source.
            result.append(PriceRow(index, sku, kind, cost, price, tuple(item[2] for item in candidates)))
        if not result:
            raise ValueError(f"{path.name} 汇总表没有售价行")
        return result
    finally:
        workbook.close()


def match_price(rows: list[PriceRow], stock_sku: str, source: dict, *, sp_no: str = "", model: str = "") -> ResolvedPrice:
    candidates = tuple(row for row in rows if key(stock_sku) in row.skus and row.source_kind == source["source_kind"])
    matched = candidates
    method = "same_price"
    if len({row.sale_price for row in matched}) > 1:
        cost = number(source["purchase_price"], f"ERP SKU={stock_sku} 来源价格")
        exact = tuple(row for row in matched if row.purchase_price == cost)
        if exact:
            matched = exact
            method = "purchase_price"
    prices = {row.sale_price for row in matched}
    if len(prices) == 1:
        return ResolvedPrice(prices.pop(), candidates, matched, method)
    context = (f"SP={sp_no}, 型号={model}, SKU={stock_sku}, 来源={source['source_kind']}, "
               f"ERP来源采购价={source.get('purchase_price')}")
    if not candidates:
        raise PriceMatchError(f"{context} 未找到汇总表售价", candidates)
    detail = "; ".join(
        f"汇总表第{row.row_number}行 原价={row.purchase_price}, 售价={row.sale_price}, "
        f"备货单候选行={','.join(str(item.row_number) for item in row.coverage)}"
        for row in candidates
    )
    raise PriceMatchError(f"{context} 售价来源不唯一: {detail}", candidates)
