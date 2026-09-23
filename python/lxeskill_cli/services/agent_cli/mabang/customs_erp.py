"""Two-stage customs workflow; ERP calculates and local files freeze the review."""
from __future__ import annotations

import hashlib
import json
import shutil
from collections import OrderedDict
from dataclasses import asdict
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook

from services.agent_cli.mabang import fill_customs_declaration as template
from services.agent_cli.mabang import download_fba_delivery_csv as delivery
from services.agent_cli.mabang.customs_prices import PriceMatchError, candidate_details, key, match_price, number, read_prices
from services.agent_cli.mabang.erp_http import ErpHttpError, request_json
from services.agent_cli.mabang.shipment_quantity_validation import _read_delivery_rows

API_PATH = "/api/v1/erp/customs-declarations/preview"
API_SCHEMA = "lxe.erp.customs-preview.v1"
PREVIEW_SCHEMA = "lxe.customs.preview.v1"


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_actual_lines(path):
    headers, rows = _read_delivery_rows(path)
    if any(name not in headers for name in ("MSKU", "MSKU发货量")):
        raise ValueError(f"{path.name} 缺少 MSKU 或 MSKU发货量列")
    quantities = {}
    for index, row in enumerate(rows, 2):
        if not any(row.values()):
            continue
        msku = key(row.get("MSKU"))
        if not msku:
            raise ValueError(f"{path.name} 第{index}行 MSKU 为空")
        quantity = number(row.get("MSKU发货量"), f"{path.name} 第{index}行 MSKU发货量")
        if quantity != quantity.to_integral_value():
            raise ValueError(f"{path.name} 第{index}行 MSKU发货量必须为整数")
        quantities[msku] = quantities.get(msku, Decimal(0)) + quantity
    if not quantities:
        raise ValueError(f"{path.name} 没有 MSKU 数据行")
    # Zero lines are deliberately absent in the packing API contract.
    return [{"msku": msku, "actual_quantity": str(quantities[msku])}
            for msku in sorted(quantities) if quantities[msku] > 0]


def _download(sp_no):
    result = delivery.run({"delivery_no": sp_no})
    if not result.get("success"):
        raise ValueError(result.get("exception") or str(result))
    return template.resolve_workspace_input(result["csv_path"])


def _erp(request):
    _, result = request_json("POST", API_PATH, operation="预览报关数量", json_payload=request)
    if result.get("response_schema") != API_SCHEMA or result.get("status") not in ("ready", "incomplete"):
        raise ValueError(f"ERP 报关响应格式无效: schema={result.get('response_schema')}, status={result.get('status')}")
    if not isinstance(result.get("shipments"), list) or not result.get("calculation_fingerprint"):
        raise ValueError("ERP 报关响应缺少完整货件或计算指纹")
    if result.get("inventory_changes_committed") is not False or result.get("snapshot_changes_committed") is not False:
        raise ValueError("ERP 报关响应未声明只读结果")
    return result


def model_summary(erp):
    groups = OrderedDict()
    for shipment in erp["shipments"]:
        for line in shipment["lines"]:
            group = (shipment["sp_no"], line["supplier_name"], line["model"])
            target = groups.setdefault(group, dict(zip(("sp_no", "supplier_name", "model"), group)) | {
                "planned_quantity": Decimal(0), "actual_quantity": Decimal(0),
                "shortage_quantity": Decimal(0), "overage_quantity": Decimal(0)})
            planned = number(line["planned_quantity"], "ERP 计划量")
            actual = number(line["actual_quantity"], "ERP 实发量")
            target["planned_quantity"] += planned
            target["actual_quantity"] += actual
            target["shortage_quantity"] += max(planned - actual, 0)
            target["overage_quantity"] += max(actual - planned, 0)
    return json.loads(_json(list(groups.values())))


def declaration_rows(prices, shipment, *, price_diagnostics=None):
    groups = OrderedDict()
    matches = []
    for line in shipment["lines"]:
        actual = number(line["actual_quantity"], "ERP SKU 实发量")
        sources = line["sources"]
        if sum((number(s["actual_quantity"], "ERP 来源实发量") for s in sources), Decimal(0)) != actual:
            raise ValueError(f"ERP SKU={line['stock_sku']} 来源实发量不守恒")
        for source in sources:
            quantity = number(source["actual_quantity"], "ERP 来源实发量")
            if quantity == 0:
                continue
            source_cost = number(source["purchase_price"], f"ERP SKU={line['stock_sku']} 来源采购价")
            diagnostic = {
                "sp_no": shipment["sp_no"], "supplier_name": line["supplier_name"],
                "model": line["model"], "stock_sku": line["stock_sku"], "product_name": line["product_name"],
                "source_id": source.get("source_id"), "source_kind": source["source_kind"],
                "source_contract_no": source.get("source_contract_no"), "source_sp_no": source.get("source_sp_no"),
                "purchase_price": source_cost, "planned_quantity": source["planned_quantity"],
                "actual_quantity": quantity,
            }
            if price_diagnostics is not None:
                price_diagnostics.append(diagnostic)
            try:
                price = match_price(prices, line["stock_sku"], source,
                                    sp_no=shipment["sp_no"], model=line["model"])
            except PriceMatchError as exc:
                diagnostic.update(status="blocked", candidates=candidate_details(exc.candidates), error=str(exc))
                raise
            diagnostic.update(status="matched", sale_price=price.sale_price, method=price.method,
                              candidates=candidate_details(price.candidates),
                              matched_summary_rows=sorted({item.row_number for item in price.matched_rows}))
            display_row = min(item.row_number for item in price.matched_rows)
            values = {
                "row_number": display_row, "source_name": line["product_name"], "model": line["model"],
                "quantity": Decimal(0), "commodity_name": line["commodity_name"],
                "sale_price": price.sale_price, "unit": line["unit"], "sku": line["stock_sku"],
                "purchase_price": source_cost,
                "source_kind": source["source_kind"],
            }
            classification = template.classify_declaration(template.SourceDeclarationRow(**values, total_price=0))
            group = (line["supplier_name"], line["model"], line["commodity_name"], line["unit"],
                     price.sale_price, classification.hs_code, classification.declaration_element)
            row = groups.setdefault(group, values)
            # These fields describe a rendered row, not a fictitious single lot.
            # Complete lot and Excel provenance stays in price_diagnostics.
            row["row_number"] = min(row["row_number"], display_row)
            if row["sku"] != line["stock_sku"]:
                row["sku"] = ""
            if row["purchase_price"] != values["purchase_price"]:
                row["purchase_price"] = None
            if row["source_kind"] != source["source_kind"]:
                row["source_kind"] = "mixed"
            row["quantity"] += quantity
            matches.append((group, diagnostic))
    rows = []
    output_indices = {}
    for group, row in sorted(groups.items(), key=lambda item: (item[1]["row_number"], item[0])):
        row["total_price"] = template._actual_total(row["quantity"], row["sale_price"])
        rows.append(row)
        output_indices[group] = len(rows)
    for group, diagnostic in matches:
        diagnostic["declaration_row_number"] = output_indices[group]
    return rows


def _freeze(path, directory, label, files):
    path = template.resolve_workspace_input(path)
    target = directory / f"{label}{path.suffix.lower()}"
    shutil.copy2(path, target)
    files.append({"name": target.name, "sha256": _sha(target), "original_path": str(path)})
    return target


def _report(directory, erp, summary, errors, price_diagnostics=()):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "型号缺货预览"
    sheet.append(["SP", "厂家", "型号", "计划量", "实发量", "缺货量", "多发量"])
    fields = ("sp_no", "supplier_name", "model", "planned_quantity", "actual_quantity", "shortage_quantity", "overage_quantity")
    for item in summary:
        sheet.append([item[f] if i < 3 else float(item[f]) for i, f in enumerate(fields)])
    detail = workbook.create_sheet("SKU数量与来源")
    detail.append(["SP", "SKU", "型号", "计划量", "实发量", "来源", "合同", "来源计划量", "来源实发量"])
    excluded = workbook.create_sheet("未纳入商品")
    excluded.append(["SP", "MSKU", "SKU", "范围", "数量"])
    for shipment in erp["shipments"]:
        for line in shipment["lines"]:
            for source in line["sources"]:
                detail.append([shipment["sp_no"], line["stock_sku"], line["model"], float(line["planned_quantity"]),
                               float(line["actual_quantity"]), source["source_kind"], source["source_contract_no"],
                               float(source["planned_quantity"]), float(source["actual_quantity"])])
        for line in shipment["excluded_lines"]:
            excluded.append([shipment["sp_no"], line["msku"], line["stock_sku"], line["tracking_mode"], float(line["actual_quantity"])])
    pricing = workbook.create_sheet("售价匹配")
    pricing.append(["SP", "厂家", "型号", "SKU", "ERP来源ID", "来源", "合同", "来源SP", "来源原价",
                    "来源实发量", "采用售价", "匹配方式", "候选汇总表及备货单行", "采用的汇总表候选行",
                    "本SP报关明细序号", "状态", "实际错误"])
    for item in price_diagnostics:
        candidates = "\n".join(
            f"汇总表{row['summary_row']}: 原价={row['purchase_price']}, 售价={row['sale_price']}; "
            f"备货单候选行={','.join(str(scope['row_number']) for scope in row['restock_candidates'])}"
            for row in item.get("candidates", [])
        )
        pricing.append([item["sp_no"], item["supplier_name"], item["model"], item["stock_sku"],
                        item["source_id"], item["source_kind"], item["source_contract_no"], item["source_sp_no"],
                        float(item["purchase_price"]), float(item["actual_quantity"]),
                        float(item["sale_price"]) if "sale_price" in item else None,
                        {"same_price": "候选售价一致", "purchase_price": "原价区分售价"}.get(item.get("method"), ""),
                        candidates, ",".join(str(row) for row in item.get("matched_summary_rows", [])),
                        item.get("declaration_row_number"), item["status"], item.get("error", "")])
    issues = workbook.create_sheet("问题")
    issues.append(["问题说明"])
    for error in errors:
        issues.append([error])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    path = directory / "quantity_review.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def preview(arguments):
    paths = template._coerce_input_paths(arguments.get("input_xlsx") or [])
    paths = [template.resolve_workspace_input(path) for path in paths]
    numbers = [template.extract_sp_no_from_filename(path) for path in paths]
    if len(numbers) != len(set(numbers)):
        raise ValueError("输入备货单包含重复 SP")
    if len(paths) > 1 and arguments.get("consignment_excel"):
        raise ValueError("多 SP 模式请按 SP 单号准备本地 WMS 装箱资料")
    root = template.resolve_workspace_input(arguments.get("output_dir") or template.DEFAULT_OUTPUT_DIR)
    directory = root / f"preview-{uuid4().hex}"
    directory.mkdir(parents=True)
    files, sources, shipments = [], [], []
    template_path = _freeze(arguments.get("template_xlsx") or "", directory, "template", files)
    for index, (path, sp) in enumerate(zip(paths, numbers)):
        price_path = _freeze(path, directory, f"prices-{index}", files)
        csv = _freeze(_download(sp), directory, f"delivery-{index}", files)
        shipments.append({"sp_no": sp, "request_id": f"customs-{_sha(csv)}", "source_file_name": csv.name,
                          "source_sha256": _sha(csv), "captured_at": datetime.now(timezone.utc).isoformat(),
                          "lines": read_actual_lines(csv)})
        sources.append({"input_name": price_path.name, "input_original": str(path)})
    request = {"shipments": shipments}
    redirected = set()
    while True:
        try:
            erp = _erp(request)
            break
        except ErpHttpError as exc:
            # At most one alias resolution per input; never loop on changing mappings.
            original = key(exc.detail.get("sp_no"))
            alias = key(exc.detail.get("required_packing_sp_no"))
            if exc.code != "packing_sp_alias_required" or original not in numbers or original in redirected or not alias.startswith("SP"):
                raise
            redirected.add(original)
            index = numbers.index(original)
            csv = _freeze(_download(alias), directory, f"delivery-alias-{index}", files)
            shipments[index].update(sp_no=alias, source_sha256=_sha(csv), source_file_name=csv.name,
                                    lines=read_actual_lines(csv), captured_at=datetime.now(timezone.utc).isoformat())
    if len(erp["shipments"]) != len(sources):
        raise ValueError("ERP 返回货件数量与输入不一致")
    if [key(s["packing_sp_no"]) for s in erp["shipments"]] != [s["sp_no"] for s in shipments]:
        raise ValueError("ERP 返回的装箱 SP 与本次发货单不一致")
    countries = {s["country"] for s in erp["shipments"]}
    if len(countries) != 1 or not countries.issubset(template.SUPPORTED_DESTINATION_COUNTRIES):
        raise ValueError(f"ERP 目的国不一致或不支持: {sorted(countries)}")
    errors, bundles, price_diagnostics = [], [], []
    for index, (source, shipment, original) in enumerate(zip(sources, erp["shipments"], paths)):
        errors.extend(str(issue["message"]) for issue in shipment["issues"] if issue["status"] == "incomplete")
        try:
            if template.extract_destination_country_from_filename(original) != shipment["country"]:
                raise ValueError(f"{original.name} 目的国与 ERP 不一致")
            prices = read_prices(directory / source["input_name"])
            rows = declaration_rows(prices, shipment, price_diagnostics=price_diagnostics)
            weight_path = template._resolve_consignment_excel_path(shipment["packing_sp_no"], arguments.get("consignment_excel"))
            weight_path = _freeze(weight_path, directory, f"weight-{index}", files)
            source["weight_name"] = weight_path.name
            weight = template.read_consignment_weight_info(directory / source["weight_name"])
            bundles.append({**source, "sp_no": shipment["sp_no"], "country": shipment["country"],
                            "rows": rows, "weight": asdict(weight)})
        except (ValueError, RuntimeError, FileNotFoundError) as exc:
            errors.append(str(exc))
    row_count = sum(len(b["rows"]) for b in bundles)
    if row_count > template.MAX_CUSTOMS_PRODUCT_ROWS:
        errors.append(f"商品数超过报关资料模板容量: {row_count} > {template.MAX_CUSTOMS_PRODUCT_ROWS}")
    workbook = template._load_workbook(template_path)
    try:
        required_sheets = {template.DECLARATION_ELEMENTS_SHEET, template.CUSTOMS_DECLARATION_SHEET,
                           *(config.sheet_name for config in template.FORMULA_SHEET_CONFIGS)}
        missing = required_sheets - set(workbook.sheetnames)
        if missing:
            raise ValueError(f"报关模板缺少 sheet: {sorted(missing)}")
        template._find_target_header_row(workbook[template.DECLARATION_ELEMENTS_SHEET])
        layout = template._find_customs_detail_layout(workbook[template.CUSTOMS_DECLARATION_SHEET])
        if layout.block_count < row_count:
            raise ValueError(f"报关单明细区容量不足: 需要 {row_count} 行，模板只有 {layout.block_count} 行")
        for config in template.FORMULA_SHEET_CONFIGS:
            template._find_formula_header_row(workbook[config.sheet_name], config)
    except ValueError as exc:
        errors.append(str(exc))
    finally:
        workbook.close()
    summary = model_summary(erp)
    report = _report(directory, erp, summary, errors, price_diagnostics)
    files.append({"name": report.name, "sha256": _sha(report), "original_path": str(report)})
    data = {"schema": PREVIEW_SCHEMA, "request": request, "erp": erp, "files": files,
            "bundles": bundles, "template_name": template_path.name, "model_summary": summary,
            "errors": errors, "can_generate": not errors and row_count > 0,
            "price_diagnostics": price_diagnostics}
    serialized = _json(data)
    digest = hashlib.sha256(serialized.encode()).hexdigest()
    preview_path = directory / f"{digest}.json"
    preview_path.write_text(serialized, encoding="utf-8")
    result = {"success": not errors, "status": "blocked" if errors else "confirmation_required" if row_count else "no_shipment",
            "can_generate": data["can_generate"], "preview_path": str(preview_path),
            "validation_report_xlsx": str(report), "model_summary": summary, "issues": errors,
            "excluded_lines": [{"sp_no": s["sp_no"], **line} for s in erp["shipments"] for line in s["excluded_lines"]],
            "row_count": row_count, "quantity_basis": "erp_preview", "inventory_changes_committed": False}
    if errors:
        result.update(exception="; ".join(errors), error={"code": "customs_preview_blocked", "message": "; ".join(errors)})
    return result


def fill(preview_path):
    path = template.resolve_workspace_input(preview_path)
    serialized = path.read_bytes()
    if hashlib.sha256(serialized).hexdigest() != path.stem:
        raise ValueError("报关预览已损坏或被修改，请重新预览")
    data = json.loads(serialized)
    if data.get("schema") != PREVIEW_SCHEMA or not data.get("can_generate"):
        raise ValueError("此预览不能生成报关资料，请解决问题后重新预览")
    for source in data["files"]:
        name = source["name"]
        if Path(name).name != name or _sha(path.parent / name) != source["sha256"]:
            raise ValueError(f"预览来源文件已改变，请重新预览: {name}")
    fresh = _erp(data["request"])
    if fresh["status"] != "ready" or fresh["calculation_fingerprint"] != data["erp"]["calculation_fingerprint"]:
        raise ValueError("ERP 计算依据已变化，请重新预览并确认")
    bundles = []
    for item in data["bundles"]:
        rows = [template.SourceDeclarationRow(**row) for row in item["rows"]]
        weight = template.ConsignmentWeightInfo(item["weight"]["excel_path"], item["weight"]["box_count"],
                                               Decimal(item["weight"]["total_gross_weight"]))
        if not rows:
            # A zero-shipped SP has no reportable boxes or weight.
            continue
        bundles.append(template.InputDeclarationBundle(path.parent / item["input_name"], item["sp_no"],
            item["country"], rows, weight,
            template.allocate_weights_by_quantity(rows, total_gross_weight=weight.total_gross_weight), []))
    return render(bundles, path.parent / data["template_name"], path.parent, data, str(path))


def render(bundles, template_path, directory, data, preview_path):
    rows = [row for bundle in bundles for row in bundle.source_rows]
    sp_nos = [bundle.sp_no for bundle in bundles]
    box_count = sum(bundle.consignment_weight_info.box_count for bundle in bundles)
    total_amount = template.calculate_total_amount(rows)
    upper = template.amount_to_chinese_upper_rmb(total_amount)
    output = directory / template._build_output_filename(sp_nos)
    temporary = output.with_suffix(".tmp.xlsx")
    shutil.copy2(template_path, temporary)
    workbook = template._load_workbook(temporary)
    try:
        sheet = workbook[template.DECLARATION_ELEMENTS_SHEET]
        header = template._find_target_header_row(sheet)
        template._clear_old_declaration_rows(sheet, header_row=header)
        notice = template._write_declaration_rows(sheet, header_row=header, rows=rows)
        template._write_customs_detail_rows(workbook[template.CUSTOMS_DECLARATION_SHEET], rows=rows,
            destination_country=bundles[0].destination_country,
            weight_allocation=template._combine_weight_allocations(bundles), box_count=box_count)
        formula_sheets = template._fill_formula_sheets(workbook, row_count=len(rows))
        template._write_uppercase_amount(workbook, upper)
        workbook.save(temporary)
    finally:
        workbook.close()
    temporary.replace(output)
    return {"success": True, "output_xlsx": str(output), "preview_path": preview_path,
            "validation_report_xlsx": str(directory / "quantity_review.xlsx"), "model_summary": data["model_summary"],
            "sp_nos": sp_nos, "destination_country": bundles[0].destination_country,
            "quantity_basis": "erp_preview", "row_count": len(rows), "box_count": box_count,
            "total_gross_weight": float(sum(b.consignment_weight_info.total_gross_weight for b in bundles)),
            "total_amount": float(total_amount), "total_amount_upper": upper,
            "formula_sheets": formula_sheets, "unmatched_count": len(notice), "notice": notice,
            "inventory_changes_committed": False, "snapshot_changes_committed": False}


def run(arguments):
    try:
        return preview(arguments)
    except Exception as exc:
        return failure(exc)


def failure(exc):
    return {"success": False, "exception": str(exc),
            "error": {"code": getattr(exc, "code", "customs_preview_failed"), "message": str(exc)}}
