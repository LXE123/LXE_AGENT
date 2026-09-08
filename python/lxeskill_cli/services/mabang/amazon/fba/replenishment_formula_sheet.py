"""Editable, per-export shipping projection. The Python decisions remain snapshots."""
from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any

from .replenishment_template import calculate_weighted_daily_sales

if TYPE_CHECKING:
    from .store_msku_replenishment import ReplenishmentRow

FINAL_SHIPPING_SHEET = "最终备货意见"
SOURCE_VALUE_COLUMNS = (
    "7天销量", "14天销量", "30天销量", "90天销量", "可售", "待入库", "预留",
    "计划入库", "在途", "待调仓", "调仓中",
)
STOCK_COLUMNS = ("可售", "待入库", "预留", "在途", "待调仓", "调仓中")
FINAL_SHIPPING_COLUMNS = (
    "MSKU", "ASIN", "本地SKU", "品名", "7天销量", "14天销量", "30天销量", "90天销量",
    "加权日销", "可售", "待入库", "预留", "计划入库", "在途", "待调仓", "调仓中",
    "FBA总库存（马帮）", "未发出货件总计", "深圳可用库存", "建议运输方式（生成时）",
    "空运补货天数", "海运补货天数", "理论空运需求", "理论海运需求",
    "扣FBA后空运量", "扣FBA后海运量", "扣未关联货件后空运量", "扣未关联货件后海运量",
    "空运建议量", "海运建议量", "单件重量(g)", "预计总重量(kg)", "备注", "深圳库存缺口",
    "输入检查", "试算提示",
)


@dataclass(frozen=True)
class ShippingFormulaInputs:
    values: dict[str, float]
    weights: tuple[float, float, float]
    minimum_sea_quantity: float


def source_number(value: Any, *, context: str) -> float:
    try:
        if value is None or isinstance(value, bool) or str(value).strip() == "":
            raise ValueError("值为空或不是数值")
        result = float(str(value).strip().replace(",", ""))
        if not math.isfinite(result) or result < 0:
            raise ValueError("数值必须有限且大于等于0")
        return result
    except (ValueError, TypeError) as exc:
        raise ValueError(f"备货公式输入无效: {context}, value={value!r}, error={exc}") from exc


def shipping_quantities(row: ReplenishmentRow) -> tuple[int, int, int]:
    """Final urgent-air, air (including companion air), and sea quantities."""
    quantity = row.replenish_quantity or 0
    if row.sheet_name not in {"空运（急发）", "空运", "海运"} or quantity <= 0:
        return 0, 0, 0
    if row.sheet_name == "海运":
        air = row.companion_air_quantity if row.companion_air_quantity is not None else 0
        sea = row.sea_quantity
        if sea is None or air < 0 or sea < 0 or air + sea != quantity:
            raise ValueError(
                f"最终发货数量拆分不一致: MSKU={row.msku}, "
                f"companion_air={air}, sea={sea}, total={quantity}"
            )
        return 0, air, sea
    return (quantity, 0, 0) if row.sheet_name == "空运（急发）" else (0, quantity, 0)


def _initial_values(row: ReplenishmentRow, missing_snapshot: bool) -> dict[str, Any]:
    inputs = row.formula_inputs
    if inputs is None:
        raise ValueError(f"缺少备货公式源数据: MSKU={row.msku}，请重新生成同源销量分析报告")
    values = inputs.values
    # Keep the same evaluation order as the formal calculation and Excel's
    # left-associative formula. sum() can cross a ceil boundary on Python 3.12.
    daily = calculate_weighted_daily_sales(
        sales_7d=values["7天销量"], sales_14d=values["14天销量"], sales_30d=values["30天销量"],
        params={"weighted_sales": dict(zip(("7d_weight", "14d_weight", "30d_weight"), inputs.weights))},
    )
    fba = sum(values[key] for key in STOCK_COLUMNS)
    is_sea = row.sheet_name == "海运"
    air_days = (row.companion_air_days or 0) if is_sea else row.replenish_days
    sea_days = row.sea_days if is_sea else 0
    if air_days is None or sea_days is None:
        raise ValueError(f"缺少试算补货天数: MSKU={row.msku}")
    gross_air = math.ceil(daily * air_days)
    gross_sea = max(0, math.ceil(daily * sea_days) - gross_air) if is_sea else 0
    air_fba = max(0, math.ceil(gross_air - fba))
    sea_fba = max(0, math.ceil(gross_sea - max(0, fba - gross_air)))
    air_net = max(0, math.ceil(air_fba - row.unlinked_quantity))
    sea_net = max(0, math.ceil(sea_fba - max(0, row.unlinked_quantity - air_fba)))
    gate = not is_sea or sea_net >= inputs.minimum_sea_quantity
    air, sea = (air_net, sea_net) if gate else (0, 0)
    if air + sea != row.replenish_quantity:
        raise ValueError(f"试算初始值与正式计算不一致: MSKU={row.msku}, formula={air + sea}, calculated={row.replenish_quantity}")
    weight = (air + sea) * row.weight_grams / 1000 if row.weight_grams is not None and row.weight_grams > 0 else None
    notes = ["可试算"]
    if row.actual_inventory is None:
        notes.append("深圳可用库存缺失或无效")
    elif air + sea > row.actual_inventory:
        notes.append("深圳库存不足")
    if weight is None:
        notes.append("单件重量缺失或无效")
    if missing_snapshot:
        notes.append("未取得同日未关联货件快照，需核实当前试算数量")
    public = {
        "A": row.msku, "B": row.asin, "C": row.local_sku, "D": row.local_sku_name or row.product_name,
        **dict(zip(("E", "F", "G", "H"), (values[k] for k in SOURCE_VALUE_COLUMNS[:4]))),
        "I": daily, **dict(zip(("J", "K", "L", "M", "N", "O", "P"), (values[k] for k in SOURCE_VALUE_COLUMNS[4:]))),
        "Q": fba, "R": row.unlinked_quantity, "S": row.actual_inventory,
        "T": "空运＋海运" if is_sea and air > 0 else row.sheet_name,
        "U": air_days, "V": sea_days, "W": gross_air, "X": gross_sea,
        "Y": air_fba, "Z": sea_fba, "AA": air_net, "AB": sea_net,
        "AC": air, "AD": sea, "AE": row.weight_grams, "AF": weight, "AG": row.remark,
        "AH": max(0, air + sea - row.actual_inventory) if row.actual_inventory is not None else None,
        "AI": "正常", "AJ": "；".join(notes),
    }
    return public


def write_formula_sheet(workbook: Any, rows: list[ReplenishmentRow], *, missing_snapshot: bool) -> dict[str, dict[str, Any]]:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    order = {"空运（急发）": 0, "空运": 1, "海运": 2}
    selected = sorted(
        (r for r in rows if any(shipping_quantities(r))),
        key=lambda r: (order[r.sheet_name], -r.replenish_quantity, r.msku, r.asin),
    )
    sheet = workbook.create_sheet(FINAL_SHIPPING_SHEET, 0)
    sheet.cell(1, 1, "黄色为输入，蓝色为本行公式；权重和最低件数直接写在公式中。海运天数含空运覆盖期；库存、未关联货件先扣空运再扣海运。仅本表试算，运输方式及其他报表保持生成时结果。")
    sheet.append(FINAL_SHIPPING_COLUMNS)
    caches: dict[str, dict[str, Any]] = {FINAL_SHIPPING_SHEET: {}}
    end = len(selected) + 2
    for r, row in enumerate(selected, 3):
        public = _initial_values(row, missing_snapshot)
        inputs = row.formula_inputs
        assert inputs is not None  # Validated by _initial_values.
        for column, value in public.items():
            cell = sheet[f"{column}{r}"]
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"
        is_sea = row.sheet_name == "海运"
        weights = tuple(repr(float(weight)) for weight in inputs.weights)
        minimum = repr(float(inputs.minimum_sea_quantity))
        valid = f'AI{r}="正常"'
        # Each stage is independently visible; no cross-sheet lookup or hidden row IDs.
        formulas = {
            "I": f'E{r}/7*{weights[0]}+F{r}/14*{weights[1]}+G{r}/30*{weights[2]}',
            "Q": f'SUM(J{r}:L{r},N{r}:P{r})',
            "AI": f'IFERROR(IF(AND(COUNT(E{r}:R{r},U{r}:V{r})=16,MIN(E{r}:R{r},U{r}:V{r})>=0,U{r}=INT(U{r}),V{r}=INT(V{r}),{"V" if is_sea else "U"}{r}>0),"正常","输入无效"),"输入无效")',
            "W": f'IF({valid},ROUNDUP(I{r}*U{r},0),"")',
            "X": f'IF({valid},MAX(0,ROUNDUP(I{r}*V{r},0)-W{r}),"")' if is_sea else f'IF({valid},0,"")',
            "Y": f'IF({valid},MAX(0,ROUNDUP(W{r}-Q{r},0)),"")',
            "Z": f'IF({valid},MAX(0,ROUNDUP(X{r}-MAX(0,Q{r}-W{r}),0)),"")',
            "AA": f'IF({valid},MAX(0,ROUNDUP(Y{r}-R{r},0)),"")',
            "AB": f'IF({valid},MAX(0,ROUNDUP(Z{r}-MAX(0,R{r}-Y{r}),0)),"")',
            "AC": f'IF({valid},IF(AB{r}>={minimum},AA{r},0),"")' if is_sea else f'IF({valid},AA{r},"")',
            "AD": f'IF({valid},IF(AB{r}>={minimum},AB{r},0),"")' if is_sea else f'IF({valid},0,"")',
            "AF": f'IFERROR(IF(AND({valid},ISNUMBER(AE{r}),AE{r}>0),(AC{r}+AD{r})*AE{r}/1000,""),"")',
            "AH": f'IFERROR(IF(AND({valid},COUNT(S{r},AC{r}:AD{r})=3),MAX(0,AC{r}+AD{r}-S{r}),""),"")',
        }
        sea_note = f'&IF(AND(AC{r}+AD{r}=0,AA{r}+AB{r}>0),"；海运最低件数条件未满足","")' if is_sea else ""
        snapshot_note = '&"；未取得同日未关联货件快照，需核实当前试算数量"' if missing_snapshot else ""
        formulas["AJ"] = (
            f'IF({valid},"可试算"'
            f'{sea_note}&IF(AC{r}+AD{r}=0,"；当前建议量为0","")'
            f'&IF(AH{r}="","；深圳可用库存缺失或无效",IF(AH{r}>0,"；深圳库存不足",""))'
            f'&IF(AF{r}="","；单件重量缺失或无效",""),'
            f'"输入无效：销量、FBA库存、未关联货件须为非负数，补货天数须为有效整数"){snapshot_note}'
        )
        for column, formula in formulas.items():
            address = f"{column}{r}"
            sheet[address] = "=" + formula
            caches[FINAL_SHIPPING_SHEET][address] = public[column]

    sheet.freeze_panes = "E3"
    sheet.sheet_view.showGridLines = False
    # A normal filter avoids Excel Table calculated-column autofill replacing
    # row-specific literal weights/thresholds when a user edits one formula.
    sheet.auto_filter.ref = f"A2:AJ{end}"
    if selected:
        numeric = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0, allow_blank=False)
        numeric.showErrorMessage = True
        numeric.error = "请输入大于等于0的数值；不要清空必需输入。"
        sheet.add_data_validation(numeric)
        for column in ("E", "F", "G", "H", "J", "K", "L", "M", "N", "O", "P", "R", "AE"):
            numeric.add(f"{column}3:{column}{end}")
        days = DataValidation(type="whole", operator="greaterThanOrEqual", formula1=0, allow_blank=False)
        days.showErrorMessage = True
        days.error = "请输入大于等于0的整数；0表示不适用。"
        sheet.add_data_validation(days)
        days.add(f"U3:V{end}")
    input_columns = {5, 6, 7, 8, 10, 11, 12, 13, 14, 15, 16, 18, 19, 21, 22, 31, 33}
    formula_columns = {9, 17, *range(23, 31), 32, 34, 35, 36}
    for index in range(1, len(FINAL_SHIPPING_COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 15
        header = sheet.cell(2, index)
        header.fill = PatternFill("solid", fgColor="243746")
        header.font = Font(bold=True, color="FFFFFF")
        header.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 15
    sheet.row_dimensions[2].height = 15
    for cells in sheet.iter_rows(min_row=3, max_row=end, max_col=len(FINAL_SHIPPING_COLUMNS)):
        sheet.row_dimensions[cells[0].row].height = 15
        for index, cell in enumerate(cells, 1):
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if index in input_columns | formula_columns:
                cell.fill = PatternFill("solid", fgColor="FFF2CC" if index in input_columns else "EAF2F8")
            if index in (9, 31, 32):
                cell.number_format = "#,##0.00"
            elif 5 <= index <= 30 and index != 20 or index == 34:
                cell.number_format = "#,##0.##"
    workbook.active = 0
    return caches


def cache_formula_values(path: Path, caches: dict[str, dict[str, Any]]) -> None:
    """Insert initial OOXML cached values after openpyxl's final metadata save."""
    import os
    import posixpath
    from .report_staging import staged_report_path
    from xml.etree import ElementTree as ET
    from zipfile import ZipFile

    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    relns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    # Preserve the default worksheet namespace used by the original Excel package.
    ET.register_namespace("", ns)
    ET.register_namespace("r", relns)
    with staged_report_path(path) as temporary_path:
        with ZipFile(path) as source:
            relationships = {r.attrib["Id"]: r.attrib["Target"] for r in ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))}
            sheets = ET.fromstring(source.read("xl/workbook.xml")).find(f"{{{ns}}}sheets")
            replacements = {}
            for sheet in sheets:
                values = caches.get(sheet.attrib["name"])
                if not values:
                    continue
                target = relationships[sheet.attrib[f"{{{relns}}}id"]]
                filename = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
                tree = ET.fromstring(source.read(filename))
                found = set()
                for cell in tree.iter(f"{{{ns}}}c"):
                    address = cell.attrib["r"]
                    if address not in values:
                        continue
                    if cell.find(f"{{{ns}}}f") is None:
                        raise ValueError(f"公式缓存目标不是公式: {sheet.attrib['name']}!{address}")
                    value = values[address]
                    if value is None or isinstance(value, str):
                        cell.attrib["t"] = "str"
                    else:
                        cell.attrib.pop("t", None)
                    cached = cell.find(f"{{{ns}}}v")
                    if cached is None:
                        cached = ET.SubElement(cell, f"{{{ns}}}v")
                    cached.text = "" if value is None else str(value)
                    found.add(address)
                if found != set(values):
                    raise ValueError(f"公式缓存目标缺失: {set(values) - found}")
                replacements[filename] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            with ZipFile(temporary_path, "w") as output:
                for info in source.infolist():
                    data = replacements[info.filename] if info.filename in replacements else source.read(info.filename)
                    output.writestr(info, data)
        # Windows rejects replacement while the source ZIP is still open.
        os.replace(temporary_path, path)
