"""Per-export full-source product catalog verification; no shared cache or implicit legacy fallback."""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from services.mabang.export_common import clean_text
from services.mabang.official_api import invalid

from .combo_sku import ComboSku, normalize_sku_key
from .sku_catalog import (
    NOT_FOUND_REASON, VERIFICATION_METHOD, LocalSkuDefinition, SkuCatalogSnapshot,
    combo_map, definition_map,
)
from .store_resolver import ID_TYPE_FBA_WAREHOUSE, ID_TYPE_SHOP

SHEET = "源数据核验信息"
VERSION = 4
TAG_COLUMNS = ("绑定核验结果", "是否通过绑定核验", "未通过原因")


class SourceVerificationError(ValueError):
    pass


def _digest(value: Any) -> str:
    def canonical(item: Any) -> Any:
        if isinstance(item, dict):
            return {key: canonical(val) for key, val in item.items()}
        if isinstance(item, (tuple, list)):
            return [canonical(val) for val in item]
        if isinstance(item, (float, int, Decimal)) and not isinstance(item, bool):
            return str(Decimal(str(item)).normalize())
        if isinstance(item, datetime):
            return item.isoformat()
        return item
    return hashlib.sha256(json.dumps(canonical(value), ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()



def product_key(record: Any) -> tuple[str, str, str, str]:
    """Stable full product identity shared by source and inventory reports."""
    def value(column: str, attribute: str) -> Any:
        return record.get(column) if isinstance(record, dict) else getattr(record, attribute)
    return (normalize_sku_key(value("MSKU", "msku")), clean_text(value("父ASIN", "parent_asin")) or "未填写父ASIN",
            clean_text(value("ASIN", "asin")), normalize_sku_key(value("本地SKU", "local_sku")))


def unverified_reasons(metadata: dict[str, Any]) -> dict[tuple[str, str, str, str], str]:
    return {tuple(item["key"]): item["reason"] for item in metadata["unverified_rows"]}


def read_report_metadata(path: str | Path) -> dict[str, Any]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return _read_info(workbook)[0]
    finally:
        workbook.close()


def classify(records: list[dict[str, Any]], skus: tuple[LocalSkuDefinition, ...]) -> list[tuple[str, bool, str]]:
    definitions = definition_map(skus)
    expected = {normalize_sku_key(record.get("本地SKU")) for record in records} - {""}
    if set(definitions) != expected:
        invalid("源数据核验", "商品查询快照未完整覆盖源表本地 SKU", {
            "missing": sorted(expected - definitions.keys()), "unexpected": sorted(definitions.keys() - expected),
        })
    tags = []
    for record in records:
        local = normalize_sku_key(record.get("本地SKU"))
        if not local:
            tags.append(("缺少本地SKU", False, "源表无本地SKU，不参与备货计算"))
        elif definitions[local].stock_type is None:
            tags.append(("未匹配", False, NOT_FOUND_REASON))
        else:
            tags.append(("已核验", True, ""))
    return tags


def read_source_local_skus(path: Path) -> list[str]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        headers, records, _ = _table(workbook)
        if "本地SKU" not in headers:
            raise SourceVerificationError("MSKU 源表缺少本地SKU列")
        return [clean_text(record.get("本地SKU")) for record in records]
    finally:
        workbook.close()


def _table(workbook: Any) -> tuple[list[str], list[dict[str, Any]], list[int]]:
    sheet = workbook.worksheets[0]
    values = sheet.iter_rows(values_only=True)
    headers = [clean_text(value) for value in next(values, ())]
    if not headers or len(headers) != len(set(headers)):
        raise SourceVerificationError("MSKU 源表表头为空或重复")
    records, indexes = [], []
    for index, values in enumerate(values, 2):
        record = dict(zip(headers, values, strict=False))
        if any(clean_text(value) for value in record.values()):
            records.append(record)
            indexes.append(index)
    return headers, records, indexes


def _write_info(workbook: Any, metadata: dict[str, Any], skus: tuple[LocalSkuDefinition, ...] = ()) -> None:
    if SHEET in workbook.sheetnames:
        del workbook[SHEET]
    sheet = workbook.create_sheet(SHEET)
    sheet.append(("kind", "json"))
    sheet.append(("metadata", json.dumps(metadata, ensure_ascii=False)))
    for sku in skus:
        sheet.append(("sku", json.dumps(sku.to_record(), ensure_ascii=False)))
    sheet.sheet_state = "hidden"


def _read_info(workbook: Any) -> tuple[dict[str, Any], tuple[LocalSkuDefinition, ...]]:
    if SHEET not in workbook.sheetnames:
        raise SourceVerificationError("文件缺少源数据核验信息；请重新下载 MSKU 源表并重跑销量和库存报表（仅支持单店单站点）")
    values = list(workbook[SHEET].iter_rows(values_only=True))
    try:
        if tuple(values[0]) != ("kind", "json") or values[1][0] != "metadata":
            raise ValueError("核验信息格式错误")
        metadata = json.loads(values[1][1])
        skus = tuple(LocalSkuDefinition.from_record(json.loads(row[1])) for row in values[2:] if row[0] == "sku")
        definition_map(skus)
        if any(row[0] != "sku" for row in values[2:]):
            raise ValueError("旧版或未知核验记录，请重新下载 MSKU 并重跑销量和库存报表")
        if metadata["scope"] != "xlsx_all":
            raise ValueError("源数据范围不是全量 XLSX")
        excluded = metadata["unverified_rows"]
        if not isinstance(excluded, list) or any(not isinstance(item, dict) or not isinstance(item.get("key"), list) or len(item["key"]) != 4 or any(not isinstance(value, str) for value in item["key"]) or not isinstance(item.get("reason"), str) or not item["reason"] for item in excluded):
            raise ValueError("未核验记录格式无效")
        if type(metadata["version"]) is not int or metadata["version"] not in (3, VERSION):
            raise ValueError(f"核验版本无效: {metadata['version']!r}，请重新下载 MSKU 并重跑销量和库存报表")
        if metadata.get("verification_method") != VERIFICATION_METHOD:
            raise ValueError("核验方式不是库存/组合商品接口，请重新下载")
        for key in ("source_fingerprint", "snapshot_id"):
            if not isinstance(metadata[key], str) or not re.fullmatch(r"[a-f0-9]{64}", metadata[key]):
                raise ValueError(f"核验指纹无效: {key}={metadata[key]!r}")
        # Version 3 records official identity. Version 4 records the web download
        # ID together with its type, without treating warehouse IDs as official sid.
        identity_fields = ("shop_id", "site") if metadata["version"] == 3 else ("store_id", "id_type")
        for key in ("store_name", "requested_store_name", "collected_at", *identity_fields):
            if not isinstance(metadata[key], str) or not metadata[key].strip():
                raise ValueError(f"核验字段无效: {key}={metadata[key]!r}")
        if metadata["version"] == VERSION and metadata["id_type"] not in (ID_TYPE_SHOP, ID_TYPE_FBA_WAREHOUSE):
            raise ValueError(f"下载店铺 ID 类型无效: {metadata['id_type']!r}")
        counts = [metadata[key] for key in ("original_row_count", "binding_verified_row_count", "binding_unverified_row_count")]
        if any(type(value) is not int or value < 0 for value in counts) or counts[0] != counts[1] + counts[2]:
            raise ValueError(f"核验数量无效: {counts!r}")
        if len(excluded) != counts[2]:
            raise ValueError("未核验记录数量不一致")
        return metadata, skus
    except (ValueError, TypeError, KeyError, IndexError) as exc:
        raise SourceVerificationError(f"源数据核验信息无效: {exc}") from exc


def annotate_source(path: Path, snapshot: SkuCatalogSnapshot, *, requested_store_name: str) -> dict[str, Any]:
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    try:
        headers, records, indexes = _table(workbook)
        if (SHEET in workbook.sheetnames or "Active核验信息" in workbook.sheetnames
                or any(column in headers for column in (*TAG_COLUMNS, "在售核验结果", "是否参与计算"))):
            raise SourceVerificationError("下载原表包含保留的源数据核验字段，拒绝覆盖")
        # The selected download shop supplies identity. XLSX store/site labels
        # are preserved for display, including regional and shared-store labels.
        tags = classify(records, snapshot.skus)
        metadata = {
            "version": VERSION, "scope": "xlsx_all", "verification_method": VERIFICATION_METHOD, "store_name": snapshot.store_name,
            "unverified_rows": [dict(key=list(product_key(record)), reason=tag[2]) for record, tag in zip(records, tags, strict=True) if not tag[1]],
            "requested_store_name": requested_store_name, "store_id": snapshot.store_id,
            "id_type": snapshot.id_type, "collected_at": datetime.now(timezone.utc).isoformat(),
            "source_fingerprint": _digest([headers, records]),
            "original_row_count": len(records), "binding_verified_row_count": sum(tag[1] for tag in tags),
            "binding_unverified_row_count": sum(not tag[1] for tag in tags),
        }
        metadata["snapshot_id"] = _digest([metadata, [item.to_record() for item in snapshot.skus]])
        sheet = workbook.worksheets[0]
        for offset, column in enumerate(TAG_COLUMNS, len(headers) + 1):
            sheet.cell(1, offset, column)
        for index, tag in zip(indexes, tags, strict=True):
            for offset, value in enumerate(tag, len(headers) + 1):
                sheet.cell(index, offset, value)
        _write_info(workbook, metadata, snapshot.skus)
        workbook.save(path)
        return metadata
    finally:
        workbook.close()


@dataclass(frozen=True)
class VerifiedSource:
    headers: list[str]
    records: list[dict[str, Any]]
    skus: tuple[LocalSkuDefinition, ...]
    metadata: dict[str, Any]

    @property
    def combo_map(self) -> dict[str, ComboSku]:
        return combo_map(self.skus)

    @property
    def counts(self) -> dict[str, int]:
        return {key: self.metadata[key] for key in ("original_row_count", "binding_verified_row_count", "binding_unverified_row_count")}


def load_verified_source(path: str | Path, *, store_name: str) -> VerifiedSource:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        metadata, skus = _read_info(workbook)
        if store_name not in (metadata["store_name"], metadata["requested_store_name"]):
            raise SourceVerificationError(f"源数据核验店铺不一致: 请求={store_name}, 快照={metadata['store_name']}")
        headers, records, _ = _table(workbook)
        if not all(column in headers for column in TAG_COLUMNS):
            raise SourceVerificationError("源表缺少 绑定标记列，请重新下载")
        raw_headers = [column for column in headers if column not in TAG_COLUMNS]
        raw_records = [{column: record.get(column) for column in raw_headers} for record in records]
        unsigned = {key: value for key, value in metadata.items() if key != "snapshot_id"}
        if (_digest([raw_headers, raw_records]) != metadata["source_fingerprint"]
                or _digest([unsigned, [item.to_record() for item in skus]]) != metadata["snapshot_id"]):
            raise SourceVerificationError("源表或 源数据核验快照发生变化，请重新下载")
        tags = classify(raw_records, skus)
        for record, expected in zip(records, tags, strict=True):
            actual = (record[TAG_COLUMNS[0]], record[TAG_COLUMNS[1]], clean_text(record[TAG_COLUMNS[2]]))
            if not isinstance(actual[1], bool) or actual != expected:
                raise SourceVerificationError("源表 绑定标记与核验快照不一致")
        verified_count = sum(tag[1] for tag in tags)
        if (len(records), verified_count, len(records) - verified_count) != tuple(metadata[key] for key in ("original_row_count", "binding_verified_row_count", "binding_unverified_row_count")):
            raise SourceVerificationError("源数据核验数量不一致")
        return VerifiedSource(headers, records, skus, metadata)
    finally:
        workbook.close()


def stamp_report(path: str | Path, metadata: dict[str, Any]) -> None:
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    try:
        _write_info(workbook, metadata)
        workbook.save(path)
    finally:
        workbook.close()


def require_matching_reports(sales: Path, inventory: Path, *, store_name: str) -> dict[str, Any]:
    from openpyxl import load_workbook
    items = []
    for path in (sales, inventory):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            metadata, _ = _read_info(workbook)
            if store_name not in (metadata["store_name"], metadata["requested_store_name"]):
                raise SourceVerificationError(f"报表源数据核验店铺不一致: {path}")
            items.append(metadata)
        finally:
            workbook.close()
    if items[0] != items[1]:
        raise SourceVerificationError("销量与库存报表的源数据快照或源数据指纹不一致，请使用同一源表重跑")
    if not items[0]["binding_verified_row_count"]:
        raise SourceVerificationError("无可计算记录：没有通过本地 SKU 绑定核验的 MSKU")
    return items[0]
