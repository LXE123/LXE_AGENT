"""Per-export Active verification; no shared cache or implicit legacy fallback."""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from typing import Any

from services.mabang.export_common import clean_text
from services.mabang.official_api import invalid

from .combo_sku import ActiveListingSnapshot, ListingSkuBinding, normalize_sku_key, select_combo_skus
from .store_sites import SITE_TO_MARKETPLACE

SHEET = "Active核验信息"
VERSION = 1
TAG_COLUMNS = ("在售核验结果", "是否参与计算", "排除原因")


class ActiveSourceError(ValueError):
    pass


def _digest(value: Any) -> str:
    def canonical(item: Any) -> Any:
        if isinstance(item, dict):
            return {key: canonical(val) for key, val in item.items()}
        if isinstance(item, (tuple, list)):
            return [canonical(val) for val in item]
        if isinstance(item, (float, int, Decimal)) and not isinstance(item, bool):
            return str(Decimal(str(item)).normalize())
        if isinstance(item, datetime):
            return item.isoformat()
        return item
    return hashlib.sha256(json.dumps(canonical(value), ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()


def classify(records: list[dict[str, Any]], bindings: tuple[ListingSkuBinding, ...]) -> list[tuple[str, bool, str]]:
    by_msku: dict[str, list[ListingSkuBinding]] = {}
    for binding in bindings:
        by_msku.setdefault(normalize_sku_key(binding.msku), []).append(binding)
    tags = []
    active_rows = []
    for record in records:
        msku = normalize_sku_key(record.get("MSKU"))
        asin = clean_text(record.get("ASIN"))
        local = normalize_sku_key(record.get("本地SKU"))
        candidates = by_msku.get(msku, [])
        if not candidates:
            tags.append(("未确认在售", False, "未匹配本店本站点的 Active Listing"))
            continue
        context = f"Active 核验 MSKU={msku} ASIN={asin} 本地SKU={local}"
        matched = [item for item in candidates if not asin or clean_text(item.asin) == asin]
        definitions = {(clean_text(item.asin), normalize_sku_key(item.local_sku), item.stock_type) for item in matched}
        if len(definitions) != 1:
            invalid(context, "ASIN 不一致或 Active 绑定存在歧义", [asdict(item) for item in candidates])
        _, binding_local, stock_type = next(iter(definitions))
        if stock_type not in (1, 2) or (local and local != binding_local):
            invalid(context, "本地 SKU 绑定变化或类型无效", [asdict(item) for item in matched])
        active_rows.append(SimpleNamespace(msku=msku, asin=asin, local_sku=local))
        tags.append(("Active", True, ""))
    # Includes conflicts where different MSKUs claim different types for one local SKU.
    select_combo_skus(active_rows, bindings)
    return tags


def _table(workbook: Any) -> tuple[list[str], list[dict[str, Any]], list[int]]:
    sheet = workbook.worksheets[0]
    values = sheet.iter_rows(values_only=True)
    headers = [clean_text(value) for value in next(values, ())]
    if not headers or len(headers) != len(set(headers)):
        raise ActiveSourceError("MSKU 源表表头为空或重复")
    records, indexes = [], []
    for index, values in enumerate(values, 2):
        record = dict(zip(headers, values, strict=False))
        if any(clean_text(value) for value in record.values()):
            records.append(record)
            indexes.append(index)
    return headers, records, indexes


def _write_info(workbook: Any, metadata: dict[str, Any], bindings: tuple[ListingSkuBinding, ...] = ()) -> None:
    if SHEET in workbook.sheetnames:
        del workbook[SHEET]
    sheet = workbook.create_sheet(SHEET)
    sheet.append(("kind", "json"))
    sheet.append(("metadata", json.dumps(metadata, ensure_ascii=False)))
    for binding in bindings:
        sheet.append(("binding", json.dumps(asdict(binding), ensure_ascii=False)))
    sheet.sheet_state = "hidden"


def _read_info(workbook: Any) -> tuple[dict[str, Any], tuple[ListingSkuBinding, ...]]:
    if SHEET not in workbook.sheetnames:
        raise ActiveSourceError("文件缺少 Active 核验信息；请重新下载 MSKU 源表并重跑销量和库存报表（欧洲整组暂不支持 Active 计算）")
    values = list(workbook[SHEET].iter_rows(values_only=True))
    try:
        if tuple(values[0]) != ("kind", "json") or values[1][0] != "metadata":
            raise ValueError("核验信息格式错误")
        metadata = json.loads(values[1][1])
        bindings = tuple(ListingSkuBinding(**json.loads(row[1])) for row in values[2:] if row[0] == "binding")
        if any(row[0] != "binding" for row in values[2:]):
            raise ValueError("未知核验记录")
        if type(metadata["version"]) is not int or metadata["version"] != VERSION:
            raise ValueError(f"核验版本无效: {metadata['version']!r}")
        for key in ("source_fingerprint", "snapshot_id"):
            if not isinstance(metadata[key], str) or not re.fullmatch(r"[a-f0-9]{64}", metadata[key]):
                raise ValueError(f"核验指纹无效: {key}={metadata[key]!r}")
        for key in ("store_name", "requested_store_name", "shop_id", "site", "collected_at"):
            if not isinstance(metadata[key], str) or not metadata[key].strip():
                raise ValueError(f"核验字段无效: {key}={metadata[key]!r}")
        counts = [metadata[key] for key in ("original_row_count", "active_row_count", "excluded_row_count")]
        if any(type(value) is not int or value < 0 for value in counts) or counts[0] != counts[1] + counts[2]:
            raise ValueError(f"核验数量无效: {counts!r}")
        return metadata, bindings
    except (ValueError, TypeError, KeyError, IndexError) as exc:
        raise ActiveSourceError(f"Active 核验信息无效: {exc}") from exc


def annotate_source(path: Path, snapshot: ActiveListingSnapshot, *, requested_store_name: str) -> dict[str, Any]:
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    try:
        headers, records, indexes = _table(workbook)
        if SHEET in workbook.sheetnames or any(column in headers for column in TAG_COLUMNS):
            raise ActiveSourceError("下载原表包含保留的 Active 核验字段，拒绝覆盖")
        for record in records:
            store = clean_text(record.get("店铺名称"))
            site = clean_text(record.get("站点")).upper()
            site = SITE_TO_MARKETPLACE.get(site, site).lower()
            if site == "uk":
                site = "gb"
            if (store and store not in (snapshot.store_name, requested_store_name)) or (site and site != snapshot.site):
                invalid("MSKU 源表店铺/站点核验", "源表包含其他店铺或站点", record)
        tags = classify(records, snapshot.bindings)
        metadata = {
            "version": VERSION, "store_name": snapshot.store_name,
            "requested_store_name": requested_store_name, "shop_id": snapshot.shop_id,
            "site": snapshot.site, "collected_at": datetime.now(timezone.utc).isoformat(),
            "source_fingerprint": _digest([headers, records]),
            "original_row_count": len(records), "active_row_count": sum(tag[1] for tag in tags),
            "excluded_row_count": sum(not tag[1] for tag in tags),
        }
        metadata["snapshot_id"] = _digest([metadata, [asdict(item) for item in snapshot.bindings]])
        sheet = workbook.worksheets[0]
        for offset, column in enumerate(TAG_COLUMNS, len(headers) + 1):
            sheet.cell(1, offset, column)
        for index, tag in zip(indexes, tags, strict=True):
            for offset, value in enumerate(tag, len(headers) + 1):
                sheet.cell(index, offset, value)
        _write_info(workbook, metadata, snapshot.bindings)
        workbook.save(path)
        return metadata
    finally:
        workbook.close()


@dataclass(frozen=True)
class ActiveSource:
    headers: list[str]
    records: list[dict[str, Any]]
    bindings: tuple[ListingSkuBinding, ...]
    metadata: dict[str, Any]

    @property
    def counts(self) -> dict[str, int]:
        return {key: self.metadata[key] for key in ("original_row_count", "active_row_count", "excluded_row_count")}


def load_active_source(path: str | Path, *, store_name: str) -> ActiveSource:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        metadata, bindings = _read_info(workbook)
        if store_name not in (metadata["store_name"], metadata["requested_store_name"]):
            raise ActiveSourceError(f"Active 核验店铺不一致: 请求={store_name}, 快照={metadata['store_name']}")
        headers, records, _ = _table(workbook)
        if not all(column in headers for column in TAG_COLUMNS):
            raise ActiveSourceError("源表缺少 Active 标记列，请重新下载")
        raw_headers = [column for column in headers if column not in TAG_COLUMNS]
        raw_records = [{column: record.get(column) for column in raw_headers} for record in records]
        unsigned = {key: value for key, value in metadata.items() if key != "snapshot_id"}
        if (_digest([raw_headers, raw_records]) != metadata["source_fingerprint"]
                or _digest([unsigned, [asdict(item) for item in bindings]]) != metadata["snapshot_id"]):
            raise ActiveSourceError("源表或 Active 核验快照发生变化，请重新下载")
        tags = classify(raw_records, bindings)
        for record, expected in zip(records, tags, strict=True):
            actual = (record[TAG_COLUMNS[0]], record[TAG_COLUMNS[1]], clean_text(record[TAG_COLUMNS[2]]))
            if not isinstance(actual[1], bool) or actual != expected:
                raise ActiveSourceError("源表 Active 标记与核验快照不一致")
        active = [record for record, tag in zip(records, tags, strict=True) if tag[1]]
        if (len(records), len(active), len(records) - len(active)) != tuple(metadata[key] for key in ("original_row_count", "active_row_count", "excluded_row_count")):
            raise ActiveSourceError("Active 核验数量不一致")
        if not active:
            raise ActiveSourceError("无符合条件的记录：源表没有确认在售的 Active MSKU")
        return ActiveSource(headers, active, bindings, metadata)
    finally:
        workbook.close()


def stamp_report(path: str | Path, metadata: dict[str, Any]) -> None:
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    try:
        _write_info(workbook, metadata)
        workbook.save(path)
    finally:
        workbook.close()


def require_matching_reports(sales: Path, inventory: Path, *, store_name: str) -> dict[str, Any]:
    from openpyxl import load_workbook
    items = []
    for path in (sales, inventory):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            metadata, _ = _read_info(workbook)
            if store_name not in (metadata["store_name"], metadata["requested_store_name"]):
                raise ActiveSourceError(f"报表 Active 核验店铺不一致: {path}")
            items.append(metadata)
        finally:
            workbook.close()
    if items[0] != items[1]:
        raise ActiveSourceError("销量与库存报表的 Active 快照或源数据指纹不一致，请使用同一源表重跑")
    if not items[0]["active_row_count"]:
        raise ActiveSourceError("无符合条件的记录：报表没有确认在售的 Active MSKU")
    return items[0]
