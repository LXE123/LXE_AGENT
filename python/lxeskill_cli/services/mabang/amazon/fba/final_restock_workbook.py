"""Build the offline 35-column final restock recommendation workbook."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
import os
from pathlib import Path
import re
import subprocess
from tempfile import TemporaryDirectory
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill
from openpyxl.utils import column_index_from_string
from openpyxl.workbook.properties import CalcProperties

from services.mabang.amazon.fba.final_restock_recommendation import final_restock_formulas
from services.mabang.amazon.fba.final_restock_settings import read_state as read_final_restock_settings
from services.mabang.amazon.fba.source_verification import (
    SHEET as SOURCE_VERIFICATION_SHEET,
    load_verified_source,
    read_report_metadata,
)
from services.mabang.amazon.fba.store_msku_actual_inventory import (
    ACTUAL_INVENTORY_QUANTITY_COLUMN,
    SOURCE_FILE_RE,
)
from services.mabang.amazon.fba.store_msku_replenishment import (
    parse_weight_grams,
    validate_unlinked_shipments_snapshot_same_day,
)
from services.mabang.amazon.fba.combo_sku import normalize_sku_key
from services.mabang.amazon.fba.unlinked_shipments import load_unlinked_shipment_quantities
from shared.datasets import dataset_dir
from shared.office.execution import run_kit_bounded


WORKSHEET_NAME = "mskulist"
OUTPUT_FILENAME_SUFFIX = "补货最新.xlsx"
OFFICE_TIMEOUT_SECONDS = 300
MAX_DIAGNOSTIC_CHARACTERS = 2000
MAX_ERROR_EXAMPLES = 20

HEADERS = (
    "店铺名称", "MSKU", "父ASIN", "ASIN", "7天销量", "14天销量", "30天销量", "90天销量",
    "销量趋势", "加权日销", "可售", "待入库", "预留", "计划入库", "在途", "运营中心转运",
    "运营中心处理中", "FBA总库存（马帮）", "可销售天数", "未发出货件总计", "建议运输方式",
    "空运补货天数", "海运补货天数", "空运计划", "空运实际", "空运总重", "海运计划",
    "海运实际", "海运总重", "单件重量(g)", "本地SKU名称", "备注", "本地库存", "本地SKU", "可供天数",
)

SOURCE_NUMBER_COLUMNS = (
    "7天销量", "14天销量", "30天销量", "90天销量", "可售", "待入库", "预留", "计划入库",
    "在途", "待调仓", "调仓中",
)
SOURCE_REQUIRED_COLUMNS = (
    "MSKU", "父ASIN", "ASIN", "本地SKU", "本地SKU名称", "备注", "单品重量(g)(cm)", "本地库存",
    *SOURCE_NUMBER_COLUMNS,
)
OUTPUT_SOURCE_COLUMNS = {
    "7天销量": "7天销量",
    "14天销量": "14天销量",
    "30天销量": "30天销量",
    "90天销量": "90天销量",
    "可售": "可售",
    "待入库": "待入库",
    "预留": "预留",
    "计划入库": "计划入库",
    "在途": "在途",
    "待调仓": "运营中心转运",
    "调仓中": "运营中心处理中",
}
FORMULA_COLUMNS = tuple(final_restock_formulas(2))
EXCEL_ERROR_VALUES = {
    "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA",
}
OUTPUT_FILENAME_INVALID_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_FONT = Font(name="Calibri", size=11, bold=False)
_NO_FILL = PatternFill(fill_type=None)
_NO_BORDER = Border()


class FinalRestockWorkbookError(ValueError):
    """The offline inputs or generated workbook cannot be trusted."""


@dataclass(frozen=True)
class FinalRestockWorkbookResult:
    store_name: str
    source_data_time: str
    output_xlsx_path: str
    row_count: int
    formula_error_count: int
    formula_error_examples: tuple[str, ...]

    @property
    def success(self) -> bool:
        return self.formula_error_count == 0

    def to_payload(self) -> dict[str, Any]:
        payload: dict[str, Any] = {
            "success": self.success,
            "store_name": self.store_name,
            "source_data_time": self.source_data_time,
            "output_xlsx_path": self.output_xlsx_path,
            "row_count": self.row_count,
            "formula_error_count": self.formula_error_count,
            "formula_error_examples": list(self.formula_error_examples),
        }
        if not self.success:
            payload["exception"] = (
                f"Office 重算发现 {self.formula_error_count} 个 Excel 错误单元格；"
                f"实际错误示例: {', '.join(self.formula_error_examples)}"
            )
        return payload


def _text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def _path(path: str | Path, *, label: str) -> Path:
    value = Path(path).expanduser()
    if not value.is_file():
        raise FileNotFoundError(f"{label}不存在: {value}")
    return value.resolve(strict=True)


def _strict_number(value: Any, *, row_number: int, column: str) -> int | float:
    if value is None or isinstance(value, bool) or not _text(value):
        raise FinalRestockWorkbookError(
            f"MSKU 源表第 {row_number} 行 {column} 未知或为空，不能按 0 计算"
        )
    try:
        number = Decimal(_text(value).replace(",", ""))
    except (InvalidOperation, ValueError) as exc:
        raise FinalRestockWorkbookError(
            f"MSKU 源表第 {row_number} 行 {column} 不是有效数值: {value!r}"
        ) from exc
    if not number.is_finite():
        raise FinalRestockWorkbookError(
            f"MSKU 源表第 {row_number} 行 {column} 必须是有限数: {value!r}"
        )
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _optional_number(value: Any, *, context: str, allow_negative: bool = False) -> int | float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise FinalRestockWorkbookError(f"{context} 不是有效数值: {value!r}")
    try:
        number = Decimal(str(value).strip().replace(",", ""))
    except (InvalidOperation, ValueError) as exc:
        raise FinalRestockWorkbookError(f"{context} 不是有效数值: {value!r}") from exc
    if not number.is_finite() or (number < 0 and not allow_negative):
        requirement = "有限数" if allow_negative else "有限非负数"
        raise FinalRestockWorkbookError(f"{context} 必须是{requirement}: {value!r}")
    return int(number) if number == number.to_integral_value() else float(number)


def _source_time(source_path: Path) -> str:
    match = SOURCE_FILE_RE.match(source_path.name)
    if not match:
        raise FinalRestockWorkbookError(
            f"MSKU 源文件名无法识别源数据时间: {source_path.name}"
        )
    return match.group("source_time")


def _unique_source_rows(records: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    if not records:
        raise FinalRestockWorkbookError("MSKU 源文件没有数据行")
    by_key: dict[str, dict[str, Any]] = {}
    for index, record in enumerate(records, start=2):
        missing = [column for column in SOURCE_REQUIRED_COLUMNS if column not in record]
        if missing:
            source_field = "（马帮字段 uq122）" if "本地库存" in missing else ""
            raise FinalRestockWorkbookError(
                f"MSKU 源表 MSKU={_text(record.get('MSKU'))!r} 缺少列: {', '.join(missing)}{source_field}"
            )
        msku = _text(record.get("MSKU"))
        key = normalize_sku_key(msku)
        if not key:
            raise FinalRestockWorkbookError(f"MSKU 源表第 {index} 行没有 MSKU")
        if key in by_key:
            previous = _text(by_key[key].get("MSKU"))
            raise FinalRestockWorkbookError(
                f"MSKU 源表存在重复 MSKU: {previous!r}、{msku!r}"
            )
        by_key[key] = record
    return records, by_key


def _load_shenzhen_rows(
    report_path: Path,
    *,
    expected_keys: set[str],
) -> dict[str, str]:
    workbook = load_workbook(report_path, read_only=True, data_only=True)
    by_key: dict[str, str] = {}
    try:
        business_sheets = [
            sheet for sheet in workbook.worksheets
            if sheet.title != SOURCE_VERIFICATION_SHEET
        ]
        if not business_sheets:
            raise FinalRestockWorkbookError("深圳库存报告没有业务工作表")
        for sheet in business_sheets:
            rows = sheet.iter_rows(values_only=True)
            headers = [_text(value) for value in next(rows, ())]
            required = ("MSKU", ACTUAL_INVENTORY_QUANTITY_COLUMN, "本地SKU", "子SKU")
            missing = [column for column in required if column not in headers]
            if missing:
                raise FinalRestockWorkbookError(
                    f"深圳库存报告工作表 {sheet.title} 缺少列: {', '.join(missing)}"
                )
            indexes = {name: headers.index(name) for name in required}
            for row_number, values in enumerate(rows, start=2):
                cells = list(values)
                if not any(_text(value) for value in cells):
                    continue
                msku = _text(cells[indexes["MSKU"]] if indexes["MSKU"] < len(cells) else None)
                key = normalize_sku_key(msku)
                if not key:
                    raise FinalRestockWorkbookError(
                        f"深圳库存报告工作表 {sheet.title} 第 {row_number} 行有数据但没有 MSKU"
                    )
                if key in by_key:
                    raise FinalRestockWorkbookError(
                        f"深圳库存报告存在重复 MSKU: {msku!r}"
                    )
                actual = cells[indexes[ACTUAL_INVENTORY_QUANTITY_COLUMN]] if indexes[ACTUAL_INVENTORY_QUANTITY_COLUMN] < len(cells) else None
                # Placeholder rows remain valid report identities; quantities never supply AG.
                _optional_number(actual, context=f"深圳库存报告 MSKU={msku} 库存", allow_negative=True)
                local_sku = cells[indexes["本地SKU"]] if indexes["本地SKU"] < len(cells) else None
                by_key[key] = _text(local_sku)
    except FinalRestockWorkbookError:
        raise
    except Exception as exc:
        raise FinalRestockWorkbookError(
            f"读取深圳库存报告失败: {report_path}, error={type(exc).__name__}: {exc}"
        ) from exc
    finally:
        workbook.close()

    missing_keys = expected_keys - by_key.keys()
    extra_keys = by_key.keys() - expected_keys
    if missing_keys or extra_keys:
        raise FinalRestockWorkbookError(
            "深圳库存报告 MSKU 覆盖与源表不一致: "
            f"缺少 {len(missing_keys)} 行 {sorted(missing_keys)[:10]!r}; "
            f"多出 {len(extra_keys)} 行 {sorted(extra_keys)[:10]!r}"
        )
    return by_key


def _shipment_quantities(path: Path, *, store_name: str, source_data_time: str) -> dict[str, int | float]:
    validate_unlinked_shipments_snapshot_same_day(path, source_data_time)
    quantities = load_unlinked_shipment_quantities(path, store_name=store_name)
    normalized: dict[str, int | float] = {}
    original_for_key: dict[str, str] = {}
    for raw_msku, quantity in quantities.items():
        msku = _text(raw_msku)
        key = normalize_sku_key(msku)
        if not key:
            raise FinalRestockWorkbookError("未关联货件快照存在没有 MSKU 的汇总行")
        if key in normalized:
            raise FinalRestockWorkbookError(
                f"未关联货件快照 MSKU 规范化后重复: {original_for_key[key]!r}、{msku!r}"
            )
        normalized[key] = _optional_number(quantity, context=f"未关联货件快照 MSKU={msku} 数量") or 0
        original_for_key[key] = msku
    return normalized


def _rows_for_workbook(
    records: list[dict[str, Any]],
    *,
    store_name: str,
    inventory_by_key: dict[str, str],
    shipment_by_key: dict[str, int | float],
) -> list[dict[str, Any]]:
    _, source_by_key = _unique_source_rows(records)
    extra_shipments = shipment_by_key.keys() - source_by_key.keys()
    if extra_shipments:
        raise FinalRestockWorkbookError(
            f"未关联货件快照含源表没有的 MSKU: {sorted(extra_shipments)[:10]!r}"
        )
    result: list[dict[str, Any]] = []
    for source_row_number, record in enumerate(records, start=2):
        msku = _text(record.get("MSKU"))
        key = normalize_sku_key(msku)
        out: dict[str, Any] = {
            "店铺名称": store_name,
            "MSKU": msku,
            "父ASIN": _text(record.get("父ASIN")),
            "ASIN": _text(record.get("ASIN")),
            "本地SKU名称": _text(record.get("本地SKU名称")),
            "备注": _text(record.get("备注")),
        }
        for source_column in SOURCE_NUMBER_COLUMNS:
            out[OUTPUT_SOURCE_COLUMNS[source_column]] = _strict_number(
                record.get(source_column), row_number=source_row_number, column=source_column
            )
        raw_weight = record.get("单品重量(g)(cm)")
        # A treats blank weights as zero, but malformed nonblank data remains an error.
        if raw_weight is None or (isinstance(raw_weight, str) and not raw_weight.strip()):
            parsed_weight = 0
        else:
            try:
                number = Decimal(str(raw_weight).strip().replace(",", ""))
            except InvalidOperation:
                parsed_weight = parse_weight_grams(raw_weight)
            else:
                parsed_weight = float(number) if number.is_finite() and number >= 0 else None
        if isinstance(raw_weight, bool) or parsed_weight is None:
            raise FinalRestockWorkbookError(
                f"AD重量解析阶段失败: MSKU={msku}，源字段=单品重量(g)(cm)，"
                f"第{source_row_number}行，原值={raw_weight!r}；原因=无法解析为有效克数"
            )
        out["单件重量(g)"] = parsed_weight
        out["未发出货件总计"] = shipment_by_key.get(key, 0)
        report_local_sku = inventory_by_key[key]
        source_local_sku = _text(record.get("本地SKU"))
        inventory_context = f"AG本地库存(uq122)，MSKU={msku}，MSKU源表第{source_row_number}行"
        out["本地库存"] = _optional_number(
            record.get("本地库存"), context=inventory_context, allow_negative=True,
        )
        if out["本地库存"] is None:
            raise FinalRestockWorkbookError(f"{inventory_context} 缺失或为空，不能按0计算")
        out["本地SKU"] = report_local_sku or source_local_sku
        result.append(out)
    return result


def _write_workbook(
    path: Path, rows: list[dict[str, Any]], *, peak_boost_enabled: bool = True, peak_boost_days: int | None = None,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = WORKSHEET_NAME
    worksheet.append(list(HEADERS))
    for row_index, row in enumerate(rows, start=2):
        for column_index, header in enumerate(HEADERS, start=1):
            value = row.get(header)
            cell = worksheet.cell(row_index, column_index, value)
            if isinstance(value, str):
                # Platform text may start with '='; only the formulas below are executable.
                cell.data_type = "s"
        for column, formula in final_restock_formulas(
            row_index, peak_boost_enabled=peak_boost_enabled, peak_boost_days=peak_boost_days,
        ).items():
            worksheet[f"{column}{row_index}"] = formula
    for row in worksheet.iter_rows():
        for cell in row:
            cell.font = _FONT
            cell.fill = _NO_FILL
            cell.border = _NO_BORDER
    for row_index in range(2, worksheet.max_row + 1):
        worksheet[f"J{row_index}"].number_format = "0.0000"
        worksheet[f"S{row_index}"].number_format = "0.00"
        worksheet[f"AI{row_index}"].number_format = "0.00"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:AI{worksheet.max_row}"
    worksheet.sheet_view.zoomScale = 85
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    try:
        workbook.save(path)
    finally:
        workbook.close()


def _office_detail(value: Any) -> str:
    text = str(value or "").strip()
    return text[:MAX_DIAGNOSTIC_CHARACTERS] if text else "Office Kit 未返回诊断文本"


def _recalculate(source: Path, output: Path) -> None:
    arguments = ["recalculate", "--input", str(source), "--output", str(output)]
    try:
        result = run_kit_bounded(arguments, timeout_seconds=OFFICE_TIMEOUT_SECONDS)
    except subprocess.TimeoutExpired as exc:
        raise FinalRestockWorkbookError(
            f"Office 重算超时（{OFFICE_TIMEOUT_SECONDS} 秒）: {_office_detail(exc.stderr or exc.stdout or exc)}"
        ) from exc
    except OSError as exc:
        raise FinalRestockWorkbookError(
            f"Office 重算执行异常: {type(exc).__name__}: {_office_detail(exc)}"
        ) from exc
    if result.returncode != 0:
        raise FinalRestockWorkbookError(
            f"Office 重算失败（退出码 {result.returncode}）: "
            f"{_office_detail(result.stderr or result.stdout)}"
        )
    if not output.is_file() or output.stat().st_size == 0:
        raise FinalRestockWorkbookError(
            f"Office 报告成功但没有生成有效文件: {output}; "
            f"输出: {_office_detail(result.stdout)}"
        )


def _inspect_recalculated(path: Path, *, rows: list[dict[str, Any]]) -> tuple[int, tuple[str, ...]]:
    formulas = load_workbook(path, read_only=True, data_only=False)
    values = load_workbook(path, read_only=True, data_only=True)
    try:
        expected_rows = len(rows)
        if formulas.sheetnames != [WORKSHEET_NAME] or values.sheetnames != [WORKSHEET_NAME]:
            raise FinalRestockWorkbookError(
                f"重算后工作表结构错误: {formulas.sheetnames!r}"
            )
        formula_sheet = formulas[WORKSHEET_NAME]
        value_sheet = values[WORKSHEET_NAME]
        if (
            formula_sheet.max_column != len(HEADERS)
            or formula_sheet.max_row != expected_rows + 1
            or value_sheet.max_column != len(HEADERS)
            or value_sheet.max_row != expected_rows + 1
        ):
            raise FinalRestockWorkbookError(
                "重算后工作簿行列数量与源表不一致: "
                f"actual={formula_sheet.max_row}x{formula_sheet.max_column}, "
                f"expected={expected_rows + 1}x{len(HEADERS)}"
            )
        error_count = 0
        examples: list[str] = []
        formula_indexes = {column: column_index_from_string(column) - 1 for column in FORMULA_COLUMNS}
        for row_number, (formula_row, value_row) in enumerate(
            zip(formula_sheet.iter_rows(), value_sheet.iter_rows(), strict=True), start=1
        ):
            if row_number > 1:
                if value_row[8].value is None:
                    raise FinalRestockWorkbookError(
                        f"Office 重算后 I{row_number} 没有公式结果，不能确认重算已执行"
                    )
                for column, index in formula_indexes.items():
                    formula_cell = formula_row[index]
                    formula = formula_cell.value
                    if formula_cell.data_type != "f" or not isinstance(formula, str) or not formula.startswith("="):
                        raise FinalRestockWorkbookError(
                            f"重算后公式丢失或改变为非公式单元格: {column}{row_number}={formula!r}"
                        )
            for formula_cell, cell in zip(formula_row, value_row, strict=True):
                if formula_cell.data_type == "f" and (row_number == 1 or formula_cell.column_letter not in FORMULA_COLUMNS):
                    raise FinalRestockWorkbookError(
                        f"重算后非计算单元格包含公式: {formula_cell.coordinate}"
                    )
                value = cell.value
                if cell.data_type == "e" or (isinstance(value, str) and value in EXCEL_ERROR_VALUES):
                    error_count += 1
                    if len(examples) < MAX_ERROR_EXAMPLES:
                        examples.append(f"{cell.coordinate}={value}")
        return error_count, tuple(examples)
    finally:
        formulas.close()
        values.close()


def _output_path(store_name: str, *, output_dir: str | Path | None) -> Path:
    clean_name = OUTPUT_FILENAME_INVALID_CHARS.sub("_", store_name).rstrip(" .")
    if not clean_name:
        raise FinalRestockWorkbookError("店铺名不能用于输出文件名")
    directory = (
        Path(output_dir)
        if output_dir is not None
        else dataset_dir("replenish_final_restock_recommendation")
    )
    directory.mkdir(parents=True, exist_ok=True)
    return directory / f"{date.today():%Y%m%d}{clean_name}{OUTPUT_FILENAME_SUFFIX}"


def generate_final_restock_recommendation(
    store_name: str,
    *,
    msku_source_xlsx_path: str | Path,
    shenzhen_inventory_report_xlsx_path: str | Path,
    unlinked_shipments_snapshot_path: str | Path,
    output_dir: str | Path | None = None,
    peak_boost_enabled: bool = False,
    peak_boost_days: int | None = None,
) -> FinalRestockWorkbookResult:
    """Validate same-batch offline inputs, then write and recalculate one workbook."""
    clean_store_name = _text(store_name)
    if not clean_store_name:
        raise FinalRestockWorkbookError("store_name 不能为空")
    settings = read_final_restock_settings()
    # A chat request is the authority for this run, never a historical saved switch.
    # Validate before touching inputs; the business formula function owns value bounds.
    final_restock_formulas(2, peak_boost_enabled=peak_boost_enabled, peak_boost_days=peak_boost_days)
    settings = {"peak_boost_enabled": peak_boost_enabled,
                "peak_boost_days": settings["peak_boost_days"] if peak_boost_days is None else peak_boost_days}
    source_path = _path(msku_source_xlsx_path, label="MSKU 源文件")
    inventory_path = _path(shenzhen_inventory_report_xlsx_path, label="深圳库存报告")
    snapshot_path = _path(unlinked_shipments_snapshot_path, label="未关联货件快照")

    verified_source = load_verified_source(source_path, store_name=clean_store_name)
    inventory_metadata = read_report_metadata(inventory_path)
    if inventory_metadata != verified_source.metadata:
        raise FinalRestockWorkbookError(
            "深圳库存报告与本次 MSKU 源文件核验批次不一致；请使用同一源文件重新导出"
        )
    source_data_time = _source_time(source_path)
    records, source_by_key = _unique_source_rows(verified_source.records)
    inventory_by_key = _load_shenzhen_rows(
        inventory_path,
        expected_keys=set(source_by_key),
    )
    shipment_by_key = _shipment_quantities(
        snapshot_path,
        store_name=clean_store_name,
        source_data_time=source_data_time,
    )
    rows = _rows_for_workbook(
        records,
        store_name=clean_store_name,
        inventory_by_key=inventory_by_key,
        shipment_by_key=shipment_by_key,
    )

    output_path = _output_path(clean_store_name, output_dir=output_dir)
    with TemporaryDirectory(prefix="final-restock-", dir=output_path.parent) as temporary_directory:
        temporary = Path(temporary_directory)
        formula_path = temporary / "final-restock-formulas.xlsx"
        recalculated_path = temporary / "final-restock-recalculated.xlsx"
        _write_workbook(formula_path, rows, **settings)
        _recalculate(formula_path, recalculated_path)
        formula_error_count, formula_error_examples = _inspect_recalculated(
            recalculated_path,
            rows=rows,
        )
        published_path = (
            output_path.with_name(f"{output_path.stem}-公式错误诊断{output_path.suffix}")
            if formula_error_count else output_path
        )
        os.replace(recalculated_path, published_path)

    return FinalRestockWorkbookResult(
        store_name=clean_store_name,
        source_data_time=source_data_time,
        output_xlsx_path=str(published_path),
        row_count=len(rows),
        formula_error_count=formula_error_count,
        formula_error_examples=formula_error_examples,
    )


__all__ = [
    "FinalRestockWorkbookError",
    "FinalRestockWorkbookResult",
    "generate_final_restock_recommendation",
]
