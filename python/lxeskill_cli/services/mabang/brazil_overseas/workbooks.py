"""Validate without changing a single byte of the platform workbook."""
import io
import os
from pathlib import Path
from zipfile import ZipFile, is_zipfile
import xlrd
from openpyxl import load_workbook
from .contracts import ALLOCATION_EXPORT_FIELDS, WAREHOUSE_NAME
from .errors import BrazilError
from .legacy_xls import workbook_stream

INVENTORY_HEADERS = {'库存SKU编号', '仓库', '仓位', '销量(7)', '销量(28)', '销量(42)', '仓位库存', '可用库存量'}
XLS_SIGNATURE = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'


def read_workbook(body, extension):
    notices = []
    if extension == 'xls' and body.startswith(XLS_SIGNATURE):
        try:
            wb = xlrd.open_workbook(file_contents=body, on_demand=True, formatting_info=True, logfile=io.StringIO())
        except xlrd.compdoc.CompDocError as exc:
            if not str(exc).startswith('Workbook corruption: seen[') or not str(exc).endswith(' == 4'):
                raise BrazilError('invalid_workbook', f'XLS 结构解析失败: {type(exc).__name__}: {exc}') from exc
            stream = workbook_stream(body)
            wb = xlrd.open_workbook(file_contents=stream, formatting_info=True, logfile=io.StringIO())
            notices.append('平台 XLS 短流尾链兼容读取；已校验流边界和 Workbook，原字节不变')
        try:
            if wb.nsheets != 1:
                raise BrazilError('invalid_sheets', f'实际工作表数 {wb.nsheets}，预期 1')
            sheet = wb.sheet_by_index(0)
            rows = [sheet.row_values(i) for i in range(sheet.nrows)]
            for row in range(sheet.nrows):
                if any(c.ctype == xlrd.XL_CELL_ERROR for c in sheet.row(row)):
                    raise BrazilError('invalid_cell',f'XLS 第 {row+1} 行包含错误单元格')
            # Expand only actual merge ranges in the validation view, never in the file.
            for r0,r1,c0,c1 in sheet.merged_cells:
                if r0 < 1 or r1 > len(rows) or c1 > sheet.ncols:
                    raise BrazilError('invalid_merge','原表合并单元格范围异常')
                for ri in range(r0,r1):
                    for ci in range(c0,c1):
                        rows[ri][ci] = rows[r0][c0]
            name = sheet.name
        finally:
            wb.release_resources()
    elif extension == 'xlsx' and is_zipfile(io.BytesIO(body)):
        with ZipFile(io.BytesIO(body)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 250_000_000 or archive.testzip():
                raise BrazilError('invalid_workbook', 'XLSX 解压大小超限或 ZIP 校验失败')
        wb = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
        try:
            if len(wb.sheetnames) != 1:
                raise BrazilError('invalid_sheets', f'实际工作表数 {len(wb.sheetnames)}，预期 1')
            rows = []
            for row in wb.active.iter_rows():
                if any(c.data_type in {'f','e'} for c in row):
                    raise BrazilError('invalid_cell', '原表含公式或错误，无法核对源值')
                rows.append([c.value for c in row])
            name = wb.active.title
        finally:
            wb.close()
    else:
        raise BrazilError('invalid_workbook', f'下载内容不是真实 {extension.upper()}；响应预览: {body.decode("utf-8",errors="replace")}')
    if not rows:
        raise BrazilError('missing_headers', '导出文件没有表头')
    headers = rows[0]
    if any(not isinstance(v,str) or not v.strip() for v in headers) or len(set(headers)) != len(headers):
        raise BrazilError('invalid_headers', f'原表头为空或重复: {headers}')
    records = [dict(zip(headers,row,strict=True)) for row in rows[1:] if any(v not in ('',None) for v in row)]
    return headers, records, name, notices


def validate(body, report, *, expected_total=None, expected_codes=None, sample_skus=()):
    extension = 'xlsx' if report == 'inventory_sales_snapshot' else 'xls'
    headers, rows, sheet, notices = read_workbook(body, extension)
    required = INVENTORY_HEADERS if extension == 'xlsx' else {n for n,_ in ALLOCATION_EXPORT_FIELDS}
    if not required.issubset(headers):
        raise BrazilError('missing_headers', f'实际表头 {headers}; 缺少 {sorted(required-set(headers))}')
    warehouse_column = '仓库' if extension == 'xlsx' else '目标仓库'
    if any(str(row[warehouse_column]).strip() != WAREHOUSE_NAME for row in rows):
        raise BrazilError('warehouse_mismatch', '导出包含非巴西海外仓或空白仓库')
    if extension == 'xlsx':
        if len(rows) != expected_total:
            raise BrazilError('coverage_mismatch', f'库存列表 {expected_total} 条，原表 {len(rows)} 行，无法确认完整覆盖')
        if not set(sample_skus).issubset({str(r['库存SKU编号']).strip() for r in rows}):
            raise BrazilError('coverage_mismatch', '库存列表样本 SKU 未全部出现在原表')
    else:
        actual = {str(r['批次编号']).strip() for r in rows}
        expected = set(expected_codes or ())
        if actual != expected:
            raise BrazilError('coverage_mismatch', f'调拨批次覆盖不符: 缺少 {len(expected-actual)}，多出 {len(actual-expected)}')
    return {'validation_notices':notices, 'row_count':len(rows), 'headers':headers, 'sheet_names':[sheet], 'extension':extension,
            'warehouses': sorted({str(r[warehouse_column]).strip() for r in rows})}


def publish(path: Path, body: bytes):
    temporary = path.with_name('.'+path.name+'.tmp')
    try:
        with temporary.open('xb') as stream:
            stream.write(body)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary,path)
    finally:
        temporary.unlink(missing_ok=True)
