from __future__ import annotations

import os
import re
import tempfile
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import is_zipfile

import xlrd
from openpyxl import Workbook, load_workbook

from .client import ZhihuiTmsClient
from .errors import ZhihuiTmsError, safe_payload_text
from .product_export import ZhihuiTmsExportResult


_XLS_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_DATE_LABEL_PATTERN = re.compile(r"^\d{8}$")
_DOWNLOAD_URL_KEYS = ("url", "downloadUrl", "fileUrl", "path")


class ZhihuiTmsDeliveryError(ZhihuiTmsError):
    """A downloaded workbook cannot be safely delivered."""


@dataclass(frozen=True)
class ZhihuiTmsArtifact:
    path: str
    kind: str
    page: int | None
    total_pages: int


@dataclass(frozen=True)
class ZhihuiTmsDeliveryResult:
    artifacts: tuple[ZhihuiTmsArtifact, ...]
    page_artifacts: tuple[ZhihuiTmsArtifact, ...]
    merged_artifact: ZhihuiTmsArtifact | None
    headers: tuple[str, ...]
    total_rows: int


@dataclass(frozen=True)
class _WorkbookRows:
    headers: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]


def _delivery_error(code: str, message: str, *, payload: Any = None) -> ZhihuiTmsDeliveryError:
    return ZhihuiTmsDeliveryError(code, message, payload=payload)


def _validate_date_label(value: str) -> str:
    label = str(value or "").strip()
    if not _DATE_LABEL_PATTERN.fullmatch(label):
        raise ValueError("date_label 必须是 YYYYMMDD")
    try:
        datetime.strptime(label, "%Y%m%d")
    except ValueError as exc:
        raise ValueError("date_label 必须是有效的 YYYYMMDD 日期") from exc
    return label


def _resolve_output_dir(output_dir: str | Path) -> Path:
    path = Path(output_dir).expanduser().resolve()
    path.mkdir(parents=True, exist_ok=True)
    return path


def _extract_download_url(pop: Any) -> str:
    if isinstance(pop, str):
        url = pop.strip()
    elif isinstance(pop, Mapping):
        url = ""
        for key in _DOWNLOAD_URL_KEYS:
            candidate = pop.get(key)
            if isinstance(candidate, str) and candidate.strip():
                url = candidate.strip()
                break
    else:
        url = ""
    if not url:
        raise _delivery_error(
            "tms_download_url_missing",
            f"智汇 TMS 导出响应缺少可下载地址: {safe_payload_text(pop)}",
            payload=pop,
        )
    return url


def _trim_row(values: Sequence[Any]) -> tuple[Any, ...]:
    row = list(values)
    while row and row[-1] in (None, ""):
        row.pop()
    return tuple(row)


def _normalize_rows(raw_rows: list[tuple[Any, ...]], *, source: str) -> _WorkbookRows:
    if not raw_rows:
        raise _delivery_error("tms_workbook_empty", f"智汇 TMS 文件没有表头: {source}")
    header_values = _trim_row(raw_rows[0])
    if not header_values or all(value in (None, "") for value in header_values):
        raise _delivery_error("tms_workbook_header_missing", f"智汇 TMS 文件表头为空: {source}")
    headers = tuple(str(value).strip() for value in header_values)
    if any(not value for value in headers):
        raise _delivery_error("tms_workbook_header_invalid", f"智汇 TMS 文件表头包含空列: {source}")

    rows: list[tuple[Any, ...]] = []
    for raw_row in raw_rows[1:]:
        row = _trim_row(raw_row)
        if len(row) > len(headers):
            raise _delivery_error(
                "tms_workbook_row_wider_than_header",
                f"智汇 TMS 文件数据列数超过表头: {source}",
            )
        rows.append(row + (None,) * (len(headers) - len(row)))
    return _WorkbookRows(headers=headers, rows=tuple(rows))


def _read_xlsx(content: bytes, *, source: str) -> _WorkbookRows:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            worksheet = workbook.active
            raw_rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    except Exception as exc:
        raise _delivery_error(
            "tms_xlsx_invalid",
            f"智汇 TMS XLSX 无法读取: {source}: {exc}",
        ) from exc
    return _normalize_rows(raw_rows, source=source)


def _read_xls(content: bytes, *, source: str) -> _WorkbookRows:
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            if workbook.nsheets < 1:
                raise ValueError("workbook has no sheets")
            worksheet = workbook.sheet_by_index(0)
            raw_rows = [tuple(worksheet.row_values(index)) for index in range(worksheet.nrows)]
        finally:
            workbook.release_resources()
    except Exception as exc:
        raise _delivery_error(
            "tms_xls_invalid",
            f"智汇 TMS XLS 无法读取: {source}: {exc}",
        ) from exc
    return _normalize_rows(raw_rows, source=source)


def _read_workbook(content: bytes, *, source: str) -> _WorkbookRows:
    if content.startswith(_XLS_SIGNATURE):
        return _read_xls(content, source=source)
    if is_zipfile(BytesIO(content)):
        return _read_xlsx(content, source=source)
    raise _delivery_error(
        "tms_workbook_signature_invalid",
        f"智汇 TMS 文件头不是有效的 XLS/XLSX: {source}",
    )


def _atomic_write_workbook(path: Path, workbook_rows: _WorkbookRows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb",
            prefix=f".{path.name}.",
            suffix=".tmp",
            dir=path.parent,
            delete=False,
        ) as temporary:
            temp_path = Path(temporary.name)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "商品"
        worksheet.append(list(workbook_rows.headers))
        for row in workbook_rows.rows:
            worksheet.append(list(row))
        workbook.save(temp_path)
        workbook.close()
        os.replace(temp_path, path)
        temp_path = None
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


def _merged_rows(headers: tuple[str, ...], pages: Sequence[_WorkbookRows]) -> _WorkbookRows:
    return _WorkbookRows(
        headers=headers,
        rows=tuple(row for workbook_rows in pages for row in workbook_rows.rows),
    )


def _write_merged_artifact(
    destination: Path,
    *,
    date_label: str,
    headers: tuple[str, ...],
    pages: Sequence[_WorkbookRows],
    total_pages: int,
    partial: bool,
) -> tuple[ZhihuiTmsArtifact, int]:
    workbook_rows = _merged_rows(headers, pages)
    prefix = "部分合并" if partial else "合并"
    path = (destination / f"智汇tms-商品-{prefix}-{date_label}.xlsx").resolve()
    _atomic_write_workbook(path, workbook_rows)
    return (
        ZhihuiTmsArtifact(
            path=str(path),
            kind="merged_partial" if partial else "merged",
            page=None,
            total_pages=total_pages,
        ),
        len(workbook_rows.rows),
    )


def _attach_partial_delivery(
    error: Exception,
    *,
    destination: Path,
    date_label: str,
    headers: tuple[str, ...] | None,
    pages: Sequence[_WorkbookRows],
    total_pages: int,
) -> None:
    """Attach one validated partial workbook when a later page fails.

    The original exception stays authoritative. A failure writing the optional
    partial artifact must never hide the observed platform error.
    """
    if headers is None or not pages or not any(workbook_rows.rows for workbook_rows in pages):
        return
    try:
        artifact, rows = _write_merged_artifact(
            destination,
            date_label=date_label,
            headers=headers,
            pages=pages,
            total_pages=total_pages,
            partial=True,
        )
    except Exception:  # noqa: BLE001 - preserve the original platform failure
        return
    setattr(error, "partial_artifacts", (artifact,))
    setattr(error, "partial_pages", len(pages))
    setattr(error, "partial_rows", rows)


def deliver_product_exports(
    client: ZhihuiTmsClient,
    export_result: ZhihuiTmsExportResult,
    *,
    output_dir: str | Path,
    date_label: str,
    on_event: Callable[[dict[str, Any]], None] | None = None,
) -> ZhihuiTmsDeliveryResult:
    label = _validate_date_label(date_label)
    destination = _resolve_output_dir(output_dir)
    total_pages = len(export_result.pages)
    if total_pages == 0:
        return ZhihuiTmsDeliveryResult(
            artifacts=(),
            page_artifacts=(),
            merged_artifact=None,
            headers=(),
            total_rows=0,
        )

    page_rows: list[_WorkbookRows] = []
    expected_headers: tuple[str, ...] | None = None
    try:
        for page in export_result.pages:
            url = _extract_download_url(page.pop)
            content, _content_type = client.download_bytes(url)
            workbook_rows = _read_workbook(content, source=url)
            if expected_headers is None:
                expected_headers = workbook_rows.headers
            elif workbook_rows.headers != expected_headers:
                raise _delivery_error(
                    "tms_workbook_header_mismatch",
                    f"智汇 TMS 第{page.page}页表头与前页不一致: {workbook_rows.headers!r} != {expected_headers!r}",
                    payload={"page": page.page, "headers": workbook_rows.headers},
                )
            page_rows.append(workbook_rows)
            if on_event is not None:
                on_event({
                    "stage": "downloaded", "page": page.page,
                    "total_pages": total_pages, "rows": len(workbook_rows.rows),
                })
    except Exception as exc:  # noqa: BLE001 - retain the observed download/schema failure
        _attach_partial_delivery(
            exc,
            destination=destination,
            date_label=label,
            headers=expected_headers,
            pages=page_rows,
            total_pages=total_pages,
        )
        raise

    assert expected_headers is not None
    merged_artifact, total_rows = _write_merged_artifact(
        destination,
        date_label=label,
        headers=expected_headers,
        pages=page_rows,
        total_pages=total_pages,
        partial=False,
    )
    if on_event is not None:
        on_event({"stage": "merged", "total_pages": total_pages, "rows": total_rows})
    return ZhihuiTmsDeliveryResult(
        artifacts=(merged_artifact,),
        page_artifacts=(),
        merged_artifact=merged_artifact,
        headers=expected_headers,
        total_rows=total_rows,
    )


__all__ = [
    "ZhihuiTmsArtifact",
    "ZhihuiTmsDeliveryError",
    "ZhihuiTmsDeliveryResult",
    "deliver_product_exports",
]
